"""계산 결과를 표(원가계산서·내역서·일위대가·단가대비표)로 바꾸고 엑셀로 만든다.

화면과 엑셀이 같은 표를 쓰도록 응답 dict에서 한 번만 만든다.
"""

from __future__ import annotations

import io
import re
from datetime import date
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

DISCLAIMER = "표준품셈 기준 금액, 검토 전 참고용"
_RANK = {"재료비": 0, "노무비": 1, "경비": 2}
_SUBTOTALS = ("노무비 계", "경비 계")


def _rate_period(version: dict, equipment: bool = False) -> str:
    if not version:
        return ""
    if equipment:
        return str(version.get("version", ""))
    title = version.get("title", "")
    year = re.search(r"20\d{2}", title)
    if year and ("상반기" in title or "하반기" in title):
        half = "상반기" if "상반기" in title else "하반기"
        return f"{year.group()} {half}"
    return " ~ ".join(part for part in (version.get("effective_from"), version.get("effective_to")) if part)


def _korean_date(value: str | None) -> str:
    if not value:
        return ""
    match = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})", value)
    return f"{int(match.group(1))}.{int(match.group(2))}.{int(match.group(3))}." if match else value


def _num(value) -> int | float | None:
    if value is None or value == "":
        return None
    number = Decimal(str(value))
    return int(number) if number == number.to_integral_value() else float(number)


def _basis_text(line: dict) -> str:
    base, rate = line.get("base") or "", line.get("rate")
    if rate is not None and line.get("status") == "산정":
        return f"{base} {line['base_amount']:,}원 × {rate}%"
    return base


def statement_rows(statement: dict | None) -> list[dict]:
    """비목별로 묶어 재료비 → 노무비 → 경비 → 순공사원가 … 도급액 순으로 늘어놓는다."""
    if not statement:
        return []
    indexed = list(enumerate(statement.get("lines", [])))

    def key(item: tuple[int, dict]):
        index, line = item
        category = line.get("category")
        if line["name"] in _SUBTOTALS:
            return (_RANK[category], 1, index)
        if category in _RANK:
            return (_RANK[category], 0, index)
        return (3, 0, index) if category else (4, 0, index)

    rows = []
    for _, line in sorted(indexed, key=key):
        status = line.get("status", "산정")
        is_total = line["name"] in ("순공사원가", "총원가", "도급액")
        rows.append({
            "name": line["name"], "category": line.get("category"),
            "basis": _basis_text(line), "amount": line.get("amount") if status != "미산정" else None,
            "status": status, "reason": line.get("reason"), "note": line.get("note"),
            "kind": "subtotal" if line["name"] in _SUBTOTALS else ("total" if is_total else "item"),
            "final": line["name"] == "도급액",
        })
    return rows


def _spec_text(inputs: list[dict]) -> str:
    values = {item["name"]: item["value"] for item in inputs}
    parts = []
    if "structure" in values:
        parts.append(values["structure"])
    if "slump_band" in values:
        parts.append(values["slump_band"])
    if "facility_type" in values or "site_type" in values:
        parts.append(f"{values.get('facility_type', '')}/{values.get('site_type', '').replace('Type-', '')}")
    if "pump_size" in values:
        boom = f"{values['pump_size']} {values.get('placement', '')}".strip()
        parts.append(boom)
    return "·".join(part for part in parts if part)


def bill_row(response: dict) -> dict | None:
    priced = response.get("priced")
    if not priced or not priced.get("reference_amounts"):
        return None
    reference = priced["reference_amounts"]
    unit_prices, amounts = priced.get("unit_prices") or {}, reference.get("subtotals") or {}
    work = response.get("work") or {}
    return {
        "name": f"{work.get('title', '')}({work.get('section_no', '')})" if work else "",
        "spec": _spec_text(response.get("inputs", [])), "unit": "㎥", "quantity": reference.get("volume"),
        "unit_price": {name: unit_prices.get(name) for name in ("재료비", "노무비", "경비")},
        "amount": {name: amounts.get(name) for name in ("재료비", "노무비", "경비")},
        "total": reference.get("total"), "partial": bool(priced.get("partial")),
    }


def unit_price_rows(priced: dict | None) -> list[dict]:
    if not priced:
        return []
    rows = []
    for key in ("lines", "equipment_lines", "cost_lines", "supply_lines"):
        for line in priced.get(key, []):
            if line.get("kind") == "equipment":
                continue  # 장비 합계 줄은 아래 구성 줄(손료·운전원)과 겹친다
            amount = line.get("amount")
            kind = line.get("kind")
            base = line.get("base")
            rate = line.get("rate")
            if kind == "rate_cost":
                unit, quantity = "식", 1
                unit_price = amount
                spec = f"노무비의 {Decimal(str(rate)) * 100:g}%" if rate is not None else ""
            else:
                unit, quantity = line.get("unit") or "", line.get("quantity")
                if "/㎥" in unit:
                    unit = unit.replace("/㎥", "")
                unit_price = line.get("unit_price")
                spec = line.get("machine_spec") or (
                    f"노임 코드 {line['rate_code']}" if line.get("rate_code") else "")
                if base and rate is not None:
                    spec = f"{base} {rate}%"
            rows.append({"category": line.get("category"), "name": line.get("name"),
                         "spec": spec, "unit": unit, "quantity": quantity,
                         "unit_price": unit_price, "amount": amount,
                         "status": line.get("status") or ("산정" if amount is not None else "미산정"),
                         "reason": line.get("reason"), "citations": line.get("citations", [])})
    return rows


def rate_rows(priced: dict | None, statement: dict | None,
              conditions: list[dict] | None = None) -> list[dict]:
    """쓴 단가 목록: 노임, 기계경비, 제비율 버전, 관급 자재."""
    if not priced:
        return []
    labor_version = priced.get("rate_version") or {}
    equipment_version = priced.get("equipment_rate_version") or {}
    labor_source = f"{labor_version.get('title', '')}({labor_version.get('publisher', '')})"
    rows = []
    for line in priced.get("lines", []) + priced.get("equipment_lines", []):
        if line.get("unit_price") is None:
            continue
        if line.get("kind") == "labor":
            rows.append({"kind": "노임", "name": line["name"], "spec": line["rate_code"],
                         "unit": labor_version.get("unit", "원/일"), "price": line["unit_price"],
                         "source": labor_source, "period": _rate_period(labor_version)})
        elif line.get("machine_code") and line.get("rate_code"):
            rows.append({"kind": "노임", "name": line["name"], "spec": line["rate_code"], "unit": "원/hr",
                         "price": line["unit_price"], "source": labor_source + " 일 노임 ÷ 8시간 환산",
                         "period": _rate_period(labor_version)})
        elif line.get("machine_code"):
            rows.append({"kind": "기계경비", "name": line["name"], "spec": line.get("machine_spec", ""),
                         "unit": "원/hr", "price": line["unit_price"],
                         "source": f"{equipment_version.get('title', '')}({equipment_version.get('publisher', '')})",
                         "period": _rate_period(equipment_version, equipment=True)})
    version = (statement or {}).get("overhead_version")
    if version:
        group = next((item.get("group") for item in (conditions or [])
                      if item.get("name") == "work_category"), "")
        rows.append({"kind": "제비율", "name": "조달청 제비율", "spec": group, "unit": "",
                     "price": None, "source": ", ".join(version.get("source_files", [])),
                     "period": version.get("effective_from", "")})
    for line in priced.get("supply_lines", []):
        rows.append({"kind": "자재", "name": line["name"], "spec": line.get("status", ""),
                     "unit": "", "price": None, "source": line.get("reason") or "",
                     "period": ""})
    return rows


def add_tables(response: dict) -> dict:
    return {"statement_rows": statement_rows(response.get("statement")), "bill": bill_row(response),
            "unit_rows": unit_price_rows(response.get("priced")),
            "rate_rows": rate_rows(response.get("priced"), response.get("statement"), response.get("conditions"))}


def _cost_statement_rows(tables: dict, statement: dict) -> list[list]:
    """원가계산서의 6개 열을 항상 같은 위치에 채운다."""
    totals = statement.get("totals", {})
    contract_amount = totals.get("contract_amount")
    statement_rows = tables.get("statement_rows", [])
    by_name = {row["name"]: row for row in statement_rows}
    output = []

    def add(major: str, minor: str, amount, basis: str = "", status: str = "산정") -> None:
        excluded = status in ("제외", "미산정")
        value = None if excluded else _num(amount)
        ratio = None
        if value is not None and contract_amount:
            ratio = float(Decimal(str(value)) / Decimal(str(contract_amount)))
        output.append([major, minor, value, ratio, basis, status if excluded else ""])

    for category in ("재료비", "노무비", "경비"):
        lines = [row for row in statement_rows if row.get("category") == category]
        for row in lines:
            minor = ("소계" if row["name"] in _SUBTOTALS else
                     "직접재료비" if category == "재료비" and row["name"] == "재료비" else row["name"])
            basis = row.get("basis") or row.get("reason") or row.get("note") or ""
            add(category, minor, row.get("amount"), basis, row.get("status", "산정"))
        if category == "재료비" and lines:
            add(category, "소계", totals.get("materials"), "직접재료비 및 재료비 항목 합계")

    summary_names = ("순공사원가", "일반관리비", "이윤", "총원가", "부가가치세", "도급액")
    for name in summary_names:
        row = by_name.get(name)
        if row:
            major = "합계" if row.get("category") == "합계" else (row.get("category") or "")
            basis = row.get("basis") or row.get("reason") or row.get("note") or ""
            if name == "부가가치세" and totals.get("total_cost") is not None:
                basis = f"총원가 {int(totals['total_cost']):,}원 × 10%"
            add(major, name, row.get("amount"), basis, row.get("status", "산정"))

    known = {row["name"] for row in statement_rows if row.get("category") in ("재료비", "노무비", "경비")}
    known.update(summary_names)
    for row in statement_rows:
        if row["name"] in known:
            continue
        status = row.get("status", "산정")
        add(row.get("category") or "경비", row["name"], row.get("amount"),
            row.get("basis") or row.get("reason") or row.get("note") or "", status)
    return output


# ---------- 엑셀 ----------

_FONT_NAME = "맑은 고딕"
_HEAD = PatternFill("solid", fgColor="E7E6E6")
_STRONG = PatternFill("solid", fgColor="FFF2CC")
_MUTED = PatternFill("solid", fgColor="F2F2F2")
_THIN = Side(style="thin", color="B7B7B7")
_DOUBLE = Side(style="double", color="595959")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_MONEY = "#,##0"
_DECIMAL = "#,##0.####"
_FOOTER = "표준품셈 기준 금액 · 검토 전 참고용 · 품셈AI"


def _work_name(response: dict) -> str:
    title = (response.get("work") or {}).get("title") or "공사비"
    return re.sub(r"\s*\(20\d{2}(?:년)?[^)]*\)\s*$", "", title).strip()


def estimate_filename(response: dict) -> str:
    """Build a filesystem-safe estimate filename from the displayed work title."""
    name = re.sub(r'[\\/:*?"<>|]', "", _work_name(response)).strip()
    return f"{name or '견적서'}_견적서.xlsx"


def _metadata(response: dict) -> str:
    tables = response.get("tables", {})
    bill = tables.get("bill") or {}
    quantity, unit = bill.get("quantity"), bill.get("unit") or "㎥"
    conditions = {item["name"]: item for item in response.get("conditions", [])}
    work_category = conditions.get("work_category", {})
    scale = conditions.get("project_scale", {}).get("value", "이 견적만")
    scale_text = "단독 공사" if scale == "이 견적만" else f"전체 공사 {int(scale):,}원"
    summary = " · ".join(part for part in (
        work_category.get("group") or work_category.get("value"),
        conditions.get("duration", {}).get("value"),
        conditions.get("contractor_type", {}).get("value"), scale_text) if part)
    qty = f" {quantity}{unit}" if quantity is not None else ""
    basis_date = response.get("basis_date") or date.today().isoformat()
    return f"공사명: {_work_name(response)}{qty} | 기준일: {basis_date} | 조건 요약: {summary}"


def _new_sheet(book: Workbook, name: str, title: str, metadata: str, widths: list[int],
               orientation: str = "portrait", first: bool = False):
    sheet = book.active if first else book.create_sheet()
    sheet.title = name
    sheet.sheet_view.showGridLines = False
    last = len(widths)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last)
    sheet.cell(1, 1, title)
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last)
    sheet.cell(2, 1, metadata)
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A4"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = orientation
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_margins.left = sheet.page_margins.right = 0.25
    sheet.page_margins.top = sheet.page_margins.bottom = 0.45
    sheet.page_margins.header = sheet.page_margins.footer = 0.2
    return sheet


def _merge_repeated(sheet, column: int, first_row: int, last_row: int) -> None:
    start = first_row
    value = sheet.cell(first_row, column).value
    for row in range(first_row + 1, last_row + 2):
        next_value = sheet.cell(row, column).value if row <= last_row else object()
        if next_value == value:
            continue
        if value not in (None, "") and row - start > 1:
            sheet.merge_cells(start_row=start, start_column=column, end_row=row - 1, end_column=column)
        start, value = row, next_value


def _footer(sheet, last_column: int) -> int:
    row = sheet.max_row + 1
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_column)
    sheet.cell(row, 1, _FOOTER)
    return row


def _format_sheet(sheet, last_column: int, header_rows: tuple[int, ...], data_start: int,
                  data_end: int, numeric_columns: tuple[int, ...] = (),
                  money_columns: tuple[int, ...] = (), ratio_columns: tuple[int, ...] = (),
                  muted_rows: tuple[int, ...] = (), bold_rows: tuple[int, ...] = (),
                  highlight_rows: tuple[int, ...] = (), footer_row: int | None = None) -> None:
    for row in range(1, sheet.max_row + 1):
        sheet.row_dimensions[row].height = max(sheet.row_dimensions[row].height or 0, 21)
        for column in range(1, last_column + 1):
            cell = sheet.cell(row, column)
            cell.font = Font(name=_FONT_NAME, size=10)
            cell.border = _BORDER
            horizontal = "right" if row >= data_start and column in numeric_columns else "left"
            cell.alignment = Alignment(horizontal=horizontal, vertical="center", wrap_text=True)
            if row in header_rows:
                cell.font = Font(name=_FONT_NAME, size=10, bold=True)
                cell.fill = _HEAD
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            elif data_start <= row <= data_end and isinstance(cell.value, (int, float)):
                if column in money_columns:
                    cell.number_format = _MONEY
                elif column in ratio_columns:
                    cell.number_format = "0.0%"
                elif column in numeric_columns:
                    cell.number_format = _DECIMAL
    sheet.row_dimensions[1].height = 29
    sheet.row_dimensions[2].height = 27
    for row in header_rows:
        sheet.row_dimensions[row].height = 27
    sheet["A1"].font = Font(name=_FONT_NAME, size=16, bold=True)
    sheet["A1"].alignment = Alignment(horizontal="left", vertical="center")
    sheet["A2"].font = Font(name=_FONT_NAME, size=10)
    for row in muted_rows:
        for column in range(1, last_column + 1):
            cell = sheet.cell(row, column)
            cell.font = Font(name=_FONT_NAME, size=10, color="808080")
            cell.fill = _MUTED
    for row in bold_rows:
        for column in range(1, last_column + 1):
            cell = sheet.cell(row, column)
            cell.font = Font(name=_FONT_NAME, size=10, bold=True)
    for row in highlight_rows:
        for column in range(1, last_column + 1):
            cell = sheet.cell(row, column)
            cell.font = Font(name=_FONT_NAME, size=10, bold=True)
            cell.fill = _STRONG
            cell.border = Border(left=_THIN, right=_THIN, top=_DOUBLE, bottom=_DOUBLE)
    if footer_row:
        sheet.row_dimensions[footer_row].height = 23
        sheet.cell(footer_row, 1).font = Font(name=_FONT_NAME, size=9, italic=True, color="666666")
        sheet.cell(footer_row, 1).alignment = Alignment(horizontal="left", vertical="center")
    sheet.print_area = f"A1:{get_column_letter(last_column)}{sheet.max_row}"


def _append_cost_rate_sources(sheet, response: dict) -> tuple[int, int]:
    statement = response.get("statement") or {}
    priced = response.get("priced") or {}
    overhead = statement.get("overhead_version") or {}
    labor = priced.get("rate_version") or {}
    equipment = priced.get("equipment_rate_version") or {}
    conditions = {item["name"]: item for item in response.get("conditions", [])}
    group = conditions.get("work_category", {}).get("group", "")
    start = sheet.max_row + 2
    sheet.merge_cells(start_row=start, start_column=1, end_row=start, end_column=6)
    sheet.cell(start, 1, "적용 요율 출처")
    header = start + 1
    sheet.append([]) if sheet.max_row < header - 1 else None
    for column, value in enumerate(("구분", "적용 기준", "출처", "기준일/적용기간"), 1):
        sheet.cell(header, column, value)
    rows = []
    if overhead:
        effective_from = _korean_date(overhead.get("effective_from"))
        rows.append(["제비율", f"조달청 {effective_from}, {group}",
                     ", ".join(overhead.get("source_files", [])), effective_from])
    if labor:
        rows.append(["노임", labor.get("title", ""), labor.get("publisher", ""), _rate_period(labor)])
    if equipment:
        rows.append(["기계경비", equipment.get("title", ""), equipment.get("publisher", ""),
                     _rate_period(equipment, equipment=True)])
    for values in rows:
        sheet.append(values + [None, None])
    return start, header


def _citation_source(citations: list[dict] | None) -> str:
    labels = []
    for citation in citations or []:
        section = citation.get("section_no") or ""
        item = citation.get("item") or ""
        page = citation.get("pdf_page")
        label = " ".join(part for part in (section, item, f"p{page}" if page else "") if part)
        if label and label not in labels:
            labels.append(label)
    return "; ".join(labels)


def _unit_export_rows(response: dict, unit_rows: list[dict]) -> tuple[list[list], int]:
    priced = response.get("priced") or {}
    subtotals = priced.get("subtotals") or {}
    unit_prices = priced.get("unit_prices") or {}
    work_title = _work_name(response)
    header = [f"제 1호표 {work_title} (㎥ 당)", None, None, None,
              None, _num(priced.get("total")), None, _num(unit_prices.get("노무비")),
              None, _num(unit_prices.get("재료비")), None, _num(unit_prices.get("경비")), None]
    rows = [header]
    category_columns = {"노무비": (7, 8), "재료비": (9, 10), "경비": (11, 12)}
    for item in unit_rows:
        status = item.get("status") or ("산정" if item.get("amount") is not None else "미산정")
        muted = status in ("제외", "미산정")
        amount = None if muted else _num(item.get("amount"))
        unit_price = None if muted else _num(item.get("unit_price"))
        note = (f"{status} · {item.get('reason') or ''}".strip(" ·") if muted else
                _citation_source(item.get("citations")))
        values = [item.get("name", ""), item.get("spec", ""), _num(item.get("quantity")),
                  item.get("unit", ""), None, amount, None, None, None, None, None, None, note]
        columns = category_columns.get(item.get("category"))
        if columns:
            values[columns[0] - 1] = unit_price
            values[columns[1] - 1] = amount
        rows.append(values)
    return rows, 5


def _basis_rows(response: dict) -> list[list]:
    result = response.get("result") or {}
    statement = response.get("statement") or {}
    basis = []
    daily = result.get("daily_volume")
    if daily:
        basis.append(["일일시공량", f"{daily['value']} {daily['unit']}", daily["formula"],
                      _citation_source(daily.get("citations")) or "; ".join(daily.get("sources", []))])
    if result.get("work_days"):
        basis.append(["작업일수", result["work_days"]["value"], result["work_days"]["formula"], ""])
    for line in result.get("lines", []):
        basis.append([f"{line['name']} 작업조·투입", f"{line['value']} {line['unit']}",
                      f"작업조 {line['crew']}" if line.get("crew") else "",
                      _citation_source(line.get("citations")) or line.get("source", "")])
    for line in result.get("unit_lines", []):
        basis.append([f"{line['name']} 1㎥당", f"{line['applied']} {line['unit']}", line["formula"],
                      f"{line['rule']}; {_citation_source(line.get('citations')) or line.get('source', '')}"])
    for item in result.get("not_calculated", []):
        basis.append([item["item"], "미산정", "", _citation_source(item.get("citations")) or item["source"]])
    for line in statement.get("lines", []):
        status = line.get("status", "산정")
        value = status if status in ("제외", "미산정") else line.get("amount")
        source = line.get("source") or {}
        source_text = ", ".join(part for part in (source.get("file"), source.get("cell")) if part)
        basis.append([line.get("name", ""), value,
                      _basis_text(line) or line.get("reason") or line.get("note") or "", source_text])
    for note in statement.get("basis_notes", []):
        basis.append(["제비율 적용 기준", "", note, ""])
    return basis


def _grouped_header(sheet, widths: int, groups: list[tuple[str, int, int]],
                    fixed: list[tuple[str, int]]) -> None:
    for label, start, end in groups:
        sheet.merge_cells(start_row=3, start_column=start, end_row=3, end_column=end)
        sheet.cell(3, start, label)
        for column, label2 in ((start, "단가"), (start + 1, "금액")):
            sheet.cell(4, column, label2)
    for label, column in fixed:
        sheet.merge_cells(start_row=3, start_column=column, end_row=4, end_column=column)
        sheet.cell(3, column, label)


def build_xlsx(response: dict) -> bytes:
    tables = response["tables"]
    statement = response.get("statement") or {}
    metadata = _metadata(response)
    book = Workbook()

    # 공사원가계산서
    widths = [16, 24, 16, 13, 48, 14]
    sheet = _new_sheet(book, "원가계산서", "공사원가계산서", metadata, widths, first=True)
    headers = ["비목(대)", "비목(중)", "금액", "구성비(%)", "산출 근거", "상태"]
    for column, value in enumerate(headers, 1):
        sheet.cell(3, column, value)
    cost_rows = _cost_statement_rows(tables, statement)
    for values in cost_rows:
        sheet.append(values)
    cost_start, cost_end = 4, 3 + len(cost_rows)
    _merge_repeated(sheet, 1, cost_start, cost_end)
    totals = {"순공사원가", "총원가", "도급액"}
    bold_rows = tuple(row for row in range(cost_start, cost_end + 1)
                      if sheet.cell(row, 2).value == "소계" or sheet.cell(row, 2).value in totals)
    excluded_rows = tuple(row for row in range(cost_start, cost_end + 1)
                          if sheet.cell(row, 6).value in ("제외", "미산정"))
    contract_rows = tuple(row for row in range(cost_start, cost_end + 1)
                          if sheet.cell(row, 2).value == "도급액")
    rate_start, rate_header = _append_cost_rate_sources(sheet, response)
    rate_end = sheet.max_row
    footer = _footer(sheet, len(widths))
    _format_sheet(sheet, len(widths), (3, rate_header), cost_start, cost_end,
                  numeric_columns=(3, 4), money_columns=(3,), ratio_columns=(4,),
                  muted_rows=excluded_rows, bold_rows=bold_rows,
                  highlight_rows=contract_rows, footer_row=footer)
    for column in range(1, 5):
        sheet.cell(rate_start, column).fill = PatternFill("solid", fgColor="D9EAD3")
        sheet.cell(rate_start, column).font = Font(name=_FONT_NAME, size=10, bold=True)
    sheet.freeze_panes = "A4"
    sheet.print_title_rows = "3:3"
    sheet.page_setup.orientation = "portrait"

    # 일위대가표: 계 → 노무비 → 재료비 → 경비
    unit_widths = [30, 24, 12, 9, 14, 14, 14, 14, 14, 14, 14, 14, 34]
    unit_sheet = _new_sheet(book, "일위대가", "일위대가표", metadata, unit_widths)
    _grouped_header(unit_sheet, len(unit_widths),
                    [("계", 5, 6), ("노무비", 7, 8), ("재료비", 9, 10), ("경 비", 11, 12)],
                    [("공 종", 1), ("규 격", 2), ("수 량", 3), ("단 위", 4), ("비 고", 13)])
    unit_rows, unit_data_start = _unit_export_rows(response, tables.get("unit_rows", []))
    for values in unit_rows:
        unit_sheet.append(values)
    unit_end = unit_sheet.max_row
    unit_total_rows = (unit_data_start,)
    unit_muted = tuple(row for row in range(unit_data_start + 1, unit_end + 1)
                        if str(unit_sheet.cell(row, 13).value or "").startswith(("제외", "미산정")))
    unit_footer = _footer(unit_sheet, len(unit_widths))
    unit_sheet.freeze_panes = "A5"
    unit_sheet.print_title_rows = "3:4"
    _format_sheet(unit_sheet, len(unit_widths), (3, 4), unit_data_start, unit_end,
                  numeric_columns=(3, 5, 6, 7, 8, 9, 10, 11, 12),
                  money_columns=(6, 8, 10, 12), muted_rows=unit_muted,
                  bold_rows=unit_total_rows, footer_row=unit_footer)

    # 공사내역서: 실무 관행에 따라 재료비 → 노무비 → 경비
    bill_widths = [30, 26, 9, 12, 14, 14, 14, 14, 14, 14, 14, 14, 34]
    bill_sheet = _new_sheet(book, "내역서", "공사내역서", metadata, bill_widths, orientation="landscape")
    _grouped_header(bill_sheet, len(bill_widths),
                    [("재료비", 5, 6), ("노무비", 7, 8), ("경비", 9, 10), ("합계", 11, 12)],
                    [("품 명", 1), ("규 격", 2), ("단위", 3), ("수량", 4), ("비고", 13)])
    bill = tables.get("bill") or {}
    bill_rows = []
    if bill:
        prices = bill.get("unit_price") or {}
        amounts = bill.get("amount") or {}
        total_unit = (response.get("priced") or {}).get("total")
        note = "제1호표 · 미산정 항목 제외 부분 합계" if bill.get("partial") else "제1호표"
        bill_rows.append([_work_name(response), bill.get("spec", ""), bill.get("unit", ""),
                          _num(bill.get("quantity")), _num(prices.get("재료비")), _num(amounts.get("재료비")),
                          _num(prices.get("노무비")), _num(amounts.get("노무비")),
                          _num(prices.get("경비")), _num(amounts.get("경비")),
                          _num(total_unit), _num(bill.get("total")), note])
        bill_rows.append(["합계", "", "", None, _num(prices.get("재료비")), _num(amounts.get("재료비")),
                          _num(prices.get("노무비")), _num(amounts.get("노무비")),
                          _num(prices.get("경비")), _num(amounts.get("경비")),
                          _num(total_unit), _num(bill.get("total")), note])
    for values in bill_rows:
        bill_sheet.append(values)
    bill_start, bill_end = 5, bill_sheet.max_row
    bill_footer = _footer(bill_sheet, len(bill_widths))
    bill_sheet.freeze_panes = "A5"
    bill_sheet.print_title_rows = "3:4"
    _format_sheet(bill_sheet, len(bill_widths), (3, 4), bill_start, bill_end,
                  numeric_columns=(4, 5, 6, 7, 8, 9, 10, 11, 12),
                  money_columns=(5, 6, 7, 8, 9, 10, 11, 12),
                  bold_rows=(bill_end,) if bill_rows else (), footer_row=bill_footer)

    # 단가대비표
    rate_widths = [16, 34, 26, 18, 16, 56, 22]
    rate_sheet = _new_sheet(book, "단가대비표", "단가대비표", metadata, rate_widths)
    rate_headers = ["구분", "품명", "규격", "단위", "적용 단가", "출처(문서·쪽)", "기준일/적용기간"]
    for column, value in enumerate(rate_headers, 1):
        rate_sheet.cell(3, column, value)
    rate_rows = [[row.get("kind", ""), row.get("name", ""), row.get("spec", ""), row.get("unit", ""),
                  _num(row.get("price")), row.get("source", ""), row.get("period", "")]
                 for row in tables.get("rate_rows", [])]
    for values in rate_rows:
        rate_sheet.append(values)
    rate_start, rate_end = 4, rate_sheet.max_row
    rate_footer = _footer(rate_sheet, len(rate_widths))
    rate_sheet.print_title_rows = "3:3"
    _format_sheet(rate_sheet, len(rate_widths), (3,), rate_start, rate_end,
                  numeric_columns=(5,), money_columns=(5,), footer_row=rate_footer)

    # 산출근거
    basis_widths = [28, 22, 54, 58]
    basis_sheet = _new_sheet(book, "산출근거", "산출근거", metadata, basis_widths)
    for column, value in enumerate(("항목", "값", "계산", "근거(품셈 절·표·쪽)"), 1):
        basis_sheet.cell(3, column, value)
    basis = _basis_rows(response)
    for values in basis:
        basis_sheet.append(values)
    basis_start, basis_end = 4, basis_sheet.max_row
    basis_footer = _footer(basis_sheet, len(basis_widths))
    basis_sheet.print_title_rows = "3:3"
    _format_sheet(basis_sheet, len(basis_widths), (3,), basis_start, basis_end,
                  numeric_columns=(2,), money_columns=(), footer_row=basis_footer)

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
