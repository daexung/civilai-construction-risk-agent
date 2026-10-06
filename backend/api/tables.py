"""계산 결과를 표(원가계산서·내역서·일위대가·단가대비표)로 바꾸고 엑셀로 만든다.

화면과 엑셀이 같은 표를 쓰도록 응답 dict에서 한 번만 만든다.
"""

from __future__ import annotations

import io
import json
import re
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from fractions import Fraction
from functools import lru_cache
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

DISCLAIMER = "표준품셈 기준 금액, 검토 전 참고용"
_RANK = {"재료비": 0, "노무비": 1, "경비": 2}
_SUBTOTALS = ("노무비 계", "경비 계")


def _rate_period(version: dict, equipment: bool = False) -> str:
    if not version:
        return ""
    if equipment:
        return str(version.get("version", ""))
    title = version.get("title", "")
    year = re.search(r"20\d{2}", title)
    if year and ("상반기" in title or "하반기" in title):
        half = "상반기" if "상반기" in title else "하반기"
        return f"{year.group()} {half}"
    return " ~ ".join(part for part in (version.get("effective_from"), version.get("effective_to")) if part)


def _korean_date(value: str | None) -> str:
    if not value:
        return ""
    match = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})", value)
    return f"{int(match.group(1))}.{int(match.group(2))}.{int(match.group(3))}." if match else value


def _num(value) -> int | float | None:
    if value is None or value == "":
        return None
    # Exact calculation outputs can be rational strings (e.g. 100/65).
    # Excel cells need a numeric approximation; the adjacent formula retains basis.
    if isinstance(value, str) and re.fullmatch(r"-?\d+/\d+", value):
        fraction = Fraction(value)
        number = Decimal(fraction.numerator) / Decimal(fraction.denominator)
    else:
        number = Decimal(str(value))
    return int(number) if number == number.to_integral_value() else float(number)


def _basis_text(line: dict) -> str:
    base, rate = line.get("base") or "", line.get("rate")
    if rate is not None and line.get("status") == "산정":
        return f"{base} {line['base_amount']:,}원 × {rate}%"
    return base


def statement_rows(statement: dict | None) -> list[dict]:
    """비목별로 묶어 재료비 → 노무비 → 경비 → 순공사원가 … 도급액 순으로 늘어놓는다."""
    if not statement:
        return []
    indexed = list(enumerate(statement.get("lines", [])))

    def key(item: tuple[int, dict]):
        index, line = item
        category = line.get("category")
        if line["name"] in _SUBTOTALS:
            return (_RANK[category], 1, index)
        if category in _RANK:
            return (_RANK[category], 0, index)
        return (3, 0, index) if category else (4, 0, index)

    rows = []
    for _, line in sorted(indexed, key=key):
        status = line.get("status", "산정")
        is_total = line["name"] in ("순공사원가", "총원가", "도급액")
        rows.append({
            "name": line["name"], "category": line.get("category"),
            "basis": _basis_text(line), "amount": line.get("amount") if status != "미산정" else None,
            "status": status, "reason": line.get("reason"), "note": line.get("note"),
            "kind": "subtotal" if line["name"] in _SUBTOTALS else ("total" if is_total else "item"),
            "final": line["name"] == "도급액",
        })
    return rows


def _spec_text(inputs: list[dict]) -> str:
    values = {item["name"]: item["value"] for item in inputs}
    parts = []
    if "structure" in values:
        parts.append(values["structure"])
    if "slump_band" in values:
        parts.append(values["slump_band"])
    if "facility_type" in values or "site_type" in values:
        parts.append(f"{values.get('facility_type', '')}/{values.get('site_type', '').replace('Type-', '')}")
    if "pump_size" in values:
        boom = f"{values['pump_size']} {values.get('placement', '')}".strip()
        parts.append(boom)
    return "·".join(part for part in parts if part)


def bill_row(response: dict) -> dict | None:
    priced = response.get("priced")
    if not priced or not priced.get("reference_amounts"):
        return None
    reference = priced["reference_amounts"]
    unit_prices, amounts = priced.get("unit_prices") or {}, reference.get("subtotals") or {}
    work = response.get("work") or {}
    unit = ((response.get("result") or {}).get("unit_basis") or {}).get("per", "1㎥")[1:]
    return {
        "name": f"{work.get('title', '')}({work.get('section_no', '')})" if work else "",
        "spec": _spec_text(response.get("inputs", [])), "unit": unit, "quantity": reference.get("volume"),
        "unit_price": {name: unit_prices.get(name) for name in ("재료비", "노무비", "경비")},
        "amount": {name: amounts.get(name) for name in ("재료비", "노무비", "경비")},
        "total": reference.get("total"), "partial": bool(priced.get("partial")),
    }


def unit_price_rows(priced: dict | None) -> list[dict]:
    if not priced:
        return []
    rows = []
    for key in ("lines", "equipment_lines", "cost_lines", "supply_lines"):
        for line in priced.get(key, []):
            if line.get("kind") == "equipment":
                continue  # 장비 합계 줄은 아래 구성 줄(손료·운전원)과 겹친다
            amount = line.get("amount")
            kind = line.get("kind")
            base = line.get("base")
            rate = line.get("rate")
            if kind == "rate_cost":
                unit, quantity = "식", 1
                unit_price = amount
                spec = f"노무비의 {Decimal(str(rate)) * 100:g}%" if rate is not None else ""
            else:
                unit, quantity = line.get("unit") or "", line.get("quantity")
                for basis_unit in ("㎥", "㎡"):
                    unit = unit.replace(f"/{basis_unit}", "")
                unit_price = line.get("unit_price")
                spec = line.get("spec") or line.get("machine_spec") or (
                    f"노임 코드 {line['rate_code']}" if line.get("rate_code") else "")
                if base and rate is not None:
                    spec = f"{base} {rate}%"
            rows.append({"category": line.get("category"), "name": line.get("name"),
                         "spec": spec, "unit": unit, "quantity": quantity,
                         "unit_price": unit_price, "amount": amount,
                         "status": line.get("status") or ("산정" if amount is not None else "미산정"),
                         "reason": line.get("reason"), "citations": line.get("citations", [])})
    return rows


def rate_rows(priced: dict | None, statement: dict | None,
              conditions: list[dict] | None = None) -> list[dict]:
    """실제 적용한 노임·기계 단가와 관급자재를 정리한다."""
    if not priced:
        return []
    labor_version = priced.get("rate_version") or {}
    equipment_version = priced.get("equipment_rate_version") or {}
    labor_source = f"{labor_version.get('title', '')}({labor_version.get('publisher', '')})"
    rows = []
    for line in priced.get("lines", []) + priced.get("equipment_lines", []):
        if line.get("unit_price") is None:
            continue
        if line.get("kind") == "labor":
            rows.append({"kind": "노임", "name": line["name"], "spec": "",
                         "unit": "원/인", "price": line["unit_price"],
                         "source": labor_source, "period": _rate_period(labor_version),
                         "note": f"일 8시간 · 노임 코드 {line.get('rate_code', '')}"})
        elif line.get("machine_code") and line.get("rate_code"):
            daily_wage = _num(Decimal(str(line["unit_price"])) * 8)
            hourly = f"{Decimal(str(line['unit_price'])):,.1f}"
            rows.append({"kind": "노임", "name": line["name"], "spec": "", "unit": "원/인",
                         "price": daily_wage, "source": labor_source,
                         "period": _rate_period(labor_version),
                         "note": f"시간당 {hourly}원 환산(÷8) · 노임 코드 {line['rate_code']}"})
        elif line.get("machine_code"):
            rows.append({"kind": "기계경비", "name": line["name"], "spec": line.get("machine_spec", ""),
                         "unit": "원/hr", "price": line["unit_price"],
                         "source": f"{equipment_version.get('title', '')}({equipment_version.get('publisher', '')})",
                         "period": _rate_period(equipment_version, equipment=True), "note": ""})
    for line in priced.get("supply_lines", []):
        if line.get("status") == "산정" and line.get("unit_price") is not None:
            rows.append({"kind": "자재", "name": line["name"], "spec": line.get("spec", ""),
                         "unit": "원/㎥", "price": _num(line["unit_price"]),
                         "source": "사용자 입력(부가세 제외)", "period": "",
                         "note": f"할증 {format((Decimal(str(line['allowance'])) * 100).normalize(), 'g')}%(품셈 1-3-1)"})
        else:
            note = "관급(발주처 지급)" if line.get("status") == "제외" else line.get("reason", "")
            rows.append({"kind": "자재", "name": line["name"], "spec": "",
                         "unit": line.get("unit", ""), "price": "-", "source": "",
                         "period": "", "note": note})
    for line in priced.get("lines", []):
        if line.get("kind") == "material":
            rows.append({"kind": "자재", "name": line["name"], "spec": "",
                         "unit": line.get("unit", "").split("/", 1)[0], "price": "-",
                         "source": _citation_source(line.get("citations")), "period": "",
                         "note": line.get("reason", "")})
    return rows


def add_tables(response: dict) -> dict:
    return {"statement_rows": statement_rows(response.get("statement")), "bill": bill_row(response),
            "unit_rows": unit_price_rows(response.get("priced")),
            "rate_rows": rate_rows(response.get("priced"), response.get("statement"), response.get("conditions"))}


def _cost_statement_rows(tables: dict, statement: dict) -> list[list]:
    """원가계산서의 실무 행을 고정된 다섯 열로 만든다."""
    totals = statement.get("totals", {})
    contract_amount = totals.get("contract_amount")
    statement_rows = tables.get("statement_rows", [])
    by_name = {row["name"]: row for row in statement_rows}
    output = []

    def add(major: str, minor: str, amount, basis: str = "", status: str = "산정") -> None:
        excluded = status in ("제외", "미산정")
        value = None if excluded else _num(amount)
        ratio = None
        if value is not None and contract_amount:
            ratio = float(Decimal(str(value)) / Decimal(str(contract_amount)))
        note = f"{status} · {basis}" if excluded and basis else status if excluded else basis
        output.append([major, minor, value, ratio, note, status if excluded else ""])

    for category in ("재료비", "노무비", "경비"):
        lines = [row for row in statement_rows if row.get("category") == category]
        for row in lines:
            minor = ("소계" if row["name"] in _SUBTOTALS else
                     "직접재료비" if category == "재료비" and row["name"] == "재료비" else row["name"])
            basis = row.get("basis") or row.get("reason") or row.get("note") or ""
            add(category, minor, row.get("amount"), basis, row.get("status", "산정"))
        if category == "재료비":
            present = {row["name"] for row in lines}
            for label in ("간접재료비(-)", "작업설·부산물 등(△)(-)"):
                if label not in present:
                    output.append([category, label, "-", None, "", ""])
        if category == "재료비" and lines:
            add(category, "소계", totals.get("materials"), "직접재료비 및 재료비 항목 합계")

    summary_names = ("순공사원가", "일반관리비", "이윤", "총원가", "부가가치세", "도급액")
    for name in summary_names:
        row = by_name.get(name)
        if row:
            major = name if name in ("순공사원가", "총원가", "도급액") else (row.get("category") or "")
            minor = "" if name in ("순공사원가", "총원가", "도급액") else name
            basis = row.get("basis") or row.get("reason") or row.get("note") or ""
            if name == "부가가치세" and totals.get("total_cost") is not None:
                basis = f"총원가 {int(totals['total_cost']):,}원 × 10%"
            add(major, minor, row.get("amount"), basis, row.get("status", "산정"))

    known = {row["name"] for row in statement_rows if row.get("category") in ("재료비", "노무비", "경비")}
    known.update(summary_names)
    for row in statement_rows:
        if row["name"] in known:
            continue
        if "1-2-9" in row["name"] or "살수 양생" in row["name"]:
            continue
        status = row.get("status", "산정")
        add(row.get("category") or "경비", row["name"], row.get("amount"),
            row.get("basis") or row.get("reason") or row.get("note") or "", status)
    return output


def _missing_cost_rows(response: dict, cost_rows: list[list]) -> list[list[str]]:
    rows: list[list[str]] = []
    seen: set[str] = set()

    def display_reason(item: str, reason: str) -> str:
        if "레미콘" in item:
            return "관급자재(발주처 지급)"
        if "연료" in item or "잡재료" in item:
            return "품셈 8-1-7 유류가격 해당 지역 가격 — 미입력"
        if "1-2-9" in item:
            return "품셈 1-2-9 — 이번 계산 범위 밖"
        if "살수 양생" in item:
            return "품셈 6-1-4 사. — 별도 계상"
        if "퇴직공제" in item:
            return "추정금액 1억 미만"
        if "산업안전보건관리비" in item:
            return "총 공사금액 2천만원 미만"
        if "공사이행보증" in item:
            return "지방계약법 시행령 제51조 대상 아님"
        if "건설기계대여대금" in item:
            return "보증 수수료 산정 방법 확인 필요"
        return reason

    def add(item: str, status: str, reason: str = "") -> None:
        if item and item not in seen:
            seen.add(item)
            rows.append([item, status, display_reason(item, reason)])

    for major, minor, _amount, _ratio, note, status in cost_rows:
        if status in ("제외", "미산정"):
            add(minor or major, status, note.removeprefix(f"{status} · ").strip())
    statement = response.get("statement") or {}
    for line in statement.get("lines", []):
        status = line.get("status", "산정")
        if status in ("제외", "미산정"):
            add(line.get("name", ""), status, line.get("reason") or "")
    result = response.get("result") or {}
    for item in result.get("not_calculated", []):
        add(item.get("item", ""), "미산정", item.get("source") or "이번 계산에서 미산정")
    return rows


# ---------- 엑셀 ----------

_FONT_NAME = "맑은 고딕"
_HEAD = PatternFill("solid", fgColor="E7E6E6")
_STRONG = PatternFill("solid", fgColor="FFF2CC")
_MUTED = PatternFill("solid", fgColor="F2F2F2")
_THIN = Side(style="thin", color="B7B7B7")
_DOUBLE = Side(style="double", color="595959")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_MONEY = "#,##0"
_DECIMAL = "#,##0.0###"
_FOOTER = "표준품셈 기준 금액 · 검토 전 참고용 · 품셈이"


def _work_name(response: dict) -> str:
    work = response.get("work") or {}
    title = work.get("title") or "공사비"
    if work.get("spec_id"):
        from backend.agent.rules.specs import load_specs
        spec = load_specs().get(work["spec_id"])
        if spec:
            title = spec["title"]
    return re.sub(r"\s*\(20\d{2}(?:년)?[^)]*\)\s*$", "", title).strip()


def estimate_filename(response: dict) -> str:
    """Brand, work and Korean download date; price basis stays in the workbook."""
    name = re.sub(r'[\\/:*?"<>|\x00-\x1f\x7f]', '', _work_name(response)).strip(' .')
    name = re.sub(r'^(?:(?:토목|건축|기계설비)\s+)?\d+-\d+-\d+\s*', '', name)
    name = re.sub(r'\s+', ' ', name)[:60].strip(' .') or '공사비'
    downloaded = datetime.now(timezone(timedelta(hours=9))).strftime('%Y%m%d')
    return f"품셈이_{name}_견적서_{downloaded}.xlsx"


def _metadata(response: dict) -> tuple[str, str]:
    tables = response.get("tables", {})
    bill = tables.get("bill") or {}
    quantity, unit = bill.get("quantity"), bill.get("unit") or "㎥"
    conditions = {item["name"]: item for item in response.get("conditions", [])}
    work_category = conditions.get("work_category", {})
    scale = conditions.get("project_scale", {}).get("value", "이 견적만")
    scale_text = "단독 공사" if scale == "이 견적만" else f"전체 공사 {int(scale):,}원"
    summary = " · ".join(part for part in (
        work_category.get("group") or work_category.get("value"),
        conditions.get("duration", {}).get("value"),
        conditions.get("contractor_type", {}).get("value"), scale_text) if part)
    qty = f" {quantity}{unit}" if quantity is not None else ""
    basis_date = response.get("basis_date") or date.today().isoformat()
    review = (response.get("result") or {}).get("review_status")
    badge = f" | {review}" if review == "AI 초안 · 검토 전" else ""
    return (f"공사명: {_work_name(response)}{qty}",
            f"기준일: {basis_date} | 조건 요약: {summary}{badge}")


def won_in_korean(amount: int) -> str:
    """정수 원 금액을 견적서 관행의 한글 금액 표기로 바꾼다."""
    number = int(amount)
    if number < 0:
        raise ValueError("금액은 0 이상이어야 합니다.")
    if number == 0:
        return "영"

    digits = ("영", "일", "이", "삼", "사", "오", "육", "칠", "팔", "구")
    small_units = ("", "십", "백", "천")
    large_units = ("", "만", "억", "조")
    parts = []
    group_index = 0
    while number:
        group = number % 10_000
        if group:
            group_text = ""
            for position in range(3, -1, -1):
                digit = group // (10 ** position) % 10
                if digit:
                    group_text += digits[digit] + small_units[position]
            parts.append(group_text + large_units[group_index])
        number //= 10_000
        group_index += 1
        if group_index >= len(large_units) and number:
            raise ValueError("조 단위보다 큰 금액은 지원하지 않습니다.")
    return "".join(reversed(parts))


def _estimate_sheet(book: Workbook, response: dict) -> None:
    """계산 결과의 최종 금액과 비목별 요약을 첫 시트에 표시한다."""
    sheet = _new_sheet(book, "견적서", "견 적 서", ("", ""), [28, 20, 52], first=True)
    sheet.freeze_panes = "A7"
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_view.showGridLines = False

    bill = (response.get("tables") or {}).get("bill") or {}
    title = _work_name(response)
    spec = bill.get("spec") or ""
    quantity, unit = bill.get("quantity"), bill.get("unit") or "㎥"
    project = " · ".join(part for part in (title, spec) if part)
    if quantity is not None:
        project = f"{project} {quantity}{unit}".strip()
    conditions = {item["name"]: item for item in response.get("conditions", [])}
    scale = conditions.get("project_scale", {}).get("value", "이 견적만")
    scale_text = "단독 공사" if scale == "이 견적만" else f"전체 공사 {int(scale):,}원"
    condition_summary = " · ".join(part for part in (
        conditions.get("work_category", {}).get("group")
        or conditions.get("work_category", {}).get("value"),
        conditions.get("duration", {}).get("value"),
        conditions.get("contractor_type", {}).get("value"), scale_text) if part)
    basis_date = response.get("basis_date") or date.today().isoformat()

    sheet["A2"] = f"공사명 : {project}"
    sheet["A3"] = f"기 준 일 : {basis_date}"
    sheet.merge_cells("A4:C4")
    review = (response.get("result") or {}).get("review_status")
    badge = " | AI 초안 · 검토 전" if review == "AI 초안 · 검토 전" else ""
    sheet["A4"] = f"조    건 : {condition_summary}{badge}"

    totals = (response.get("statement") or {}).get("totals") or {}
    contract = totals.get("contract_amount")
    amount_text = ""
    if contract is not None:
        amount_text = f"견적금액 : 일금 {won_in_korean(contract)}원정 (₩{int(contract):,})   ※ 부가가치세 포함"
    sheet.merge_cells("A6:C6")
    sheet["A6"] = amount_text
    sheet["A6"].font = Font(name=_FONT_NAME, size=13, bold=True)
    sheet["A6"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet["A6"].border = Border(left=_DOUBLE, right=_DOUBLE, top=_DOUBLE, bottom=_DOUBLE)
    sheet.row_dimensions[6].height = 34

    header_row = 8
    for column, label in enumerate(("구 분", "금 액(원)", "비 고"), 1):
        cell = sheet.cell(header_row, column, label)
        cell.font = Font(name=_FONT_NAME, size=10, bold=True)
        cell.fill = _HEAD
        cell.alignment = Alignment(horizontal="center", vertical="center")
    material_excluded = any(row.get("kind") == "자재" and "관급" in (row.get("note") or "")
                            for row in (response.get("tables") or {}).get("rate_rows", []))
    ready_mix = next((line for line in (response.get("priced") or {}).get("supply_lines", [])
                      if line.get("status") == "산정" and line.get("unit_price") is not None), None)
    entries = (
        ("재료비", "materials", ""),
        ("노무비", "labor", "직접 + 간접 노무비"),
        ("경비", "expenses", "장비·보험료·기타경비 등"),
        ("일반관리비", "management", ""),
        ("이윤", "profit", ""),
        ("공급가액", "total_cost", ""),
        ("부가가치세", "vat", "10%"),
        ("합 계", "contract_amount", ""),
    )
    if ready_mix:
        price_text = f"{Decimal(str(ready_mix['unit_price'])):,.0f}"
        entries = tuple((label, key,
                         f"레미콘 사급(사용자 입력 {price_text}원/㎥)" if label == "재료비" else note)
                        for label, key, note in entries)
    elif material_excluded:
        entries = tuple((label, key, "레미콘 관급 제외" if label == "재료비" else note)
                        for label, key, note in entries)
    row_by_label = {}
    for offset, (label, key, note) in enumerate(entries, 1):
        row = header_row + offset
        row_by_label[label] = row
        sheet.cell(row, 1, label)
        value = totals.get(key)
        sheet.cell(row, 2, _num(value))
        sheet.cell(row, 3, note)
        sheet.cell(row, 2).number_format = _MONEY
        for column in range(1, 4):
            sheet.cell(row, column).border = _BORDER
            sheet.cell(row, column).alignment = Alignment(
                horizontal="right" if column == 2 else "left", vertical="center", wrap_text=True)

    for label in ("공급가액", "합 계"):
        for column in range(1, 4):
            sheet.cell(row_by_label[label], column).font = Font(name=_FONT_NAME, size=10, bold=True)
    total_row = row_by_label["합 계"]
    for column in range(1, 4):
        sheet.cell(total_row, column).fill = _STRONG
        sheet.cell(total_row, column).border = Border(left=_THIN, right=_THIN, top=_DOUBLE, bottom=_DOUBLE)

    missing = [item for item in _missing_cost_rows(response, _cost_statement_rows(
        (response.get("tables") or {}), response.get("statement") or {})) if item[1] == "미산정"]
    next_row = total_row + 2
    if missing:
        item, _status, reason = missing[0]
        suffix = f" 외 {len(missing) - 1}건" if len(missing) > 1 else ""
        sheet.merge_cells(start_row=next_row, start_column=1, end_row=next_row, end_column=3)
        sheet.cell(next_row, 1, f'빠진 항목: {item}({reason}){suffix} — 원가계산서 "빠진 항목" 참조')
        sheet.cell(next_row, 1).alignment = Alignment(vertical="center", wrap_text=True)
        next_row += 2
    sheet.merge_cells(start_row=next_row, start_column=1, end_row=next_row, end_column=3)
    sheet.cell(next_row, 1, _FOOTER)
    sheet.cell(next_row, 1).font = Font(name=_FONT_NAME, size=9, italic=True, color="666666")
    sheet.cell(next_row, 1).alignment = Alignment(horizontal="center", vertical="center")
    sheet.print_area = f"A1:C{next_row}"
    sheet.print_title_rows = "1:3"
    for row in range(1, next_row + 1):
        for column in range(1, 4):
            cell = sheet.cell(row, column)
            if cell.font.name != _FONT_NAME:
                cell.font = Font(name=_FONT_NAME, size=10, bold=cell.font.bold,
                                 italic=cell.font.italic, color=cell.font.color)
    sheet["A1"].font = Font(name=_FONT_NAME, size=20, bold=True)
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 38
    sheet.row_dimensions[2].height = 25
    sheet.row_dimensions[3].height = 22
    sheet.row_dimensions[4].height = 25
    sheet.row_dimensions[header_row].height = 24
    for column, width in enumerate((28, 20, 52), 1):
        sheet.column_dimensions[get_column_letter(column)].width = width


def _new_sheet(book: Workbook, name: str, title: str, metadata: tuple[str, str], widths: list[int],
               orientation: str = "portrait", first: bool = False):
    sheet = book.active if first else book.create_sheet()
    sheet.title = name
    sheet.sheet_view.showGridLines = False
    last = len(widths)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last)
    sheet.cell(1, 1, title)
    project_line, conditions_line = metadata
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last)
    sheet.cell(2, 1, project_line)
    sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=last)
    sheet.cell(3, 1, conditions_line)
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A5"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = orientation
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_margins.left = sheet.page_margins.right = 0.25
    sheet.page_margins.top = sheet.page_margins.bottom = 0.45
    sheet.page_margins.header = sheet.page_margins.footer = 0.2
    return sheet


def _merge_repeated(sheet, column: int, first_row: int, last_row: int) -> None:
    start = first_row
    value = sheet.cell(first_row, column).value
    for row in range(first_row + 1, last_row + 2):
        next_value = sheet.cell(row, column).value if row <= last_row else object()
        if next_value == value:
            continue
        if value not in (None, "") and row - start > 1:
            sheet.merge_cells(start_row=start, start_column=column, end_row=row - 1, end_column=column)
        start, value = row, next_value


def _footer(sheet, last_column: int) -> int:
    row = sheet.max_row + 1
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_column)
    sheet.cell(row, 1, _FOOTER)
    return row


def _format_sheet(sheet, last_column: int, header_rows: tuple[int, ...], data_start: int,
                  data_end: int, numeric_columns: tuple[int, ...] = (),
                  money_columns: tuple[int, ...] = (), ratio_columns: tuple[int, ...] = (),
                  muted_rows: tuple[int, ...] = (), bold_rows: tuple[int, ...] = (),
                  highlight_rows: tuple[int, ...] = (), footer_row: int | None = None) -> None:
    for row in range(1, sheet.max_row + 1):
        for column in range(1, last_column + 1):
            cell = sheet.cell(row, column)
            cell.font = Font(name=_FONT_NAME, size=10)
            cell.border = _BORDER
            horizontal = "right" if row >= data_start and column in numeric_columns else "left"
            cell.alignment = Alignment(horizontal=horizontal, vertical="center", wrap_text=True)
            if row in header_rows:
                cell.font = Font(name=_FONT_NAME, size=10, bold=True)
                cell.fill = _HEAD
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            elif data_start <= row <= data_end and isinstance(cell.value, (int, float)):
                if column in money_columns:
                    cell.number_format = _MONEY
                elif column in ratio_columns:
                    cell.number_format = "0.0%"
                elif column in numeric_columns:
                    cell.number_format = _DECIMAL
    sheet.row_dimensions[1].height = 29
    sheet.row_dimensions[2].height = 22
    sheet.row_dimensions[3].height = 22
    for row in header_rows:
        sheet.row_dimensions[row].height = 27
    sheet["A1"].font = Font(name=_FONT_NAME, size=16, bold=True)
    sheet["A1"].alignment = Alignment(horizontal="left", vertical="center")
    sheet["A2"].font = Font(name=_FONT_NAME, size=10)
    sheet["A3"].font = Font(name=_FONT_NAME, size=10)
    for row in muted_rows:
        for column in range(1, last_column + 1):
            cell = sheet.cell(row, column)
            cell.font = Font(name=_FONT_NAME, size=10, color="808080")
            cell.fill = _MUTED
    for row in bold_rows:
        for column in range(1, last_column + 1):
            cell = sheet.cell(row, column)
            cell.font = Font(name=_FONT_NAME, size=10, bold=True)
    for row in highlight_rows:
        for column in range(1, last_column + 1):
            cell = sheet.cell(row, column)
            cell.font = Font(name=_FONT_NAME, size=10, bold=True)
            cell.fill = _STRONG
            cell.border = Border(left=_THIN, right=_THIN, top=_DOUBLE, bottom=_DOUBLE)
    if footer_row:
        sheet.row_dimensions[footer_row].height = 23
        sheet.cell(footer_row, 1).font = Font(name=_FONT_NAME, size=9, italic=True, color="666666")
        sheet.cell(footer_row, 1).alignment = Alignment(horizontal="left", vertical="center")
    sheet.print_area = f"A1:{get_column_letter(last_column)}{sheet.max_row}"


@lru_cache(maxsize=1)
def _overhead_rates() -> dict:
    from backend.paths import ROOT
    path = ROOT / "data" / "rates" / "overhead_rates.json"
    return json.loads(path.read_text(encoding="utf-8"))


def _overhead_rate_entries(version_id: str, work: str) -> dict[str, dict]:
    version = _overhead_rates().get("versions", {}).get(version_id, {}).get(work, {})
    entries: dict[str, dict] = {}

    def collect(value) -> None:
        if isinstance(value, dict):
            cell = value.get("cell")
            if cell:
                entries[cell] = value
            for child in value.values():
                collect(child)
        elif isinstance(value, list):
            for child in value:
                collect(child)

    collect(version)
    return entries


def _applied_rate_range(entry: dict, line: dict, work_category: str) -> str:
    parts = []
    if entry.get("scale"):
        parts.append(entry["scale"])
    if entry.get("period"):
        parts.append(entry["period"])
    if entry.get("contractor"):
        parts.append("종합" if entry["contractor"] == "종합건설업" else "전문")
    if entry.get("grade"):
        parts.append(entry["grade"])
    if entry.get("kind"):
        parts.append(work_category)
    if parts:
        return "·".join(parts)
    return f"{line.get('base') or '공통 요율'} 기준"


def _append_cost_rate_sources(sheet, response: dict) -> tuple[int, int]:
    statement = response.get("statement") or {}
    priced = response.get("priced") or {}
    overhead = statement.get("overhead_version") or {}
    labor = priced.get("rate_version") or {}
    equipment = priced.get("equipment_rate_version") or {}
    conditions = {item["name"]: item for item in response.get("conditions", [])}
    work_category = conditions.get("work_category", {}).get("value", "")
    work = "architecture" if work_category in ("재개발·재건축", "주택 신축", "주택 외 건축") else "civil"
    group = "건축" if work == "architecture" else "토목"
    start = sheet.max_row + 2
    sheet.merge_cells(start_row=start, start_column=1, end_row=start, end_column=5)
    sheet.cell(start, 1, "적용 요율 출처")
    header = start + 1
    sheet.append([]) if sheet.max_row < header - 1 else None
    for column, value in enumerate(("항목", "요율", "적용 구간", "출처"), 1):
        sheet.cell(header, column, value)
    sheet.merge_cells(start_row=header, start_column=4, end_row=header, end_column=5)
    rows = []
    if overhead:
        effective_from = _korean_date(overhead.get("effective_from"))
        rate_entries = _overhead_rate_entries(overhead.get("id", ""), work)
        for line in statement.get("lines", []):
            source = line.get("source") or {}
            if (line.get("name") == "부가가치세" or line.get("status") != "산정"
                    or line.get("rate") is None or line.get("amount") is None
                    or not str(source.get("file") or "").startswith("data/raw/")):
                continue
            entry = rate_entries.get(source.get("cell"), {})
            rows.append([line["name"], f"{line['rate']}%",
                         _applied_rate_range(entry, line, work_category),
                         f"조달청 제비율 {effective_from}, {group}"])
        if any(line.get("name") == "부가가치세" for line in statement.get("lines", [])):
            rows.append(["부가가치세", "10%", "-", "부가가치세법"])
    if labor:
        rows.append(["노임", "-", _rate_period(labor),
                     f"{labor.get('title', '')}({labor.get('publisher', '')})"])
    if equipment:
        rows.append(["기계경비", "-", _rate_period(equipment, equipment=True),
                     f"{equipment.get('title', '')}({equipment.get('publisher', '')})"])
    for values in rows:
        sheet.append(values + [None])
        row = sheet.max_row
        sheet.merge_cells(start_row=row, start_column=4, end_row=row, end_column=5)
    return start, header


def _citation_source(citations: list[dict] | None) -> str:
    labels = []
    for citation in citations or []:
        section = citation.get("section_no") or ""
        item = citation.get("item") or ""
        page = citation.get("pdf_page")
        label = " ".join(part for part in (section, item if item != "표" else "", f"p{page}" if page else "") if part)
        if label and label not in labels:
            labels.append(label)
    return "; ".join(labels)


def _unit_citation_source(citations: list[dict] | None) -> str:
    labels = []
    for citation in citations or []:
        section = citation.get("section_no") or ""
        if section in ("1-2-2",) or "노임단가" in (citation.get("section_title") or ""):
            continue
        page = citation.get("pdf_page")
        label = " ".join(part for part in (section, f"p{page}" if page else "") if part)
        if label and label not in labels:
            labels.append(label)
    return "; ".join(labels)


def _unit_export_rows(response: dict, unit_rows: list[dict]) -> tuple[list[list], int]:
    priced = response.get("priced") or {}
    subtotals = priced.get("subtotals") or {}
    unit_prices = priced.get("unit_prices") or {}
    work_title = _work_name(response)
    basis_unit = ((response.get("result") or {}).get("unit_basis") or {}).get("per", "1㎥")[1:]
    header = [f"제 1호표 {work_title} ({basis_unit} 당)", None, None, None,
              None, _num(priced.get("total")), None, _num(unit_prices.get("노무비")),
              None, _num(unit_prices.get("재료비")), None, _num(unit_prices.get("경비")), None]
    rows = [header]
    category_columns = {"노무비": (7, 8), "재료비": (9, 10), "경비": (11, 12)}
    for item in unit_rows:
        status = item.get("status") or ("산정" if item.get("amount") is not None else "미산정")
        muted = status in ("제외", "미산정")
        amount = None if muted else _num(item.get("amount"))
        unit_price = None if muted else _num(item.get("unit_price"))
        note = (f"{status} · {item.get('reason') or ''}".strip(" ·") if muted else
                _unit_citation_source(item.get("citations")))
        if item.get("category") == "노무비" and item.get("spec", "").startswith("노임 코드"):
            note = "; ".join(part for part in (note, item["spec"]) if part)
        if "펌프차 운전원" in item.get("name", ""):
            note = "; ".join(part for part in (note, "일 283,323원 ÷ 8시간") if part)
        if item.get("name") == "레미콘(사급)":
            note = "; ".join(part for part in (note, "사용자 입력 단가") if part)
        spec = "" if item.get("category") == "노무비" else item.get("spec", "")
        values = [item.get("name", ""), spec, _num(item.get("quantity")),
                  item.get("unit", ""), None, amount, None, None, None, None, None, None, note]
        columns = category_columns.get(item.get("category"))
        if columns:
            values[columns[0] - 1] = unit_price
            values[columns[1] - 1] = amount
        rows.append(values)
    return rows, 6


def _basis_rows(response: dict) -> list[list]:
    result = response.get("result") or {}
    statement = response.get("statement") or {}
    priced = response.get("priced") or {}
    basis_unit = ((result.get("unit_basis") or {}).get("per") or "1㎥")[1:]
    basis = []
    daily = result.get("daily_volume")
    if daily:
        basis.append([f"일일시공량 ({daily['unit']})", _num(daily["value"]), daily["formula"],
                      _citation_source(daily.get("citations")) or "; ".join(daily.get("sources", []))])
        for citation in daily.get("citations", []):
            factor = citation.get("row")
            if factor in ("f1", "f2"):
                basis.append([f"보정계수 {factor}", _num(citation.get("value")),
                              citation.get("column") or "", _citation_source([citation])])
    if result.get("work_days"):
        basis.append(["작업일수 (일)", _num(result["work_days"]["value"]), result["work_days"]["formula"], ""])
    for line in result.get("lines", []):
        basis.append([f"{line['name']} 작업조·투입 ({line['unit']})", _num(line["value"]),
                      f"작업조 {line['crew']}" if line.get("crew") else "",
                      _citation_source(line.get("citations")) or line.get("source", "")])
    for line in result.get("unit_lines", []):
        basis.append([f"{line['name']} 1{basis_unit}당 ({line['unit']})", _num(line["applied"]), line["formula"],
                      f"{line['rule']}; {_citation_source(line.get('citations')) or line.get('source', '')}"])
        for adjustment in line.get("adjustments", []):
            basis.append([f"{line['name']} {adjustment['종류']}", adjustment["값"],
                          adjustment["원문 인용"], _citation_source(adjustment.get("citations"))])
    for memo in result.get("adjustment_memos", []):
        basis.append(["할증 참고", "", memo, "자동 계산 없음"])
    for line in priced.get("supply_lines", []):
        if line.get("status") == "산정":
            quantity = _num(line.get("quantity"))
            unit_price = Decimal(str(line["unit_price"]))
            formula = f"{unit_price:,.0f}원/㎥ × {quantity}㎥"
            source = "; ".join(part for part in (
                _unit_citation_source(line.get("citations")), "사용자 입력 단가") if part)
        else:
            quantity = None
            formula = f"{line.get('status')} · {line.get('reason') or ''}"
            source = _unit_citation_source(line.get("citations"))
        basis.append([f"{line['name']} 1㎥당 (㎥)", quantity, formula, source])
    labor_total = (statement.get("totals") or {}).get("labor")
    for line in priced.get("cost_lines", []):
        if line.get("kind") != "rate_cost":
            continue
        rate = Decimal(str(line.get("rate") or 0)) * 100
        base_amount = line.get("base_amount")
        if base_amount is None and line.get("base") == "labor_subtotal":
            base_amount = labor_total
        base_label = "노무비 계" if line.get("base") == "labor_subtotal" else str(line.get("base") or "기준액")
        basis.append([f"{line['name']} 기준액 (원)", _num(base_amount),
                      f"{base_label} {_num(base_amount):,}원 × {rate.normalize():g}%" if base_amount is not None else "",
                      _unit_citation_source(line.get("citations"))])
    notes = statement.get("basis_notes", [])
    if notes:
        basis.append(["적용 기준", "", "", ""])
        for index, note in enumerate(notes[:2]):
            label = "공사 규모 판정" if index == 0 else "안전관리비 기초액"
            basis.append([label, "", note, ""])
    return basis


def _grouped_header(sheet, widths: int, groups: list[tuple[str, int, int]],
                    fixed: list[tuple[str, int]]) -> None:
    for label, start, end in groups:
        sheet.merge_cells(start_row=4, start_column=start, end_row=4, end_column=end)
        sheet.cell(4, start, label)
        for column, label2 in ((start, "단가"), (start + 1, "금액")):
            sheet.cell(5, column, label2)
    for label, column in fixed:
        sheet.merge_cells(start_row=4, start_column=column, end_row=5, end_column=column)
        sheet.cell(4, column, label)


def build_xlsx(response: dict) -> bytes:
    tables = response["tables"]
    statement = response.get("statement") or {}
    metadata = _metadata(response)
    book = Workbook()
    _estimate_sheet(book, response)

    # 공사원가계산서
    widths = [16, 24, 16, 13, 56]
    sheet = _new_sheet(book, "원가계산서", "공사원가계산서", metadata, widths)
    headers = ["비목", "구분", "금액", "구성비", "비고(산출 근거)"]
    for column, value in enumerate(headers, 1):
        sheet.cell(4, column, value)
    cost_rows = _cost_statement_rows(tables, statement)
    for values in cost_rows:
        sheet.append(values[:4] + [values[4]])
        if values[0] in ("순공사원가", "총원가", "도급액") and not values[1]:
            row = sheet.max_row
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    cost_start, cost_end = 5, 4 + len(cost_rows)
    _merge_repeated(sheet, 1, cost_start, cost_end)
    totals = {"순공사원가", "총원가", "도급액"}
    bold_rows = tuple(row for row in range(cost_start, cost_end + 1)
                      if sheet.cell(row, 2).value == "소계" or sheet.cell(row, 1).value in totals)
    excluded_rows = tuple(row for row in range(cost_start, cost_end + 1)
                          if cost_rows[row - cost_start][5] in ("제외", "미산정"))
    contract_rows = tuple(row for row in range(cost_start, cost_end + 1)
                          if sheet.cell(row, 1).value == "도급액")
    missing_rows = _missing_cost_rows(response, cost_rows)
    missing_start = sheet.max_row + 2
    sheet.merge_cells(start_row=missing_start, start_column=1, end_row=missing_start, end_column=5)
    sheet.cell(missing_start, 1, "빠진 항목")
    missing_header = missing_start + 1
    for column, label in enumerate(("항목", "상태", "이유"), 1):
        sheet.cell(missing_header, column, label)
    for item, status, reason in missing_rows:
        sheet.append([item, status, reason, None, None])
        current = sheet.max_row
        sheet.merge_cells(start_row=current, start_column=3, end_row=current, end_column=5)
    missing_end = sheet.max_row
    rate_start, rate_header = _append_cost_rate_sources(sheet, response)
    footer = _footer(sheet, len(widths))
    _format_sheet(sheet, len(widths), (4, missing_header, rate_header), cost_start, cost_end,
                  numeric_columns=(3, 4), money_columns=(3,), ratio_columns=(4,),
                  muted_rows=excluded_rows, bold_rows=bold_rows,
                  highlight_rows=contract_rows, footer_row=footer)
    for column in range(1, 5):
        sheet.cell(rate_start, column).fill = PatternFill("solid", fgColor="D9EAD3")
        sheet.cell(rate_start, column).font = Font(name=_FONT_NAME, size=10, bold=True)
    sheet.freeze_panes = "A5"
    sheet.print_title_rows = "4:4"
    sheet.page_setup.orientation = "portrait"

    # 공사내역서: 실무 관행에 따라 재료비 → 노무비 → 경비
    bill_widths = [30, 26, 9, 12, 14, 14, 14, 14, 14, 14, 14, 14, 34]
    bill_sheet = _new_sheet(book, "내역서", "공사내역서", metadata, bill_widths, orientation="landscape")
    _grouped_header(bill_sheet, len(bill_widths),
                    [("재료비", 5, 6), ("노무비", 7, 8), ("경비", 9, 10), ("합계", 11, 12)],
                    [("품 명", 1), ("규 격", 2), ("단위", 3), ("수량", 4), ("비고", 13)])
    bill = tables.get("bill") or {}
    bill_rows = []
    if bill:
        prices = bill.get("unit_price") or {}
        amounts = bill.get("amount") or {}
        total_unit = (response.get("priced") or {}).get("total")
        note = "제1호표"
        bill_rows.append([_work_name(response), bill.get("spec", ""), bill.get("unit", ""),
                          _num(bill.get("quantity")), _num(prices.get("재료비")), _num(amounts.get("재료비")),
                          _num(prices.get("노무비")), _num(amounts.get("노무비")),
                          _num(prices.get("경비")), _num(amounts.get("경비")),
                          _num(total_unit), _num(bill.get("total")), note])
        bill_rows.append(["합계", "", "", None, None, _num(amounts.get("재료비")),
                          None, _num(amounts.get("노무비")),
                          None, _num(amounts.get("경비")),
                          None, _num(bill.get("total")), ""])
    for values in bill_rows:
        bill_sheet.append(values)
    bill_start, bill_end = 6, bill_sheet.max_row
    bill_note_row = None
    if bill and bill.get("partial"):
        bill_note_row = bill_sheet.max_row + 1
        bill_sheet.merge_cells(start_row=bill_note_row, start_column=1,
                               end_row=bill_note_row, end_column=len(bill_widths))
        bill_sheet.cell(bill_note_row, 1, "미산정 항목 제외 부분 합계")
    bill_footer = _footer(bill_sheet, len(bill_widths))
    bill_sheet.freeze_panes = "A6"
    bill_sheet.print_title_rows = "4:5"
    _format_sheet(bill_sheet, len(bill_widths), (4, 5), bill_start, bill_end,
                  numeric_columns=(4, 5, 6, 7, 8, 9, 10, 11, 12),
                  money_columns=(5, 6, 7, 8, 9, 10, 11, 12),
                  bold_rows=(bill_end,) if bill_rows else (), footer_row=bill_footer)
    if bill_note_row:
        bill_sheet.cell(bill_note_row, 1).font = Font(name=_FONT_NAME, size=9, italic=True, color="666666")

    # 일위대가표: 계 → 노무비 → 재료비 → 경비
    unit_widths = [30, 24, 12, 9, 14, 14, 14, 14, 14, 14, 14, 14, 34]
    unit_sheet = _new_sheet(book, "일위대가", "일위대가표", metadata, unit_widths)
    _grouped_header(unit_sheet, len(unit_widths),
                    [("계", 5, 6), ("노무비", 7, 8), ("재료비", 9, 10), ("경 비", 11, 12)],
                    [("공 종", 1), ("규 격", 2), ("수 량", 3), ("단 위", 4), ("비 고", 13)])
    unit_rows, unit_data_start = _unit_export_rows(response, tables.get("unit_rows", []))
    for values in unit_rows:
        unit_sheet.append(values)
    unit_end = unit_sheet.max_row
    unit_total_rows = (unit_data_start,)
    unit_muted = tuple(row for row in range(unit_data_start + 1, unit_end + 1)
                        if str(unit_sheet.cell(row, 13).value or "").startswith(("제외", "미산정")))
    unit_note_row = unit_end + 1
    unit_sheet.merge_cells(start_row=unit_note_row, start_column=1,
                           end_row=unit_note_row, end_column=len(unit_widths))
    unit_sheet.cell(unit_note_row, 1, "금액 0.1원 미만 버림(품셈 1-2-2)")
    unit_footer = _footer(unit_sheet, len(unit_widths))
    unit_sheet.freeze_panes = "A6"
    unit_sheet.print_title_rows = "4:5"
    _format_sheet(unit_sheet, len(unit_widths), (4, 5), unit_data_start, unit_end,
                  numeric_columns=(3, 5, 6, 7, 8, 9, 10, 11, 12),
                  money_columns=(6, 8, 10, 12), muted_rows=unit_muted,
                  bold_rows=unit_total_rows, footer_row=unit_footer)
    unit_sheet.cell(unit_note_row, 1).font = Font(name=_FONT_NAME, size=9, italic=True, color="666666")
    for row in range(unit_data_start + 1, unit_end + 1):
        if unit_sheet.cell(row, 1).value == "펌프차 운전원":
            unit_sheet.cell(row, 7).number_format = "#,##0.0"

    # 단가대비표
    rate_widths = [16, 34, 24, 14, 16, 52, 20, 44]
    rate_sheet = _new_sheet(book, "단가대비표", "단가대비표", metadata, rate_widths)
    rate_headers = ["구분", "품명", "규격", "단위", "적용 단가", "단가 근거", "적용 기간", "비고"]
    for column, value in enumerate(rate_headers, 1):
        rate_sheet.cell(4, column, value)
    rate_rows = [[row.get("kind", ""), row.get("name", ""), row.get("spec", ""), row.get("unit", ""),
                  "-" if row.get("price") == "-" else _num(row.get("price")),
                  row.get("source", ""), row.get("period", ""), row.get("note", "")]
                 for row in tables.get("rate_rows", [])]
    for values in rate_rows:
        rate_sheet.append(values)
    rate_start, rate_end = 5, rate_sheet.max_row
    rate_footer = _footer(rate_sheet, len(rate_widths))
    rate_sheet.print_title_rows = "4:4"
    _format_sheet(rate_sheet, len(rate_widths), (4,), rate_start, rate_end,
                  numeric_columns=(5,), money_columns=(5,), footer_row=rate_footer)

    # 산출근거
    basis_widths = [28, 22, 54, 58]
    basis_sheet = _new_sheet(book, "산출근거", "산출근거", metadata, basis_widths)
    for column, value in enumerate(("항목", "값", "계산", "근거(품셈 절·표·쪽)"), 1):
        basis_sheet.cell(4, column, value)
    basis = _basis_rows(response)
    for values in basis:
        basis_sheet.append(values)
    basis_start, basis_end = 5, basis_sheet.max_row
    basis_footer = _footer(basis_sheet, len(basis_widths))
    basis_sheet.print_title_rows = "4:4"
    _format_sheet(basis_sheet, len(basis_widths), (4,), basis_start, basis_end,
                  numeric_columns=(2,), money_columns=(), footer_row=basis_footer)

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
