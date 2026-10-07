"""여러 공종 묶음 견적: 항목 분리, 항목별 질문, 간접비 1회 계산, API·Excel을 오프라인으로 검사한다."""

from __future__ import annotations

import io
import os
import sys
from contextlib import nullcontext
from pathlib import Path
from unittest.mock import patch

os.environ["AGENT_OFFLINE"] = "1"
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from backend.agent.nodes.bundle import bundle, split_items  # noqa: E402
from backend.api.main import app  # noqa: E402
from backend.agent.tools.calc.cost_statement import calculate_cost_statement  # noqa: E402

CLIENT = TestClient(app, headers={"X-Guest-Session": "offline-bundle-session-" + "x" * 32})
QUERY = "철근구조물 레미콘 인력운반 타설 150㎥, 무근구조물 레미콘 인력운반 타설 50㎥ 견적, 도로공사"
ANSWERS = {"work": "6-1-1", "scattered_small_volume": False, "concrete_supply": "관급"}


def main() -> int:
    checks = []
    checks.append(("B-split1 쉼표·와 분리", split_items("레미콘 타설 100㎥와 거푸집 50㎡, 공사기간 6개월")
                   == ["레미콘 타설 100m3", "거푸집 50m2 공사기간 6개월"]))
    checks.append(("B-split2 천 단위 쉼표·치수는 항목 아님",
                   split_items("옹벽 높이 3m, 레미콘 1,500㎥ 타설") == ["옹벽 높이 3m 레미콘 1,500m3 타설"]))
    checks.append(("B-split3 앞 조건은 첫 항목에", split_items("도로공사, 레미콘 100㎥ 그리고 거푸집 20㎡")
                   == ["도로공사 레미콘 100m3", "거푸집 20m2"]))

    start = CLIENT.post("/api/chat", json={"message": QUERY}).json()
    thread = start["thread_id"]
    second = CLIENT.post("/api/chat", json={"thread_id": thread, "answers": ANSWERS}).json()
    done = CLIENT.post("/api/chat", json={"thread_id": thread, "answers": ANSWERS}).json()
    checks.append(("B-ask 항목별로 묻는다", start["status"] == "MISSING_INFO" and start["message"].startswith("1번 항목")
                   and second["status"] == "MISSING_INFO" and second["message"].startswith("2번 항목")
                   and len(second["items"]) == 1))

    items = done["items"]
    labor = sum(int(item["priced"]["reference_amounts"]["subtotals"]["노무비"]) for item in items)
    expenses = sum(int(item["priced"]["reference_amounts"]["subtotals"]["경비"]) for item in items)
    lines = done["statement"]["lines"]
    direct_labor = next(line["amount"] for line in lines if line["name"] == "직접노무비")
    checks.append(("B-sum 직접비 합산·간접비 1회", done["status"] in ("OK", "PARTIAL") and len(items) == 2
                   and direct_labor == labor
                   and next(line["amount"] for line in lines if line["name"] == "직접경비") == expenses
                   and sum(line["name"] == "간접노무비" for line in lines) == 1
                   and done["statement"]["conditions"]["work_category"] == "도로"))
    # 각 1,500만원이면 따로는 안전관리비 제외(2천만원 미만), 묶으면 3,000만원이라 산정된다.
    conditions = {"work_category": "도로", "duration": "1~6개월", "contractor_type": "종합건설업",
                  "project_scale": "이 견적만"}
    priced = {"reference_amounts": {"subtotals": {"재료비": "0", "노무비": "15000000", "경비": "0"},
                                    "total": "15000000"}, "rate_version": {}, "unpriced": [], "excluded": []}
    alone = calculate_cost_statement(priced, conditions, "2026-10-06")
    together = bundle({"items": [{"query": "a", "priced": priced}, {"query": "b", "priced": priced}],
                       "inputs": conditions, "basis_date": "2026-10-06"})
    safety = lambda statement: next(line["status"] for line in statement["lines"] if line["name"] == "산업안전보건관리비")
    checks.append(("B-scale 묶은 규모로 제비율·적용 기준 판정", safety(alone) == "제외"
                   and safety(together["statement"]) == "산정" and "30,000,000원" in together["statement"]["basis_notes"][0]))
    checks.append(("B-answer 묶음 설명", done["answer"].startswith("2개 공종을 한 견적서로")
                   and len(done["tables"]["bills"]) == 2 and all(item["reason"] == "" for item in items)))

    changed = CLIENT.post("/api/chat", json={"thread_id": thread, "conditions": {"contractor_type": "전문건설업"}}).json()
    checks.append(("B-conditions 조건 변경 재계산", changed["statement"]["conditions"]["contractor_type"] == "전문건설업"
                   and len(changed["items"]) == 2
                   and changed["statement"]["totals"]["contract_amount"] != done["statement"]["totals"]["contract_amount"]))

    export = CLIENT.get(f"/api/export/{thread}.xlsx")
    book = load_workbook(io.BytesIO(export.content))
    bill_names = [row[0] for row in book["내역서"].iter_rows(min_row=6, values_only=True)]
    unit_headers = [row[0] for row in book["일위대가"].iter_rows(min_row=6, values_only=True)
                    if str(row[0] or "").startswith("제 ")]
    checks.append(("B-excel 내역서 2행+합계·호표 2개", export.status_code == 200
                   and "%EC%99%B8%201%EA%B1%B4" in export.headers["content-disposition"]  # '외 1건'
                   and bill_names[:3] == [bill_names[0], bill_names[1], "합계"]
                   and [header[:4] for header in unit_headers] == ["제 1호", "제 2호"]))

    too_many = CLIENT.post("/api/chat", json={"message": "레미콘 10㎥, 거푸집 20㎡, 철근 3톤, 비계 40㎡ 견적"}).json()
    checks.append(("B-limit 4개 이상은 안내", too_many["status"] == "BLOCKED" and "3개" in too_many["message"]))
    single = CLIENT.post("/api/chat", json={"message": "철근구조물 레미콘 인력운반 타설 150㎥"}).json()
    checks.append(("B-single 단일 공종은 그대로", single["items"] == [] and single["status"] == "MISSING_INFO"
                   and not single["message"].startswith("1번 항목")))

    for name, ok in checks:
        print(("PASS " if ok else "FAIL ") + name)
    passed = sum(ok for _, ok in checks)
    print(f"통과 {passed} / 전체 {len(checks)}")
    return 0 if passed == len(checks) else 1


if __name__ == "__main__":
    with patch('backend.api.usage_limits.processing', return_value=nullcontext((None, 'fixture', False))), \
            patch('backend.api.usage_limits.consume', return_value={}), \
            patch('backend.api.usage_limits.status', return_value={}):
        raise SystemExit(main())
