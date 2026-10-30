"""Focused offline checks for generated calculation module boundaries."""
import copy
import json
import tempfile
import unittest
from fractions import Fraction
from unittest.mock import patch
from pathlib import Path

from backend.agent.tools.calc.generated import generated_daily_crew
from pipeline.calculation_modules import (ALLOWED_OPERATIONS, ResponseTruncatedError, author_candidate,
                                         SEMANTIC_AUTHORING_SCHEMA, author_semantic_candidate,
                                         assemble_semantic_candidate, validate_semantic_authoring,
                                         validate_candidate, run_candidate_case,
                                         run_candidate_case_detailed, main)

ROOT = Path(__file__).resolve().parents[1]
DRAFTS = ROOT / "data" / "drafts" / "specs"


def draft(path):
    return json.loads((DRAFTS / path).read_text(encoding="utf-8"))["draft"]


def test_daily_crew_non_numeric_source_cells_fail_closed_with_citation():
    from backend.agent.tools.calc.daily_crew import adjusted_daily_crew

    cases = [
        ("유지관리/1-2-2.json",
         {"work_condition": "인력 시공_낙엽수", "dbm_size": "31∼41㎝ 미만",
          "quantity": "1", "aerial_lift_spec": "고소작업차"}, "p915-t0", "-"),
        ("유지관리/4-3-1.json",
         {"work_type": "보호통", "diameter": "ø65~80", "quantity": "1",
          "parallel_replacement": False}, "p981-t0", "-"),
    ]
    for path, inputs, table_id, source_value in cases:
        result = adjusted_daily_crew(draft(path), inputs)
        assert result["status"] == "unresolvable", result
        assert "표 셀에 계산 가능한 수치가 없음" in result["reason"], result
        assert source_value in result["reason"], result
        assert table_id in json.dumps(result["citations"], ensure_ascii=False), result


def test_contract(executor="json_interpreter", policies=None, case_id="independent-1"):
    policy_note_id = (__import__("backend.agent.tools.calc.policies", fromlist=["POLICIES"]).POLICIES[
        next(iter(policies))]["source_note_id"] if policies else "n1")
    return {"contract_version": "1.0",
        "source": {"edition": "2026", "pdf_sha256": "a" * 64, "pages": [193],
                   "citations": [{"page": 193, "section_no": "6-3-1"}],
                "note_rule_map": [{"note_id": policy_note_id, "rule_ids": ["r1"]}]},
        "site_inputs": [
            {"name": name, "type": kind, "unit": unit, "source": "field fact"}
            for name, kind, unit in (
                ("height", "positive_rational", "m"), ("crane_used", "boolean", None),
                ("lifting_equipment_used", "boolean", None), ("quantity", "positive_rational", "ea"),
                ("girder_type", "enum", None), ("direct_lifting", "boolean", None),
                ("special_condition", "boolean", None))],
        "application": {"conditions": ["height > 7"], "target": {"name": "daily_output", "unit": "m3/day"},
                        "order": ["rule r1"], "system_policies": list(policies or []),
                        "policy_versions": ({next(iter(policies)): "1"} if policies else {}),
                        "policy_sources": ({next(iter(policies)): {"edition": "2026", "section_no": "6-3-1",
                            "pdf_page": 193, "printed_page": 137,
                            "text": "수직고가 7m를 초과하는 경우 매 3m마다 시공량을 9%까지 감한다."}} if policies else {})},
        "executor": {"kind": executor, "module_version": "1.0"},
        "verification": {"cases": [{"id": case_id, "citation": "PDF p.193"}]}}


def source_backed_formwork_candidate(spec):
    from pipeline.calculation_contract import build_source_bundle
    from pipeline.calculation_modules import collect_known_input_fields
    from backend.agent.tools.calc.policies import allowed_policy_ids
    pdf = next((ROOT / "data/raw/standard_estimation").glob("*.pdf"))
    source, bundle = build_source_bundle(pdf, ROOT / "data/processed/chunks.all.jsonl",
        ROOT / "data/processed/page_map.json", "공통", "6-3-1")
    policy_id = "formwork-6-3-1-height-started-interval-compound-v1"
    note_id = next(item["evidence_id"] for item in bundle["calculation_notes"]
                   if item["kind"] == "table_note")
    semantics = {
        "conditions": ["height above 7 m without lifting equipment", "roof slab under 20 degrees",
                       "continuous crane use"],
        "handled_condition_map": [
            {"condition": item["condition"], "rule_ids": [rule_id]}
            for item, rule_id in zip(spec.get("blocked", []),
                                     ("height-policy-ref", "roof-slope-reduction", "continuous-crane-no-output-deduction"))],
        "target": {"name": "daily output", "unit": "m3/day"},
        "order": ["height-policy-ref", "roof-slope-reduction", "continuous-crane-no-output-deduction"],
        "inputs": [],
        "rules": [
            {"id": "roof-slope-reduction", "when": {"is_roof_slab": True},
             "operation": "multiply_output", "value": "83/100", "source_note_ids": [note_id],
             "reason": "The table note reduces output by 17% for roof slabs under 20 degrees."},
            {"id": "continuous-crane-no-output-deduction", "when": {"crane_used": True},
             "operation": "multiply_output", "value": "1", "source_note_ids": [note_id],
             "reason": "The table note requires no daily-output deduction when a continuous crane is used."},
        ],
        "policy_references": [{"id": "height-policy-ref", "when": {"height": {"gt": 7},
            "crane_used": False, "lifting_equipment_used": False}, "policy_id": policy_id,
            "source_note_ids": [note_id], "reason": "Reference the exact registered source-note policy."}],
        "note_interpretations": [{"note_id": item["evidence_id"],
            "rule_ids": (["height-policy-ref", "roof-slope-reduction",
                          "continuous-crane-no-output-deduction"] if item["evidence_id"] == note_id else []),
            "excluded_reason": ("Base complexity selection and added crane resource costing are separate from daily-output modifiers."
                                if item["evidence_id"] == note_id else "This note is outside the focused output-adjustment rules.")}
            for item in bundle["calculation_notes"]],
    }
    cases = [{"id": "source-backed-boundary", "spec": spec}]
    known_conditions = {item["condition"]: item for item in spec.get("blocked", [])}
    known_fields = collect_known_input_fields([{"spec": spec}])
    return assemble_semantic_candidate(semantics, source, cases,
        allowed_system_policies=allowed_policy_ids("공통", "6-3-1"),
        known_blocked_conditions=known_conditions, known_input_fields=known_fields)


class GeneratedCalculationTests(unittest.TestCase):
    def test_semantic_schema_assembly_policy_reference_and_real_interpreter(self):
        from pipeline.calculation_contract import build_source_bundle
        pdf = next((ROOT / "data/raw/standard_estimation").glob("*.pdf"))
        source, bundle = build_source_bundle(pdf, ROOT / "data/processed/chunks.all.jsonl",
            ROOT / "data/processed/page_map.json", "공통", "6-3-1")
        spec = draft("공통/6-3-1.json")
        policy_id = "formwork-6-3-1-height-started-interval-compound-v1"
        source_note_id = __import__("backend.agent.tools.calc.policies", fromlist=["POLICIES"]).POLICIES[policy_id]["source_note_id"]
        self.assertEqual(source_note_id, next(item["evidence_id"] for item in bundle["calculation_notes"]
                                               if item["kind"] == "table_note"))
        table_note = next(item for item in bundle["calculation_notes"] if item["kind"] == "table_note")
        semantic = {"conditions": ["수직고 7m 초과", "지붕 슬래브 설치(경사도 20도 미만)"],
            "target": {"name": "daily output", "unit": "m3/day"}, "order": ["height", "roof"],
            "inputs": [],
            "handled_condition_map": [],
            "rules": [
                {"id": "roof", "when": {"is_roof_slab": True},
                 "operation": "multiply_output", "value": "83/100", "source_note_ids": [source_note_id],
                 "reason": "Official source states 17% reduction."},
            ],
            "policy_references": [{"id": "height", "when": {"height": {"gt": 7}, "crane_used": False, "lifting_equipment_used": False},
                 "policy_id": policy_id, "source_note_ids": [source_note_id], "reason": "System assumption reference."}],
            "note_interpretations": [{"note_id": item["evidence_id"],
                "rule_ids": (["height", "roof"] if item["evidence_id"] == source_note_id else []),
                "excluded_reason": ("The table's other crew/equipment and base-complexity treatments are outside this daily-output modifier contract."
                                    if item["evidence_id"] == source_note_id else "Other source note is outside the height/roof output modifiers in this focused fixture.")}
                for item in bundle["calculation_notes"]]}
        semantic["handled_condition_map"] = [
            {"condition": spec["blocked"][0]["condition"], "rule_ids": ["height"]},
            {"condition": spec["blocked"][1]["condition"], "rule_ids": ["roof"]}]
        known_conditions = {item["condition"]: item for item in spec["blocked"]}
        known_fields = __import__("pipeline.calculation_modules", fromlist=["collect_known_input_fields"]).collect_known_input_fields([{"spec": spec}])
        self.assertEqual(validate_semantic_authoring(semantic, source, {policy_id}, known_conditions, known_fields), [])
        self.assertEqual(set(SEMANTIC_AUTHORING_SCHEMA["required"]),
                         {"conditions", "handled_condition_map", "target", "order", "inputs", "rules", "policy_references", "note_interpretations"})
        cases = [
            {"id": "author-7.001", "spec": spec, "inputs": {"volume": "100", "complexity": "복잡", "height": "7.001", "is_roof_slab": False, "crane_used": False, "lifting_equipment_used": False}, "allowed_system_policies": [policy_id], "expected": {"status": "computed", "daily_volume_m3": "27.3"}},
            {"id": "author-40.001", "spec": spec, "inputs": {"volume": "100", "complexity": "복잡", "height": "40.001", "is_roof_slab": False, "crane_used": False, "lifting_equipment_used": False}, "allowed_system_policies": [policy_id], "expected": {"status": "computed", "daily_volume_m3": "9.67426462240814347997043"}},
            {"id": "author-roof", "spec": spec, "inputs": {"volume": "100", "complexity": "복잡", "height": "7.001", "is_roof_slab": True, "crane_used": False, "lifting_equipment_used": False}, "allowed_system_policies": [policy_id], "expected": {"status": "computed", "daily_volume_m3": "22.659"}},
        ]
        candidate = assemble_semantic_candidate(semantic, source, cases,
            allowed_system_policies={policy_id}, known_blocked_conditions=known_conditions,
            known_input_fields=known_fields)
        self.assertTrue(candidate["id"].startswith("generated-"))
        self.assertEqual(set(candidate["handled_conditions"]), {spec["blocked"][0]["condition"], spec["blocked"][1]["condition"]})
        self.assertNotIn(spec["blocked"][2]["condition"], candidate["handled_conditions"])
        self.assertEqual(candidate["calculation_contract"]["source"]["pages"], [
            {"pdf_page": item["pdf_page"], "printed_page": item["printed_page"]}
            for item in bundle["full_section_pdf_pages"]])
        self.assertEqual(candidate["calculation_contract"]["verification"]["cases"],
                         [{"id": c["id"], "citation": "trusted evaluation case"} for c in cases])
        policy_rule = next(rule for rule in candidate["rules"] if rule.get("system_policy_ref") == policy_id)
        self.assertEqual(policy_rule["operation"], "apply_system_policy")
        self.assertNotIn("value_expr", policy_rule)
        candidate["sections"] = [{"division": spec["division"], "section_no": "6-3-1"}]
        for case in cases:
            self.assertEqual(run_candidate_case(candidate, case), [], case["id"])
        representation_case = copy.deepcopy(cases[0])
        representation_case["expected"]["daily_volume_m3"] = "27.300"
        self.assertEqual(run_candidate_case_detailed(candidate, representation_case)["status"], "passed")
        execution_case = copy.deepcopy(cases[0])
        execution_case["inputs"]["height"] = "not-a-height"
        failed_execution = run_candidate_case_detailed(candidate, execution_case)
        self.assertEqual(failed_execution["status"], "execution_failed")
        self.assertEqual(failed_execution["execution"]["status"], "rejected")
        self.assertTrue(failed_execution["execution"]["reason"])
        self.assertEqual(run_candidate_case(candidate, execution_case)[0].split(":", 1)[0],
                         "execution failed")
        from backend.agent.tools.calc.generated import generated_daily_crew
        observed = generated_daily_crew(spec, cases[1]["inputs"], labor_only=True, module_override=candidate)
        adjustment = observed["provenance"]["daily_volume_m3"]["output_adjustments"][0]
        self.assertEqual(adjustment["official_source"]["pdf_page"], 193)
        self.assertEqual(adjustment["official_source"]["printed_page"], 137)
        self.assertIn("System-only convention", adjustment["system_assumption"])
        roof_result = generated_daily_crew(spec, cases[2]["inputs"], labor_only=True, module_override=candidate)
        self.assertEqual([entry.get("system_policy_ref") or entry.get("reason")
                          for entry in roof_result["provenance"]["daily_volume_m3"]["output_adjustments"]],
                         [policy_id, "Official source states 17% reduction."])
        prompts = []
        with tempfile.TemporaryDirectory() as temp:
            generated = author_semantic_candidate(source, cases, model=lambda prompt:
                (prompts.append(prompt) or json.dumps(semantic, ensure_ascii=False)),
                max_repairs=0, allowed_system_policies={policy_id}, output_dir=Path(temp), attempt_log=None)
            self.assertEqual(generated["status"], "candidate", generated)
            self.assertTrue(Path(generated["path"]).is_file())
        self.assertNotIn("27.3", prompts[0])
        self.assertNotIn("10.631060024624333494473", prompts[0])
        self.assertIn("policy_references", prompts[0])
        self.assertEqual(ALLOWED_OPERATIONS.__contains__("apply_system_policy"), True)

    def test_saved_bad_model_response_replay_rejects_contract_and_policy_formula(self):
        raw_path = ROOT / "evals/runtime_validation/gemini_6_1_4_2026-10-17/module-json-raw/generation-4.json"
        if not raw_path.exists(): self.skipTest("saved paid response is absent")
        text = json.loads(raw_path.read_text(encoding="utf-8"))["response_text"]
        parsed = json.loads(text)
        contract = parsed.get("calculation_contract", {})
        self.assertIsInstance(contract.get("source", {}).get("note_rule_map"), dict)
        self.assertTrue(any("expected_multiplier" in str(item) for item in contract.get("verification", {}).get("cases", [])))
        rules = parsed.get("rules", []) or parsed.get("calculation_contract", {}).get("executor", {}).get("rules", [])
        self.assertTrue(any("height" in rule.get("id", "") and "multiply_output" in rule
                            for rule in rules))
        errors = validate_candidate(parsed, allowed_system_policies={
            "formwork-6-3-1-height-started-interval-compound-v1"}, require_contract=True)
        self.assertTrue(errors)
        self.assertTrue(any("note_rule_map" in error for error in errors))
        height = next(rule for rule in rules if "height" in rule.get("id", ""))
        self.assertIn("ceil", json.dumps(height))
        from pipeline.calculation_modules import _expr
        legacy_factor = _expr(height["multiply_output"], {"vertical_height": "40.001"})
        self.assertEqual(legacy_factor, Fraction(-8, 100))
        semantic_raw = ROOT / "evals/runtime_validation/gemini_6_1_4_2026-10-17/module-json-raw/generation-7.json"
        if semantic_raw.exists():
            from evals.run_gemini_6_1_4 import _module_source
            _, bundle, _, _, author_cases, _, _ = _module_source()
            old_semantic = json.loads(json.loads(semantic_raw.read_text(encoding="utf-8"))["response_text"])
            labels = {item["condition"]: item for case in author_cases
                      for item in case.get("spec", {}).get("blocked", [])}
            field_catalog = __import__("pipeline.calculation_modules", fromlist=["collect_known_input_fields"]).collect_known_input_fields(author_cases)
            replay_errors = validate_semantic_authoring(old_semantic, bundle,
                {"formwork-6-3-1-height-started-interval-compound-v1"}, labels, field_catalog)
            self.assertTrue(any("additional property" in error for error in replay_errors))

    def test_semantic_policy_reference_must_not_include_a_formula(self):
        source = json.dumps({"calculation_notes": [{"note_id": "n"}]})
        semantic = {"conditions": ["height"], "handled_condition_map": [], "target": {"name": "daily", "unit": "m3/day"},
            "order": ["r"], "inputs": [], "rules": [], "policy_references": [{"id": "r", "when": {"height": {"gt": 7},
            "crane_used": False, "lifting_equipment_used": False},
            "policy_id": "formwork-6-3-1-height-started-interval-compound-v1", "value_expr": {"number": 0.91},
            "source_note_ids": ["n"], "reason": "x"}],
            "note_interpretations": [{"note_id": "n", "rule_ids": ["r"], "excluded_reason": ""}]}
        errors = validate_semantic_authoring(semantic, source,
            {"formwork-6-3-1-height-started-interval-compound-v1"})
        self.assertTrue(any("additional property" in item for item in errors))

    def test_installed_vertex_genai_sdk_accepts_semantic_json_schema(self):
        from google.genai import types
        config = types.GenerateContentConfig(response_mime_type="application/json",
                                             response_json_schema=SEMANTIC_AUTHORING_SCHEMA)
        self.assertIsNotNone(config.response_json_schema)
        self.assertEqual(config.response_mime_type, "application/json")

    def test_identical_semantic_contract_failure_stops_repair_loop(self):
        from pipeline.calculation_contract import build_source_bundle
        from backend.agent.tools.calc.policies import allowed_policy_ids
        from pipeline.calculation_modules import collect_known_input_fields
        pdf = next((ROOT / "data/raw/standard_estimation").glob("*.pdf"))
        source, bundle = build_source_bundle(pdf, ROOT / "data/processed/chunks.all.jsonl",
            ROOT / "data/processed/page_map.json", "공통", "6-3-1")
        spec = draft("공통/6-3-1.json")
        policy_id = "formwork-6-3-1-height-started-interval-compound-v1"
        note_id = next(item["evidence_id"] for item in bundle["calculation_notes"]
                       if item["kind"] == "table_note")
        response = {
            "conditions": ["height above 7 m without lifting equipment", "roof slab under 20 degrees"],
            "handled_condition_map": [
                {"condition": spec["blocked"][0]["condition"], "rule_ids": ["height-ref"]},
                {"condition": spec["blocked"][1]["condition"], "rule_ids": ["roof"]}],
            "target": {"name": "daily output", "unit": "m3/day"}, "order": ["height-ref", "roof"], "inputs": [],
            "rules": [{"id": "roof", "when": {"is_roof_slab": True}, "operation": "multiply_output",
                       "value": "83/100", "source_note_ids": [note_id], "reason": "Source note reduction."}],
            "policy_references": [{"id": "height-ref", "when": {"height": {"gt": 7},
                "crane_used": False, "lifting_equipment_used": False}, "policy_id": policy_id,
                "source_note_ids": [note_id], "reason": "Reference scoped system policy."}],
            "note_interpretations": [{"note_id": item["evidence_id"],
                "rule_ids": (["height-ref", "roof"] if item["evidence_id"] == note_id else []),
                "excluded_reason": ("Focused output adjustment scope." if item["evidence_id"] != note_id else "")}
                for item in bundle["calculation_notes"]]}
        # An identical schema defect is returned twice; preflight must pass and
        # the repair loop must stop without spending a third callback.
        response["policy_references"][0]["operation"] = "multiply_output"
        response["policy_references"][0]["value"] = "91/100"
        cases = [{"id": "authoring-fixture", "spec": spec, "inputs": {"volume": "100",
            "complexity": "복잡", "height": "7.001", "is_roof_slab": False,
            "crane_used": False, "lifting_equipment_used": False}, "expected": {
                "status": "computed", "daily_volume_m3": "27.3"}, "allowed_system_policies": [policy_id]}]
        calls = []
        with tempfile.TemporaryDirectory() as temp:
            result = author_semantic_candidate(source, cases,
                model=lambda prompt: (calls.append(prompt) or json.dumps(response)),
                max_repairs=2, allowed_system_policies=allowed_policy_ids("공통", "6-3-1"),
                output_dir=Path(temp), attempt_log=None)
        self.assertEqual(len(calls), 2)
        self.assertEqual(result["attempts"], 2)
        self.assertEqual(result["status"], "failed")
        self.assertEqual(result["exit_reason"], "identical semantic failure repeated; no same-error retry")
        self.assertEqual(result["model_calls"], 2)
        self.assertEqual(len(result["attempt_log"]), 2)
        self.assertFalse(list(Path(temp).glob("*.json")))

    def test_semantic_bad_policy_source_is_blocked_before_callback(self):
        from pipeline.calculation_contract import build_source_bundle
        from backend.agent.tools.calc.policies import allowed_policy_ids
        pdf = next((ROOT / "data/raw/standard_estimation").glob("*.pdf"))
        source, _ = build_source_bundle(pdf, ROOT / "data/processed/chunks.all.jsonl",
            ROOT / "data/processed/page_map.json", "공통", "6-3-1")
        broken = json.loads(source)
        broken["pdf_sha256"] = "0" * 64
        calls = []
        with tempfile.TemporaryDirectory() as temp:
            result = author_semantic_candidate(json.dumps(broken), [{"id": "source-preflight", "spec": draft("공통/6-3-1.json")}],
                model=lambda prompt: (calls.append(prompt) or "{}"), max_repairs=2,
                allowed_system_policies=allowed_policy_ids("공통", "6-3-1"),
                output_dir=Path(temp), attempt_log=None)
        self.assertEqual(calls, [])
        self.assertEqual(result["model_calls"], 0)
        self.assertEqual(result["exit_reason"], "source_preflight_failed")
        self.assertEqual(result["status"], "failed")
        self.assertFalse(list(Path(temp).glob("*.json")))
        self.assertFalse(list(Path(temp).glob("*.json")))

    def test_semantic_wrong_calculation_is_repaired_without_reference_value_leak(self):
        from pipeline.calculation_contract import build_source_bundle
        from backend.agent.tools.calc.policies import allowed_policy_ids
        pdf = next((ROOT / "data/raw/standard_estimation").glob("*.pdf"))
        source, bundle = build_source_bundle(pdf, ROOT / "data/processed/chunks.all.jsonl",
            ROOT / "data/processed/page_map.json", "공통", "6-3-1")
        spec = draft("공통/6-3-1.json")
        cases = json.loads((ROOT / "evals/calculation_modules/formwork_authoring_cases.json").read_text(encoding="utf-8"))
        for case in cases:
            case["spec"] = spec
        policy_id = "formwork-6-3-1-height-started-interval-compound-v1"
        note_id = next(item["evidence_id"] for item in bundle["calculation_notes"] if item["kind"] == "table_note")
        from pipeline.calculation_modules import collect_known_input_fields
        known = collect_known_input_fields(cases)
        def response(roof_factor):
            return {"conditions": ["roof slab under 20 degrees", "height above 7 m without lifting equipment"],
                "handled_condition_map": [{"condition": item["condition"], "rule_ids": [rid]}
                    for item, rid in zip(spec["blocked"], ("height-ref", "roof", "crane"))],
                "target": {"name": "daily output", "unit": "m3/day"},
                "order": ["height-ref", "roof", "crane"], "inputs": [],
                "rules": [{"id": "roof", "when": {"is_roof_slab": True}, "operation": "multiply_output",
                    "value": roof_factor, "source_note_ids": [note_id], "reason": "Source table note roof adjustment."},
                    {"id": "crane", "when": {"crane_used": True}, "operation": "multiply_output",
                    "value": "1", "source_note_ids": [note_id], "reason": "Source note crane handling."}],
                "policy_references": [{"id": "height-ref", "when": {"height": {"gt": 7}, "crane_used": False,
                    "lifting_equipment_used": False}, "policy_id": policy_id, "source_note_ids": [note_id],
                    "reason": "Scoped system policy."}],
                "note_interpretations": [{"note_id": item["evidence_id"],
                    "rule_ids": (["height-ref", "roof", "crane"] if item["evidence_id"] == note_id else []),
                    "excluded_reason": ("Focused rules only." if item["evidence_id"] != note_id else "")}
                    for item in bundle["calculation_notes"]]}
        calls = []
        replies = [response("82/100"), response("83/100")]
        with tempfile.TemporaryDirectory() as temp:
            result = author_semantic_candidate(source, cases,
                model=lambda prompt: (calls.append(prompt) or json.dumps(replies[len(calls)-1], ensure_ascii=False)),
                max_repairs=2, allowed_system_policies=allowed_policy_ids("공통", "6-3-1"),
                output_dir=Path(temp), attempt_log=None)
            self.assertEqual(result["status"], "candidate", result)
            self.assertEqual(result["attempts"], 2)
            self.assertEqual(len(calls), 2)
            self.assertIn("authoring case auth-height-roof: output field computed result mismatch failed", calls[1])
            self.assertNotIn("22.659", calls[1])
            self.assertNotIn("10.631060024624333494473", calls[1])
            self.assertTrue(Path(result["path"]).is_file())
            saved = json.loads(Path(result["path"]).read_text(encoding="utf-8"))
            self.assertEqual(saved["review_status"], "candidate_unreviewed")
            self.assertEqual(result["model_calls"], 2)
            self.assertEqual(result["exit_reason"], "candidate_unreviewed")
            self.assertEqual(result["attempt_log"][0]["feedback_delivered"], "")
            self.assertIn("auth-height-roof", result["attempt_log"][1]["feedback_delivered"])
    def test_truncated_generation_is_not_retried_as_semantic_repair(self):
        calls = []
        def model(_prompt):
            calls.append(True)
            raise ResponseTruncatedError("finish_reason=MAX_TOKENS")
        with tempfile.TemporaryDirectory() as temp:
            result = author_candidate("source", [{"id": "private-case", "inputs": {},
                "expected": {"daily_volume_m3": "123.45"}}], model=model,
                case_runner=lambda *_: [], output_dir=Path(temp), attempt_log=None)
        self.assertEqual(len(calls), 1)
        self.assertEqual(result["status"], "failed")
        self.assertIn("ResponseTruncatedError", result["errors"][0])

    def test_roof_reduces_daily_output_not_quantity_and_keeps_source_immutable(self):
        spec = draft("공통/6-3-1.json")
        original = copy.deepcopy(spec)
        result = generated_daily_crew(spec, {
            "volume": "100", "complexity": "복잡", "height": "7.001",
            "is_roof_slab": True, "crane_used": True, "lifting_equipment_used": True,
        }, labor_only=True)
        self.assertEqual(result["status"], "computed")
        self.assertEqual(result["daily_volume_m3"], "24.9")
        self.assertEqual(result["work_days"], "1000/249")
        self.assertEqual(spec, original)
        trace = result["provenance"]["generated_module"]
        self.assertEqual(trace["rules"][0]["source"], "2026 건설공사 표준품셈 PDF 193쪽(인쇄 137쪽), p193-t0 비고")
        base = result["provenance"]["daily_volume_m3"]["base_output"]
        self.assertEqual(base["value"], "30")
        self.assertEqual(result["provenance"]["daily_volume_m3"]["output_adjustments"][0]["after"], "24.9")

    def test_height_boundaries_and_unresolved_policy(self):
        spec = next(json.loads(path.read_text(encoding="utf-8"))["draft"] for path in DRAFTS.rglob("6-3-1.json") if json.loads(path.read_text(encoding="utf-8"))["draft"].get("division") == "\uacf5\ud1b5")
        candidate = source_backed_formwork_candidate(spec)
        common = {"volume": "100", "complexity": "\ubcf5\uc7a1", "is_roof_slab": False}
        stale_registered = generated_daily_crew(spec, {**common, "height": "8", "crane_used": False,
            "lifting_equipment_used": False}, labor_only=True)
        self.assertEqual(stale_registered["status"], "rejected")
        self.assertIn("not authorized", stale_registered.get("reason", ""))
        for height, expected in (("7", "30"), ("7.001", "27.3"), ("10", "27.3"),
                                 ("10.001", "24.843"), ("13", "24.843"), ("13.001", "22.60713"),
                                 ("39.999", "10.631060024624333494473"),
                                 ("40", "10.631060024624333494473"),
                                 ("40.001", "9.67426462240814347997043")):
            result = generated_daily_crew(spec, {**common, "height": height, "crane_used": False,
                "lifting_equipment_used": False}, labor_only=True, module_override=candidate)
            self.assertEqual(result["status"], "computed", height)
            self.assertEqual(result["daily_volume_m3"], expected, height)
            self.assertEqual(result["provenance"]["daily_volume_m3"]["base_output"]["value"], "30")
        for height in ("7.001", "10", "10.001", "13"):
            result = generated_daily_crew(spec, {**common, "height": height, "crane_used": False,
                "lifting_equipment_used": True}, labor_only=True, module_override=candidate)
            self.assertEqual(result["daily_volume_m3"], "30")
        for height in ("8", "13"):
            result = generated_daily_crew(spec, {**common, "height": height, "crane_used": True,
                "lifting_equipment_used": False}, labor_only=True, module_override=candidate)
            self.assertEqual(result["status"], "computed")
            self.assertEqual(result["daily_volume_m3"], "30")
        combined = generated_daily_crew(spec, {**common, "height": "7.001", "is_roof_slab": True,
            "crane_used": False, "lifting_equipment_used": False}, labor_only=True, module_override=candidate)
        self.assertEqual(combined["daily_volume_m3"], "22.659")
        self.assertEqual(combined["provenance"]["daily_volume_m3"]["base_output"]["value"], "30")
        self.assertEqual(len(combined["provenance"]["daily_volume_m3"]["output_adjustments"]), 2)

    def test_height_candidate_contract_and_production_interpreter(self):
        spec = copy.deepcopy(draft("공통/6-3-1.json"))
        policy_ids = {"formwork-6-3-1-height-started-interval-compound-v1"}
        policy_note_id = __import__("backend.agent.tools.calc.policies", fromlist=["POLICIES"]).POLICIES[
            next(iter(policy_ids))]["source_note_id"]
        candidate = {"id": "height-candidate", "version": "1.0", "source_hash": "a" * 64,
            "calculation_contract": test_contract(policies=policy_ids),
            "sections": [{"division": spec["division"], "section_no": "6-3-1"}],
            "system_policy_ids": list(policy_ids),
            "policy_note_ids": {next(iter(policy_ids)): policy_note_id},
            "policy_versions": {"formwork-6-3-1-height-started-interval-compound-v1": "1"},
            "policy_sources": {"formwork-6-3-1-height-started-interval-compound-v1":
                __import__("backend.agent.tools.calc.policies", fromlist=["POLICIES"]).POLICIES[
                    "formwork-6-3-1-height-started-interval-compound-v1"]["source"]},
            "inputs": [{"name": "lifting_equipment_used", "type": "boolean",
                        "allowed_values": [True, False], "unit": None, "required": True,
                        "ask": "Is lifting equipment used?", "source": "source"}],
            "handled_conditions": ["수직고 7m 초과"], "rules": [{
            "id": "r1", "when": {"height": {"gt": 7}, "crane_used": False,
                         "lifting_equipment_used": False},
                "operation": "apply_system_policy",
                "system_policy_ref": "formwork-6-3-1-height-started-interval-compound-v1",
                "source": "2026 Standard Construction Estimate, common 6-3-1, PDF p.193 (printed p.137)"}]}
        legacy_errors = validate_candidate(candidate, allowed_system_policies=policy_ids)
        self.assertTrue(any("stable PDF/table evidence locator" in error for error in legacy_errors))
        candidate = source_backed_formwork_candidate(spec)
        self.assertEqual(validate_candidate(candidate, allowed_system_policies=policy_ids, require_contract=True), [])
        self.assertNotIn("height_interval_deduction", ALLOWED_OPERATIONS)
        bad_units = copy.deepcopy(candidate)
        height_rule = next(rule for rule in bad_units["rules"] if rule.get("system_policy_ref") == next(iter(policy_ids)))
        height_rule["operation"] = "multiply_output"
        height_rule["value_expr"] = {"sub": [{"number": 1}, {"mul": [{"number": "0.09"},
            {"ceil": {"div": [{"sub": [{"input": "height"}, {"number": 7}]}, {"number": 3}]}}]}]}
        self.assertTrue(run_candidate_case(bad_units, {"spec": spec, "inputs": {
            "height": "8", "volume": "1", "complexity": "복잡", "is_roof_slab": False,
            "crane_used": False, "lifting_equipment_used": False},
            "allowed_system_policies": sorted(policy_ids), "expected": {"status": "computed"}}))
        self.assertTrue(validate_candidate({**candidate, "inputs": [{"name": "height_deduction_percent",
            "type": "percentage", "required": True, "ask": "감산률을 입력해 주세요."}]},
            allowed_system_policies=policy_ids))
        spec["quantity_model"]["params"]["module_id"] = candidate["id"]
        for height, expected in (("7.001", "27.3"), ("40", "10.631060024624333494473"),
                                 ("40.001", "9.67426462240814347997043")):
            case = {"spec": spec, "inputs": {"volume": "100", "complexity": "복잡",
                    "height": height, "is_roof_slab": False, "crane_used": False,
                    "lifting_equipment_used": False}, "allowed_system_policies": sorted(policy_ids),
                    "expected": {"status": "computed", "daily_volume_m3": expected}}
            self.assertEqual(run_candidate_case(candidate, case), [], height)
        broken = copy.deepcopy(candidate)
        next(rule for rule in broken["rules"] if rule.get("system_policy_ref")).pop("system_policy_ref")
        self.assertTrue(validate_candidate(broken))
        deleted_rule = copy.deepcopy(candidate)
        deleted_rule["rules"] = []
        self.assertTrue(validate_candidate(deleted_rule, allowed_system_policies=policy_ids, require_contract=True))
        prompts = []
        prompt_case = {"id": "independent-1", "spec": spec, "inputs": {"volume": "100", "complexity": "복잡",
            "height": "7.001", "is_roof_slab": False, "crane_used": False,
            "lifting_equipment_used": False}, "allowed_system_policies": sorted(policy_ids),
            "expected": {"status": "computed", "daily_volume_m3": "27.3"}}
        with tempfile.TemporaryDirectory() as temp:
            author_candidate("evidence", [prompt_case], model=lambda prompt: (prompts.append(prompt) or json.dumps(candidate)),
                             case_runner=run_candidate_case, allowed_system_policies=policy_ids,
                             output_dir=Path(temp), attempt_log=None)
        self.assertIn("Do not invent system policy", prompts[0])
        self.assertIn("note_rule_map is a LIST", prompts[0])
        self.assertIn("(1 - 9/100)^k", prompts[0])
        self.assertIn("do not add inputs, outputs, expected values", prompts[0])
        self.assertNotIn("height_interval_deduction", prompts[0])
        self.assertNotIn("height_input", prompts[0])
        self.assertNotIn("27.3", prompts[0])

    def test_registered_module_reaches_gate_instead_of_stale_blocked_if(self):
        from backend.agent.nodes.gate import gate
        spec = draft("공통/6-3-1.json")
        with patch("backend.agent.nodes.gate.load_specs", return_value={spec["id"]: spec}):
            result = gate({"spec_id": spec["id"], "inputs": {
                "height": "10", "crane_used": False, "is_roof_slab": False,
            }})
        self.assertNotEqual(result.get("status"), "BLOCKED")

    def test_registered_module_does_not_bypass_unrelated_block(self):
        from backend.agent.tools.calc.generated import check_registered_blocked
        spec = draft("공통/6-3-1.json")
        spec["blocked"].append({"condition": "separate safety rule", "blocked_if": {
            "input": "crane_used", "op": "==", "value": True}, "reason": "separate", "source": "test"})
        blocked = check_registered_blocked(spec, {"crane_used": True})
        self.assertEqual(blocked["reason"], "separate")

    def test_per_unit_opaque_source_denominators_require_exact_unit_identity(self):
        from backend.agent.tools.calc.per_unit import conversion, parse_basis

        won = chr(51064)
        unit_device = chr(45824)
        unit_lot = chr(54596) + chr(51648)
        unit_sheet = chr(47560)
        self.assertEqual(parse_basis(won + "/" + unit_device), (1, unit_device))
        self.assertEqual(parse_basis(won + "/1,000" + unit_sheet), (1000, unit_sheet))
        self.assertEqual(conversion(unit_device, unit_device), Fraction(1))
        self.assertEqual(conversion(unit_lot, unit_device), None)
        self.assertEqual(parse_basis(won), None)  # a numerator without a basis is not guessed
        self.assertEqual(conversion(unit_device, None), None)

    def test_steel_girder_addition_and_subtraction(self):
        spec = draft("토목/5-3-2.json")
        base = {"girder_type": "개구제형", "quantity": "8", "location": "육상",
                "crane_size": "100~300ton", "special_condition": False}
        plus = generated_daily_crew(spec, {**base, "direct_lifting": True}, labor_only=True)
        minus = generated_daily_crew(spec, {**base, "direct_lifting": False, "special_condition": True}, labor_only=True)
        self.assertEqual(plus["daily_volume_m3"], "5")
        self.assertEqual(minus["daily_volume_m3"], "3")

    def test_existing_crew_addition_and_equipment_time_calculator(self):
        spec = json.loads((ROOT / "backend/agent/rules/specs/common/6-1-4_pump.json").read_text(encoding="utf-8"))
        values = {field["name"]: field["allowed_values"][0]
                  for field in spec["inputs"]
                  if isinstance(field.get("allowed_values"), list) and field["name"] != "concrete_supply"}
        values["volume"] = "100"
        base = __import__("backend.agent.tools.calc.daily_crew", fromlist=["adjusted_daily_crew"]).adjusted_daily_crew(
            spec, values, labor_only=True)
        rule = spec["crew_rules"][0]
        trigger, value = next(iter(rule["when"].items()))
        changed = __import__("backend.agent.tools.calc.daily_crew", fromlist=["adjusted_daily_crew"]).adjusted_daily_crew(
            spec, {**values, trigger: value}, labor_only=True)
        trade, increase = next(iter(rule["change"].items()))
        self.assertEqual(base["status"], "computed")
        self.assertEqual(changed["status"], "computed")
        self.assertEqual(Fraction(changed["person_days"][trade]) - Fraction(base["person_days"][trade]),
                         Fraction(changed["work_days"]) * int(increase))
        self.assertEqual(changed["equipment_days"]["콘크리트펌프차"], changed["work_days"])

    def test_bad_candidate_and_bounded_repair_never_promotes(self):
        self.assertTrue(validate_candidate("this is not JSON"))
        spec = draft("토목/5-3-2.json")
        candidate = {"id": "example", "version": "1.0", "source_hash": "a" * 64,
                     "sections": [{"division": spec["division"], "section_no": spec["section_no"]}],
                     "calculation_contract": test_contract(case_id="independent-1"),
                     "handled_conditions": ["지조립 없이 직접 인양하여 가설하는 경우"],
                     "rules": [{"id": "r1", "when_expr": {"gt": [{"input": "quantity"}, {"quantity": {"value": "4", "unit": "ea"}}]},
                                "operation": "add_output",
                                "value_expr": {"mul": [
                                    {"floor": {"div": [{"input": "quantity"}, {"quantity": {"value": "4", "unit": "ea"}}]}},
                                    {"quantity": {"value": "1", "unit": "m3/day"}}]},
                                "source": "independent test"}]}
        responses = iter([json.dumps({"id": "bad"}), json.dumps(candidate)])
        calls = []
        def model(prompt):
            calls.append(prompt)
            return next(responses)
        case = {"id": "independent-1", "spec": spec, "inputs": {"girder_type": "개구제형", "quantity": "8", "location": "육상",
                                            "direct_lifting": False, "special_condition": False,
                                            "crane_size": "100~300ton"},
                "expected": {"status": "computed", "daily_volume_m3": "6"}}
        with tempfile.TemporaryDirectory() as temp:
            result = author_candidate("source bundle", [case], model=model,
                                      case_runner=run_candidate_case, output_dir=Path(temp), attempt_log=None)
            self.assertEqual(result["status"], "candidate", result)
            saved = json.loads(Path(result["path"]).read_text(encoding="utf-8"))
            self.assertEqual(result["status"], "candidate")
            self.assertEqual(result["attempts"], 2)
            self.assertEqual(saved["review_status"], "candidate_unreviewed")
            self.assertNotEqual(Path(result["path"]).parent.name, "approved")
        self.assertEqual(len(calls), 2)
        unsafe = {"id": "unsafe", "source_hash": "source", "rules": [
            {"when": {"flag": True}, "operation": "python_exec", "value": "__import__('os')", "source": "x"}]}
        self.assertTrue(validate_candidate(unsafe))

    def test_candidate_path_and_output_are_not_overwritten(self):
        candidate = {"id": "..\\approved\\evil", "source_hash": "x", "rules": [
            {"when": {"x": True}, "operation": "multiply_output", "value": "1", "source": "x"}]}
        self.assertTrue(validate_candidate(candidate))
        safe = {**candidate, "id": "safe"}
        with tempfile.TemporaryDirectory() as temp:
            target = Path(temp) / "safe.json"
            target.write_text("keep", encoding="utf-8")
            response = author_candidate("source", [{}], model=lambda _: json.dumps(safe),
                                        case_runner=lambda *_: [], output_dir=Path(temp), attempt_log=None)
            self.assertEqual(response["status"], "failed")
            self.assertEqual(target.read_text(encoding="utf-8"), "keep")

    def test_cli_defaults_to_offline_prepare(self):
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "evidence.txt"
            cases = Path(temp) / "cases.json"
            source.write_text("source facts", encoding="utf-8")
            cases.write_text("[]", encoding="utf-8")
            with patch("sys.stdout", new_callable=__import__("io").StringIO) as stdout:
                status = main(["--source", str(source), "--cases", str(cases)])
            self.assertEqual(status, 0)
            self.assertIn('"call_model": false', stdout.getvalue())

    def test_repair_limit_is_three_and_repeated_failure_stays_unregistered(self):
        calls = []
        def model(_prompt):
            calls.append(1)
            return "{}"
        with tempfile.TemporaryDirectory() as temp:
            result = author_candidate("source", [{}], model=model, max_repairs=3,
                                      case_runner=lambda *_: [], output_dir=Path(temp), attempt_log=None)
            self.assertEqual(result["status"], "failed")
            self.assertEqual(result["attempts"], 4)
            self.assertEqual(len(calls), 4)
            self.assertEqual(list(Path(temp).glob("*.json")), [])

    def test_python_candidate_runner_requires_isolated_resource_limited_docker(self):
        from pipeline.python_sandbox import docker_run_command, run_generated_python, SandboxUnavailable
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            candidate, harness, case_file = root / "candidate.py", root / "harness.py", root / "case.json"
            for path in (candidate, harness, case_file):
                path.write_text("pass", encoding="utf-8")
            cmd = docker_run_command("python:3.12@sha256:" + "a" * 64, candidate, harness, case_file, root)
            for flag in ("--network=none", "--read-only", "--cap-drop=ALL", "--memory=256m",
                         "--pids-limit=32", "--cpus=1", "--user=65534:65534"):
                self.assertIn(flag, cmd)
            with self.assertRaises(ValueError):
                docker_run_command("python:latest", candidate, harness, case_file, root)
        with patch("pipeline.python_sandbox.subprocess.run", side_effect=FileNotFoundError("docker")):
            with self.assertRaises(SandboxUnavailable):
                run_generated_python("def calculate(spec, inputs): return inputs", {}, {},
                                     image="python:3.12@sha256:" + "a" * 64)

    def test_isolated_python_payload_excludes_expected_values_and_rejects_constant_stub(self):
        from pipeline.python_sandbox import run_generated_python, validate_python_candidate
        source = "def calculate(spec, inputs, labor_only=False):\n    return {'status': 'computed', 'value': inputs['x']}\n"
        self.assertFalse(validate_python_candidate(source))
        self.assertTrue(validate_python_candidate("def calculate(spec, inputs):\n    return {'value': 17}\n"))
        seen = {}
        def fake_run(command, **kwargs):
            payload_path = next(Path(arg.split('src=', 1)[1].split(',dst=', 1)[0])
                                for arg in command if isinstance(arg, str) and 'src=' in arg and 'case.json' in arg)
            seen.update(json.loads(payload_path.read_text(encoding='utf-8')))
            class Result:
                returncode = 0
                stdout = '{"status":"ok"}\n'
                stderr = ""
            return Result()
        def dispatch(command, **kwargs):
            if command[1:3] == ["image", "inspect"]:
                return type("I", (), {"returncode": 0})()
            return fake_run(command, **kwargs)
        with patch("pipeline.python_sandbox.subprocess.run", side_effect=dispatch):
            result = run_generated_python(source, {"quantity_model": {}}, {"x": 2}, image="python:3.12@sha256:"+"a"*64)
        self.assertEqual(result["status"], "ok")
        self.assertEqual(seen, {"spec": {"quantity_model": {}}, "inputs": {"x": 2}})
        self.assertNotIn("expected", seen)

    def test_python_mock_repair_feedback_hides_reference_values(self):
        from pipeline.calculation_modules import author_python_candidate
        constant = "def calculate(spec, inputs, labor_only=False):\n    return {'value': 14}\n"
        repaired = "def calculate(spec, inputs, labor_only=False):\n    return {'value': inputs['x'] * 2}\n"
        prompts = []
        cases = [{"id": "private-expected", "spec": {}, "inputs": {"x": 7}, "expected": {"value": 14}}]
        def model(prompt):
            prompts.append(prompt)
            return constant if len(prompts) == 1 else repaired
        def sandbox(source, spec, inputs, **kwargs):
            if "'value': 14" in source:
                return {"status": "ok", "result": {"value": 14}}
            return {"status": "ok", "result": {"value": inputs["x"] * 2}}
        with tempfile.TemporaryDirectory() as temp, patch("pipeline.python_sandbox.run_generated_python", side_effect=sandbox):
            result = author_python_candidate("source evidence", cases, model=model,
                sandbox_image="python@sha256:"+"a"*64, output_dir=Path(temp), attempt_log=None)
            candidate = Path(result["path"])
            metadata = json.loads(candidate.with_suffix(".candidate.json").read_text(encoding="utf-8"))
            import hashlib
            code_hash = hashlib.sha256(candidate.read_bytes()).hexdigest()
        self.assertEqual(result["status"], "candidate")
        self.assertEqual(result["attempts"], 2)
        self.assertNotIn("14", prompts[0])
        self.assertNotIn("14", prompts[1])
        self.assertEqual(metadata["code_sha256"], code_hash)

    def test_python_generation_contract_matches_original_section_and_notes(self):
        from pipeline.calculation_modules import author_python_candidate
        from pipeline.calculation_contract import CONTRACT_VERSION
        contract = {"contract_version": CONTRACT_VERSION,
            "source": {"edition": "2026", "pdf_sha256": "a"*64, "division": "common",
                "section_no": "probe", "pages": [1], "citations": [{"page": 1}],
                "note_rule_map": [{"note_id": "note-1", "rule_ids": ["rule-1"]}]},
            "site_inputs": [{"name": "x", "type": "positive_rational", "unit": "ea", "source": "site"}],
            "application": {"conditions": [], "target": "value", "order": ["rule-1"], "system_policies": []},
            "executor": {"kind": "isolated_python", "module_version": "1"},
            "verification": {"cases": [{"id": "author-1", "citation": "page 1"}]}}
        source = "def calculate(spec, inputs, labor_only=False):\n    return {'status': 'computed', 'value': inputs['x']}\n"
        contract["verification"]["cases"].append({"id": "author-2", "citation": "page 1"})
        response = json.dumps({"code": source, "calculation_contract": contract})
        cases = [
            {"id": "author-1", "spec": {}, "inputs": {"x": 8}, "expected": {"value": 8}},
            {"id": "author-2", "spec": {}, "inputs": {"x": 9}, "expected": {"value": 9}},
        ]
        def sandbox(code, spec, inputs, **kwargs):
            return {"status": "ok", "result": {"status": "computed", "value": inputs["x"]}}
        withheld = [{"id": "heldout-1", "spec": {}, "inputs": {"x": 10},
                     "expected": {"status": "computed", "value": 10}}]
        with tempfile.TemporaryDirectory() as temp, patch("pipeline.python_sandbox.run_generated_python", side_effect=sandbox):
            result = author_python_candidate("source", cases, model=lambda _: response,
                sandbox_image="python@sha256:"+"a"*64, max_repairs=0, output_dir=Path(temp),
                evaluation_cases=withheld,
                require_contract=True, source_pdf_hash="a"*64, source_note_ids={"note-1"},
                division="common", section_no="probe", attempt_log=None)
            saved_contract = json.loads(Path(result["path"]).with_suffix(".contract.json").read_text(encoding="utf-8"))
            candidate = Path(result["path"])
            metadata = json.loads(candidate.with_suffix(".candidate.json").read_text(encoding="utf-8"))
            import hashlib
            code_hash = hashlib.sha256(candidate.read_bytes()).hexdigest()
        self.assertEqual(result["status"], "candidate", result)
        self.assertEqual(saved_contract, contract)
        self.assertEqual(metadata["code_sha256"], code_hash)

    def test_contracted_python_refuses_single_input_or_constant_expected_suite(self):
        from pipeline.calculation_modules import author_python_candidate
        base = {"id": "one", "spec": {}, "inputs": {"x": 1}, "expected": {"value": 4}}
        called = []
        with self.assertRaisesRegex(ValueError, "vary site inputs"):
            author_python_candidate("source", [base, {**base, "id": "two"}],
                model=lambda prompt: called.append(prompt), sandbox_image="python@sha256:"+"a"*64,
                evaluation_cases=[{"id": "heldout", "spec": {}, "inputs": {"x": 3},
                                   "expected": {"value": 4}}], require_contract=True,
                attempt_log=None)
        self.assertEqual(called, [])

    def test_registered_python_module_routes_estimate_inputs_to_sandbox(self):
        from backend.agent.tools.calc import generated
        from backend.agent.nodes.compute import calculator_for
        source = "def calculate(spec, inputs, labor_only=False):\n    return {'status': 'computed'}\n"
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            registry_path = root / "approved" / "python-registry.json"
            code_path = root / "approved" / "python" / "calc.py"
            code_path.parent.mkdir(parents=True)
            code_path.write_text(source, encoding="utf-8")
            contract = {"contract_version": "1.0",
                "source": {"edition": "2026", "pdf_sha256": "a"*64, "pages": [1],
                           "citations": [{"page": 1}], "note_rule_map": []},
                "site_inputs": [{"name": "x", "type": "positive_rational", "unit": "ea", "source": "site"}],
                "application": {"conditions": [], "target": "daily output", "order": [], "system_policies": []},
                "executor": {"kind": "isolated_python", "module_version": "1"},
                "verification": {"cases": [{"id": "t1", "citation": "page 1"}]}}
            import hashlib
            module = {"id": "calc", "version": "1", "sections": [{"division": "common", "section_no": "x"}],
                "inputs": ["x"], "source_hash": "a"*64, "code_sha256": hashlib.sha256(code_path.read_bytes()).hexdigest(),
                "file": "calc.py", "sandbox_image": "python@sha256:"+"b"*64,
                "review_status": "approved_for_runtime", "calculation_contract": contract}
            registry_path.write_text(json.dumps({"registry_version": 1, "modules": [module]}), encoding="utf-8")
            spec = {"division": "common", "section_no": "x", "inputs": [{"name": "x"}]}
            with patch.object(generated, "MODULE_DIR", root), patch.object(generated, "PYTHON_REGISTRY", registry_path), \
                    patch("pipeline.python_sandbox.run_generated_python", return_value={"status": "ok", "result": {"status": "computed"}}) as run:
                self.assertEqual(calculator_for(spec), generated.generated_python_daily_crew)
                result = generated.generated_python_daily_crew(spec, {"x": 2})
            self.assertEqual(result["status"], "computed")
            self.assertEqual(run.call_args.args[1:], (spec, {"x": 2}))

    def test_docker_python_authoring_withheld_evaluation_registers_and_estimates(self):
        """Exercise the real pinned container across candidate, final eval, registry and estimate paths."""
        import hashlib
        import subprocess
        from pipeline.calculation_contract import CONTRACT_VERSION
        from pipeline.calculation_modules import author_python_candidate, register_python_candidate
        from backend.agent.tools.calc import generated
        from backend.agent.nodes.compute import calculator_for

        image = "python@sha256:a6e34c598f2467ed0e9a8d349809fcd8b5c603269512df273a0bb1784edc11b1"
        try:
            available = subprocess.run(["docker", "image", "inspect", image],
                                       capture_output=True, timeout=10).returncode == 0
        except (OSError, subprocess.TimeoutExpired):
            available = False
        if not available:
            self.skipTest("digest-pinned Docker image is unavailable")

        source = ("def calculate(spec, inputs, labor_only=False):\n"
                  "    return {'status': 'computed', 'daily_volume_m3': str(inputs['quantity'] * 3)}\n")
        contract = {"contract_version": CONTRACT_VERSION,
            "source": {"edition": "2026", "pdf_sha256": "a" * 64, "division": "common",
                "section_no": "probe", "pages": [1], "citations": [{"page": 1}],
                "note_rule_map": []},
            "site_inputs": [{"name": "quantity", "type": "positive_rational", "unit": "m3", "source": "site"}],
            "application": {"conditions": [], "target": {"name": "daily_volume_m3", "unit": "m3/day"},
                "order": ["daily output = 3 × quantity"], "system_policies": []},
            "executor": {"kind": "isolated_python", "module_version": "1"},
            "verification": {"cases": [{"id": "author-1", "citation": "PDF p.1"},
                                         {"id": "author-2", "citation": "PDF p.1"}]}}
        response = json.dumps({"code": source, "calculation_contract": contract})
        author = [
            {"id": "author-1", "spec": {}, "inputs": {"quantity": 7},
             "expected": {"status": "computed", "daily_volume_m3": "21"}},
            {"id": "author-2", "spec": {}, "inputs": {"quantity": 8},
             "expected": {"status": "computed", "daily_volume_m3": "24"}},
        ]
        heldout = [{"id": "heldout-1", "spec": {}, "inputs": {"quantity": 17},
                    "expected": {"status": "computed", "daily_volume_m3": "51"}}]

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            candidate_result = author_python_candidate("test evidence", author, model=lambda _: response,
                sandbox_image=image, max_repairs=0, output_dir=root / "candidates", evaluation_cases=heldout,
                require_contract=True, source_pdf_hash="a" * 64, source_note_ids=set(),
                division="common", section_no="probe", attempt_log=None)
            self.assertEqual(candidate_result["status"], "candidate", candidate_result)
            candidate = Path(candidate_result["path"])
            final = candidate_result["final_evaluation"]
            evaluation_record = {**final, "status": "passed", "code_sha256": hashlib.sha256(candidate.read_bytes()).hexdigest()}
            candidate.with_suffix(".evaluation.json").write_text(json.dumps(evaluation_record), encoding="utf-8")
            registry = root / "approved" / "python-registry.json"
            register_python_candidate(candidate, registry, module_id="probe", version="1",
                sections=[{"division": "common", "section_no": "probe"}], inputs=["quantity"],
                source_hash="a" * 64, contract=contract, sandbox_image=image)
            code_target = registry.parent / "python" / "probe.py"
            estimate_spec = {"division": "common", "section_no": "probe",
                             "inputs": [{"name": "quantity"}],
                             "quantity_model": {"name": "adjusted_daily_crew",
                                                "params": {"module_id": "probe"}}}
            with patch.object(generated, "MODULE_DIR", root), \
                    patch.object(generated, "PYTHON_REGISTRY", registry):
                self.assertEqual(calculator_for(estimate_spec), generated.generated_python_daily_crew)
                estimate = generated.generated_python_daily_crew(estimate_spec, {"quantity": 17}, labor_only=True)
            self.assertEqual(estimate["status"], "computed", estimate)
            self.assertEqual(estimate["daily_volume_m3"], "51")
            self.assertEqual(estimate["provenance"]["generated_module"]["code_sha256"],
                             hashlib.sha256(code_target.read_bytes()).hexdigest())

    def test_source_bundle_keeps_full_section_pages_tables_notes_and_references(self):
        import fitz
        from pipeline.calculation_contract import build_source_bundle
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            pdf = fitz.open()
            for text in ("section body", "referenced clause"):
                page = pdf.new_page()
                page.insert_text((72, 72), text)
            pdf_path = root / "source.pdf"
            pdf.save(pdf_path)
            chunks_path = root / "chunks.jsonl"
            chunks_path.write_text("\n".join(json.dumps(item) for item in [
                {"id": "table-1", "kind": "table", "division": "common", "section_no": "1-1",
                 "pages": [1], "source": {"page": 1}, "values": {"daily": {"output": "10"}},
                 "references": [{"division": "common", "section_no": "2-1"}]},
                {"id": "note-1", "kind": "note", "division": "common", "section_no": "1-1",
                 "pages": [1], "source": {"page": 1}, "text": "reduce daily output by 5%"},
                {"id": "ref-1", "kind": "paragraph", "division": "common", "section_no": "2-1",
                 "pages": [2], "source": {"page": 2}, "text": "referenced clause"}]) + "\n", encoding="utf-8")
            page_map_path = root / "page-map.json"
            page_map_path.write_text(json.dumps({"pages": {"1": {"printed_page": 9}, "2": {"printed_page": 10}}}),
                                     encoding="utf-8")
            serialized, payload = build_source_bundle(pdf_path, chunks_path, page_map_path, "common", "1-1")
            self.assertEqual([item["pdf_page"] for item in payload["full_section_pdf_pages"]], [1, 2])
            self.assertEqual(payload["full_section_pdf_pages"][0]["printed_page"], 9)
            self.assertEqual(payload["parsed_section_tables_notes"][0]["values"]["daily"]["output"], "10")
            self.assertIn("note-1", payload["calculation_notes"][0]["legacy_aliases"])
            self.assertEqual(payload["referenced_chunks"][0]["id"], "ref-1")
            self.assertEqual(json.loads(serialized)["pdf_sha256"], payload["pdf_sha256"])

    def test_authoring_and_withheld_case_catalogs_are_separate_and_cover_required_types(self):
        base = ROOT / "evals" / "calculation_modules"
        authoring = json.loads((base / "authoring_cases.json").read_text(encoding="utf-8"))
        final = json.loads((base / "final_evaluation_cases.json").read_text(encoding="utf-8"))
        self.assertFalse({case["id"] for case in authoring} & {case["id"] for case in final})
        author_text = json.dumps(authoring, ensure_ascii=False)
        final_kinds = {case.get("kind") for case in final}
        for marker in ("ratio-roof-output", "interval-formwork-40m", "crew-addition",
                       "output-increase-decrease", "equipment-time", "cost-combination"):
            self.assertIn(marker, author_text)
        self.assertTrue({"boundary", "interaction", "crew_and_equipment", "output_delta",
                         "cost_combination"} <= final_kinds)

    def test_formwork_candidate_case_files_resolve_site_values_and_are_disjoint(self):
        from pipeline.calculation_modules import load_case_file
        base = ROOT / "evals" / "calculation_modules"
        authoring = load_case_file(base / "formwork_authoring_cases.json")
        evaluation = load_case_file(base / "formwork_final_evaluation_cases.json")
        self.assertTrue(all(case["inputs"]["complexity"] in case["spec"]["inputs"][1]["allowed_values"]
                            for case in authoring + evaluation))
        self.assertFalse({case["id"] for case in authoring} & {case["id"] for case in evaluation})
        self.assertTrue(all(case["spec"]["section_no"] == "6-3-1" for case in authoring + evaluation))

    def test_height_answer_to_estimate_and_partial_xlsx_preserves_other_item(self):
        from io import BytesIO
        from openpyxl import load_workbook
        from backend.agent.estimate import tools
        from backend.agent.estimate.state import choose_work
        from backend.agent.rules.specs import load_specs
        from backend.api.tables import build_xlsx

        spec = next(item for item in load_specs().values()
                    if item["division"] == "공통" and item["section_no"] == "6-3-1")
        session = tools.new_session({"items": [
            {"item_id": "i1", "request_text": "합판거푸집 100㎡",
             "quantity": {"value": "100", "unit": "㎡"}},
            {"item_id": "i2", "request_text": "다른 항목"},
        ]})
        item = session["items"]["i1"]
        choose_work(item, {"confirmed": True, "decision": "chosen", "section_no": "6-3-1",
                           "section": spec["title"]}, spec)
        session["items"]["i2"]["explicit"] = {"marker": {"value": "preserve", "source": "test"}}
        session["items"]["i2"]["status"] = "UNSUPPORTED"
        session["items"]["i2"]["reason"] = "별도 테스트 항목"
        session["items"]["i2"]["result_revision"] = session["items"]["i2"]["input_revision"]
        answers = {"complexity": "복잡", "height": "8", "is_roof_slab": False, "crane_used": False}
        tools.set_conditions(session, "", values={name: {"value": value} for name, value in answers.items()},
                             source="answer")
        pending = tools.compute_labor(session)
        self.assertEqual(pending["status"], "needs_input")
        self.assertEqual(pending["missing"][0]["field"], "lifting_equipment_used")
        tools.set_conditions(session, "", values={"lifting_equipment_used": {"value": False}}, source="answer")
        labor = tools.compute_labor(session)
        self.assertEqual(labor["status"], "ok", labor.get("missing"))
        self.assertEqual(labor["data"]["daily_volume_m3"], "27.3")
        self.assertEqual(labor["data"]["work_days"], "1000/273")
        self.assertEqual(labor["data"]["quantity"], "100")
        self.assertNotIn("height_deduction_percent", item["conditions"])
        self.assertNotIn("height_deduction_source", item["conditions"])
        from backend.api.main import _computed_result_out
        from backend.api.tables import build_xlsx, statement_rows
        height_priced = tools.estimate_cost(session)
        self.assertEqual(height_priced["status"], "ok")
        exported_result = _computed_result_out(item["computed_result"], spec, spec.get("review", ""))
        self.assertEqual(exported_result["daily_volume"]["value"], "27.3")
        self.assertEqual(exported_result["daily_volume"]["adjustments"][0]["after"], "27.3")
        height_response = {"work": {"title": spec["title"], "spec_id": spec["id"]},
            "tables": {"bill": {"quantity": "100", "unit": "?", "spec": "??"},
                       "statement_rows": statement_rows(session["statement"]), "unit_rows": [], "rate_rows": []},
            "statement": session["statement"], "priced": item["priced_result"],
            "result": exported_result, "conditions": [], "inputs": [], "basis_date": "2026-10-17"}
        height_book = load_workbook(BytesIO(build_xlsx(height_response)), data_only=False)
        height_cells = [cell.value for sheet in height_book for row in sheet.iter_rows() for cell in row if cell.value]
        self.assertTrue(any("27.3" in str(value) for value in height_cells))
        self.assertTrue(any("formwork-6-3-1-height-started-interval-compound-v1" in str(value)
                            or "9%" in str(value) for value in height_cells))
        self.assertTrue(any("PDF p.193 (printed p.137)" in str(value) for value in height_cells))
        tools.set_conditions(session, "", values={"crane_used": {"value": True}}, source="answer")
        labor = tools.compute_labor(session)
        tools.set_conditions(session, "", values={"lifting_equipment_used": {"value": True}}, source="answer")
        labor = tools.compute_labor(session)
        self.assertEqual(labor["data"]["daily_volume_m3"], "30")  # 양중장비 사용 시 높이 감산 없음
        priced = tools.estimate_cost(session)
        self.assertEqual(priced["status"], "ok")
        self.assertEqual(priced["data"]["statement"]["status"], "PARTIAL")
        self.assertEqual(session["items"]["i2"]["explicit"]["marker"]["value"], "preserve")

        response = {"work": {"title": spec["title"], "spec_id": spec["id"]},
                    "tables": {"bill": {"quantity": "100", "unit": "㎡", "spec": "복잡"},
                               "statement_rows": [], "unit_rows": [], "rate_rows": []},
                    "statement": session["statement"], "priced": item["priced_result"],
                    "result": {key: item["computed_result"][key]
                               for key in ("unit_lines", "not_calculated", "unit_basis")},
                    "conditions": [], "inputs": [],
                    "basis_date": "2026-10-17"}
        workbook = load_workbook(BytesIO(build_xlsx(response)), data_only=False)
        cells = [cell.value for sheet in workbook for row in sheet.iter_rows() for cell in row if cell.value]
        self.assertTrue(any("상시 크레인 장비비" in str(value) for value in cells))

    def test_two_calculated_items_reprice_aggregate_and_export(self):
        from io import BytesIO
        from openpyxl import load_workbook
        from backend.agent.estimate import tools
        from backend.agent.estimate.aggregate import aggregate
        from backend.agent.estimate.state import choose_work
        from backend.agent.rules.specs import load_specs
        from backend.api.tables import build_xlsx, statement_rows

        spec = next(item for item in load_specs().values()
                    if item["division"] == "토목" and item["section_no"] == "5-3-2")
        session = tools.new_session({"items": [
            {"item_id": "girder-a", "request_text": "강재거더 A 8개"},
            {"item_id": "girder-b", "request_text": "강재거더 B 12개"},
        ]})
        answers = {"girder_type": "개구제형", "location": "육상", "direct_lifting": False,
                   "special_condition": False, "crane_size": "100~300ton"}
        for item_id, quantity in (("girder-a", "8"), ("girder-b", "12")):
            session["item_order"] = [item_id] + [x for x in session["item_order"] if x != item_id]
            session["active_item_id"] = item_id
            item = session["items"][item_id]
            choose_work(item, {"confirmed": True, "decision": "chosen", "section_no": "5-3-2",
                               "section": spec["title"]}, spec)
            tools.set_conditions(session, "", values={**{key: {"value": value} for key, value in answers.items()},
                                                       "quantity": {"value": quantity}}, source="answer")
            tools.estimate_cost(session)
            self.assertEqual(item["status"], "PRICED")

        session["item_order"] = ["girder-a", "girder-b"]
        session["active_item_id"] = "girder-a"
        combined = aggregate(session)
        self.assertEqual(set(combined["aggregate_result"]["included"]), {"girder-a", "girder-b"})
        previous_b = copy.deepcopy(session["items"]["girder-b"]["priced_result"])
        previous_total = combined["aggregate_result"]["direct_total"]

        tools.set_conditions(session, "", values={"quantity": {"value": "16"}}, source="answer")
        updated = tools.estimate_cost(session)
        self.assertIn(updated["status"], ("ok", "error"))  # error may report overall partial/not-ready status
        self.assertEqual(session["items"]["girder-b"]["priced_result"], previous_b)
        final = aggregate(session)
        self.assertEqual(final["aggregate_result"]["included"], ["girder-a", "girder-b"])
        self.assertNotEqual(final["aggregate_result"]["direct_total"], previous_total)

        response = {"work": {"title": spec["title"], "spec_id": spec["id"]},
                    "tables": {"bill": {"quantity": "28", "unit": "ea", "spec": "개구제형"},
                               "statement_rows": statement_rows(final["statement"]),
                               "unit_rows": [], "rate_rows": []},
                    "statement": final["statement"],
                    "priced": session["items"]["girder-a"]["priced_result"],
                    "result": {key: session["items"]["girder-a"]["computed_result"].get(key)
                               for key in ("unit_lines", "not_calculated", "unit_basis")},
                    "conditions": [], "inputs": [], "basis_date": "2026-10-17"}
        workbook = load_workbook(BytesIO(build_xlsx(response)), data_only=False)
        self.assertIn("원가계산서", workbook.sheetnames)
        self.assertTrue(any("도급액" in str(cell.value) for sheet in workbook for row in sheet.iter_rows()
                            for cell in row if cell.value))


if __name__ == "__main__":
    unittest.main()
