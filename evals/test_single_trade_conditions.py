"""도장 횟수·방수 종속 조건의 계산/저장/엑셀 회귀. 외부 호출 없음."""
import copy
import io
import os
import unittest
from fractions import Fraction

os.environ['AGENT_OFFLINE'] = '1'
os.environ['AGENT_LLM'] = 'off'
os.environ['KNOWLEDGE_RETRIEVAL'] = 'legacy'

from backend.agent.rules.specs import load_specs
from backend.agent.nodes.fill import extract_inputs
from backend.agent.tools.calc.per_unit import per_unit
from backend.agent.estimate import tools, dialogue
from backend.agent.nodes.select import decide
from backend.agent.rules.specs import specs_by_section


def spec(no):
    return next(s for s in load_specs().values() if s['division'] == '건축' and s['section_no'] == no)


def session(no, text):
    s = tools.new_estimate(text, '2026-10-01')
    target = spec(no)
    result = tools.find_work(s, hits=[{'section_no': no, 'division': '건축', 'rank': 1,
                                     'section': f"{no} {target['title']}", 'page': 693,
                                     'text': target['title']}])
    assert result['status'] == 'ok', result
    return s


def answer(s, **values):
    return tools.set_conditions(s, '', values={k: {'value': v} for k, v in values.items()}, source='answer')


class SingleTradeConditions(unittest.TestCase):
    def test_explicit_paint_method_selection(self):
        specs = specs_by_section()
        cases = {
            '수성페인트 롤러칠': '11-2-2',
            '수성페인트 붓칠': '11-2-1',
            '수성페인트 뿜칠': '11-2-3',
        }
        for query, expected in cases.items():
            hits = [{'section_no': no, 'division': '건축', 'rank': rank, 'section': no,
                     'chunk_id': f'{no}-chunk', 'score': 10 / rank}
                    for rank, no in enumerate(('11-2-1', '11-2-2', '11-2-3'), 1)]
            result = decide(hits, specs, query=query)
            self.assertEqual(result['decision'], 'chosen', query)
            self.assertEqual(result['selected']['section_no'], expected, query)
            self.assertEqual(result['selection_source'], 'explicit_name_and_method')

        s = tools.new_estimate('수성페인트 롤러칠')
        hits = [{'section_no': no, 'division': '건축', 'rank': rank, 'section': no,
                 'chunk_id': f'{no}-chunk', 'score': 10 / rank}
                for rank, no in enumerate(('11-2-1', '11-2-2', '11-2-3'), 1)]
        result = tools.find_work(s, hits=hits)
        self.assertEqual(result['data']['spec_id'], spec('11-2-2')['id'])
        self.assertEqual(tools._item(s)['selection']['source'], 'explicit_name_and_method')
        self.assertEqual(tools._item(s)['selection']['section_no'], '11-2-2')

    def test_ambiguous_or_negative_method_keeps_confirmation(self):
        specs = specs_by_section()
        hits = [{'section_no': no, 'division': '건축', 'rank': rank, 'section': no}
                for rank, no in enumerate(('11-2-1', '11-2-2', '11-2-3'), 1)]
        for query in ('수성페인트 칠하기', '수성페인트 롤러칠 또는 붓칠',
                      '수성페인트 롤러칠 말고 붓칠'):
            result = decide(hits, specs, query=query)
            self.assertNotEqual(result['decision'], 'chosen', query)

    def test_explicit_candidate_requires_trade_and_method_match(self):
        specs = copy.deepcopy(specs_by_section())
        oil = copy.deepcopy(spec('11-2-2'))
        oil.update(id='test/oil-paint-roller', section_no='11-2-5',
                   title='유성페인트 롤러칠', work='유성페인트 롤러칠')
        specs[('건축', '11-2-5')] = [oil]

        def select_for(query, sections):
            hits = [{'section_no': section, 'division': '건축', 'rank': rank,
                     'section': section, 'chunk_id': f'{section}-chunk'}
                    for rank, section in enumerate(sections, 1)]
            return decide(hits, specs, query=query)

        # A matching method with the wrong paint type is not an exact candidate.
        for query, sections in (
                ('수성페인트 롤러칠 2회 1㎡', ('11-2-5', '11-2-1')),
                ('유성페인트 롤러칠 2회 1㎡', ('11-2-2', '11-2-1'))):
            result = select_for(query, sections)
            self.assertEqual(result['decision'], 'provisional', query)

        # A matching candidate may be selected even when a different paint type
        # is also present among the search candidates.
        result = select_for('수성페인트 롤러칠 2회 1㎡', ('11-2-5', '11-2-2', '11-2-1'))
        self.assertEqual(result['decision'], 'chosen')
        self.assertEqual(result['selected']['section_no'], '11-2-2')

        # Generic paint type and missing method stay as clarification questions.
        result = select_for('페인트 롤러칠 1㎡', ('11-2-2', '11-2-5'))
        self.assertEqual(result['decision'], 'provisional')
        result = select_for('수성페인트 칠하기 1㎡', ('11-2-1', '11-2-2'))
        self.assertEqual(result['decision'], 'provisional')

    def test_same_work_name_across_divisions_keeps_confirmation(self):
        specs = dict(specs_by_section())
        duplicate = copy.deepcopy(specs[("건축", "11-2-2")][0])
        duplicate["division"] = "토목"
        duplicate["section_no"] = "5-2-4"
        duplicate["id"] = "duplicate-other-division"
        specs[("토목", "5-2-4")] = [duplicate]
        hits = [
            {'section_no': no, 'division': division, 'rank': rank, 'section': no}
            for rank, (division, no) in enumerate((('건축', '11-2-2'), ('토목', '5-2-4')), 1)
        ]
        result = decide(hits, specs, query='수성페인트 롤러칠')
        self.assertNotEqual(result['decision'], 'chosen')

    def test_paint_original_failure_and_repeat(self):
        s = spec('11-2-2')
        old = copy.deepcopy(s)
        old['inputs'] = [f for f in old['inputs'] if f['name'] != 'coat_count']
        old['quantity_model']['params'].pop('repeat_input')
        baseline = per_unit(old, {'quantity': '1', 'is_ceiling': False}, labor_only=True)
        self.assertEqual(Fraction(baseline['unit_lines'][0]['exact']), Fraction('0.012'))
        for count in (1, 2, 3):
            r = per_unit(s, {'quantity': '1', 'is_ceiling': False, 'coat_count': count}, labor_only=True)
            self.assertEqual(r['status'], 'computed')
            self.assertEqual(Fraction(r['unit_lines'][0]['exact']), Fraction('0.012') * count)
            self.assertEqual(Fraction(r['unit_lines'][1]['exact']), Fraction('0.002') * count)
            self.assertIn(f'× {count}회', r['unit_lines'][0]['formula'])
            self.assertTrue(any('1회 칠하는 기준' in (c.get('quote') or '') for c in r['unit_lines'][0]['citations']))

    def test_repeat_extraction_and_missing(self):
        s = spec('11-2-2')
        self.assertEqual(extract_inputs('롤러칠 2회 1㎡', s)[0]['coat_count'], 2)
        self.assertNotIn('coat_count', extract_inputs('롤러칠 1.5회', s)[0])
        self.assertNotIn('coat_count', extract_inputs('롤러칠 -1회', s)[0])
        self.assertIn('coat_count', extract_inputs('1회 또는 2회', s)[1])
        r = per_unit(s, {'quantity': '1', 'is_ceiling': False}, labor_only=True)
        self.assertIn('coat_count', r['missing'])
        for bad in (0, -1, True, '1.5'):
            r = per_unit(s, {'quantity': '1', 'is_ceiling': False, 'coat_count': bad}, labor_only=True)
            self.assertEqual(r['status'], 'rejected')

    def test_interior_waterproof_without_exterior_fields(self):
        s = spec('6-4-3')
        values = {'quantity': '1', 'thickness': '10㎜이하', 'is_exterior': False}
        old = copy.deepcopy(s)
        for f in old['inputs']:
            f.pop('when', None)
        self.assertEqual(per_unit(old, values, labor_only=True)['status'], 'ask')
        r = per_unit(s, values, labor_only=True)
        self.assertEqual(r['status'], 'computed')
        self.assertEqual([Fraction(l['exact']) for l in r['unit_lines']], [Fraction('.047'), Fraction('.035')])
        # 과거 외벽 입력이 남아 있어도 실내 계산에 쓰지 않는다.
        stale = per_unit(s, {**values, 'floor_count': 0, 'building_height': '0'}, labor_only=True)
        self.assertEqual(stale['unit_lines'], r['unit_lines'])

    def test_exterior_still_requires_conditions(self):
        s = spec('6-4-3')
        base = {'quantity': '1', 'thickness': '10㎜이하', 'is_exterior': True}
        self.assertEqual(per_unit(s, base, labor_only=True)['missing'], ['floor_count'])
        for zero in (0, '0'):
            self.assertEqual(per_unit(s, {**base, 'floor_count': zero}, labor_only=True)['missing'], ['building_height'])
        self.assertNotEqual(per_unit(s, {**base, 'floor_count': 5}, labor_only=True)['status'], 'computed')

    def test_numeric_answers_and_private_trace(self):
        s = session('6-4-3', '실내 바닥 방수모르타르 1㎡')
        for zero in (0, '0'):
            r = answer(s, is_exterior=True, floor_count=zero)
            self.assertFalse(r['rejected'])
            dialogue._trace_selected_conditions({'session': s})
            # 실제 질문/조건 재구성 경로도 거친다.
            tools.compute_labor(s)
            dialogue._trace_selected_conditions({'session': s})
        bad = answer(s, floor_count='-1')
        self.assertIn('0 이상의 정수', bad['rejected']['floor_count'])
        self.assertNotIn('0 / 이', bad['rejected']['floor_count'])

    def test_waterproof_cost_and_wall_switch(self):
        s = session('6-4-3', '실내 바닥 방수모르타르 1㎡')
        answer(s, thickness='10㎜이하')
        result = tools.estimate_cost(s)
        self.assertEqual(result['status'], 'ok')
        self.assertNotIn('floor_count', tools._item(s)['conditions'])
        self.assertNotIn('building_height', tools._item(s)['conditions'])
        answer(s, is_exterior=True)
        result = tools.compute_labor(s)
        self.assertIn('floor_count', [q['field'] for q in result['missing']])
        answer(s, floor_count='0')
        result = tools.compute_labor(s)
        self.assertIn('building_height', [q['field'] for q in result['missing']])
        answer(s, is_exterior=False)
        self.assertEqual(tools.estimate_cost(s)['status'], 'ok')
        self.assertNotIn('floor_count', tools._item(s)['conditions'])

    def test_recalculation_and_excel(self):
        from backend.api.main import _build_response
        from backend.api.dialogue_service import legacy_state
        from backend.api.tables import build_xlsx
        from openpyxl import load_workbook
        s = session('11-2-2', '실내 벽 수성페인트 롤러칠 2회 1㎡ 바탕 제외')
        self.assertFalse(answer(s, is_ceiling=False)['rejected'])
        result = tools.estimate_cost(s)
        self.assertEqual(result['status'], 'ok')
        item = tools._item(s)
        self.assertEqual(item['conditions']['coat_count'], 2)
        key = item['labor_key']
        state = legacy_state({'session': s, 'goal': 'cost', 'pending': [], 'turn': {'tools': ['estimate_cost']}})
        response = _build_response('offline-conditions', state)
        book = load_workbook(io.BytesIO(build_xlsx(response)), data_only=False)
        values = [str(c.value) for sheet in book for row in sheet for c in row if c.value is not None]
        self.assertTrue(any('2회 칠' in v for v in values))
        self.assertTrue(any('× 2회' in v for v in values))
        answer(s, coat_count='1')
        self.assertIsNone(tools.current_estimate(s))
        result = tools.estimate_cost(s)
        self.assertEqual(result['status'], 'ok')
        self.assertNotEqual(item['labor_key'], key)
        self.assertEqual(Fraction(item['computed_result']['unit_lines'][0]['exact']), Fraction('.012'))


if __name__ == '__main__':
    unittest.main()
