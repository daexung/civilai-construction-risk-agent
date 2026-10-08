"""AGENT_MODE=tools 비회원 대화를 API부터 Excel까지 검사한다(고정 후보 + 대본 LLM, 외부 호출 없음).

- 고정 후보: 검색 결과를 운영에서 나온 후보(6-1-4, 6-1-1, 1-6-2)로 고정한다. 실제 검색 대화는 check_dialogue_search.py.
- 대본 LLM: 행동(JSON)은 정해 둔 순서로, 응답 문장은 서버가 준 facts로 만든다. 값·근거 검증은 서버 그대로 거친다.
- "세제곱미터"는 대본도 사용자 문장 그대로의 근거를 보내므로, 서버 단위 파서가 읽지 못하면 실패한다.
"""

from __future__ import annotations

import io
import json
import os
import sys
import threading
import time
from contextlib import nullcontext
from fractions import Fraction
from pathlib import Path
from unittest.mock import MagicMock, patch

os.environ.update(AGENT_OFFLINE="1", AGENT_LLM="on", AGENT_MODE="tools")
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

import backend.api.main as api_main  # noqa: E402
from backend.agent.estimate import dialogue, reply_check, tools  # noqa: E402
from backend.agent.nodes.compute import CALCULATORS  # noqa: E402
from backend.agent.rules.specs import load_specs  # noqa: E402
from backend.agent.tools.llm import client  # noqa: E402
from backend.api import chat_storage, dialogue_service  # noqa: E402

HITS = [{"rank": rank, "section_no": no, "division": division, "section": f"{no} {title}", "page": page,
         "table_id": None, "text": title, "chunk_id": f"fixed-{rank}"}
        for rank, (no, division, title, page) in enumerate([("6-1-4", "공통", "콘크리트 펌프차 타설", 186),
                                                             ("6-1-1", "공통", "레디믹스트콘크리트 타설", 185),
                                                             ("1-6-2", "토목", "표층 인력포설", 120)], 1)]
FORM_HITS = [{"rank": rank, "section_no": no, "division": "공통", "section": f"{no} {title}", "page": page,
              "table_id": None, "text": title, "chunk_id": f"form-{rank}"}
             for rank, (no, title, page) in enumerate([("6-3-1", "합판거푸집 설치 및 해체", 204),
                                                       ("6-3-4", "문양거푸집(판넬) 설치 및 해체", 209),
                                                       ("6-3-2", "강재거푸집 설치 및 해체", 206)], 1)]
PUMP = {"pump_size": "32m", "structure": "철근", "slump_band": "15㎝", "facility_type": "Type-Ⅱ",
        "site_type": "Type-Ⅱ", "placement": "붐", "vibrator_used": True, "reset_status": "없음"}
SECRET = "offline-dialogue-session-" + "x" * 32


class ScriptedLLM:
    """행동은 queue 순서대로, 응답 문장은 facts로 만든다. calls에 호출 기록을 남긴다."""

    def __init__(self):
        self.actions: list[dict] = []
        self.calls: list[str] = []
        self.reply_override = None
        self.reply_raise = False

    def __call__(self, prompt, system, response_schema=None, **kwargs):
        data = json.loads(prompt)
        if response_schema is dialogue.REPLY_SCHEMA:
            self.calls.append("reply")
            if self.reply_raise:
                raise client.LLMUnavailable("scripted reply failure", attempts=1)
            if self.reply_override:
                return json.dumps(self.reply_override(data), ensure_ascii=False)
            facts = {fact["id"]: fact for fact in data["facts"]}
            fact = facts.get("contract_amount") or facts.get("line0")
            if not fact:
                return json.dumps({"text": "아래 질문을 확인해 주세요."}, ensure_ascii=False)
            phrase = f"{fact['item']} {reply_check.display(fact['value'])}{fact['unit']}"
            return json.dumps({"text": f"{phrase}입니다."},
                              ensure_ascii=False)
        self.calls.append("agent")
        action = self.actions.pop(0) if self.actions else {"action": "reply", "reason": "끝"}
        if "__raise__" in action:
            raise client.LLMUnavailable("scripted agent failure", attempts=1)
        if "__raw__" in action:
            return action["__raw__"]
        return json.dumps(action, ensure_ascii=False)

    def plan(self, *tools_and_args):
        self.actions = [{"action": "call_tool", "tool": name, "args": args, "reason": "대본"} for name, args in tools_and_args]


def vague_after_ok(reply: dict, thread_id: str, before: tuple) -> bool:
    """모호한 요청 뒤: 확인 질문(scope)만 나오고 세션(물량·조건 포함)은 그대로."""
    state = dialogue.load(dialogue.build_dialogue_graph(api_main.GRAPH.checkpointer), thread_id)
    current = tools._item(state["session"])
    after = (state["session"]["estimate_id"], current.get("request_text"), current.get("quantity"),
             (current.get("selection") or {}).get("section_no"), current["conditions"].get("pump_size"))
    return [q["name"] for q in reply["questions"]] == ["scope"] and after == before


def check_free_card_answers(checks: list) -> None:
    """실제 API·체크포인트, 사용량/LLM/회원 DB만 모의 처리한다."""
    llm, usage = ScriptedLLM(), {"used": 0}
    original_free = dialogue_service.free_card_answer
    member_lock = [False]

    def status(*args):
        return {"used": usage["used"], "remaining": 5 - usage["used"],
                "service_remaining": 500 - usage["used"]}

    def consume(*args):
        usage["used"] += 1
        return status()

    def locked_free(graph, thread, request):
        assert api_main._GUESTS[thread]["lock"].locked() if thread in api_main._GUESTS else member_lock[0]
        return original_free(graph, thread, request)

    with patch.dict(os.environ, {"AGENT_MODE": "tools", "AGENT_LLM": "on"}), \
            patch("backend.api.usage_limits.processing", return_value=nullcontext((None, "fixture", False))), \
            patch("backend.api.usage_limits.consume", side_effect=consume), \
            patch("backend.api.usage_limits.status", side_effect=status), \
            patch.object(dialogue_service, "free_card_answer", side_effect=locked_free), \
            patch.object(tools, "retrieve", return_value={"hits": HITS}) as retrieval, \
            patch.object(client, "generate", side_effect=llm) as generate, \
            patch.object(api_main, "warmup_client", return_value=True), \
            TestClient(api_main.app, headers={"X-Guest-Session": SECRET}) as http:
        def post(body):
            result = http.post("/api/chat", json=body)
            assert result.status_code == 200, result.text
            return result.json()

        def start():
            llm.plan(("find_work", {}), ("compute_labor", {}))
            return post({"message": "콘크리트 타설 1세제곱미터 품셈 알려줘"})

        def card(previous, values):
            values = {name: value for name, value in values.items() if name not in {"placement", "vibrator_used", "reset_status"} or name in
                      {q["field"] for q in dialogue.load(api_main.DIALOGUE, previous["thread_id"])["pending"]}}
            return {"thread_id": previous["thread_id"], "answers": values,
                    "refs": {q["field"]: q["ref"] for q in dialogue.load(api_main.DIALOGUE, previous["thread_id"])["pending"] if q["field"] in values}}

        first = start()
        before = status()
        generate.reset_mock(); retrieval.reset_mock()
        body = {**card(first, {"work": "6-1-4"}), "request_id": "41111111-1111-4111-8111-111111111111"}
        second = post(body)
        state = dialogue.load(api_main.DIALOGUE, first["thread_id"])
        checks.append(("C1 비회원 질문 1회 → work 카드 무료, usage·서비스 잔여 유지, LLM·검색 0, template",
                       before["used"] == 1 and second["usage"] == before and state["free_card_answers"] == 1
                       and generate.call_count == retrieval.call_count == 0 and second["answer_source"] == "template"))
        replay = post(body)
        checks.append(("C2 같은 request_id 무료 카드 재전송: 같은 응답·무료 횟수 1 유지",
                       replay == second and dialogue.load(api_main.DIALOGUE, first["thread_id"])["free_card_answers"] == 1))
        stale = post(card(first, {"work": "6-1-4"}))
        checks.append(("C3 옛 ref는 차감 1·기존 미반영 안내", stale["usage"]["used"] == before["used"] + 1
                       and "지금 확인하는 질문" in stale["message"]))
        generate.reset_mock(); retrieval.reset_mock()
        full = post(card(stale, PUMP))
        checks.append(("C4 조건 카드 무료·LLM·검색 0·품 계산 완료",
                       full["status"] == "COMPUTED" and full["usage"] == stale["usage"]
                       and generate.call_count == retrieval.call_count == 0))
        fresh = start()
        generate.reset_mock(); before = usage["used"]
        mixed = post({**card(fresh, {"work": "6-1-4"}), "message": "이 조건으로 품 알려줘"})
        checks.append(("C5 answers+message 차감 1·LLM 경로 유지", usage["used"] == before + 1 and generate.call_count > 0))
        before = usage["used"]
        post({"thread_id": full["thread_id"], "conditions": {"duration": "7~12개월"}})
        checks.append(("C6 conditions 재계산은 차감 1 유지", usage["used"] == before + 1))
        current = start()
        current = post(card(current, {"work": "6-1-4"}))
        before = usage["used"]
        for _ in range(29):
            current = post(card(current, {"pump_size": "잘못된 값"}))
        saved = dialogue.load(api_main.DIALOGUE, current["thread_id"])
        checks.append(("C7 현재 ref의 잘못된 값 반복도 무료 횟수에 포함: 30회까지 차감 0",
                       saved["free_card_answers"] == 30 and usage["used"] == before))
        current = post(card(current, {"pump_size": "잘못된 값"}))
        checks.append(("C8 31번째 카드부터 차감 1", usage["used"] == before + 1
                       and dialogue.load(api_main.DIALOGUE, current["thread_id"])["free_card_answers"] == 30))
        llm.plan(("find_work", {}), ("compute_labor", {}))
        restarted = post({"thread_id": current["thread_id"], "message": "콘크리트 타설 1세제곱미터 품셈", "restart": True})
        checks.append(("C9 restart 새 견적은 무료 횟수 0", dialogue.load(api_main.DIALOGUE, restarted["thread_id"])["free_card_answers"] == 0))
        retry_first = start()
        retry_body = {**card(retry_first, {"work": "6-1-4"}), "request_id": "42222222-2222-4222-8222-222222222222"}
        before = usage["used"]
        with patch.object(tools, "compute_labor", side_effect=RuntimeError("free card failure")):
            failed = post(retry_body)
        failure_state = dialogue.load(api_main.DIALOGUE, retry_first["thread_id"])
        recovered = post(retry_body)
        checks.append(("C10 실패한 무료 턴은 횟수·상태 유지, 같은 ID 재시도도 무료",
                       failed["status"] == "ERROR" and failure_state["free_card_answers"] == 0
                       and usage["used"] == before and recovered["status"] == "MISSING_INFO"
                       and dialogue.load(api_main.DIALOGUE, recovered["thread_id"])["free_card_answers"] == 1))
        # 실제 현재 scope ref여도 새 견적 시작은 유료이며 횟수는 0으로 초기화한다.
        scope_state = dialogue.load(api_main.DIALOGUE, recovered["thread_id"])
        scope_state["pending_revision"] += 1
        question = dialogue._confirm_scope(scope_state, "콘크리트 타설 2세제곱미터 품셈")["missing"][0]
        question["ref"] = dialogue._ref(question, scope_state["pending_revision"])
        scope_state["pending"] = [question]
        api_main.DIALOGUE.update_state(dialogue.config(recovered["thread_id"]), scope_state)
        scope_response = dialogue_service.restore(api_main.DIALOGUE, recovered["thread_id"])
        before = usage["used"]
        new_estimate = post(card(scope_response, {"scope": dialogue.SCOPE_NEW}))
        checks.append(("C10b 현재 ref의 새 견적 시작 카드도 차감 1·무료 횟수 0 초기화",
                       usage["used"] == before + 1
                       and dialogue.load(api_main.DIALOGUE, new_estimate["thread_id"])["free_card_answers"] == 0))
        recovered = post(card(new_estimate, {"work": "6-1-4"}))
        # 무료 조건의 경계: 누락·추가·다른 ref, conditions/restart, 새 검색 카드, legacy.
        valid = card(recovered, {"pump_size": "32m"})
        for label, changed in [
            ("refs 누락", {"refs": {}}), ("answer ref 누락", {"answers": {"pump_size": "32m", "structure": "철근"}}),
            ("추가 ref", {"refs": {**valid["refs"], "unknown": "other"}}),
            ("빈 conditions도 제외", {"conditions": {}}), ("restart 제외", {"restart": True}),
            ("새 견적 시작 카드 제외", {"answers": {"scope": dialogue.SCOPE_NEW}, "refs": {"scope": "scope@1@1"}}),
        ]:
            checks.append((f"C11 {label}: 무료 판정 false",
                           not original_free(api_main.DIALOGUE, valid["thread_id"], {**valid, **changed})))
        # 회원은 같은 실행 경로·MemorySaver와 DB transaction/lock/append_pair를 모의 확인한다.
        member, conversation = "43333333-3333-4333-8333-333333333333", "44444444-4444-4444-8444-444444444444"
        api_main.DIALOGUE.update_state(dialogue.config(conversation), dialogue.load(api_main.DIALOGUE, recovered["thread_id"]))
        conn = MagicMock(); conn.__enter__.return_value = conn
        duplicate = [None]
        def execute(sql, args):
            if "pg_advisory_xact_lock" in sql:
                member_lock[0] = True
            result = MagicMock()
            result.fetchone.return_value = duplicate[0] if "SELECT payload" in sql else {"user_id": member}
            return result
        conn.execute.side_effect = execute
        payload = api_main.ChatRequest(conversation_id=conversation, thread_id=conversation,
                                       request_id="45555555-5555-4555-8555-555555555555",
                                       answers=valid["answers"], refs=valid["refs"])
        before = usage["used"]; generate.reset_mock(); retrieval.reset_mock()
        with patch.object(chat_storage, "connection", return_value=conn), \
                patch.object(chat_storage, "owned"), patch.object(chat_storage, "append_pair") as append, \
                patch.object(api_main, "PostgresSaver", return_value=api_main.GRAPH.checkpointer):
            member_response = api_main._member_chat(payload, member, (None, "fixture", True))
            member_count = dialogue.load(api_main.DIALOGUE, conversation)["free_card_answers"]
            duplicate[0] = {"payload": member_response}
            member_replay = api_main._member_chat(payload, member, (None, "fixture", True))
        checks.append(("C12 회원 lock 안 무료 판정·메시지 저장 1회·중복 재전송·LLM/검색 0 (DB mock)",
                       member_lock[0] and usage["used"] == before and append.call_count == 1
                       and member_response["usage"]["used"] == before and member_replay == member_response
                       and dialogue.load(api_main.DIALOGUE, conversation)["free_card_answers"] == member_count
                       and generate.call_count == retrieval.call_count == 0))
        with patch.object(api_main, "_uses_dialogue", return_value=False), \
                patch.object(api_main, "_chat_response", return_value={"thread_id": valid["thread_id"], "status": "MISSING_INFO"}):
            before = usage["used"]
            post(valid)
        checks.append(("C13 legacy 카드도 기존처럼 차감", usage["used"] == before + 1))


def check_one_question(checks: list) -> None:
    """실제 API 응답의 카드 하나만 제출한다. 전체 pending은 별도로 검증한다."""
    llm, consumed = ScriptedLLM(), []
    with patch.dict(os.environ, {"AGENT_MODE": "tools", "AGENT_LLM": "off"}), \
            patch("backend.api.usage_limits.processing", return_value=nullcontext((None, "fixture", False))), \
            patch("backend.api.usage_limits.consume", side_effect=lambda *args: consumed.append(1) or {"used": len(consumed)}), \
            patch("backend.api.usage_limits.status", side_effect=lambda *args: {"used": len(consumed)}), \
            patch.object(tools, "retrieve", return_value={"hits": HITS}) as retrieval, \
            patch.object(client, "generate", side_effect=llm) as generate, \
            patch.object(api_main, "warmup_client", return_value=True), \
            TestClient(api_main.app, headers={"X-Guest-Session": SECRET}) as http:
        def post(body):
            result = http.post("/api/chat", json=body)
            assert result.status_code == 200, result.text
            return result.json()

        def answer(previous, value):
            question = previous["questions"][0]
            return post({"thread_id": previous["thread_id"], "answers": {question["name"]: value},
                         "refs": {question["name"]: question["ref"]}})

        def start(message="콘크리트 펌프차 타설 비용 알려줘"):
            first = post({"message": message})
            return answer(first, "6-1-4")

        def state(reply):
            return dialogue.load(api_main.DIALOGUE, reply["thread_id"])

        expected = ["pump_size", "structure", "slump_band", "facility_type", "site_type", "volume"]
        current = start()
        checks.append(("O1 공종 선택 후 화면 붐 길이 1개·remaining 6·전체 pending 보존·물량 맨 뒤",
                       [q["name"] for q in current["questions"]] == ["pump_size"]
                       and current["questions_remaining"] == 6
                       and [q["field"] for q in state(current)["pending"]] == expected
                       and current["questions"][0]["ref"] == state(current)["pending"][0]["ref"]))
        sequence, counts = [], []
        before = len(consumed); generate.reset_mock(); retrieval.reset_mock()
        values = {**PUMP, "volume": "1", "concrete_supply": "관급"}
        while current["questions"]:
            assert len(sequence) < 12
            sequence.append(current["questions"][0]["name"])
            counts.append(current["questions_remaining"])
            assert len(current["questions"]) == 1 and current["questions_remaining"] == len(state(current)["pending"])
            current = answer(current, values[sequence[-1]])
        print("ONE_QUESTION_SEQUENCE", sequence, "remaining", counts, "final", current["status"])
        checks.append(("O2 화면 카드로 하나씩: 물량 → 가격 → 결과, 매번 무료·LLM/검색 0",
                       sequence == [*expected, "concrete_supply"] and counts == [6,5,4,3,2,1,1]
                       and current["estimate_current"] and current["questions_remaining"] == 0
                       and len(consumed) == before and generate.call_count == retrieval.call_count == 0))
        for mode in ("rule", "llm"):
            text_start = start()
            estimate_id = state(text_start)["session"]["estimate_id"]
            generate.reset_mock()
            llm.plan(("set_conditions", {"values": {"pump_size": {"value": "41m", "evidence": "41m"},
                                                    "structure": {"value": "철근", "evidence": "철근"},
                                                    "slump_band": {"value": "15㎝", "evidence": "슬럼프 15"}}}),
                     ("estimate_cost", {}))
            with patch.dict(os.environ, {"AGENT_LLM": "on" if mode == "llm" else "off"}):
                changed = post({"thread_id": text_start["thread_id"], "message": "41m 철근 슬럼프 15"})
            saved = state(changed); item = tools._item(saved["session"])
            checks.append((f"O3 {mode}: 글로 세 조건 동시 반영·견적 유지·남은 질문 1개 표시",
                           saved["session"]["estimate_id"] == estimate_id
                           and all(item["conditions"].get(k) == v for k,v in {"pump_size":"41m", "structure":"철근", "slump_band":"15㎝"}.items())
                           and [q["field"] for q in saved["pending"]] == expected[3:]
                           and [q["name"] for q in changed["questions"]] == ["facility_type"]
                           and changed["questions_remaining"] == 3))
            if mode == "llm":
                prompts = [(json.loads(call.args[0]), call.kwargs.get("response_schema")) for call in generate.call_args_list]
                agent = [data for data,schema in prompts if schema is dialogue.ACTION_SCHEMA]
                reply = [data for data,schema in prompts if schema is dialogue.REPLY_SCHEMA]
                checks.append(("O4 LLM 행동 입력은 전체 pending 6개, 답변 입력만 첫 질문 1개",
                               len(agent[0]["session"]["pending_questions"]) == 6
                               and reply[-1]["pending_questions"] == [saved["pending"][0]["ask"]]))
        hidden_start = start(); estimate_id = state(hidden_start)["session"]["estimate_id"]
        hidden = post({"thread_id": hidden_start["thread_id"], "message": "철근이야"})
        saved = state(hidden)
        checks.append(("O5 화면에 안 보인 구조 답: 철근 반영·견적 유지·붐 길이가 여전히 첫 질문",
                       tools._item(saved["session"])["conditions"].get("structure") == "철근"
                       and saved["session"]["estimate_id"] == estimate_id
                       and [q["name"] for q in hidden["questions"]] == ["pump_size"]
                       and hidden["questions_remaining"] == 5))
        both_start = start(); estimate_id = state(both_start)["session"]["estimate_id"]
        both = post({"thread_id": both_start["thread_id"], "message": "41m에 철근이야"})
        saved = state(both); item = tools._item(saved["session"])
        checks.append(("O6 '41m에 철근이야'는 대기 답으로 반영·새 견적/범위 질문 아님",
                       saved["session"]["estimate_id"] == estimate_id and item["conditions"].get("pump_size") == "41m"
                       and item["conditions"].get("structure") == "철근"
                       and both["questions"][0]["name"] == "slump_band" and both["questions_remaining"] == 4))
        mismatch = start("콘크리트 타설 1m 비용 알려줘")
        checks.append(("O7 단위 불일치 물량은 맨 앞·불일치 사유 포함",
                       mismatch["questions"][0]["name"] == "volume"
                       and state(mismatch)["pending"][0]["field"] == "volume"
                       and "맞지 않아요" in (mismatch["questions"][0].get("reason") or "")))
        before = len(consumed)
        stale = post({"thread_id": hidden["thread_id"], "answers": {"pump_size":"32m"},
                      "refs": {"pump_size": hidden_start["questions"][0]["ref"]}})
        checks.append(("O8 옛 ref 거부·차감 유지·붐 길이 미반영",
                       "반영하지 않았어요" in stale["message"] and len(consumed) == before + 1
                       and "pump_size" not in tools._item(state(stale)["session"])["conditions"]))
        door_hits = [{**HITS[0], "section_no":"10-1-7", "division":"건축", "section":"10-1-7 자동문 설치"}]
        with patch.object(tools, "retrieve", return_value={"hits":door_hits}):
            door = post({"message":"자동문 설치 비용 알려줘"})
            done = answer(door,"3")
        checks.append(("O9 조건 1개 자동문: 질문 1개·remaining 1, 답 후 기존 품 결과",
                       [q["name"] for q in door["questions"]] == ["quantity"] and door["questions_remaining"] == 1
                       and done["status"] in ("OK", "PARTIAL") and done["questions_remaining"] == 0))
        private = start("콘크리트 타설 1세제곱미터 비용 알려줘")
        while private["questions"][0]["name"] != "concrete_supply":
            private = answer(private,PUMP[private["questions"][0]["name"]])
        child = answer(private,"사급")
        checks.append(("O10 when 의존: 부모 관급/사급 전 단가 질문 없음, 사급 선택 후 단가만",
                       private["questions_remaining"] == 1 and child["questions"][0]["name"] == "ready_mix_price"
                       and child["questions_remaining"] == 1))


def check_per_unit_dialogue(checks):
    with patch.dict(os.environ, {"AGENT_MODE":"tools", "AGENT_LLM":"off"}), \
            patch("backend.api.usage_limits.processing", return_value=nullcontext((None,"fixture",False))), \
            patch("backend.api.usage_limits.consume", return_value={}), \
            patch("backend.api.usage_limits.status", return_value={}), \
            patch.object(tools,"retrieve", return_value={"hits":HITS}), \
            patch.object(client,"generate", side_effect=AssertionError("real LLM forbidden")), \
            patch.object(api_main,"warmup_client", return_value=True), \
            TestClient(api_main.app, headers={"X-Guest-Session":SECRET}) as http:
        def post(body):
            r=http.post("/api/chat",json=body); assert r.status_code==200,r.text; return r.json()
        def answer(previous,value):
            q=previous["questions"][0]
            return post({"thread_id":previous["thread_id"],"answers":{q["name"]:value},"refs":{q["name"]:q["ref"]}})
        def pump(text):
            current=post({"message":text}); current=answer(current,"6-1-4")
            sequence=[]
            while current["questions"]:
                q=current["questions"][0]; sequence.append(q["name"])
                assert q["name"] in PUMP,sequence
                current=answer(current,PUMP[q["name"]])
            print("PER_UNIT_SEQUENCE", sequence)
            return current,sequence
        current,sequence=pump("펌프차 타설 품 알려줘")
        result=current.get("result") or {}
        checks.append(("U7 API 물량 없는 작업조 품: 총량 없음·후속 안내·Excel 차단",
                       current["status"]=="COMPUTED" and result.get("per_unit_only") and result["unit_lines"]
                       and result["work_days"] is None and not result["lines"]
                       and current["message"].endswith("총 인원·작업일수가 필요하면 물량을 알려 주세요.")
                       and not current["estimate_current"]
                       and http.get(f"/api/export/{current['thread_id']}.xlsx").status_code==404))
        full=post({"thread_id":current["thread_id"],"message":"100㎥야"})
        checks.append(("U8 API 나중의 100㎥: 작업일수·인일 계산", full["status"]=="COMPUTED"
                       and full["result"]["work_days"] and full["result"]["lines"] and not full["result"].get("per_unit_only")))
        empty,_=pump("펌프차 타설 품 알려줘")
        cost=post({"thread_id":empty["thread_id"],"message":"비용 계산해줘"})
        checks.append(("U9 API 품→비용: 물량 질문·원가/Excel 차단",
                       cost["questions"][0]["name"]=="volume" and not cost["estimate_current"] and cost["statement"] is None))
        with patch.object(tools,"retrieve",return_value={"hits":[{**HITS[0],"section_no":"10-1-7","division":"건축","section":"10-1-7 자동문 설치"}]}):
            door=post({"message":"자동문 설치 품 알려줘"})
        checks.append(("U10 조건은 물량뿐인 명세: 질문 0개 즉시 단위당 품", not door["questions"]
                       and door["result"]["per_unit_only"] and not door["result"]["lines"]))
        one,order=pump("콘크리트 타설 1m^3일 때 품 알려줘")
        expected=["pump_size","structure","slump_band","facility_type","site_type"]
        assumption="가정: 붐 타설, 진동기 사용, 재셋팅 없음 — 다르면 말씀해 주세요."
        checks.append(("B1 물량 1m^3 총량·질문 5개 순서·가정 표시",
                       order==sequence==expected and one["result"]["work_days"] and not one["result"]["per_unit_only"]
                       and one["result"]["assumptions"]==result["assumptions"]==assumption
                       and assumption in one["message"] and assumption in current["message"]))
        defaults={e["name"]:e for e in current["inputs"] if e["source"]=="기본값"}
        checks.append(("B2 기본값 입력 출처 3개·다른 명세 기본값 없음",
                       {k:e["value"] for k,e in defaults.items()}=={"placement":"붐","vibrator_used":True,"reset_status":"없음"}
                       and sum("default" in f for s in load_specs().values() for f in s["inputs"])==3))
        changed=post({"thread_id":current["thread_id"],"message":"배관 타설에 진동기 안 써"})
        values={e["name"]:e for e in changed["inputs"]}
        checks.append(("B3 글 답 배관·진동기 미사용 우선·가정에서 해당 두 항목 제외",
                       values["placement"]["value"]=="배관" and values["vibrator_used"]["value"] is False
                       and values["placement"]["source"]==values["vibrator_used"]["source"]=="답변"
                       and changed["result"]["assumptions"]=="가정: 재셋팅 없음 — 다르면 말씀해 주세요."
                       and "가정: 재셋팅 없음" in changed["message"]))
        same,_=pump("펌프차 타설 품 알려줘")
        same_state=dialogue.load(api_main.DIALOGUE,same["thread_id"])
        old_key=tools._item(same_state["session"])["labor_key"]
        same=post({"thread_id":same["thread_id"],"message":"붐 타설이야"})
        same_state=dialogue.load(api_main.DIALOGUE,same["thread_id"])
        checks.append(("B7 기본값과 같은 값을 명시: 캐시는 유지·출처/가정 갱신",
                       tools._item(same_state["session"])["labor_key"]==old_key
                       and next(e for e in same["inputs"] if e["name"]=="placement")["source"]=="답변"
                       and same["result"]["assumptions"]=="가정: 진동기 사용, 재셋팅 없음 — 다르면 말씀해 주세요."))
        reset=post({"thread_id":changed["thread_id"],"message":"재셋팅 있어"})
        checks.append(("B4 재셋팅 있음 글 답은 기본값 대신 반영·기존 보류 유지",
                       reset["status"]=="BLOCKED" and next(e for e in reset["inputs"] if e["name"]=="reset_status")["value"]=="있음"
                       and "가정:" not in reset["message"]))
        priced=post({"thread_id":one["thread_id"],"message":"비용 계산해줘"})
        priced=answer(priced,"관급")
        book=load_workbook(io.BytesIO(http.get(f"/api/export/{priced['thread_id']}.xlsx").content),data_only=True)
        basis=list(book["산출근거"].values)
        checks.append(("B5 비용·원가·Excel 산출근거: 기본값 입력/출처·가정 포함",
                       priced["estimate_current"] and assumption in priced["result"]["assumptions"]
                       and assumption in priced["message"] and any(assumption in n for n in priced["statement"]["basis_notes"])
                       and all(any(e["label"]==row[0] and str(e["value"]) in (row[2] or "") and row[3]=="기본값" for row in basis) for e in defaults.values())))
        reference=tools.new_estimate("콘크리트 타설 1㎥",priced["basis_date"])
        tools.find_work(reference,hits=HITS); tools.set_conditions(reference,"",work="6-1-4",source="answer")
        tools.set_conditions(reference,"",values={k:{"value":v} for k,v in {**PUMP,"concrete_supply":"관급"}.items()},source="answer")
        explicit=tools.estimate_cost(reference)
        checks.append(("B6 기본값과 같은 조건을 명시한 경우 전체 금액 동일",
                       explicit["data"]["statement"]["totals"]==priced["statement"]["totals"]
                       and [(l["name"],l["exact"]) for l in explicit["data"]["unit_lines"]]==[(l["name"],l["exact"]) for l in one["result"]["unit_lines"]]
                       and not tools._item(reference)["computed_result"]["assumptions"]))
        saved=dialogue.load(api_main.DIALOGUE,empty["thread_id"])
        # 대본 답변이 후속 안내를 생략해도 서버가 붙이며, 입력에도 전달한다.
        saved["pending"]=[]
        item=tools._item(saved["session"])
        log=[{"tool":"compute_labor","result":tools.compute_labor(saved["session"],allow_per_unit=True)}]
        prompts=[]
        def generate(prompt,*args,**kwargs):
            prompts.append(json.loads(prompt)); return json.dumps({"text":"품을 계산했어요."})
        reply=dialogue._reply(saved,log,[],"품 알려줘",dialogue.Budget(),generate,True)
        checks.append(("U11 모의 LLM 입력·최종 답변에 단위당 후속 안내", prompts[0]["quantity_hint"]
                       and reply["text"].endswith(prompts[0]["quantity_hint"])
                       and prompts[0]["assumptions"]==assumption and assumption in reply["text"]))


def main() -> int:
    checks = []
    check_per_unit_dialogue(checks)
    llm = ScriptedLLM()
    consumed = []
    with patch("backend.api.usage_limits.processing", return_value=nullcontext((None, "fixture", False))), \
            patch("backend.api.usage_limits.status", return_value={}), \
            patch("backend.api.usage_limits.consume", side_effect=lambda quota, *request: consumed.append(1) or {}), \
            patch.object(tools, "retrieve", return_value={"hits": HITS}), \
            patch.object(client, "generate", llm), patch.object(api_main, "warmup_client", return_value=True), \
            TestClient(api_main.app, headers={"X-Guest-Session": SECRET}) as http:
        def post(body: dict, expect: int = 200) -> dict:
            result = http.post("/api/chat", json=body)
            assert result.status_code == expect, (result.status_code, result.text[:500])
            return result.json()

        def answer(previous: dict, values: dict, refs: dict | None = None) -> dict:
            if refs is None:
                names = {q["field"] for q in stored(previous["thread_id"])["pending"]}
                values = {name: value for name, value in values.items()
                          if name not in {"placement", "vibrator_used", "reset_status"} or name in names}
                refs = {q["field"]: q["ref"] for q in stored(previous["thread_id"])["pending"] if q["field"] in values}
            return post({"thread_id": previous["thread_id"], "answers": values, "refs": refs})

        def stored(thread_id: str) -> dict:
            return dialogue.load(dialogue.build_dialogue_graph(api_main.GRAPH.checkpointer), thread_id)

        # 1. 품셈 질문: 물량은 사용자 문장 그대로의 근거("1세제곱미터"), 공종만 묻는다
        llm.plan(("find_work", {}), ("set_conditions", {"quantity": {"value": "1", "unit": "㎥", "evidence": "1세제곱미터"}}),
                 ("compute_labor", {}))
        t1 = post({"message": "콘크리트 타설 1세제곱미터 품셈 알려줘"})
        thread = t1["thread_id"]
        item = lambda: tools._item(stored(thread)["session"])  # noqa: E731
        checks.append(("D1 공종만 질문, 물량 1(세제곱미터 근거를 서버가 읽음)",
                       [q["name"] for q in t1["questions"]] == ["work"] and item()["quantity"]["value"] == "1"
                       and all(q.get("ref") for q in t1["questions"])))

        # 2. 공종 → 품 조건만(가격 조건·물량 없음) → 품과 근거
        llm.plan(("compute_labor", {}))
        t2 = answer(t1, {"work": "6-1-4"})
        checks.append(("D2 품 조건만 질문(관급/사급·물량 없음)", set(q["field"] for q in stored(thread)["pending"]) == set(PUMP) - {"placement", "vibrator_used", "reset_status"}
                       and [q["name"] for q in t2["questions"]] == ["pump_size"] and t2["questions_remaining"] == 5))
        llm.plan(("compute_labor", {}))
        t2b = answer(t2, PUMP)
        direct = CALCULATORS["adjusted_daily_crew"](next(s for s in load_specs().values() if s["section_no"] == "6-1-4"),
                                                    {**PUMP, "volume": "1"}, labor_only=True)
        checks.append(("D2 1㎥ 품 결과(계산기와 같음)·원가 없음·카드 답 템플릿",
                       t2b["status"] == "COMPUTED" and item()["computed_result"]["unit_lines"] == direct["unit_lines"]
                       and t2b["statement"] is None and not t2b["estimate_current"]
                       and t2b["answer_source"] == "template" and "콘크리트공" in t2b["answer"]))

        # 3. 같은 조건으로 100세제곱미터 비용: 가격 조건만 묻는다
        llm.plan(("set_conditions", {"quantity": {"value": "100", "unit": "㎥", "evidence": "100세제곱미터"}}),
                 ("estimate_cost", {}))
        t3 = post({"thread_id": thread, "message": "같은 조건으로 100세제곱미터 비용 계산해줘"})
        export_pending = http.get(f"/api/export/{thread}.xlsx").status_code
        checks.append(("D3 가격 조건만 질문, 품 조건 유지, 물량 100, Excel은 아직 404",
                       [q["name"] for q in t3["questions"]] == ["concrete_supply"]
                       and item()["quantity"]["value"] == "100"
                       and all(item()["conditions"].get(name) is not None for name in PUMP) and export_pending == 404))

        # 23. 대기 질문 보존: 관급/사급 질문 대기 중 조회(단위당 인원·근거·원문 검색)는 질문과 ref를 그대로 둔다.
        #     아래 D4가 t3 카드(원래 ref)로 답해 정상 반영되는지까지 본다.
        card = [(q["name"], q["ref"]) for q in t3["questions"]]
        kept_cards = []
        llm.plan(("compute_labor", {}))
        kept_cards.append(post({"thread_id": thread, "message": "콘크리트공은 1세제곱미터당 몇 명이야?"}))
        llm.plan(("explain_basis", {}))
        kept_cards.append(post({"thread_id": thread, "message": "할증 기준은 어디서 나와?"}))
        llm.plan(("search_standard", {"query": "펌프차 타설 할증"}))
        kept_cards.append(post({"thread_id": thread, "message": "펌프차 타설 할증 원문 찾아줘"}))
        with patch.dict(os.environ, {"AGENT_LLM": "off"}):  # 규칙 경로도 같다
            kept_cards.append(post({"thread_id": thread, "message": "콘크리트공은 1세제곱미터당 몇 명이야?"}))
            kept_cards.append(post({"thread_id": thread, "message": "할증 기준은 어디서 나와?"}))
        checks.append(("D23 대기 질문 중 조회(단위당 인원·근거·원문 검색, LLM·규칙)는 질문·ref 유지, 물량 그대로",
                       all([(q["name"], q["ref"]) for q in reply["questions"]] == card and reply["status"] == "MISSING_INFO"
                           for reply in kept_cards)
                       and [(q["field"], q["ref"]) for q in stored(thread)["pending"]] == card
                       and item()["quantity"]["value"] == "100"))

        # 4. 가격 조건 → 원가계산서(도구만으로 계산한 합계와 같음)
        llm.plan(("estimate_cost", {}))
        t4 = answer(t3, {"concrete_supply": "관급"})
        reference = tools.new_estimate("콘크리트 타설", None)
        tools.find_work(reference, hits=HITS)
        tools.set_conditions(reference, "", work="6-1-4", source="answer")
        tools.set_conditions(reference, "", values={k: {"value": v} for k, v in {**PUMP, "concrete_supply": "관급"}.items()},
                             source="answer")
        reference["basis_date"] = t4["basis_date"]
        tools.set_conditions(reference, "100㎥", quantity={"value": "100", "unit": "㎥", "evidence": "100㎥"})
        expected = tools.estimate_cost(reference)["data"]["statement"]["totals"]
        work_days_100 = Fraction(item()["computed_result"]["work_days"])
        checks.append(("D4 원가계산서 최신·합계가 서버 도구 계산과 같음(조회 뒤 원래 카드 ref로 답해도 반영)",
                       t4["status"] in ("OK", "PARTIAL") and t4["estimate_current"]
                       and "반영하지 않았어요" not in (t4["answer"] or "")
                       and t4["statement"]["totals"] == expected))

        # 6. 근거 질문: 조건·견적 유지
        llm.plan(("explain_basis", {}))
        t6 = post({"thread_id": thread, "message": "할증 기준은 어디서 나와?"})
        checks.append(("D6 근거 질문 뒤에도 조건·견적 유지(최신)",
                       t6["estimate_current"] and t6["statement"]["totals"] == t4["statement"]["totals"]
                       and item()["quantity"]["value"] == "100"))

        # 5. 물량 변경 → 재계산 → Excel(최신 합계와 같음)
        llm.plan(("set_conditions", {"quantity": {"value": "300", "unit": "㎥", "evidence": "300세제곱미터"}}),
                 ("estimate_cost", {}))
        t5 = post({"thread_id": thread, "message": "300세제곱미터로 바꿔줘"})
        workbook = load_workbook(io.BytesIO(http.get(f"/api/export/{thread}.xlsx").content), data_only=True)
        amount = t5["statement"]["totals"]["contract_amount"]
        cells = [cell.value for sheet in workbook for row in sheet.iter_rows() for cell in row]
        checks.append(("D5 물량 300 재계산, 금액 변경, Excel에 최신 도급액",
                       t5["estimate_current"] and amount != t4["statement"]["totals"]["contract_amount"]
                       and Fraction(item()["computed_result"]["work_days"]) == 3 * work_days_100
                       and amount in cells))

        # 7. 새로고침: 저장소(체크포인터)에서 새 그래프로 읽어 마지막 응답·최신 견적 복원
        fresh_graph = dialogue.build_dialogue_graph(api_main.GRAPH.checkpointer)
        restored = dialogue_service.restore(fresh_graph, thread)
        same = {k: v for k, v in restored.items() if k not in ("answer_id", "timing", "usage", "guest_session")} == \
               {k: v for k, v in t5.items() if k not in ("answer_id", "timing", "usage", "guest_session")}
        reloaded = dialogue.load(fresh_graph, thread)
        after_reload = load_workbook(io.BytesIO(http.get(f"/api/export/{thread}.xlsx").content), data_only=True)
        checks.append(("D7 저장소에서 새로 읽은 상태로 대화·최신 견적 복원, 이어서 Excel도 같은 금액",
                       same and restored["estimate_current"] and tools.current_estimate(reloaded["session"]) is not None
                       and amount in [c.value for s in after_reload for r in s.iter_rows() for c in r]))

        # 16. 조회 질문은 상태를 바꾸지 않는다(300㎥ 견적). '1세제곱미터당'은 단위당 기준이지 물량 변경이 아니다
        def snapshot() -> tuple:
            current = item()
            book = load_workbook(io.BytesIO(http.get(f"/api/export/{thread}.xlsx").content), data_only=True)
            return (current["quantity"]["value"], dict(current["conditions"]),
                    tools.current_estimate(stored(thread)["session"]) is not None,
                    amount in [c.value for s in book for r in s.iter_rows() for c in r])

        kept = ("300", dict(item()["conditions"]), True, True)
        lookup = "콘크리트공은 1세제곱미터당 몇 명이야?"
        with patch.dict(os.environ, {"AGENT_LLM": "off"}):
            rule_lookup = post({"thread_id": thread, "message": lookup})
        checks.append(("D16a 규칙 경로: 단위당 인원 질문은 물량·조건·최신 견적·Excel 금액 유지",
                       snapshot() == kept and stored(thread)["turn"]["tools"] == ["compute_labor"]
                       and "콘크리트공" in rule_lookup["answer"]))
        llm.actions = [{"__raise__": True}]
        failed_lookup = post({"thread_id": thread, "message": lookup})
        checks.append(("D16b LLM 첫 호출 실패 → 규칙 경로: 같은 결과",
                       snapshot() == kept and stored(thread)["turn"]["tools"] == ["compute_labor"]
                       and failed_lookup["answer"] == rule_lookup["answer"]))
        llm.plan(("set_conditions", {"quantity": {"value": "1", "unit": "㎥", "evidence": "1세제곱미터"}}), ("estimate_cost", {}))
        post({"thread_id": thread, "message": lookup})
        checks.append(("D16c LLM이 조회 질문에 set_conditions를 불러도 서버가 바꾸지 않음", snapshot() == kept))
        with patch.dict(os.environ, {"AGENT_LLM": "off"}):
            changed = post({"thread_id": thread, "message": "1세제곱미터로 바꿔줘"})
        checks.append(("D16d '1세제곱미터로 바꿔줘'는 물량 변경·재계산",
                       item()["quantity"]["value"] == "1" and changed["estimate_current"]
                       and changed["statement"]["totals"]["contract_amount"] != amount))
        llm.actions = [{"action": "call_tool", "tool": "set_conditions", "reason": "대본",
                        "args": {"quantity": {"value": "300", "unit": "㎥", "evidence": "300세제곱미터"}}},
                       {"__raise__": True}]
        after_failure = post({"thread_id": thread, "message": "300세제곱미터로 바꿔줘"})
        # 규칙 경로로 다시 적용하지 않는다(set_conditions 1회). 바뀐 입력은 서버가 목표(비용)에 맞게 다시 계산한다.
        checks.append(("D16e 도구 실행 뒤 LLM 실패: 적용한 결과를 규칙 경로로 다시 바꾸지 않음(set_conditions 1회), 바뀐 물량으로 재계산",
                       item()["quantity"]["value"] == "300" and stored(thread)["turn"]["tools"] == ["set_conditions", "estimate_cost"]
                       and after_failure["estimate_current"]))
        llm.plan(("estimate_cost", {}))
        post({"thread_id": thread, "message": "다시 계산해줘"})
        checks.append(("D16f 재계산 뒤 300㎥ 견적·Excel 금액 원래대로", snapshot() == kept))

        # 17. 거부 안내: 거부된 값(25m)은 '적용하지 않음'으로 설명할 때만 숫자 검증 통과
        llm.reply_override = lambda data: {"text": "25m는 지원하지 않아 기존 32m를 유지했어요."}
        llm.plan(("set_conditions", {"values": {"pump_size": {"value": "25m", "evidence": "25m"}}}))
        refused_ok = post({"thread_id": thread, "message": "펌프차 붐을 25m로 바꿔줘"})
        llm.reply_override = lambda data: {"text": "25m로 계산했어요."}
        llm.plan(("set_conditions", {"values": {"pump_size": {"value": "25m", "evidence": "25m"}}}))
        refused_bad = post({"thread_id": thread, "message": "펌프차 붐을 25m로 바꿔줘"})
        llm.reply_override = None
        checks.append(("D17 '25m는 지원하지 않아 32m 유지'는 통과, '25m로 계산'은 거부, 32m·견적 유지",
                       refused_ok["answer_source"] == "llm" and "32m를 유지" in refused_ok["answer"]
                       and refused_bad["answer_source"] == "template" and "25" in refused_bad["llm_info"]["rejected"]
                       and snapshot() == kept))

        # 18. 작업일수 표시: 정확값(30/13) 유지, 화면·답변은 '약 2.31'. 검증은 정확값과 그 반올림 결과만 허용
        def days_reply(text: str) -> dict:
            llm.reply_override = lambda data: {"text": text}
            llm.plan(("estimate_cost", {}))
            try:
                return post({"thread_id": thread, "message": "다시 계산해줘"})
            finally:
                llm.reply_override = None
        good_days = days_reply("작업일수 약 2.31일이에요.")
        bad_days = [days_reply(text)["answer_source"] for text in ("작업일수 약 2.32일이에요.", "작업일수 2.3일이에요.")]
        days_book = load_workbook(io.BytesIO(http.get(f"/api/export/{thread}.xlsx").content), data_only=True)
        days_cells = [c.value for s in days_book for r in s.iter_rows() for c in r]
        checks.append(("D18 작업일수 정확값 30/13 유지·화면 '약 2.31'·답변 2.31만 허용(2.32·2.3 거부)·Excel 계산값 그대로",
                       good_days["result"]["work_days"]["value"] == "30/13"
                       and good_days["result"]["work_days"]["display"] == "약 2.31"
                       and item()["computed_result"]["work_days"] == "30/13"
                       and good_days["answer_source"] == "llm" and bad_days == ["template", "template"]
                       and any(isinstance(v, float) and abs(v - 30 / 13) < 1e-12 for v in days_cells)
                       and snapshot() == kept))

        # 20. 실제 LLM 평가에서 본 순서: 같은 물량을 계산 뒤 한 번 더 보내도 최신 견적을 버리지 않는다
        same = {"quantity": {"value": "300", "unit": "㎥", "evidence": "300세제곱미터"}}
        llm.plan(("set_conditions", same), ("estimate_cost", {}), ("set_conditions", {**same, "values": None}))
        repeated = post({"thread_id": thread, "message": "300세제곱미터로 바꿔줘"})
        checks.append(("D20 같은 물량 재설정(계산 뒤 반복 호출)은 최신 견적·Excel 금액 유지",
                       repeated["estimate_current"] and stored(thread)["turn"]["tools"][-1] == "set_conditions"
                       and repeated["statement"]["totals"]["contract_amount"] == amount and snapshot() == kept))

        # 21. 실제 LLM 평가에서 거부됐던 맞는 문장: 기준 문서명·공종 번호는 서버 문자열이라 허용
        llm.reply_override = lambda data: {"text": "콘크리트공 0.0308인/㎥입니다. 일당 시공량 130㎥/일 기준입니다. 자세한 내용은 "
                                                   "2026 건설공사 표준품셈 공통부문 6-1-4 콘크리트 펌프차 타설을 참고하시기 바랍니다."}
        llm.plan(("compute_labor", {}))
        cited = post({"thread_id": thread, "message": lookup})
        llm.reply_override = None
        checks.append(("D21 기준 문서명(2026 건설공사 표준품셈)·공종 번호가 든 실제 문장은 통과, 견적 유지",
                       cited["answer_source"] == "llm" and snapshot() == kept))

        # 22. 실제 LLM 평가에서 거부됐던 맞는 문장: 인용의 표 소제목('2. 인력편성')·개정 연도가 든 절 제목은 허용
        llm.reply_override = lambda data: {"text": "콘크리트공 0.0308 인/㎥이 소요됩니다. 근거는 2026 건설공사 표준품셈 공통부문 "
                                                   "6-1-4 콘크리트 펌프차 타설의 2. 인력편성 표입니다."}
        llm.plan(("compute_labor", {}))
        table_named = post({"thread_id": thread, "message": lookup})
        llm.reply_override = lambda data: {"text": "할증 기준은 6-1-4 콘크리트 펌프차 타설('08, '09, '17, '22, '24, ‘25년 보완)에서 "
                                                   "확인할 수 있습니다."}
        llm.plan(("explain_basis", {}))
        section_named = post({"thread_id": thread, "message": "할증 기준은 어디서 나와?"})
        llm.reply_override = lambda data: {"text": "콘크리트공 4인이 필요합니다."}  # 인용 값(4)은 여전히 근거 없는 숫자
        llm.plan(("compute_labor", {}))
        cited_value = post({"thread_id": thread, "message": lookup})
        llm.reply_override = None
        checks.append(("D22 인용 소제목·절 제목(개정 연도)은 통과, 인용 속 값(4인)은 거부, 견적 유지",
                       table_named["answer_source"] == "llm" and section_named["answer_source"] == "llm"
                       and cited_value["answer_source"] == "template" and snapshot() == kept))

        # 19. 대기 질문의 문장 답은 조회 질문으로 막지 않는다('?'가 붙어도)
        def pending_flow(rule: bool, pump_text: str, volume_text: str) -> tuple:
            env = {"AGENT_LLM": "off"} if rule else {}
            with patch.dict(os.environ, env):
                llm.plan(("find_work", {}), ("compute_labor", {}))
                first = post({"message": "콘크리트 펌프차 타설 품 알려줘"})
                llm.plan(("compute_labor", {}))
                asked = answer(first, {"work": "6-1-4"})
                pending_thread = first["thread_id"]
                names = {q["name"] for q in asked["questions"]}
                llm.plan(("set_conditions", {"values": {"pump_size": {"value": "32m", "evidence": "32m"}}}), ("compute_labor", {}))
                post({"thread_id": pending_thread, "message": pump_text})
                llm.plan(("set_conditions", {"quantity": {"value": "100", "unit": "㎥", "evidence": "100세제곱미터"}}),
                         ("compute_labor", {}))
                post({"thread_id": pending_thread, "message": volume_text})
            current = tools._item(stored(pending_thread)["session"])
            return names, current["conditions"].get("pump_size"), (current.get("quantity") or {}).get("value")

        rule_plain = pending_flow(True, "32m야", "100세제곱미터야")
        rule_asked = pending_flow(True, "붐은 32m야?", "물량은 100세제곱미터야?")
        llm_asked = pending_flow(False, "붐은 32m야?", "물량은 100세제곱미터야?")
        checks.append(("D19 붐 길이·물량을 물은 뒤 '32m야'·'100세제곱미터야'(물음표 포함, 규칙·LLM)는 반영",
                       all("pump_size" in names and pump == "32m" and volume == "100"
                           for names, pump, volume in (rule_plain, rule_asked, llm_asked))))

        # 8. 오래된 답: 같은 필드·같은 질문이라도 예전 revision의 답은 적용하지 않는다
        llm.plan(("find_work", {}), ("set_conditions", {"work": "6-1-4"}), ("set_conditions", {
            "quantity": {"value": "10", "unit": "㎥", "evidence": "10세제곱미터"},
            "values": {name: {"value": value, "evidence": text} for name, (value, text) in {
                "pump_size": ("32m", "32m"), "structure": ("철근", "철근콘크리트"), "slump_band": ("15㎝", "15cm"),
                "placement": ("붐", "붐타설"), "vibrator_used": (True, "진동기 사용"), "reset_status": ("없음", "재셋팅 없음"),
            }.items()}}), ("estimate_cost", {}))
        s1 = post({"message": "철근콘크리트 펌프차 32m 붐타설 슬럼프 15cm 진동기 사용 재셋팅 없음 10세제곱미터 비용"})
        llm.plan(("estimate_cost", {}))
        s1b = answer(s1, {"facility_type": "Type-Ⅱ", "site_type": "Type-Ⅱ"})
        old_ref = {q["name"]: q["ref"] for q in s1b["questions"]}
        llm.plan(("set_conditions", {"quantity": {"value": "20", "unit": "㎥", "evidence": "20세제곱미터"}}), ("estimate_cost", {}))
        s2 = post({"thread_id": s1["thread_id"], "message": "같은 조건으로 20세제곱미터 비용"})
        new_ref = {q["name"]: q["ref"] for q in s2["questions"]}
        llm.plan(("estimate_cost", {}))
        stale = answer(s2, {"concrete_supply": "관급"}, refs=old_ref)
        llm.plan(("estimate_cost", {}))
        # 거부한 턴도 같은 질문을 새 revision으로 다시 낸다. 현재 질문은 바로 직전 응답의 ref다.
        fresh = answer(stale, {"concrete_supply": "관급"})
        checks.append(("D8 같은 필드라도 예전 질문(revision)의 답은 미적용, 현재 질문 답은 적용",
                       [q["name"] for q in s1b["questions"]] == ["concrete_supply"] == [q["name"] for q in s2["questions"]]
                       and old_ref != new_ref and stale["status"] == "MISSING_INFO" and "반영하지 않았어요" in stale["answer"]
                       and stale["questions"][0]["ref"] not in (old_ref["concrete_supply"], new_ref["concrete_supply"])
                       and fresh["estimate_current"]))

        # 24. 공종이 실제로 바뀌면 대기 질문을 다시 정하고, 예전 카드 ref의 답은 거부한다
        llm.plan(("find_work", {}), ("set_conditions", {"work": "6-1-4"}), ("set_conditions", {
            "quantity": {"value": "10", "unit": "㎥", "evidence": "10세제곱미터"}}), ("estimate_cost", {}))
        w1 = post({"message": "펌프차 콘크리트 타설 10세제곱미터 비용"})
        llm.plan(("estimate_cost", {}))
        w2 = answer(w1, {name: value for name, value in PUMP.items() if name in {q["field"] for q in stored(w1["thread_id"])["pending"]}})
        w2_ref = {q["name"]: q["ref"] for q in w2["questions"]}
        llm.plan(("set_conditions", {"work": "6-1-1"}), ("estimate_cost", {}))
        switched = post({"thread_id": w1["thread_id"], "message": "레디믹스트콘크리트 타설로 바꿔줘"})
        llm.plan(("estimate_cost", {}))
        old_card = answer(switched, {"concrete_supply": "관급"}, refs={"concrete_supply": w2_ref.get("concrete_supply")})
        checks.append(("D24 공종이 바뀐 뒤 예전 카드(ref)의 답은 거부",
                       "concrete_supply" in w2_ref and tools._item(stored(w1["thread_id"])["session"])["selection"]["section_no"] == "6-1-1"
                       and "반영하지 않았어요" in old_card["answer"]
                       and tools._item(stored(w1["thread_id"])["session"])["conditions"].get("concrete_supply") != "관급"))

        # 25. 품 계산 완료 → 같은 물량으로 비용 요청: 입력 변경이 없어도 새 필수 질문(관급/사급)을 낸다.
        #     같은 질문을 다시 보여줄 때는 ref를 바꾸지 않고, 그 카드로 답하면 견적·Excel까지 완료된다.
        llm.plan(("find_work", {}), ("set_conditions", {"quantity": {"value": "100", "unit": "㎥", "evidence": "100세제곱미터"}}),
                 ("compute_labor", {}))
        c1 = post({"message": "콘크리트 타설 100세제곱미터 품 알려줘"})
        llm.plan(("compute_labor", {}))
        c2 = answer(c1, {"work": "6-1-4"})
        llm.plan(("compute_labor", {}))
        c3 = answer(c2, PUMP)
        llm.plan(("estimate_cost", {}))
        c4 = post({"thread_id": c1["thread_id"], "message": "같은 조건으로 비용도 계산해줘"})
        llm.plan(("estimate_cost", {}))
        c5 = post({"thread_id": c1["thread_id"], "message": "비용 계산해줘"})  # 같은 질문을 다시 보여줌
        llm.plan(("compute_labor", {}))
        c6 = post({"thread_id": c1["thread_id"], "message": "콘크리트공은 1세제곱미터당 몇 명이야?"})  # 조회: 질문 유지
        llm.plan(("estimate_cost", {}))
        c7 = answer(c4, {"concrete_supply": "관급"})  # 처음 받은 카드로 답한다
        cost_export = http.get(f"/api/export/{c1['thread_id']}.xlsx")
        cost_cells = [c.value for s in load_workbook(io.BytesIO(cost_export.content), data_only=True)
                      for r in s.iter_rows() for c in r] if cost_export.status_code == 200 else []
        refs_of = lambda reply: [(q["name"], q["ref"]) for q in reply["questions"]]  # noqa: E731
        checks.append(("D25 품 완료 → 같은 물량 비용 요청은 관급/사급 질문 생성, 같은 질문 재표시·조회는 ref 유지, 카드 답으로 견적·Excel 완료",
                       c3["status"] == "COMPUTED" and [q["name"] for q in c4["questions"]] == ["concrete_supply"]
                       and refs_of(c5) == refs_of(c4) == refs_of(c6)
                       and c7["estimate_current"] and c7["statement"]["totals"]["contract_amount"] in cost_cells))

        # 25b. 같은 흐름을 규칙 경로(LLM 끔·실패 대체)로: 새 견적으로 초기화하지 않고 같은 결과
        with patch.dict(os.environ, {"AGENT_LLM": "off"}):
            r1 = post({"message": "콘크리트 타설 100세제곱미터 품 알려줘"})
            r2 = answer(r1, {"work": "6-1-4"})
            r3 = answer(r2, PUMP)
            r4 = post({"thread_id": r1["thread_id"], "message": "같은 조건으로 비용도 계산해줘"})
            r5 = post({"thread_id": r1["thread_id"], "message": "콘크리트공은 1세제곱미터당 몇 명이야?"})
            r7 = answer(r4, {"concrete_supply": "관급"})
        rule_item = tools._item(stored(r1["thread_id"])["session"])
        checks.append(("D25b 규칙 경로: 품 완료 → 같은 조건 비용 요청은 견적 유지(초기화 없음)·관급/사급 질문, 조회는 ref 유지, 카드 답으로 견적 완료",
                       r3["status"] == "COMPUTED" and [q["name"] for q in r4["questions"]] == ["concrete_supply"]
                       and refs_of(r5) == refs_of(r4) and rule_item["quantity"]["value"] == "100"
                       and rule_item["conditions"].get("pump_size") == "32m" and r7["estimate_current"]
                       and r7["statement"]["totals"]["contract_amount"] == c7["statement"]["totals"]["contract_amount"]))

        # 26. 품 완료 → 공종을 말하지 않은 비용 요청('비용 계산해줘')은 지금 견적으로(특정 표현 없이도).
        #     명확한 다른 공종은 새 견적, 모호하면 상태를 보존하고 확인 질문.
        def labor_done() -> str:
            llm.plan(("find_work", {}), ("set_conditions", {"quantity": {"value": "100", "unit": "㎥",
                                                                         "evidence": "100세제곱미터"}}), ("compute_labor", {}))
            first = post({"message": "콘크리트 타설 100세제곱미터 품 알려줘"})
            llm.plan(("compute_labor", {}))
            second = answer(first, {"work": "6-1-4"})
            llm.plan(("compute_labor", {}))
            answer(second, PUMP)
            return first["thread_id"]

        def kept(thread_id: str) -> tuple:
            current = tools._item(stored(thread_id)["session"])
            return (stored(thread_id)["session"]["estimate_id"], current["selection"].get("section_no"),
                    (current.get("quantity") or {}).get("value"), current["conditions"].get("pump_size"))

        def outcome(reply: dict, thread_id: str, before: tuple) -> tuple:
            return [q["name"] for q in reply["questions"]], kept(thread_id) == before

        llm_thread = labor_done(); llm_before = kept(llm_thread)
        llm.plan(("estimate_cost", {}))
        by_llm = outcome(post({"thread_id": llm_thread, "message": "비용 계산해줘"}), llm_thread, llm_before)
        fail_thread = labor_done(); fail_before = kept(fail_thread)
        llm.actions = [{"__raise__": True}]
        by_rule = outcome(post({"thread_id": fail_thread, "message": "비용 계산해줘"}), fail_thread, fail_before)
        guard_thread = labor_done(); guard_before = kept(guard_thread)
        llm.plan(("find_work", {}), ("estimate_cost", {}))  # LLM이 새 견적을 시작하려 해도 서버가 막는다
        by_guard = outcome(post({"thread_id": guard_thread, "message": "비용 계산해줘"}), guard_thread, guard_before)
        checks.append(("D26a 품 완료 → '비용 계산해줘': 공종·조건·물량 유지, 가격 질문 생성(LLM·LLM 실패→규칙·LLM의 find_work 차단 모두 같음)",
                       by_llm == by_rule == by_guard == (["concrete_supply"], True)))

        new_thread = labor_done(); new_before = kept(new_thread)
        with patch.dict(os.environ, {"AGENT_LLM": "off"}):
            new_rule = post({"thread_id": new_thread, "message": "거푸집 설치 비용 알려줘"})
        llm_new_thread = labor_done(); llm_new_before = kept(llm_new_thread)
        llm.plan(("find_work", {}), ("estimate_cost", {}))
        new_llm = post({"thread_id": llm_new_thread, "message": "거푸집 설치 비용 알려줘"})
        checks.append(("D26b 명확한 다른 공종('거푸집 설치 비용')은 새 견적 흐름(규칙·LLM)",
                       kept(new_thread)[0] != new_before[0] and [q["name"] for q in new_rule["questions"]] == ["work"]
                       and kept(llm_new_thread)[0] != llm_new_before[0] and [q["name"] for q in new_llm["questions"]] == ["work"]))

        vague_thread = labor_done(); vague_before = kept(vague_thread)
        with patch.dict(os.environ, {"AGENT_LLM": "off"}):
            vague = post({"thread_id": vague_thread, "message": "콘크리트 거푸집 비용 알려줘"})
            stay = answer(vague, {"scope": dialogue.SCOPE_CURRENT})
        vague_new = labor_done()
        with patch.dict(os.environ, {"AGENT_LLM": "off"}):
            asked_new = post({"thread_id": vague_new, "message": "콘크리트 거푸집 비용 알려줘"})
            vague_new_before = kept(vague_new)
            restarted = answer(asked_new, {"scope": dialogue.SCOPE_NEW})
        checks.append(("D26c 모호한 요청('콘크리트 거푸집 비용')은 상태 보존·확인 질문 → '지금 견적'은 가격 질문, '새 견적'은 새 흐름",
                       [q["name"] for q in vague["questions"]] == ["scope"] and kept(vague_thread) == vague_before
                       and [q["name"] for q in stay["questions"]] == ["concrete_supply"] and kept(vague_thread) == vague_before
                       and kept(vague_new)[0] != vague_new_before[0] and [q["name"] for q in restarted["questions"]] == ["work"]))

        # 27. 공종 판정은 조건·물량 변경보다 먼저: 새 공종 문장의 물량을 지금 견적에 넣지 않는다(규칙, LLM 실패→규칙)
        def via(mode: str, thread_id: str, message: str) -> dict:
            if mode == "rule":
                with patch.dict(os.environ, {"AGENT_LLM": "off"}):
                    return post({"thread_id": thread_id, "message": message})
            llm.actions = [{"__raise__": True}]
            return post({"thread_id": thread_id, "message": message})

        def session_of(thread_id: str) -> tuple:
            state = stored(thread_id)
            current = tools._item(state["session"])
            return (state["session"]["estimate_id"], current.get("request_text"), current.get("quantity"),
                    (current.get("selection") or {}).get("section_no"), current["conditions"].get("pump_size"))

        order = {}
        for mode in ("rule", "llm_failed"):
            door_thread = labor_done(); door_before = session_of(door_thread)
            door = via(mode, door_thread, "자동문 3개소 설치 비용")
            door_after = session_of(door_thread)
            form_thread = labor_done(); form_before = session_of(form_thread)
            via(mode, form_thread, "거푸집 100㎡ 비용")
            form_after = session_of(form_thread)
            same_thread = labor_done(); same_before = session_of(same_thread)
            same = via(mode, same_thread, "콘크리트 200세제곱미터 비용")
            same_after = session_of(same_thread)
            vague_thread = labor_done(); vague_before = session_of(vague_thread)
            vague = via(mode, vague_thread, "콘크리트 거푸집 100㎡ 비용 알려줘")
            order[mode] = {
                "1 자동문: 새 세션·3개소, 콘크리트 계산 없음": door_after[0] != door_before[0] and "자동문" in door_after[1]
                and (door_after[2] or {}).get("value") == "3" and door["statement"] is None
                and stored(door_thread)["turn"]["tools"][:2] == ["find_work", "set_conditions"],
                "2 거푸집: 새 세션·100㎡": form_after[0] != form_before[0] and "거푸집" in form_after[1]
                and (form_after[2] or {}).get("value") == "100",
                "3 콘크리트 200㎥: 같은 세션·물량 갱신·조건 유지": same_after[0] == same_before[0]
                and same_after[2]["value"] == "200" and same_after[3:] == same_before[3:]
                and [q["name"] for q in same["questions"]] == ["concrete_supply"],
                "4 모호+물량: 확인 질문, 기존 상태 그대로": vague_after_ok(vague, vague_thread, vague_before)}
        checks.append(("D27 공종 판정을 조건 변경보다 먼저(규칙·LLM 실패→규칙): 자동문 3개소·거푸집 100㎡는 새 흐름, "
                       "콘크리트 200㎥는 물량 갱신, 모호하면 물량이 있어도 확인 전 상태 유지",
                       all(all(result.values()) for result in order.values())))
        if not all(all(result.values()) for result in order.values()):
            print("  D27", order)

        # 27b. LLM이 새 공종 문장의 물량을 지금 견적에 넣으려 해도(set_conditions → estimate_cost) 서버가 막는다
        guard_door = labor_done(); guard_door_before = session_of(guard_door)
        llm.plan(("set_conditions", {"quantity": {"value": "3", "unit": "개소", "evidence": "3개소"}}), ("estimate_cost", {}))
        guarded = post({"thread_id": guard_door, "message": "자동문 3개소 설치 비용"})
        checks.append(("D27b LLM이 다른 공종 물량을 지금 견적에 넣으려 하면 거부, 콘크리트 견적 그대로",
                       session_of(guard_door) == guard_door_before and guarded["statement"] is None
                       and stored(guard_door)["turn"]["trace"][0]["status"] == "rejected"))

        # 28. 새 견적의 물량은 LLM이 set_conditions를 생략해도 서버가 보존한다(공종 미확정 → 선택 → 조건 답 후에도 유지).
        #     물량이 여러 개면 고르지 않고 묻는다. LLM 경로·LLM 실패→규칙·규칙 경로가 같은 상태.
        form_answers = {"complexity": "보통", "height": "3", "is_roof_slab": False, "crane_used": False}

        def fail_or_plan(mode: str) -> None:
            """다음 턴의 LLM: llm=compute_labor만 고름, llm_failed=매 턴 실패, llm_reply=도구 없이 바로 답함."""
            if mode == "llm":
                llm.plan(("compute_labor", {}))
            elif mode == "llm_failed":
                llm.actions = [{"__raise__": True}]
            elif mode == "llm_reply":
                llm.actions = []

        def form_flow(mode: str, message: str = "거푸집 100㎡ 비용") -> dict:
            out = {}
            with patch.object(tools, "retrieve", return_value={"hits": FORM_HITS}),                     patch.dict(os.environ, {"AGENT_LLM": "off"} if mode == "rule" else {}):
                if mode in ("llm", "llm_reply"):
                    llm.plan(("find_work", {}))  # set_conditions를 부르지 않는다
                elif mode == "llm_failed":
                    llm.actions = [{"__raise__": True}]
                first = post({"message": message})
                thread_id = first["thread_id"]
                out["start"] = (tools._item(stored(thread_id)["session"]).get("quantity") or {})
                fail_or_plan(mode)
                chosen = answer(first, {"work": "6-3-1"})
                out["asked"] = [q["field"] for q in stored(thread_id)["pending"]]
                out["shown"] = [q["name"] for q in chosen["questions"]]
                out["selected"] = (tools._item(stored(thread_id)["session"]).get("quantity") or {})
                fail_or_plan(mode)
                values = {k: v for k, v in form_answers.items() if k in out["asked"]}
                done = answer(chosen, values) if values else chosen
                item_after = tools._item(stored(thread_id)["session"])
                out["after"] = (item_after.get("quantity") or {})
                out["volume"] = item_after["conditions"].get("volume")
                out["goal"] = stored(thread_id)["goal"]
                out["status"] = done["status"]
            return out

        flows = {mode: form_flow(mode) for mode in ("llm", "llm_reply", "llm_failed", "rule")}
        hundred = lambda q: q.get("value") == "100" and tools.unit_key(q.get("unit")) == tools.unit_key("㎡")  # noqa: E731
        checks.append(("D28 거푸집 100㎡: find_work만 불러도 물량 보존, 공종 선택·조건 답 후에도 100㎡ 유지(LLM·LLM 바로 답함·LLM 실패·규칙 동일)",
                       all(hundred(f["start"]) and hundred(f["selected"]) and hundred(f["after"]) and f["volume"] == "100"
                           and "volume" not in f["asked"] and f["goal"] == "cost" for f in flows.values())
                       and len({(f["status"], tuple(f["asked"])) for f in flows.values()}) == 1))
        if not checks[-1][1]:
            print("  D28", flows)
        two = form_flow("rule", "거푸집 100㎡ 200㎡ 비용")
        checks.append(("D28b 물량이 여러 개면 고르지 않고 공종 선택 뒤 물량을 묻는다", two["start"] == {} and two["asked"][-1] == "volume" and len(two["shown"]) == 1))

        # 29. 비용으로 시작한 요청은 LLM이 compute_labor만 골라도 비용 흐름을 이어 간다(가격 질문 → 견적).
        def cost_flow(mode: str) -> tuple:
            with patch.dict(os.environ, {"AGENT_LLM": "off"} if mode == "rule" else {}):
                if mode in ("llm", "llm_reply"):
                    llm.plan(("find_work", {}), ("compute_labor", {}))
                elif mode == "llm_failed":
                    llm.actions = [{"__raise__": True}]
                first = post({"message": "콘크리트 타설 100세제곱미터 비용 알려줘"})
                fail_or_plan(mode)
                second = answer(first, {"work": "6-1-4"})
                fail_or_plan(mode)
                third = answer(second, PUMP)
                fail_or_plan(mode)
                fourth = answer(third, {"concrete_supply": "관급"})
            return ([q["name"] for q in third["questions"]], fourth["estimate_current"],
                    (fourth["statement"] or {}).get("totals", {}).get("contract_amount"), stored(first["thread_id"])["goal"])

        costs = {mode: cost_flow(mode) for mode in ("llm", "llm_reply", "llm_failed", "rule")}
        checks.append(("D29 비용 요청: 조건 카드 답 뒤 LLM이 compute_labor만 고르거나 바로 답해도 관급/사급 질문 → 견적까지(LLM 실패·규칙 동일)",
                       all(c[0] == ["concrete_supply"] and c[1] and c[3] == "cost" for c in costs.values())
                       and len({c[2] for c in costs.values()}) == 1))
        if not checks[-1][1]:
            print("  D29", costs)

        # 30. 실제 LLM 평가(S5)에서 본 순서: 공종이 바로 정해진 비용 요청에 LLM이 set_conditions(work)만 부르고 답함.
        #     서버가 명세 단위(㎥)의 물량 300을 보존하고(32m·15cm 제외), 계산으로 이어 남은 질문 → 견적까지 간다.
        s5_message = "철근콘크리트 300세제곱미터를 32m 붐 펌프차로 타설 비용 알려줘. 슬럼프 15cm, 진동기 사용, 재셋팅 없음"
        with patch.object(tools, "retrieve", return_value={"hits": HITS[:1]}):
            llm.plan(("find_work", {}), ("set_conditions", {"work": "6-1-4"}))
            p1 = post({"message": s5_message})
        p_item = tools._item(stored(p1["thread_id"])["session"])
        llm.actions = []  # 카드 답 뒤에도 LLM이 바로 답한다
        p2 = answer(p1, {"facility_type": "Type-Ⅱ", "site_type": "Type-Ⅱ"})
        llm.actions = []
        p3 = answer(p2, {"concrete_supply": "관급"})
        checks.append(("D30 공종 바로 확정 + LLM이 계산 없이 답해도: 물량 300㎥ 보존·남은 질문 → 관급 → 견적(목표 비용 유지)",
                       p_item["quantity"]["value"] == "300" and p_item["conditions"].get("pump_size") == "32m"
                       and [q["name"] for q in p1["questions"]] == ["facility_type"] and p1["questions_remaining"] == 2
                       and [q["name"] for q in p2["questions"]] == ["concrete_supply"] and p3["estimate_current"]
                       and stored(p1["thread_id"])["goal"] == "cost"))
        if not checks[-1][1]:
            print("  D30", p_item.get("quantity"), [q["name"] for q in p1["questions"]], p1["status"], [q["name"] for q in p2["questions"]], p3.get("estimate_current"))

        # 9. LLM이 근거 검증을 우회하지 못함: 근거와 다른 값, source 주입은 무시된다
        before = item()["quantity"]["value"]
        llm.plan(("set_conditions", {"quantity": {"value": "999", "unit": "㎥", "evidence": "500세제곱미터"},
                                     "source": "answer"}), ("set_conditions", {"values": {"concrete_supply": {"value": "사급"}},
                                                                                "source": "answer"}))
        bypass = post({"thread_id": thread, "message": "500세제곱미터면 어때?"})
        checks.append(("D9 근거와 다른 물량·근거 없는 값은 거부(LLM이 source 지정 불가)",
                       item()["quantity"]["value"] == before and item()["conditions"].get("concrete_supply") == "관급"
                       and bypass["estimate_current"]))

        # 10. 응답 문장 검증: 근거 없는 금액·단위 불일치면 LLM 문장을 버리고 템플릿
        llm.reply_override = lambda data: {"text": "도급액 1,000원입니다."}
        llm.plan(("estimate_cost", {}))
        wrong = post({"thread_id": thread, "message": "다시 계산해줘"})
        llm.reply_override = lambda data: {"text": "노무비가 많이 들어요 12,345원"}
        llm.plan(("estimate_cost", {}))
        unclaimed = post({"thread_id": thread, "message": "다시 계산해줘"})
        llm.reply_override = None
        checks.append(("D10 사실과 다른 숫자·근거 없는 숫자 문장은 템플릿으로 대체",
                       wrong["answer_source"] == "template" and "값 불일치" in wrong["llm_info"]["rejected"]
                       and unclaimed["answer_source"] == "template" and unclaimed["llm_info"]["rejected"]
                       and reply_check.display(str(t5["statement"]["totals"]["contract_amount"])) in wrong["answer"]))

        # 11. 예산: 도구 4회·LLM 시도 6회 상한
        llm.plan(*[("explain_basis", {})] * 10)
        capped = post({"thread_id": thread, "message": "근거를 계속 보여줘"})
        info = capped["llm_info"]
        checks.append(("D11 턴 예산(도구 ≤4, LLM 시도 ≤6) 안에서 끝남",
                       info["tool_calls"] <= 4 and info["llm_attempts"] <= 6 and capped["estimate_current"]))
        llm.actions = []

        # 12. 중복 요청: 같은 request_id는 한 번만 실행·차감, 다른 내용이면 409
        runs = []
        original_turn = dialogue.run_turn

        def slow_turn(*args, **kwargs):
            runs.append(1)
            time.sleep(0.3)
            return original_turn(*args, **kwargs)

        body = {"thread_id": thread, "message": "근거 보여줘", "request_id": "11111111-1111-4111-8111-111111111111"}
        consumed.clear()
        with patch.object(dialogue, "run_turn", slow_turn):
            llm.plan(("explain_basis", {}))
            results = []
            workers = [threading.Thread(target=lambda: results.append(http.post("/api/chat", json=body))) for _ in range(2)]
            [worker.start() for worker in workers]
            [worker.join() for worker in workers]
            again = http.post("/api/chat", json=body)
            conflict = http.post("/api/chat", json={**body, "message": "다른 내용"})
        answers = {result.json()["answer_id"] for result in [*results, again]}
        checks.append(("D12 동시·재전송 중복 요청은 실행 1회·차감 1회·같은 응답, 다른 내용은 409",
                       [r.status_code for r in results] == [200, 200] and again.status_code == 200
                       and len(answers) == 1 and len(runs) == 1 and len(consumed) == 1 and conflict.status_code == 409))

        # 15. 실제 LLM 형식: 조건을 [{field, value, evidence}] 목록으로, 참/거짓은 문자열로 보낸다
        llm.plan(("set_conditions", {"values": [{"field": "concrete_supply", "value": "사급", "evidence": "사급으로"},
                                                {"field": "vibrator_used", "value": "true", "evidence": "진동기 사용"}]}),
                 ("estimate_cost", {}))
        listed = post({"thread_id": thread, "message": "레미콘 사급으로, 진동기 사용"})
        checks.append(("D15 목록 형식 조건·문자열 참/거짓을 근거 대조 후 반영", item()["conditions"].get("concrete_supply") == "사급"
                       and item()["conditions"].get("vibrator_used") is True
                       and [q["name"] for q in listed["questions"]] == ["ready_mix_price"]))

        # 14. 비회원 대화 삭제는 새 흐름 상태(dlg: thread)도 지운다
        deleted = http.delete(f"/api/guest/conversations/{thread}").status_code
        checks.append(("D14 비회원 대화 삭제 시 대화 상태도 삭제", deleted == 204 and stored(thread) is None
                       and http.get(f"/api/export/{thread}.xlsx").status_code == 404))

    check_free_card_answers(checks)
    check_one_question(checks)

    # 13. 설정 off: 기존 흐름
    os.environ["AGENT_MODE"] = ""
    with patch("backend.api.usage_limits.processing", return_value=nullcontext((None, "fixture", False))), \
            patch("backend.api.usage_limits.status", return_value={}), \
            patch("backend.api.usage_limits.consume", return_value={}), \
            patch.object(client, "generate", side_effect=AssertionError("LLM called")), \
            patch.dict(os.environ, {"AGENT_LLM": "off"}), \
            TestClient(api_main.app, headers={"X-Guest-Session": SECRET}) as http:
        legacy = http.post("/api/chat", json={"message": "철근콘크리트 벽체 260㎥ 펌프차로 타설 비용"}).json()
    checks.append(("D13 AGENT_MODE 꺼짐이면 기존 그래프 응답(ref 없음)",
                   legacy["status"] == "MISSING_INFO" and "estimate_current" not in legacy
                   and all("ref" not in q for q in legacy["questions"])))

    for name, passed in checks:
        print(f"{'PASS' if passed else 'FAIL'} {name}")
    print(f"통과 {sum(passed for _, passed in checks)} / 전체 {len(checks)}")
    return 0 if all(passed for _, passed in checks) else 1


if __name__ == "__main__":
    raise SystemExit(main())
