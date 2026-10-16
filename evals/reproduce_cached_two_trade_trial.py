"""One bounded local HTTP run for two exact, calculator-backed cached eval questions."""
from __future__ import annotations

import io
import json
import os
import subprocess
import sys
import time
import urllib.error
import urllib.request
import zipfile
from pathlib import Path

PORT = 8012
QUESTIONS = {
    "waterproofing": "시멘트 액체방수를 바닥에 300㎡ 바르려면 인력 품은 어떻게 잡아야 해?",
    "meter": "건물 기계실 배관 ø100㎜에 건물용 적산열량계를 설치하면 개당 배관공은 몇 인이야?",
}
TARGETS = {"waterproofing": "6-4-1", "meter": "6-2-2"}
ANSWERS = {
    "waterproofing": {"area": ("300", "㎡"), "surface_type": ("바닥", None)},
    # The exact question asks for the per-meter quantity; 100 mm belongs to the printed 100~125 mm band.
    "meter": {"meter_size": ("ø100~125mm", None), "quantity": ("1", "개")},
}
DEADLINE: float | None = None


def request(opener, method: str, url: str, body: dict | None = None,
            headers: dict[str, str] | None = None):
    remaining = 6.0 if DEADLINE is None else DEADLINE - time.monotonic()
    if remaining <= 0:
        raise TimeoutError("52-second reproduction deadline reached")
    data = json.dumps(body, ensure_ascii=False).encode("utf-8") if body is not None else None
    request_obj = urllib.request.Request(url, data=data, headers=headers or {}, method=method)
    try:
        with opener.open(request_obj, timeout=min(6.0, remaining)) as response:
            return response.status, response.read(), dict(response.headers)
    except urllib.error.HTTPError as error:
        return error.code, error.read(), dict(error.headers)


def decoded(body: bytes) -> dict:
    return json.loads(body.decode("utf-8"))


def scope_decisions(analysis: dict) -> list[dict]:
    actions = {"requested": "include_work", "background": "keep_background",
               "condition": "keep_condition", "excluded": "keep_excluded"}
    decisions = []
    for span in analysis["scope_spans"]:
        classification = span["classification"]
        action = actions.get(classification)
        decision = {"span_id": span["span_id"], "action": action}
        if classification == "unconfirmed":
            work_ids = span.get("work_ids", [])
            if len(work_ids) == 1:
                work = next(row["work"] for row in analysis["work_items"] if row["id"] == work_ids[0])
            elif span["text"] == QUESTIONS["meter"] and "건물용 적산열량계" in span["text"]:
                # This is the evaluator-author's explicit scope choice for this exact request.
                work = "건물용 적산열량계"
            else:
                raise RuntimeError(f"No explicit supported scope choice for {span['span_id']}: {span['text']}")
            decision.update(action="include_work", work=work)
        elif action is None:
            raise RuntimeError(f"Unexpected scope classification: {classification}")
        decisions.append(decision)
    return decisions


def main() -> int:
    global DEADLINE
    root = Path(__file__).resolve().parents[1]
    base = f"http://127.0.0.1:{PORT}"
    opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))
    deadline = time.monotonic() + 52
    DEADLINE = deadline
    stdout_path = root / ".trial-two-trade.server.stdout.log"
    stderr_path = root / ".trial-two-trade.server.stderr.log"
    server_code = (
        "from backend.shared import embedding; from backend.knowledge import embed_index; "
        "deny=lambda *a,**k: (_ for _ in ()).throw(AssertionError('external embedding call')); "
        "embedding.embed_texts=deny; embed_index.make_client=deny; "
        f"import uvicorn; uvicorn.run('backend.api.item_labor_trial:app', host='127.0.0.1', port={PORT}, log_level='warning')"
    )
    env = os.environ.copy()
    env.update({"LOCAL_TRIAL_ITEM_LABOR": "on", "AGENT_OFFLINE": "1", "AGENT_LLM": "off",
                "KNOWLEDGE_ASSET_DIR": r"C:\Users\user\Desktop\PROJECTS\civilai-construction-risk-agent\data\knowledge"})
    with stdout_path.open("wb") as server_stdout, stderr_path.open("wb") as server_stderr:
        server = subprocess.Popen([sys.executable, "-B", "-u", "-c", server_code], cwd=root,
                                  env=env, stdout=server_stdout, stderr=server_stderr)
        try:
            while time.monotonic() < deadline and server.poll() is None:
                try:
                    status, _, _ = request(opener, "GET", base + "/api/item-labor/health")
                    if status == 200:
                        print("HEALTH_HTTP 200", flush=True)
                        break
                except (OSError, TimeoutError):
                    time.sleep(0.2)
            else:
                print(f"SERVER_NOT_READY pid={server.pid} exit={server.poll()}", flush=True)
                return 10

            source = " 그리고 ".join(QUESTIONS.values())
            status, body, _ = request(opener, "POST", base + "/api/item-labor/analyze",
                                      {"text": source}, {"Content-Type": "application/json"})
            print("ANALYZE_HTTP", status, flush=True)
            if status != 200:
                print(body.decode("utf-8", "replace")[:1600], flush=True)
                return 11
            trial = decoded(body)
            analysis = trial["analysis"]
            print("SCOPE_CHOICES", json.dumps({
                "spans": [{k: span.get(k) for k in ("span_id", "classification", "text", "work_ids", "condition_ids")}
                          for span in analysis["scope_spans"]],
                "decisions": scope_decisions(analysis),
            }, ensure_ascii=False), flush=True)
            decisions = scope_decisions(analysis)
            headers = {"Content-Type": "application/json", "X-Trial-Owner": trial["owner_token"]}
            url = f"{base}/api/item-labor/{trial['trial_id']}"
            status, body, _ = request(opener, "POST", url + "/confirm", {"decisions": decisions}, headers)
            print("CONFIRM_HTTP", status, flush=True)
            if status != 200:
                print(body.decode("utf-8", "replace")[:1600], flush=True)
                return 12
            confirmed = decoded(body)
            items = confirmed.get("items", [])
            if len(items) != 2:
                print("STOP: expected two confirmed items", json.dumps(items, ensure_ascii=False), flush=True)
                return 13
            item_by_text = {item["request_text"]: item for item in items}
            if set(item_by_text) != set(QUESTIONS.values()):
                print("STOP: exact question text not retained", json.dumps(list(item_by_text), ensure_ascii=False), flush=True)
                return 14
            print("CONFIRMED_ITEMS", json.dumps([{"item_id": x["item_id"], "request_text": x["request_text"],
                    "scope_status": x.get("scope_status")} for x in items], ensure_ascii=False), flush=True)

            calculated = {}
            for key, exact_question in QUESTIONS.items():
                item = item_by_text[exact_question]
                item_url = f"{url}/items/{item['item_id']}"
                status, body, _ = request(opener, "POST", item_url + "/search", {}, headers)
                print(f"SEARCH_HTTP {key}", status, flush=True)
                if status != 200:
                    print(body.decode("utf-8", "replace")[:1600], flush=True)
                    return 15
                result = decoded(body)
                item = next(x for x in result["items"] if x["item_id"] == item["item_id"])
                expected_section = TARGETS[key]
                candidates = item.get("candidates", [])
                print(f"SEARCH_RESULT {key}", json.dumps({
                    "request_text": item["request_text"], "search_status": item.get("search_status"),
                    "search_info": item.get("search_info"),
                    "candidates": [{k: candidate.get(k) for k in ("spec_id", "section_no", "section", "has_spec")}
                                   for candidate in candidates],
                    "selected_specification": item.get("selected_specification"),
                    "questions": [{k: question.get(k) for k in ("question_id", "version", "field", "ask", "unit")}
                                  for question in item.get("questions", [])]}, ensure_ascii=False), flush=True)
                if item.get("search_status") != "COMPLETE" or not (item.get("search_info") or {}).get("query_cache_hit"):
                    print("STOP: exact query was not a local cache hit", key, flush=True)
                    return 16
                candidate = next((c for c in candidates if str(c.get("section_no")) == expected_section
                                  and c.get("has_spec") and c.get("spec_id")), None)
                selected = item.get("selected_specification") or {}
                if candidate is None and str(selected.get("section_no")) != expected_section:
                    print("STOP: the exact target section has no calculator spec; no substitute selected", key, flush=True)
                    return 17

                # Answer only the candidate the exact question and the pinned spec support.
                for _ in range(6):
                    questions = item.get("questions", [])
                    if not questions:
                        break
                    question = questions[0]
                    field = question.get("field")
                    if field == "work":
                        chosen = candidate or next((c for c in candidates if str(c.get("section_no")) == expected_section), None)
                        if not chosen:
                            raise RuntimeError(f"No matching candidate for work confirmation: {key}")
                        value, unit = chosen["section"], None
                    else:
                        pair = ANSWERS[key].get(field)
                        if pair is None:
                            raise RuntimeError(f"No source-supported answer for {key}.{field}; refusing to guess")
                        value, unit = pair
                    answer_body = {"question_id": question["question_id"], "version": question["version"], "value": value}
                    if unit:
                        answer_body["unit"] = unit
                    status, body, _ = request(opener, "POST", item_url + "/answer", answer_body, headers)
                    print(f"ANSWER_HTTP {key}", status, json.dumps(answer_body, ensure_ascii=False), flush=True)
                    if status != 200:
                        print(body.decode("utf-8", "replace")[:1600], flush=True)
                        return 18
                    result = decoded(body)
                    item = next(x for x in result["items"] if x["item_id"] == item["item_id"])
                else:
                    raise RuntimeError(f"Too many unanswered fields for {key}")

                if item.get("selected_specification", {}).get("section_no") != expected_section:
                    print("STOP: selected calculator differs from the exact target", key,
                          json.dumps(item.get("selected_specification"), ensure_ascii=False), flush=True)
                    return 19
                if item.get("status") != "CALCULATED":
                    print("STOP: item calculation did not complete", key,
                          json.dumps({"status": item.get("status"), "questions": item.get("questions"),
                                      "labor": item.get("labor")}, ensure_ascii=False), flush=True)
                    return 20
                calculated[key] = item
                print(f"CALCULATED {key}", json.dumps({"request_text": item["request_text"],
                    "selected_specification": item["selected_specification"], "quantity": item.get("quantity"),
                    "labor": item.get("labor"), "search_info": item.get("search_info")}, ensure_ascii=False), flush=True)

            if len(calculated) != 2:
                raise RuntimeError("both exact trade calculations were not completed")
            status, body, _ = request(opener, "POST", url + "/estimate",
                                      {"basis_date": "2026-10-16", "common_conditions": {}}, headers)
            print("ESTIMATE_HTTP", status, flush=True)
            if status != 200:
                print(body.decode("utf-8", "replace")[:1800], flush=True)
                return 21
            estimate = decoded(body)
            estimate_items = estimate.get("items", [])
            print("ESTIMATE_RESULT", json.dumps({"status": estimate.get("status"),
                "item_count": len(estimate_items), "aggregate": estimate.get("aggregate"),
                "items": [{k: item.get(k) for k in ("item_id", "request_text", "status", "selected_specification", "quantity")}
                          for item in estimate_items]}, ensure_ascii=False), flush=True)
            if len(estimate_items) != 2 or len(estimate.get("aggregate", {}).get("included", [])) != 2:
                print("STOP: combined estimate does not include both calculated trades", flush=True)
                return 22

            status, xlsx, _ = request(opener, "GET", url + "/estimate.xlsx", headers=headers)
            print("XLSX_HTTP", status, flush=True)
            if status != 200 or not zipfile.is_zipfile(io.BytesIO(xlsx)):
                print(xlsx[:1500].decode("utf-8", "replace") if status != 200 else "Invalid XLSX archive", flush=True)
                return 23
            output = root / ".trial" / "multi-trade-two-calculated-items.xlsx"
            output.parent.mkdir(parents=True, exist_ok=True)
            output.write_bytes(xlsx)
            with zipfile.ZipFile(io.BytesIO(xlsx)) as archive:
                if archive.testzip() is not None:
                    raise RuntimeError("XLSX archive CRC check failed")
            from openpyxl import load_workbook

            workbook = load_workbook(output, read_only=True, data_only=True)
            try:
                sheetnames = workbook.sheetnames
                if not sheetnames:
                    raise RuntimeError("XLSX contains no sheets")
            finally:
                workbook.close()
            print("XLSX_OPENPYXL_OK", json.dumps({"path": str(output), "bytes": len(xlsx),
                "sheets": sheetnames}, ensure_ascii=False), flush=True)
            print("PASS exact two-trade cached searches, calculator-backed answers, aggregate, XLSX; external embeddings blocked", flush=True)
            return 0
        finally:
            if server.poll() is None:
                server.terminate()
                try:
                    server.wait(timeout=4)
                except subprocess.TimeoutExpired:
                    server.kill()
                    server.wait(timeout=2)


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"REPRODUCTION_FAILED: {type(exc).__name__}: {exc}", file=sys.stderr, flush=True)
        raise SystemExit(30)
