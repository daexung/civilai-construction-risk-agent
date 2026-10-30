"""Bounded, recorded local API validation for paint-coat and indoor-waterproof cases.

This runner sends only the two synthetic test questions and the relevant local
specification excerpts through the Studio provider. It never records credentials
or member data and never starts the API lifespan or connects to its database.
"""
from __future__ import annotations

import contextlib
import io
import json
import os
import re
import sys
import time
import traceback
from datetime import datetime, timezone
from pathlib import Path
from uuid import uuid4

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
RUN_DIR = ROOT / "evals" / "runtime_validation" / "runs"
RESUME_ID = sys.argv[1] if len(sys.argv) > 1 else None
if RESUME_ID and not re.fullmatch(r"studio-single-trade-\d{8}T\d{6}Z-[0-9a-f]{8}", RESUME_ID):
    raise SystemExit("invalid_execution_id")
RUN_ID = RESUME_ID or f"studio-single-trade-{datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%SZ')}-{uuid4().hex[:8]}"
EVENT_PATH = RUN_DIR / f"{RUN_ID}.jsonl"
SUMMARY_PATH = RUN_DIR / f"{RUN_ID}.summary.json"
MODEL = "gemini-3.5-flash-lite"
PROVIDER = "studio"
MAX_CALLS = 10
MAX_USD = 0.20
MAX_OUTPUT_TOKENS = 256
# Official Gemini Developer API standard per-token rates; preflight treats each
# UTF-8 byte as one input token to reserve conservatively before sending.
INPUT_USD_PER_TOKEN = 0.30 / 1_000_000
OUTPUT_USD_PER_TOKEN = 2.50 / 1_000_000

state = {
    "calls": 0,
    "reserved_usd": 0.0,
    "measured_usd": 0.0,
    "input_tokens": 0,
    "output_tokens": 0,
    "logging_failed": False,
    "case_id": None,
    "case_calls": {},
    "prior_measurement_unknown": False,
    "model_events": [],
}


class RecordFailure(RuntimeError):
    pass


def safe_append(event: dict) -> None:
    """Durably append one sanitized JSONL event or block subsequent calls."""
    try:
        RUN_DIR.mkdir(parents=True, exist_ok=True)
        line = json.dumps(event, ensure_ascii=False, separators=(",", ":"), default=str) + "\n"
        with EVENT_PATH.open("a", encoding="utf-8", newline="\n") as stream:
            stream.write(line)
            stream.flush()
            os.fsync(stream.fileno())
    except Exception as exc:
        state["logging_failed"] = True
        raise RecordFailure(type(exc).__name__) from None


def now() -> str:
    return datetime.now(timezone.utc).isoformat()


def model_prompt_fields(prompt: str) -> dict:
    """Persist only the synthetic user query and already selected test conditions."""
    try:
        value = json.loads(prompt)
    except Exception:
        return {"input_json": False}
    session = value.get("session") if isinstance(value, dict) else None
    if not isinstance(session, dict):
        session = {}
    return {
        "input_json": isinstance(value, dict),
        "user_message": value.get("user_message") if isinstance(value, dict) else None,
        "session_conditions": session.get("conditions"),
        "session_quantity": session.get("quantity"),
        "session_work": session.get("work"),
        "session_fields": session.get("fields"),
    }


def condition_action(model_output) -> dict:
    if not isinstance(model_output, dict):
        return {}
    args = model_output.get("args")
    if not isinstance(args, dict):
        return {}
    raw = args.get("values")
    if isinstance(raw, list):
        return {str(item.get("field")): item.get("value")
                for item in raw if isinstance(item, dict) and item.get("field") in {
                    "coat_count", "is_ceiling", "thickness", "is_exterior", "floor_count", "building_height"}}
    if isinstance(raw, dict):
        return {name: item.get("value") for name, item in raw.items()
                if name in {"coat_count", "is_ceiling", "thickness", "is_exterior", "floor_count", "building_height"}
                and isinstance(item, dict)}
    return {}


def guarded_generate(prompt, system, **kwargs):
    if state["logging_failed"]:
        raise RecordFailure("prior_record_write_failed")
    if state["calls"] >= MAX_CALLS:
        safe_append({"event": "preflight_rejected", "execution_id": RUN_ID, "case_id": state["case_id"],
                     "timestamp": now(), "reason": "call_limit", "call_limit": MAX_CALLS})
        raise RecordFailure("call_limit")
    if state["case_calls"].get(state["case_id"], 0) >= 3:
        safe_append({"event": "preflight_rejected", "execution_id": RUN_ID, "case_id": state["case_id"],
                     "timestamp": now(), "reason": "case_call_limit", "case_call_limit": 3})
        raise RecordFailure("case_call_limit")
    input_bytes = len((prompt + system).encode("utf-8"))
    if input_bytes > 50_000:
        safe_append({"event": "preflight_rejected", "execution_id": RUN_ID, "case_id": state["case_id"],
                     "timestamp": now(), "reason": "prompt_size_limit", "input_bytes": input_bytes})
        raise RecordFailure("prompt_size_limit")
    reserved = input_bytes * INPUT_USD_PER_TOKEN + MAX_OUTPUT_TOKENS * OUTPUT_USD_PER_TOKEN
    if state["reserved_usd"] + reserved > MAX_USD:
        safe_append({"event": "preflight_rejected", "execution_id": RUN_ID, "case_id": state["case_id"],
                     "timestamp": now(), "reason": "estimated_cost_limit", "reserved_usd": reserved,
                     "cumulative_reserved_usd": state["reserved_usd"], "cost_limit_usd": MAX_USD})
        raise RecordFailure("estimated_cost_limit")

    attempt = state["calls"] + 1
    state["calls"] = attempt
    state["case_calls"][state["case_id"]] = state["case_calls"].get(state["case_id"], 0) + 1
    state["reserved_usd"] += reserved
    input_fields = model_prompt_fields(prompt)
    start_record = {
        "event": "attempt_start", "execution_id": RUN_ID, "provider": PROVIDER, "model": MODEL,
        "timestamp": now(), "attempt": attempt, "case_id": state["case_id"],
        "input_bytes": input_bytes, "output_token_limit": MAX_OUTPUT_TOKENS,
        "reserved_usd": round(reserved, 9), "input_fields": input_fields,
    }
    safe_append(start_record)

    from backend.agent.tools.llm import client as llm
    original_generate = state["original_generate"]
    kwargs["max_attempts"] = 1  # Disable all automatic retries for the actual provider request.
    kwargs["retry_delays"] = []
    kwargs["model"] = MODEL
    started = time.monotonic()
    try:
        response = original_generate(prompt, system, **kwargs)
    except BaseException as exc:
        from backend.agent.tools.llm import client as llm
        key = llm._env_value("GEMINI_API_KEY") or ""
        project = llm._env_value("GOOGLE_CLOUD_PROJECT") or ""
        record = {"event": "attempt_error", "execution_id": RUN_ID, "provider": PROVIDER,
                  "model": MODEL, "timestamp": now(), "attempt": attempt, "case_id": state["case_id"],
                  "error_type": type(exc).__name__, "elapsed_ms": round((time.monotonic() - started) * 1000),
                  "safe_error": llm._safe_error(exc, [key, project]), "automatic_retry": False}
        safe_append(record)
        state["model_events"].append(record)
        raise

    usage = response.usage or {}
    input_tokens = usage.get("prompt_token_count")
    output_tokens = (usage.get("candidates_token_count") or 0) + (usage.get("thoughts_token_count") or 0)
    measured = None
    if isinstance(input_tokens, int):
        measured = input_tokens * INPUT_USD_PER_TOKEN + output_tokens * OUTPUT_USD_PER_TOKEN
        state["input_tokens"] += input_tokens
        state["output_tokens"] += output_tokens
        state["measured_usd"] += measured
    raw_text = response.text or ""
    try:
        decoded = json.loads(raw_text)
    except Exception:
        decoded = None
    record = {
        "event": "attempt_response", "execution_id": RUN_ID, "provider": PROVIDER, "model": MODEL,
        "timestamp": now(), "attempt": attempt, "case_id": state["case_id"],
        "elapsed_ms": round((time.monotonic() - started) * 1000),
        "raw_response": raw_text,
        "token_usage": usage,
        "estimated_usd_from_usage": round(measured, 9) if measured is not None else None,
        "response_condition_fields": condition_action(decoded),
        "response_json": decoded,
        "automatic_retry": False,
    }
    safe_append(record)
    state["model_events"].append(record)
    if state["measured_usd"] > MAX_USD:
        safe_append({"event": "measured_cost_limit_crossed", "execution_id": RUN_ID,
                     "timestamp": now(), "measured_usd": state["measured_usd"], "cost_limit_usd": MAX_USD})
        raise RecordFailure("measured_cost_limit_crossed")
    return response


def local_processing(*args, **kwargs):
    @contextlib.contextmanager
    def context():
        yield (None, "isolated-local-validation", False)
    return context()


def reduced_response(data: dict) -> dict:
    keep_fields = {"coat_count", "is_ceiling", "thickness", "is_exterior", "floor_count", "building_height"}
    return {
        "status": data.get("status"), "estimate_current": data.get("estimate_current"),
        "inputs": {entry.get("name"): entry.get("value") for entry in (data.get("inputs") or [])
                   if entry.get("name") in keep_fields},
        "conditions": data.get("conditions"), "result": data.get("result"),
        "priced": data.get("priced"), "statement": data.get("statement"),
        "questions": [{"name": q.get("name"), "ask": q.get("ask")} for q in (data.get("questions") or [])],
        "llm_info": data.get("llm_info"), "thread_id": data.get("thread_id"),
        "answer_source": data.get("answer_source"), "message": data.get("message"),
    }


def stored_state(thread_id: str) -> dict:
    from backend.api import main as api
    from backend.agent.estimate import dialogue, tools
    state_value = dialogue.load(api.DIALOGUE, thread_id)
    if not state_value or not state_value.get("session"):
        return {}
    item = tools._item(state_value["session"])
    return {"spec_id": item.get("selected_spec_id"), "conditions": item.get("conditions"),
            "quantity": item.get("quantity"), "computed_result": item.get("computed_result"),
            "priced_result": item.get("priced_result"),
            "current_estimate": bool(tools.current_estimate(state_value["session"]))}


def excel_check(response: dict) -> dict:
    from backend.api import dialogue_service, main as api
    from backend.api.tables import build_xlsx
    from openpyxl import load_workbook
    try:
        export_response = dialogue_service.export_response(api.DIALOGUE, response["thread_id"])
    except Exception as exc:
        return {"export_error_type": type(exc).__name__}
    workbook = load_workbook(io.BytesIO(build_xlsx(export_response)), data_only=False)
    coat_cells = []
    contract_cells = []
    unpriced_words = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            values = [cell.value for cell in row]
            for value in values:
                if isinstance(value, str) and ("회 칠" in value or ("×" in value and "회" in value)):
                    coat_cells.append(value)
                if isinstance(value, str) and "미산정" in value:
                    unpriced_words.append(value)
            for index, value in enumerate(values[:-2]):
                if isinstance(value, str) and "도급액" in value:
                    contract_cells.append(values[index + 2])
    statement = response.get("statement") or {}
    contract_amount = (statement.get("totals") or {}).get("contract_amount")
    return {"http_status": 200, "sheets": workbook.sheetnames,
            "coat_cells": coat_cells[:20], "xlsx_contract_amounts": contract_cells,
            "response_contract_amount": contract_amount,
            "contract_amount_match": contract_amount is not None and any(
                str(value).replace(",", "") == str(contract_amount).replace(",", "") for value in contract_cells),
            "xlsx_unpriced_mentions": unpriced_words[:10]}


def run_case(case_id: str, message: str, answer_values: dict) -> dict:
    from backend.api import main as api
    from datetime import date
    thread_id = uuid4().hex
    state["case_id"] = case_id
    safe_append({"event": "case_start", "execution_id": RUN_ID, "timestamp": now(), "case_id": case_id,
                 "test_question": message})
    request = api.ChatRequest(thread_id=thread_id, message=message, basis_date=date(2026, 10, 1))
    data = api._dialogue_chat(request)
    turns = 1
    unanswered = []
    while data.get("questions") and turns <= 5:
        question = data["questions"][0]
        name = question["name"]
        if name not in answer_values:
            unanswered.append(name)
            break
        request = api.ChatRequest(thread_id=thread_id, answers={name: answer_values[name]},
            refs={name: question["ref"]}, basis_date=date(2026, 10, 1))
        data = api._dialogue_chat(request)
        turns += 1
    saved = stored_state(data["thread_id"])
    exported = excel_check(data) if data.get("estimate_current") else {"http_status": 404}
    event = {"event": "case_result", "execution_id": RUN_ID, "timestamp": now(), "case_id": case_id,
             "turns": turns, "response": reduced_response(data), "server_saved": saved,
             "excel": exported, "unanswered_questions": unanswered}
    safe_append(event)
    return event


def main() -> int:
    RUN_DIR.mkdir(parents=True, exist_ok=True)
    previous_summary = None
    if RESUME_ID:
        if not EVENT_PATH.is_file() or not SUMMARY_PATH.is_file():
            raise SystemExit("resume_record_missing")
        existing = [json.loads(line) for line in EVENT_PATH.read_text(encoding="utf-8").splitlines() if line.strip()]
        previous_summary = json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))
        start = next((row for row in existing if row.get("event") == "run_start"), {})
        if (start.get("provider"), start.get("model"), start.get("call_limit"), start.get("cost_limit_usd")) != (
                PROVIDER, MODEL, MAX_CALLS, MAX_USD):
            raise SystemExit("resume_configuration_mismatch")
        state["calls"] = sum(1 for row in existing if row.get("event") == "attempt_start")
        state["reserved_usd"] = float(previous_summary.get("reserved_usd") or 0)
        state["prior_measurement_unknown"] = previous_summary.get("measured_usd") is None
        for row in existing:
            if row.get("event") == "attempt_start" and row.get("case_id"):
                name = row["case_id"]
                state["case_calls"][name] = state["case_calls"].get(name, 0) + 1
        if state["calls"] >= MAX_CALLS or state["reserved_usd"] >= MAX_USD:
            raise SystemExit("resume_budget_exhausted")
        safe_append({"event": "run_resume", "execution_id": RUN_ID, "timestamp": now(),
                     "calls_already_started": state["calls"], "reserved_usd_before_resume": state["reserved_usd"],
                     "automatic_retries": False})
    else:
        safe_append({"event": "run_start", "execution_id": RUN_ID, "timestamp": now(),
                     "provider": PROVIDER, "model": MODEL, "call_limit": MAX_CALLS,
                     "cost_limit_usd": MAX_USD, "automatic_retries": False,
                     "max_output_tokens_per_call": MAX_OUTPUT_TOKENS,
                     "network_mode": "Google AI Studio generation only; retrieval is local BM25",
                     "database_persistence": False})

    # The override exists only in this Python process; no shell/profile or .env file is edited.
    os.environ.update({"AGENT_MODE": "tools", "AGENT_LLM": "on", "LLM_PROVIDER": PROVIDER,
                       "LLM_MODEL": MODEL, "AGENT_OFFLINE": "1", "KNOWLEDGE_RETRIEVAL": "legacy",
                       "LLM_MAX_OUTPUT_TOKENS": str(MAX_OUTPUT_TOKENS), "APP_ENVIRONMENT": "local"})
    from unittest.mock import patch
    from backend.agent.tools.llm import client as llm
    from backend.api import main as api, usage_limits, answer_runs
    from backend.shared import answer_trace
    from backend.agent.estimate import dialogue as dlg
    safe_append({"event": "preflight_stage", "execution_id": RUN_ID, "timestamp": now(),
                 "stage": "imports_complete"})

    if llm.provider_name() != PROVIDER or llm.model_name() != MODEL:
        safe_append({"event": "preflight_error", "execution_id": RUN_ID, "timestamp": now(),
                     "reason": "provider_model_mismatch"})
        return 2
    if not llm._env_value("GEMINI_API_KEY"):
        safe_append({"event": "preflight_error", "execution_id": RUN_ID, "timestamp": now(),
                     "reason": "studio_credential_missing"})
        return 2
    if api._LIFESPAN_ACTIVE:
        safe_append({"event": "preflight_error", "execution_id": RUN_ID, "timestamp": now(),
                     "reason": "api_lifespan_active"})
        return 2
    safe_append({"event": "preflight_stage", "execution_id": RUN_ID, "timestamp": now(),
                 "stage": "provider_model_credential_and_no_lifespan_verified"})
    state["original_generate"] = llm._generate

    with patch.object(llm, "_generate", side_effect=guarded_generate), \
         patch.object(usage_limits, "processing", local_processing), \
         patch.object(usage_limits, "consume", return_value={}), \
         patch.object(usage_limits, "status", return_value={}), \
         patch.object(answer_trace, "persist_start", lambda: None), \
         patch.object(answer_trace, "finalize", lambda *args, **kwargs: None):
        safe_append({"event": "preflight_stage", "execution_id": RUN_ID, "timestamp": now(),
                     "stage": "in_process_dialogue_handler_ready"})
        try:
            cases = [
                run_case("paint_2_coats_wall_1sqm",
                    "수성페인트 롤러칠 2회, 1㎡, 벽면(천장 아님). 비용을 계산해줘.",
                    {"coat_count": 2, "is_ceiling": False, "quantity": "1", "volume": "1"}),
                run_case("waterproof_indoor_10mm_1sqm",
                    "실내 바닥 방수모르타르 10㎜, 1㎡. 비용을 계산해줘.",
                    {"thickness": "10㎜이하", "is_exterior": False, "quantity": "1", "volume": "1"}),
            ]
        except BaseException as exc:
            safe_append({"event": "run_error", "execution_id": RUN_ID, "timestamp": now(),
                         "error_type": type(exc).__name__, "automatic_retries": False})
            raise

    current_cases = cases
    all_cases = (previous_summary.get("cases") or []) + current_cases if previous_summary else current_cases
    final = {"execution_id": RUN_ID, "event_path": str(EVENT_PATH), "summary_path": str(SUMMARY_PATH),
             "provider": PROVIDER, "model": MODEL, "calls": state["calls"], "call_limit": MAX_CALLS,
             "reserved_usd": round(state["reserved_usd"], 9),
             "measured_usd": round(state["measured_usd"], 9)
                 if state["input_tokens"] and not state["prior_measurement_unknown"] else None,
             "measured_usd_this_resume": round(state["measured_usd"], 9) if state["input_tokens"] else None,
             "input_tokens": state["input_tokens"], "output_tokens": state["output_tokens"],
             "cost_status": "unknown" if state["prior_measurement_unknown"] or not state["input_tokens"]
                 else "usage_based_estimate",
             "cases": all_cases, "cases_this_segment": current_cases, "automatic_retries": False}
    try:
        temporary = SUMMARY_PATH.with_suffix(".tmp")
        with temporary.open("w", encoding="utf-8", newline="\n") as stream:
            json.dump(final, stream, ensure_ascii=False, indent=2, default=str)
            stream.flush()
            os.fsync(stream.fileno())
        temporary.replace(SUMMARY_PATH)
        safe_append({"event": "run_finish", "execution_id": RUN_ID, "timestamp": now(),
                     "calls": state["calls"], "reserved_usd": final["reserved_usd"],
                     "measured_usd": final["measured_usd"], "input_tokens": state["input_tokens"],
                     "output_tokens": state["output_tokens"], "summary_path": str(SUMMARY_PATH)})
    except Exception as exc:
        state["logging_failed"] = True
        print(json.dumps({"execution_id": RUN_ID, "summary_write_error": type(exc).__name__}), flush=True)
        return 3
    print(json.dumps({"execution_id": RUN_ID, "calls": state["calls"],
        "reserved_usd": final["reserved_usd"], "measured_usd": final["measured_usd"],
        "case_statuses": [{"case_id": c.get("case_id"),
                           "status": (c.get("response") or {}).get("status"),
                           "estimate_current": (c.get("response") or {}).get("estimate_current")}
                          for c in current_cases],
        "event_path": str(EVENT_PATH), "summary_path": str(SUMMARY_PATH)},
        ensure_ascii=False, default=str), flush=True)
    return 0


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    try:
        sys.exit(main())
    except SystemExit:
        raise
    except BaseException as exc:
        # Do not print arbitrary provider exception text; it may contain credentials.
        try:
            safe_append({"event": "run_fatal_error", "execution_id": RUN_ID, "timestamp": now(),
                         "error_type": type(exc).__name__, "calls_started": state["calls"],
                         "automatic_retries": False})
        except Exception:
            pass
        print(json.dumps({"execution_id": RUN_ID, "fatal_error_type": type(exc).__name__,
                          "event_path": str(EVENT_PATH)}, ensure_ascii=False), flush=True)
        raise
