"""AGENT_MODE=tools 비회원 대화를 실제 검색 색인(오프라인)과 규칙 정책(LLM 끔)으로 검사한다.

check_dialogue.py(고정 후보·대본 LLM)와 달리 공종 후보를 실제 검색이 고른다. 검색이 기대 공종을 못 찾으면
그 항목은 FAIL로 남기고, 나머지 흐름은 고정 후보로 다시 돌려 "고정 후보 대체"로 따로 표시한다(성공으로 세지 않는다).
"""

from __future__ import annotations

import io
import os
import sys
from contextlib import nullcontext
from pathlib import Path
from unittest.mock import patch

os.environ.update(AGENT_OFFLINE="1", AGENT_LLM="off", AGENT_MODE="tools")
# 서비스 기본 색인(index_config.json + 서비스 부문)으로 검사한다. 6장 전용 평가 색인(INDEX_CONFIG)은 경쟁 공종이 적어
# 6-1-4를 바로 확정하므로 공종 질문이 나오지 않는다(검색 오류가 아님).
os.environ.pop("INDEX_CONFIG", None)
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

import backend.api.main as api_main  # noqa: E402
from backend.agent.estimate import dialogue, tools  # noqa: E402
from backend.agent.nodes.retrieve import get_search  # noqa: E402
from evals.check_dialogue import HITS, PUMP, SECRET  # noqa: E402

EXPECTED = "6-1-4"


def run(http, label: str) -> list[tuple[str, bool]]:
    checks = []

    def post(body: dict) -> dict:
        result = http.post("/api/chat", json=body)
        assert result.status_code == 200, (result.status_code, result.text[:300])
        return result.json()

    def answer(previous: dict, values: dict) -> dict:
        return post({"thread_id": previous["thread_id"], "answers": values,
                     "refs": {q["field"]: q["ref"] for q in dialogue.load(api_main.DIALOGUE, previous["thread_id"])["pending"] if q["field"] in values}})

    t1 = post({"message": "콘크리트 타설 1세제곱미터 품셈 알려줘"})
    thread = t1["thread_id"]
    work = next((q for q in t1["questions"] if q["name"] == "work"), None)
    found = bool(work) and any(EXPECTED in choice for choice in work["choices"] or [])
    checks.append((f"{label} S1 실제 검색 후보에 {EXPECTED} 포함", found))
    if not found:
        print(f"  {label} 후보:", work["choices"] if work else t1["status"],
              "| 검색 방식·경고:", t1.get("search"))
        return checks
    t2 = answer(t1, {"work": EXPECTED})
    checks.append((f"{label} S2 품 조건만 질문", set(q["field"] for q in dialogue.load(api_main.DIALOGUE, thread)["pending"]) == set(PUMP)
                   and [q["name"] for q in t2["questions"]] == ["pump_size"] and t2["questions_remaining"] == len(PUMP)))
    t2b = answer(t2, PUMP)
    checks.append((f"{label} S2 1㎥ 품 결과(물량 1)", t2b["status"] == "COMPUTED"
                   and any(row["name"] == "volume" and row["value"] == "1" for row in t2b["inputs"])))
    t3 = post({"thread_id": thread, "message": "같은 조건으로 100세제곱미터 비용 계산해줘"})
    checks.append((f"{label} S3 가격 조건만 질문", [q["name"] for q in t3["questions"]] == ["concrete_supply"]))
    t4 = answer(t3, {"concrete_supply": "관급"})
    checks.append((f"{label} S4 원가계산서 최신", t4["estimate_current"] and t4["status"] in ("OK", "PARTIAL")))
    t6 = post({"thread_id": thread, "message": "할증 기준은 어디서 나와?"})
    checks.append((f"{label} S6 근거 질문 뒤 견적 유지", t6["estimate_current"]
                   and t6["statement"]["totals"] == t4["statement"]["totals"]))
    t5 = post({"thread_id": thread, "message": "300세제곱미터로 바꿔줘"})
    amount = t5["statement"]["totals"]["contract_amount"]
    book = load_workbook(io.BytesIO(http.get(f"/api/export/{thread}.xlsx").content), data_only=True)
    checks.append((f"{label} S5 재계산 금액이 Excel에", amount != t4["statement"]["totals"]["contract_amount"]
                   and amount in [c.value for s in book for r in s.iter_rows() for c in r]))
    return checks


def main() -> int:
    index, method, warning = get_search()
    print(f"색인: chunks {len(index.chunks)}개, 방식 {method}, 경고 {warning}")
    with patch("backend.api.usage_limits.processing", return_value=nullcontext((None, "fixture", False))), \
            patch("backend.api.usage_limits.status", return_value={}), \
            patch("backend.api.usage_limits.consume", return_value={}), \
            TestClient(api_main.app, headers={"X-Guest-Session": SECRET}) as http:
        checks = run(http, "[실제 검색]")
        searched = checks[0][1]
        if not searched:
            print("  실제 검색 실패 → 아래는 고정 후보 대체 결과(성공으로 세지 않음)")
            with patch.object(tools, "retrieve", return_value={"hits": HITS}):
                fallback = run(http, "[고정 후보 대체]")
            for name, passed in fallback:
                print(f"{'PASS' if passed else 'FAIL'} {name}")
    for name, passed in checks:
        print(f"{'PASS' if passed else 'FAIL'} {name}")
    print(f"통과 {sum(passed for _, passed in checks)} / 전체 {len(checks)} (실제 검색 기준)")
    return 0 if all(passed for _, passed in checks) else 1


if __name__ == "__main__":
    raise SystemExit(main())
