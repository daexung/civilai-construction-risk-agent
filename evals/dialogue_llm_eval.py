"""실제 LLM·실제 검색으로 AGENT_MODE=tools 대화를 평가한다(수동 실행, CI 제외).

- DB는 로컬 Supabase(127.0.0.1:54322)만 쓴다. 운영 DB·운영 API를 쓰지 않는다.
- 키는 --env-file에서 LLM·임베딩 키 3개만 프로세스 안으로 읽는다(파일 복사·출력 없음).
- 호출마다 SDK가 돌려준 토큰 사용량으로 비용을 계산하고, 다음 호출로 상한을 넘을 수 있으면 막는다.
- 실패는 고정 후보·대본으로 바꾸지 않고 그대로 기록한다.

사용: python evals/dialogue_llm_eval.py --env-file <원래 폴더의 .env> --cap-usd 0.10
     --focus: 준비 턴(공종·조건·100㎥·관급)은 LLM 없이(규칙 경로, 검색 임베딩만 유료) 돌리고 FOCUS 3턴만 LLM으로 평가
"""

from __future__ import annotations

import argparse
import io
import json
import os
import sys
import time
from datetime import date
from pathlib import Path
from unittest.mock import patch
from uuid import uuid4

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
ALLOWED_KEYS = ("VERTEX_API_KEY", "GOOGLE_CLOUD_PROJECT", "GEMINI_API_KEY")
# Gemini API 가격 페이지(2026-10-07 UTC 갱신) 표준 유료 단가, 100만 토큰당 달러.
PRICE = {"input": 0.30, "output": 2.50, "embedding": 0.20, "checked": "2026-10-07",
         "source": "https://ai.google.dev/gemini-api/docs/pricing"}
FOCUS = ("300세제곱미터로 바꿔줘", "콘크리트공은 1세제곱미터당 몇 명이야?", "펌프차 붐을 25m로 바꿔줘")
NEXT_CALL_RESERVE_USD = 0.004  # 입력 4천·출력 1천 토큰 상한 가정


def load_keys(env_file: Path) -> None:
    for line in env_file.read_text(encoding="utf-8").splitlines():
        name, sep, value = line.partition("=")
        if sep and name.strip() in ALLOWED_KEYS and value.strip():
            os.environ[name.strip()] = value.strip().strip('"').strip("'")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--env-file", required=True)
    parser.add_argument("--cap-usd", type=float, default=0.10)
    parser.add_argument("--focus", action="store_true")
    args = parser.parse_args()
    os.environ.update(AGENT_MODE="tools", AGENT_LLM="on",
                      CHAT_DATABASE_URL="postgresql://postgres:postgres@127.0.0.1:54322/postgres",
                      QUOTA_HASH_SECRET="local-llm-eval-only")
    os.environ.pop("AGENT_OFFLINE", None)
    load_keys(Path(args.env_file))

    from fastapi import HTTPException
    from fastapi.testclient import TestClient
    from openpyxl import load_workbook

    import backend.api.main as api_main
    from backend.agent.estimate import dialogue, tools
    from backend.agent.tools.llm import client
    from backend.agent.tools.search import vector
    from backend.api import chat_storage
    from backend.agent.nodes.retrieve import get_search

    ledger = {"llm_calls": [], "embedding_calls": 0, "embedding_chars": 0, "blocked": 0}

    def spent() -> float:
        llm = sum(call["usd"] for call in ledger["llm_calls"])
        # 임베딩은 SDK 사용량이 없어 글자 수를 토큰 수 상한으로 본다(한글 1자 ≤ 1토큰 가정).
        return llm + ledger["embedding_chars"] * PRICE["embedding"] / 1e6

    real_get_client = client._get_client

    def metered_client(provider, key):
        real = real_get_client(provider, key)

        class Models:
            def generate_content(self, model, contents, config):
                if spent() + NEXT_CALL_RESERVE_USD > args.cap_usd:
                    ledger["blocked"] += 1
                    raise client.LLMUnavailable("평가 비용 상한 도달")
                response = real.models.generate_content(model=model, contents=contents, config=config)
                usage = response.usage_metadata
                prompt = usage.prompt_token_count or 0
                output = (usage.candidates_token_count or 0) + (getattr(usage, "thoughts_token_count", 0) or 0)
                ledger["llm_calls"].append({"model": model, "input_tokens": prompt, "output_tokens": output,
                                            "usd": prompt * PRICE["input"] / 1e6 + output * PRICE["output"] / 1e6})
                return response

        class Wrapped:
            models = Models()
        return Wrapped()

    real_embed = vector.embed_texts

    def metered_embed(client_obj, inputs, **kwargs):
        ledger["embedding_calls"] += 1
        ledger["embedding_chars"] += sum(len(str(text)) for text in inputs)
        return real_embed(client_obj, inputs, **kwargs)

    member = str(uuid4())
    conversation = str(uuid4())
    turns: list[dict] = []
    with chat_storage.connection() as conn:
        conn.execute("INSERT INTO auth.users(id,aud,role,email) VALUES (%s,'authenticated','authenticated',%s)",
                     (member, f"llm-eval-{member}@example.invalid"))
    model = client.model_name()

    def identity(token):
        if token is None:
            return None
        if token == "Bearer eval":
            return member
        raise HTTPException(401, "invalid")

    try:
        with patch.object(client, "_get_client", metered_client), patch.object(vector, "embed_texts", metered_embed), \
                patch.object(chat_storage, "identity", side_effect=identity), \
                TestClient(api_main.app, raise_server_exceptions=False) as http:
            headers = {"Authorization": "Bearer eval"}
            last: dict = {}

            def send(label: str, body: dict) -> dict:
                nonlocal last
                started = time.monotonic()
                result = http.post("/api/chat", headers=headers, json={
                    "conversation_id": conversation, "request_id": str(uuid4()),
                    **({"thread_id": conversation} if turns else {}), **body})
                payload = result.json() if result.headers.get("content-type", "").startswith("application/json") else {}
                with chat_storage.connection() as conn:
                    from langgraph.checkpoint.postgres import PostgresSaver
                    state = dialogue.load(dialogue.build_dialogue_graph(PostgresSaver(conn)), conversation) or {}
                item = tools._item(state["session"]) if state.get("session") else {}
                turns.append({
                    "turn": len(turns) + 1, "user": label, "http": result.status_code, "llm": os.environ["AGENT_LLM"],
                    "seconds": round(time.monotonic() - started, 1),
                    "status": payload.get("status"), "answer": payload.get("answer") or payload.get("message"),
                    "answer_source": payload.get("answer_source"), "tools": (state.get("turn") or {}).get("tools"),
                    "trace": (state.get("turn") or {}).get("trace"),
                    "llm_info": payload.get("llm_info"),
                    "questions": [{"name": q["name"], "choices": q.get("choices")} for q in payload.get("questions") or []],
                    "work": (item.get("selection") or {}).get("section"),
                    "quantity": (item.get("quantity") or {}).get("value"),
                    "pump_size": (item.get("conditions") or {}).get("pump_size"),
                    "contract_amount": ((payload.get("statement") or {}).get("totals") or {}).get("contract_amount"),
                    "work_days": ((payload.get("result") or {}).get("work_days") or {}).get("value"),
                    "work_days_display": ((payload.get("result") or {}).get("work_days") or {}).get("display"),
                    "estimate_current": payload.get("estimate_current"),
                    "embedding_calls_so_far": ledger["embedding_calls"],
                    "search_fallback": getattr(get_search()[0], "error", None),
                    "usd_so_far": round(spent(), 5)})
                last = payload
                return payload

            def answer(label: str, wanted: dict) -> dict:
                questions = {q["name"]: q for q in last.get("questions") or []}
                values = {name: value for name, value in wanted.items() if name in questions}
                if not values:
                    turns.append({"turn": len(turns) + 1, "user": label, "skipped": "서버가 이 질문을 내지 않음",
                                  "asked": list(questions)})
                    return last
                if "work" in values:
                    choice = next((c for c in questions["work"]["choices"] or [] if values["work"] in c), None)
                    if choice is None:
                        turns.append({"turn": len(turns) + 1, "user": label, "skipped": "후보에 6-1-4 없음",
                                      "choices": questions["work"]["choices"]})
                        return last
                return send(label, {"answers": values, "refs": {name: questions[name]["ref"] for name in values}})

            if args.focus:
                os.environ["AGENT_LLM"] = "off"
            send("콘크리트 타설 1세제곱미터 품셈 알려줘", {"message": "콘크리트 타설 1세제곱미터 품셈 알려줘"})
            answer("(버튼) 공종: 6-1-4 펌프차 타설", {"work": "6-1-4"})
            answer("(버튼) 품 조건: 32m·철근·15㎝·Type-Ⅱ·Type-Ⅱ·붐·진동기 사용·재셋팅 없음",
                   {"pump_size": "32m", "structure": "철근", "slump_band": "15㎝", "facility_type": "Type-Ⅱ",
                    "site_type": "Type-Ⅱ", "placement": "붐", "vibrator_used": True, "reset_status": "없음"})
            for message in ("같은 조건으로 100세제곱미터 비용 계산해줘",):
                send(message, {"message": message})
            answer("(버튼) 관급", {"concrete_supply": "관급"})
            messages = ("할증 기준은 어디서 나와?", "300세제곱미터로 바꿔줘", "콘크리트공은 1세제곱미터당 몇 명이야?",
                        "레미콘을 사급으로 바꾸면 얼마야?", "레미콘 단가는 세제곱미터당 9만원이야")
            if args.focus:
                # 준비가 끝난 상태(32m·100㎥·최신 견적)에서만 유료 턴을 보낸다. 아니면 LLM을 켜지 않고 끝낸다.
                ready = (str(turns[-1].get("quantity")) == "100" and turns[-1].get("estimate_current")
                         and turns[-1].get("pump_size") == "32m")
                messages = FOCUS if ready else ()
                if not ready:
                    turns.append({"turn": len(turns) + 1, "skipped": "준비 턴이 견적까지 가지 않아 유료 턴을 보내지 않음"})
                os.environ["AGENT_LLM"] = "on"
            for message in messages:
                send(message, {"message": message})
            export = http.get(f"/api/export/{conversation}.xlsx", headers=headers)
            excel_amounts = []
            if export.status_code == 200:
                book = load_workbook(io.BytesIO(export.content), data_only=True)
                excel_amounts = [c.value for s in book for r in s.iter_rows() for c in r if isinstance(c.value, int)]
    finally:
        with chat_storage.connection() as conn:
            conn.execute("DELETE FROM public.conversations WHERE user_id=%s", (member,))
            conn.execute("DELETE FROM auth.users WHERE id=%s", (member,))
    final_amount = next((t["contract_amount"] for t in reversed(turns) if t.get("contract_amount")), None)
    index, method, warning = get_search()
    report = {"model": model, "price": PRICE, "cap_usd": args.cap_usd, "focus": args.focus,
              "search": {"method": method, "warning": warning, "last_used": getattr(index, "used", None),
                         "last_error": getattr(index, "error", None)}, "turns": turns,
              "excel": {"status": export.status_code, "contains_latest_amount": final_amount in excel_amounts,
                        "latest_amount": final_amount, "matching_cells": excel_amounts.count(final_amount)},
              "cost": {"usd_total": round(spent(), 6), "llm_calls": len(ledger["llm_calls"]),
                       "input_tokens": sum(c["input_tokens"] for c in ledger["llm_calls"]),
                       "output_tokens": sum(c["output_tokens"] for c in ledger["llm_calls"]),
                       "embedding_calls": ledger["embedding_calls"], "blocked_by_cap": ledger["blocked"]}}
    out = ROOT / "evals/results" / f"dialogue_llm_eval_{date.today().isoformat()}_{int(time.time())}.json"
    out.write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2, default=str))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
