"""고정 시나리오를 한 버전의 코드로 실행해 비교용 결과를 JSON으로 남긴다.

check_baseline.py가 버전마다 별도 프로세스로 호출한다. --root의 backend만 import하므로
main 기준 환경과 현재 작업 트리의 import 경로·캐시·환경설정이 섞이지 않는다.
운영 DB·유료 API는 쓰지 않는다(오프라인 단어 검색, LLM 꺼짐, 사용량 가짜 값).

사용: python evals/baseline/run_scenarios.py --root <코드 폴더> --scenarios <json> --out <json>
      [--capture-checkpoints <파일>] [--resume-checkpoints <파일>]
"""

from __future__ import annotations

import argparse
import base64
import hashlib
import io
import json
import os
import sys
from contextlib import nullcontext
from pathlib import Path
from unittest.mock import patch

BASIS_DATE = "2026-10-06"
DATA_FILES = ["data/processed/chunks.all.jsonl", "data/processed/chunks.jsonl", "data/processed/page_map.json",
              "data/rates/overhead_rates.json", "backend/agent/tools/search/index_config.json"]


def _fixed_environment() -> dict:
    for name in ("INDEX_CONFIG", "CHAT_DATABASE_URL", "GEMINI_API_KEY", "VERTEX_API_KEY", "SUPABASE_URL",
                 "SUPABASE_SECRET_KEY", "QUOTA_HASH_SECRET", "APP_ENV"):
        os.environ.pop(name, None)
    os.environ.update({"AGENT_OFFLINE": "1", "AGENT_LLM": "off", "PYTHONHASHSEED": "0"})
    return {"AGENT_OFFLINE": "1", "AGENT_LLM": "off", "INDEX_CONFIG": "(기본 index_config.json)",
            "basis_date": BASIS_DATE, "llm": "off", "database": "none(사용량 가짜 값)"}


def _citation_ids(items) -> list:
    return [c.get("internal_id") or c.get("chunk_id") for c in items or []]


def _normalize(response: dict) -> dict:
    """비교에 쓰는 값만 남긴다. 매번 바뀌는 id·시간은 뺀다."""
    result = response.get("result") or {}
    priced = response.get("priced") or {}
    statement = response.get("statement") or {}
    qa = response.get("qa") or {}
    return {
        "status": response.get("status"), "route": response.get("route"), "route_source": response.get("route_source"),
        "message": response.get("message"), "answer": response.get("answer"),
        "work": {k: (response.get("work") or {}).get(k) for k in ("spec_id", "title", "confirmed")} if response.get("work") else None,
        "questions": [{k: q.get(k) for k in ("name", "ask", "choices", "default", "reason", "optional")}
                      for q in response.get("questions", [])],
        "inputs": [{k: i.get(k) for k in ("name", "value", "source")} for i in response.get("inputs", [])],
        "conditions": [{k: c.get(k) for k in ("name", "value", "source")} for c in response.get("conditions", [])],
        "evidence": [e.get("section") for e in response.get("evidence", [])],
        "qa": {"conclusion": qa.get("conclusion"), "sections": [c.get("section") for c in qa.get("comparisons", [])],
               "citations": [c.get("chunk_id") for c in qa.get("citations", [])]} if qa else None,
        "result": {
            "unit_basis": result.get("unit_basis"),
            "unit_lines": [{k: line.get(k) for k in ("kind", "name", "unit", "applied", "exact")}
                           for line in result.get("unit_lines", [])],
            "lines": [{k: line.get(k) for k in ("kind", "name", "value", "unit", "crew")} for line in result.get("lines", [])],
            "daily_volume": (result.get("daily_volume") or {}).get("value"),
            "work_days": (result.get("work_days") or {}).get("value"),
            "not_calculated": [n.get("item") for n in result.get("not_calculated", [])],
            "citations": sorted({cid for line in result.get("unit_lines", []) for cid in _citation_ids(line.get("citations"))}),
        } if result else None,
        "priced": {
            "status": priced.get("status"), "partial": priced.get("partial"),
            "unit_prices": priced.get("unit_prices"), "subtotals": priced.get("subtotals"), "total": priced.get("total"),
            "reference_amounts": priced.get("reference_amounts"),
            "lines": [{k: line.get(k) for k in ("name", "unit_price", "amount", "rate_code")} for line in priced.get("lines", [])],
            "unpriced": [u.get("name") for u in priced.get("unpriced", [])],
            "excluded": [u.get("name") for u in priced.get("excluded", [])],
            "rate_version": (priced.get("rate_version") or {}).get("id"),
        } if priced else None,
        "statement": {
            "status": statement.get("status"), "totals": statement.get("totals"),
            "lines": [[line.get("name"), line.get("amount"), line.get("status")] for line in statement.get("lines", [])],
        } if statement else None,
        "bill": (response.get("tables") or {}).get("bill"),
        "items": response.get("items") or [],  # 현재 브랜치에만 있는 필드. 단일 공종에서는 비어 있어야 한다.
    }


def _auto_answers(response: dict, answers: dict) -> dict:
    """두 버전에 같은 규칙으로 답한다: 시나리오에 적힌 값 → 질문 기본값 → 첫 선택지."""
    chosen = {}
    for question in response.get("questions", []):
        name, choices = question["name"], question.get("choices")
        options = choices if isinstance(choices, list) else []
        if name == "work":
            wanted = answers.get("work")
            match = next((c for c in options if wanted and wanted in str(c)), None)
            chosen[name] = match or wanted or question.get("default") or (options[0] if options else "")
        elif name in answers:
            chosen[name] = answers[name]
        elif question.get("default") is not None:
            chosen[name] = question["default"]
        else:
            chosen[name] = options[0] if options else answers.get("_free", "10")
    return chosen


def _excel(content: bytes) -> dict:
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(content))
    return {name: [[cell for cell in row if cell not in (None, "")]
                   for row in book[name].iter_rows(values_only=True) if any(c not in (None, "") for c in row)]
            for name in book.sheetnames}


def run_scenario(client, scenario: dict) -> list[dict]:
    headers = {"X-Guest-Session": "baseline-" + hashlib.sha256(scenario["id"].encode()).hexdigest()[:40]}
    answers = scenario.get("answers", {})
    steps, thread = [], None

    def post(body: dict) -> dict:
        nonlocal thread
        body = {**body, "basis_date": BASIS_DATE}
        if thread:
            body["thread_id"] = thread
        response = client.post("/api/chat", json=body, headers=headers)
        data = response.json() if response.status_code == 200 else {"http_status": response.status_code,
                                                                      "detail": response.json().get("detail")}
        thread = data.get("thread_id", thread)
        return data

    for op in scenario["steps"]:
        kind, value = next(iter(op.items()))
        if kind == "message":
            data = post({"message": value})
            steps.append({"op": op, "response": _normalize(data) if "http_status" not in data else data})
        elif kind in ("until_done", "answer_once"):
            for _ in range(6 if kind == "until_done" else 1):
                last = steps[-1]["raw"] if steps and "raw" in steps[-1] else None
                if last is None or last.get("status") != "MISSING_INFO":
                    break
                sent = _auto_answers(last, answers)
                data = post({"answers": sent})
                steps.append({"op": {"answers": sent}, "response": _normalize(data) if "http_status" not in data else data,
                              "raw": data})
        elif kind == "text":
            data = post({"message": value})
            steps.append({"op": op, "response": _normalize(data) if "http_status" not in data else data})
        elif kind == "restart":
            data = post({"message": value, "restart": True})
            steps.append({"op": op, "response": _normalize(data) if "http_status" not in data else data})
        elif kind == "conditions":
            data = post({"conditions": value})
            steps.append({"op": op, "response": _normalize(data) if "http_status" not in data else data})
        elif kind == "export":
            response = client.get(f"/api/export/{thread}.xlsx", headers=headers)
            steps.append({"op": op, "excel": _excel(response.content) if response.status_code == 200
                          else {"http_status": response.status_code}})
            continue
        if kind in ("message", "text", "restart", "conditions"):
            steps[-1]["raw"] = data
    for step in steps:
        step.pop("raw", None)
    return steps


def _checkpoint_io(capture: Path | None, resume: Path | None) -> dict:
    """합성 질문으로 만든 체크포인트를 저장(main)하거나, 저장된 체크포인트로 재개(현재 버전)한다."""
    from collections import defaultdict
    from langgraph.checkpoint.memory import MemorySaver
    from langgraph.types import Command
    from backend.agent.graph import build_graph
    from backend.agent.state import new_state

    query = "철근콘크리트 벽체 260㎥ 32m 펌프차로 타설 비용"   # 합성 테스트 질문(실사용자 데이터 아님)
    answers = {"work": "6-1-4", "pump_size": "32m", "slump_band": "15㎝", "facility_type": "Type-Ⅱ",
               "site_type": "Type-Ⅱ", "placement": "붐", "vibrator_used": True, "reset_status": "없음",
               "concrete_supply": "관급", "work_category": "기타 토목공사", "duration": "1~6개월",
               "contractor_type": "종합건설업", "project_scale": "이 견적만"}
    config = {"configurable": {"thread_id": "legacy-synthetic"}}
    out = {}
    if capture:
        saver = MemorySaver()
        graph = build_graph(saver)
        waiting = graph.invoke(new_state(query, BASIS_DATE), config)
        records = []
        for item in reversed(list(saver.list(config))):
            records.append({
                "config": item.config["configurable"], "parent": (item.parent_config or {}).get("configurable"),
                "checkpoint": base64.b64encode(saver.serde.dumps_typed(item.checkpoint)[1]).decode(),
                "checkpoint_type": saver.serde.dumps_typed(item.checkpoint)[0],
                "metadata": json.loads(json.dumps(item.metadata, default=str)),
                "writes": [[task, channel, *(lambda t: [t[0], base64.b64encode(t[1]).decode()])(saver.serde.dumps_typed(value))]
                           for task, channel, value in (item.pending_writes or [])],
            })
        capture.parent.mkdir(parents=True, exist_ok=True)
        capture.write_text(json.dumps(
            {"synthetic_query": query, "note": "합성 테스트 질문으로 만든 '질문 대기 중' 체크포인트. 실사용자 데이터 없음.",
             "records": records}, ensure_ascii=False, indent=1), encoding="utf-8")
        out["captured_next"] = list(graph.get_state(config).next)
        finished = graph.invoke(Command(resume=answers), config)
        out["resumed_totals"] = (finished.get("statement") or {}).get("totals")
    if resume:
        saved = json.loads(resume.read_text(encoding="utf-8"))
        saver = MemorySaver()
        for record in saved["records"]:
            target = {"configurable": {**record["config"]}}
            if record["parent"]:
                target["configurable"]["checkpoint_id"] = record["parent"]["checkpoint_id"]
            else:
                target["configurable"].pop("checkpoint_id", None)
            checkpoint = saver.serde.loads_typed((record["checkpoint_type"], base64.b64decode(record["checkpoint"])))
            stored = saver.put(target, checkpoint, record["metadata"], checkpoint["channel_versions"])
            grouped = defaultdict(list)
            for task, channel, kind, value in record["writes"]:
                grouped[task].append((channel, saver.serde.loads_typed((kind, base64.b64decode(value)))))
            for task, values in grouped.items():
                saver.put_writes(stored, values, task)
        graph = build_graph(saver)
        try:
            out["loaded_next"] = list(graph.get_state(config).next)
            finished = graph.invoke(Command(resume=answers), config)
            out["resumed_status"] = finished.get("status")
            out["resumed_totals"] = (finished.get("statement") or {}).get("totals")
        except Exception as exc:  # 호환 실패도 결과로 남긴다
            out["error"] = f"{type(exc).__name__}: {str(exc)[:300]}"
    return out


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", required=True)
    parser.add_argument("--scenarios", required=True)
    parser.add_argument("--out", required=True)
    parser.add_argument("--capture-checkpoints")
    parser.add_argument("--resume-checkpoints")
    args = parser.parse_args()
    root = Path(args.root).resolve()
    environment = _fixed_environment()
    os.chdir(root)
    sys.path.insert(0, str(root))

    import backend
    from fastapi.testclient import TestClient
    from backend.api.main import app

    scenarios = json.loads(Path(args.scenarios).read_text(encoding="utf-8"))["scenarios"]
    results = {}
    patches = [patch("backend.api.usage_limits.processing", return_value=nullcontext((None, "baseline", False))),
               patch("backend.api.usage_limits.consume", return_value={}),
               patch("backend.api.usage_limits.status", return_value={}),
               patch("backend.api.usage_limits.read", return_value={})]
    for item in patches:
        item.start()
    client = TestClient(app)
    for scenario in scenarios:
        try:
            results[scenario["id"]] = run_scenario(client, scenario)
        except Exception as exc:
            results[scenario["id"]] = [{"error": f"{type(exc).__name__}: {str(exc)[:300]}"}]
    checkpoints = _checkpoint_io(Path(args.capture_checkpoints) if args.capture_checkpoints else None,
                                 Path(args.resume_checkpoints) if args.resume_checkpoints else None)
    import langgraph, fastapi  # noqa: E401
    repo = Path(__file__).resolve().parents[2]
    relative = lambda path: Path(os.path.relpath(path, repo)).as_posix()  # 공개 저장소에 로컬 절대 경로를 남기지 않는다
    meta = {"code_root": relative(root), "backend_package": relative(Path(backend.__file__).parent), "environment": environment,
            "python": sys.version.split()[0],
            "data_sha256": {f: hashlib.sha256((root / f).read_bytes()).hexdigest() for f in DATA_FILES}}
    Path(args.out).write_text(json.dumps({"meta": meta, "results": results, "checkpoints": checkpoints},
                                         ensure_ascii=False, indent=1, default=str), encoding="utf-8")
    print(f"scenarios={len(results)} out={args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
