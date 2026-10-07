"""새 공종별 견적 흐름의 서비스 계층·채팅 API 연결(단계 C2-1)을 오프라인으로 검사한다.

운영 DB·유료 API를 쓰지 않는다. 사용량은 가짜 값이며, 차감 호출 횟수만 센다.
회원 경로(PostgreSQL 저장·가져오기)는 로컬 Supabase가 없어 저장소를 흉내 낸 부분만 확인한다.
"""

from __future__ import annotations

import io
import os
import sys
from contextlib import nullcontext
from pathlib import Path
from unittest.mock import MagicMock, patch

os.environ["AGENT_OFFLINE"] = "1"
os.environ["AGENT_LLM"] = "off"
os.environ.pop("INDEX_CONFIG", None)
os.environ.pop("ESTIMATE_PLAN_INPUT", None)
os.environ.pop("ESTIMATE_MAX_ITEMS", None)
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

import backend.api.main as api  # noqa: E402
from backend.api import chat_storage, estimate_service, runtime  # noqa: E402
from backend.agent.estimate.state import max_items_setting  # noqa: E402

BASIS = "2026-10-06"
REBAR = "철근구조물 레미콘 인력운반 타설 150㎥"
PLAIN = "무근구조물 레미콘 인력운반 타설 50㎥"
READY = {"scattered_small_volume": False, "concrete_supply": "관급"}
HEADERS = {"X-Guest-Session": "estimate-service-check-" + "x" * 32}


def plan(*requests) -> dict:
    return {"basis_date": BASIS, "items": [{"request_text": text} for text in requests]}


def answer_all(response: dict, work: str = "6-1-1") -> dict:
    """화면처럼 질문 name을 키로 답한다(name을 그대로 돌려보냄)."""
    answers = {}
    for question in response["questions"]:
        field = question["name"].split("@")[0].split(":")[1]
        if field == "work":
            answers[question["name"]] = next(c for c in question["choices"] if work in c)
        else:
            answers[question["name"]] = READY.get(field, (question["choices"] or ["10"])[0])
    return answers


def main() -> int:
    checks: list[tuple[str, bool, str]] = []
    consume = MagicMock(return_value={})
    patches = [patch("backend.api.usage_limits.processing", return_value=nullcontext((None, "fixture", False))),
               patch("backend.api.usage_limits.consume", consume),
               patch("backend.api.usage_limits.status", return_value={})]
    for item in patches:
        item.start()
    client = TestClient(api.app, headers=HEADERS)

    # 1) 공종 수 상한 설정
    default_limit = max_items_setting()
    with patch.dict(os.environ, {"ESTIMATE_MAX_ITEMS": "0"}):
        try:
            max_items_setting()
            bad_setting = "허용됨"
        except ValueError as exc:
            bad_setting = str(exc)
    checks.append(("L1 상한 기본값 3, 잘못된 설정값은 거부", default_limit == 3 and "1 이상" in bad_setting, bad_setting))

    # 2) estimate_plan 입력은 서버 로컬 검사 설정에서만
    before = consume.call_count
    denied = client.post("/api/chat", json={"estimate_plan": plan(REBAR, PLAIN), "debug": True, "test_mode": True})
    with patch.dict(os.environ, {"ESTIMATE_PLAN_INPUT": "local", "APP_ENV": "production", "ALLOWED_ORIGINS": "https://example.com"}):
        production_enabled = runtime.estimate_plan_input_enabled()
        try:
            runtime.validate_production()
            production_rejected = False
        except ValueError as exc:
            production_rejected = "ESTIMATE_PLAN_INPUT" in str(exc)
    checks.append(("G1 설정이 없으면 거부(클라이언트 값으로 켤 수 없음), 운영에서는 설정해도 꺼짐",
                   denied.status_code == 422 and consume.call_count == before and not production_enabled and production_rejected,
                   f"응답 {denied.status_code} {denied.json().get('detail')}, 사용량 차감 {consume.call_count - before}회"))

    os.environ["ESTIMATE_PLAN_INPUT"] = "local"
    four = plan(REBAR, PLAIN, "자동문 3개소 설치", "도어록 50개소 설치")
    over = client.post("/api/chat", json={"estimate_plan": four})
    with patch.dict(os.environ, {"ESTIMATE_MAX_ITEMS": "4"}):
        allowed = client.post("/api/chat", json={"estimate_plan": four})
    allowed_body = allowed.json() if allowed.status_code == 200 else {}
    checks.append(("L2 기본 상한 3에서 4개는 범위 조정 요청, 설정 4이면 4개 처리",
                   over.status_code == 422 and "3개 공종까지" in over.json()["detail"]
                   and allowed.status_code == 200 and len(allowed_body.get("items", [])) == 4
                   and sorted(q["name"].split("@")[0] for q in allowed_body["questions"])[:1] == ["i1:work"],
                   f"기본 {over.status_code} / 설정 4: 공종 {len(allowed_body.get('items', []))}개, 질문 {len(allowed_body.get('questions', []))}개"))

    # 3) 비회원 채팅 API: 시작 → 답(질문 name 그대로) → 완료. 사용량은 요청마다 1회
    before = consume.call_count
    first = client.post("/api/chat", json={"estimate_plan": plan(REBAR, PLAIN)}).json()
    thread = first["thread_id"]
    names = [q["name"] for q in first["questions"]]
    second = client.post("/api/chat", json={"thread_id": thread, "answers": answer_all(first)}).json()
    done = client.post("/api/chat", json={"thread_id": thread, "answers": answer_all(second)}).json()
    checks.append(("A1 질문 name에 공종 ID·버전, 같은 대화에서 중단·재개 후 통합 견적",
                   all("@" in name and name.split(":")[0] in ("i1", "i2") for name in names)
                   and first["questions"][0]["ask"].startswith("[1번 ") and done["estimate_status"] == "PARTIAL"
                   and done["statement"]["totals"]["contract_amount"] == 25423718 and done["thread_id"] == thread
                   and [item["item_id"] for item in done["items"]] == ["i1", "i2"],
                   f"도급액 {done['statement']['totals']['contract_amount']:,}, 질문 예 {names[0]}"))
    checks.append(("U1 사용량은 채팅 요청마다 1회(공종 내부 계산은 차감 없음)", consume.call_count - before == 3,
                   f"요청 3회 → 차감 {consume.call_count - before}회"))

    # 4) 오래된 답·끝난 견적의 답: 상태 유지 + 안내
    late = client.post("/api/chat", json={"thread_id": thread, "answers": {names[0]: "anything"}}).json()
    checks.append(("A2 끝난 견적에 이전 질문 답: 금액·상태 유지, 안내", late["estimate_status"] == "PARTIAL"
                   and late["statement"]["totals"]["contract_amount"] == 25423718 and "반영하지 않았어요" in late["message"],
                   late["message"][:40]))

    # 5) 공통 조건 변경(현장 조건 반영 막대와 같은 요청) → 원가계산서만
    session_before = estimate_service.snapshot(api.ESTIMATE_GRAPH, thread).values["estimate"]
    changed = client.post("/api/chat", json={"thread_id": thread, "conditions": {"contractor_type": "전문건설업"}}).json()
    session_after = estimate_service.snapshot(api.ESTIMATE_GRAPH, thread).values["estimate"]
    checks.append(("C1 공통 조건 변경: 공종 결과 유지, 원가계산서만 갱신", changed["statement"]["totals"]["contract_amount"] == 25409021
                   and session_after["items"] == session_before["items"]
                   and session_after["common_conditions"]["contractor_type"]["source"] == "answer",
                   f"도급액 → {changed['statement']['totals']['contract_amount']:,}"))

    # 6) 서비스 계층 공종 수정(set_quantity·set_explicit): 같은 견적의 그 공종과 합계만
    updated = estimate_service.update_item(api.ESTIMATE_GRAPH, thread, "i1", quantity={"value": "200", "unit": "㎥"})
    session_updated = estimate_service.snapshot(api.ESTIMATE_GRAPH, thread).values["estimate"]
    checks.append(("I1 공종 수정: i1 물량 200, i2 결과 유지, 같은 견적 ID", session_updated["estimate_id"] == session_before["estimate_id"]
                   and session_updated["items"]["i2"]["priced_result"] == session_before["items"]["i2"]["priced_result"]
                   and session_updated["items"]["i1"]["priced_result"]["reference_amounts"]["volume"] == "200"
                   and updated["statement"]["totals"]["contract_amount"] != changed["statement"]["totals"]["contract_amount"],
                   f"도급액 → {updated['statement']['totals']['contract_amount']:,}"))
    # 6-1-1은 사급이어도 레미콘 단가를 묻지 않고 미산정으로 남긴다(질문 없이 그 공종만 다시 계산).
    finished = estimate_service.update_item(api.ESTIMATE_GRAPH, thread, "i2", conditions={"concrete_supply": "사급"})
    session_explicit = estimate_service.snapshot(api.ESTIMATE_GRAPH, thread).values["estimate"]
    i2_after = session_explicit["items"]["i2"]
    checks.append(("I2 공종 조건 수정(set_explicit): i2만 다시 계산, i1 결과 유지",
                   i2_after["conditions"]["concrete_supply"] == {"value": "사급", "source": "answer", "explicit": True}
                   and i2_after["result_revision"] == i2_after["input_revision"] > session_updated["items"]["i2"]["input_revision"]
                   and session_explicit["items"]["i1"]["priced_result"] == session_updated["items"]["i1"]["priced_result"]
                   and finished["estimate_status"] in ("COMPLETE", "PARTIAL") and not finished["questions"],
                   f"{finished['estimate_status']} 도급액 {finished['statement']['totals']['contract_amount']:,}"))

    # 7) Excel: 같은 대화 주소에서 같은 집계
    export = client.get(f"/api/export/{thread}.xlsx")
    sheets = load_workbook(io.BytesIO(export.content)) if export.status_code == 200 else None
    estimate_text = " ".join(str(c) for row in sheets["견적서"].iter_rows(values_only=True) for c in row if c) if sheets else ""
    checks.append(("X1 Excel: 마지막 견적과 같은 금액", export.status_code == 200
                   and f"₩{finished['statement']['totals']['contract_amount']:,}" in estimate_text
                   and "%EC%99%B8%201%EA%B1%B4" in export.headers.get("content-disposition", ""),
                   f"{export.status_code} 도급액 {finished['statement']['totals']['contract_amount']:,}"))

    # 8) 같은 대화의 기존/새 상태 구분
    legacy_start = client.post("/api/chat", json={"message": REBAR}).json()
    mixed = legacy_start["thread_id"]
    legacy_values = api.GRAPH.get_state({"configurable": {"thread_id": mixed}}).values
    estimate_in_mixed = client.post("/api/chat", json={"thread_id": mixed, "estimate_plan": plan(REBAR, PLAIN)}).json()
    legacy_after = api.GRAPH.get_state({"configurable": {"thread_id": mixed}}).values
    first_estimate_id = estimate_in_mixed["estimate_id"]
    restart = client.post("/api/chat", json={"thread_id": mixed, "message": PLAIN, "restart": True}).json()
    legacy_after_restart = api.GRAPH.get_state({"configurable": {"thread_id": mixed}}).values
    checks.append(("F1 새 견적의 다시 보내기: 저장된 계획으로 새 견적 흐름 재실행(공종 누락 없음), 기존 상태 유지",
                   legacy_start.get("flow") is None and estimate_in_mixed["flow"] == "estimate"
                   and legacy_after["query"] == legacy_values["query"]
                   and restart.get("flow") == "estimate" and restart["estimate_id"] != first_estimate_id
                   and [item["item_id"] for item in restart["items"]] == ["i1", "i2"]
                   and sorted(q["name"].split("@")[0] for q in restart["questions"]) == ["i1:work", "i2:work"]
                   and legacy_after_restart["query"] == legacy_values["query"]
                   and estimate_service.latest_flow(api.GRAPH, api.ESTIMATE_GRAPH, mixed) == "estimate",
                   f"새 견적 ID {restart['estimate_id'][:8]}… (이전 {first_estimate_id[:8]}…), 공종 {[i['item_id'] for i in restart['items']]}"))
    cross = client.post("/api/chat", json={"thread_id": mixed, "answers": {"work": "6-1-1", "concrete_supply": "관급"}}).json()
    cross_session = estimate_service.snapshot(api.ESTIMATE_GRAPH, mixed).values["estimate"]
    checks.append(("F2 기존 흐름 형식의 답은 새 견적에 적용하지 않고 안내", cross["flow"] == "estimate"
                   and all(item["selected_spec_id"] == "" and item["explicit"] == {} for item in cross_session["items"].values())
                   and sorted(q["name"].split("@")[0] for q in cross["questions"]) == ["i1:work", "i2:work"]
                   and "이전 질문" in cross["message"], cross["message"][:40]))
    legacy_only = client.post("/api/chat", json={"message": REBAR}).json()
    legacy_thread = legacy_only["thread_id"]
    foreign = client.post("/api/chat", json={"thread_id": legacy_thread, "answers": {restart["questions"][0]["name"]: "x"}}).json()
    resumed = client.post("/api/chat", json={"thread_id": legacy_thread, "answers": {"work": "6-1-1"}}).json()
    legacy_restart = client.post("/api/chat", json={"thread_id": legacy_thread, "message": PLAIN, "restart": True}).json()
    checks.append(("F3 기존 대화: 새 견적 형식 답은 적용되지 않고, 재개·다시 보내기는 기존 흐름",
                   foreign.get("flow") is None and foreign["status"] == "MISSING_INFO"
                   and [q["name"] for q in foreign["questions"]] == [q["name"] for q in legacy_only["questions"]]
                   and resumed.get("flow") is None and legacy_restart.get("flow") is None
                   and estimate_service.snapshot(api.ESTIMATE_GRAPH, legacy_thread) is None,
                   f"재개 {resumed['status']}, 다시 보내기 {legacy_restart['status']}"))

    # 9) 삭제·만료·탈퇴: 두 체크포인트 키를 함께
    def has_state(key):
        return bool(list(api.GRAPH.checkpointer.list({"configurable": {"thread_id": key}})))
    keys = chat_storage.checkpoint_threads(mixed)
    existed = [has_state(key) for key in keys]
    deleted = client.delete(f"/api/guest/conversations/{mixed}")
    checks.append(("D1 비회원 대화 삭제: 기존·새 상태 모두 삭제", existed == [True, True] and deleted.status_code == 204
                   and not any(has_state(key) for key in keys), f"삭제 전 {existed}"))
    with patch.object(api, "_GUEST_TTL", -1):
        api._prune_guests()
    checks.append(("D2 비회원 만료 정리: 기존·새 상태 모두 삭제", not any(has_state(key) for key in chat_storage.checkpoint_threads(thread)),
                   "만료 처리 후 남은 상태 없음"))
    recorded = []
    fake_saver = MagicMock()
    fake_saver.delete_thread.side_effect = recorded.append
    fake_conn = MagicMock()
    fake_conn.__enter__.return_value = fake_conn
    fake_conn.transaction.return_value = nullcontext()
    with patch.object(chat_storage, "connection", return_value=fake_conn), \
            patch.object(chat_storage, "owned", return_value={}), \
            patch.object(chat_storage, "PostgresSaver", return_value=fake_saver):
        chat_storage.delete_conversation("conv-1", "user-1")
    member_deleted = list(recorded)
    recorded.clear()
    from backend.api import accounts
    import inspect
    removal = inspect.getsource(accounts)
    checks.append(("D3 회원 대화 삭제(저장소 흉내): 두 키 삭제 호출, 탈퇴 코드도 같은 목록 사용",
                   member_deleted == ["conv-1", "conv-1#estimate"] and "checkpoint_threads(str(conversation_id))" in removal,
                   f"삭제 호출 {member_deleted}"))

    for item in patches:
        item.stop()
    os.environ.pop("ESTIMATE_PLAN_INPUT", None)
    for name, ok, detail in checks:
        print(("PASS " if ok else "FAIL ") + name + (f"  — {detail}" if detail else ""))
    passed = sum(bool(ok) for _, ok, _ in checks)
    print(f"통과 {passed} / 전체 {len(checks)}")
    return 0 if passed == len(checks) else 1


if __name__ == "__main__":
    raise SystemExit(main())
