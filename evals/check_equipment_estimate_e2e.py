"""로더·굴착기 장비 견적 끝까지: 조건 → Q → 장비시간 → 금액 → 다중 합계 → 수정 → 복원 → Excel.

검색은 모의(모델·임베딩·서버 없음). 기대값은 원문 표 숫자와 요금 파일에서 따로 계산한다.
"""
from __future__ import annotations

import json
import os
import sys
from collections import defaultdict
from decimal import ROUND_DOWN, Decimal
from fractions import Fraction as F
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.environ.update(AGENT_MODE="tools", AGENT_LLM="off", AGENT_OFFLINE="1")

from openpyxl import load_workbook  # noqa: E402

from backend.agent.estimate import dialogue, tools  # noqa: E402
from backend.agent.rules.specs import load_specs  # noqa: E402
from backend.agent.tools.source import citation  # noqa: E402
from backend.api import dialogue_service, tables  # noqa: E402

SPECS = {s["id"]: s for s in load_specs().values()}
LOADER = SPECS["2026-정오표1차/공통/8-2-5/로더"]
EXC = SPECS["2026-정오표1차/공통/8-2-3/굴착기"]
TEXT = "로더 상차 1000㎥ + 굴착기 굴착 500㎥ 비용 견적"
ANS = {"loader_type": "타이어식", "machine_spec": "0.95㎥", "volume_basis": "자연상태", "L": "1.25",
       "t1_condition": "보통인경우", "t1_method": "산적상태에서 담을 때", "e_soil_type": "모래, 사질토",
       "e_state": "자연상태", "e_condition": "보통", "is_trench": False, "l": "8", "fuel_price": "1500",
       "quantity": "1000", "soil_type": "모래,사질토", "soil_state": "자연상태", "site_condition": "보통",
       "is_trenching": False, "is_underwater": False, "is_residential_area": False, "swing_angle": "90"}


def hit(query: str) -> dict:
    spec = LOADER if "로더" in query else EXC
    return {"rank": 1, "section_no": spec["section_no"], "division": spec["division"],
            "section": f"{spec['section_no']} {spec['title']}", "page": 193, "table_id": None,
            "text": spec["title"], "chunk_id": "mock", "ref": spec["id"]}


def answer_value(q: dict, spec_id: str | None):
    name = q["name"].split("::", 1)[-1]
    if name in ANS:
        value = ANS[name]
        if name == "quantity" and spec_id == EXC["id"]:
            value = "500"
        if q.get("choices") and value not in q["choices"]:
            value = next((c for c in q["choices"] if str(c).replace(" ", "") == str(value).replace(" ", "")),
                         q["choices"][0])
        return value
    if q.get("choices"):
        return False if name.startswith("is_") else q["choices"][0]
    return "100"


def drive(graph, tid, r, limit=60):
    for _ in range(limit):
        if not r["questions"]:
            return r
        q = r["questions"][0]
        state = dialogue.load(graph, tid)
        item_id = q["name"].split("::", 1)[0]
        spec_id = ((state or {}).get("session", {}).get("items", {}).get(item_id) or {}).get("selected_spec_id")
        r = dialogue_service.respond(graph, tid, {"answers": {q["name"]: answer_value(q, spec_id)},
                                                  "refs": {q["name"]: q["ref"]}})
    raise AssertionError(("stuck", r["questions"][:1], r.get("answer")))


def trunc(value, places="0.1") -> Decimal:
    return Decimal(value).quantize(Decimal(places), rounding=ROUND_DOWN)


def expected_unit(q: F, machine: dict, wage_daily: int, fuel_price: int) -> dict:
    """원문 표 숫자에서 따로 구한 1㎥당 일위대가(재료비=연료·잡재료, 노무비=운전원, 경비=손료)."""
    hours = Decimal((1 / q).numerator) / Decimal((1 / q).denominator)
    material = trunc(hours * Decimal(machine["fuel_l_per_hr"]) * fuel_price
                     * (1 + Decimal(machine["misc_pct_of_fuel"]) / 100))
    labor = trunc(hours * machine["operator_per_day"] * wage_daily / 8)
    expense = trunc(hours * machine["hourly_depreciation"])
    return {"재료비": int(trunc(material, "1")), "노무비": int(trunc(labor, "1")), "경비": int(trunc(expense, "1"))}


def main() -> None:
    rates = json.loads((ROOT / "data/rates/equipment_rates.json").read_text(encoding="utf-8"))["machines"]
    versions = json.loads((ROOT / "data/rates/labor_rates.json").read_text(encoding="utf-8"))["versions"]
    wage = int(next(v for v in versions if v["id"].startswith("2026H2"))["rates"]["1048"]["daily"])
    # 원문 표(손으로 옮김): 로더 타이어식 0.95㎥ · t1=9(보통·산적) · m=1.8 · l=8 · t2=14 · K=1.2 · E 자연상태 보통 0.55 · L=1.25
    cm_loader = F("1.8") * 8 + 9 + 14
    q_loader = 3600 * F("0.95") * F("1.2") * (1 / F("1.25")) * F("0.55") / cm_loader
    # 굴착기 0.12㎥ · 구분 0.12~0.4 · 90도 cm=15 · K=1.10 · E 모래 자연상태 보통 0.70 · L=1.25
    q_exc = 3600 * F("0.12") * F("1.10") * (1 / F("1.25")) * F("0.70") / 15
    want = {"loader": (expected_unit(q_loader, rates["0302-0095"], wage, 1500), 1000),
            "exc": (expected_unit(q_exc, rates["0201-0012"], wage, 1500), 500)}

    graph = dialogue.build_dialogue_graph()
    tid = "equipment-e2e"
    with patch.object(tools, "retrieve", side_effect=lambda s: {"hits": [hit(s["query"])], "search_info": {
            "method": "mock", "api_calls": 0, "warnings": []}}), \
            patch.object(citation, "_pages", return_value=defaultdict(lambda: {"printed_page": None, "division": None})):
        r = dialogue_service.respond(graph, tid, {"message": TEXT})
        r = dialogue_service.respond(graph, tid, {
            "answers": {q["name"]: "계산에 포함" for q in r["questions"]}, "refs": {q["name"]: q["ref"] for q in r["questions"]}})
        r = drive(graph, tid, r)
        state = dialogue.load(graph, tid)
        items = state["session"]["items"]
        ids = {it["selected_spec_id"]: i for i, it in items.items()}
        loader_id, exc_id = ids[LOADER["id"]], ids[EXC["id"]]
        for item_id in (loader_id, exc_id):
            assert items[item_id]["status"] == "PRICED", (item_id, items[item_id]["status"])
            assert items[item_id]["priced_result"]["status"] == "OK", items[item_id]["priced_result"]["unpriced"]
        # 1. 조건 → Q → 장비시간 (정확한 분수)
        for item_id, q, name in ((loader_id, q_loader, "로더"), (exc_id, q_exc, "굴착기")):
            lines = {l["name"]: l["exact"] for l in items[item_id]["computed_result"]["unit_lines"]}
            assert F(lines["시간당 작업량 Q"]) == q and F(lines[name]) == 1 / q, (name, lines, q)
        print("PASS Q and equipment hours match source-table expectation (loader 4104/85 m3/hr, excavator 17.7408 m3/hr)")

        # 2. 단가·금액: 1㎥당 일위대가와 물량 곱, 합계
        def check_amounts(st, qty_loader=1000, qty_exc=500):
            its = st["session"]["items"]
            total = {"재료비": 0, "노무비": 0, "경비": 0}
            for item_id, key, qty in ((loader_id, "loader", qty_loader), (exc_id, "exc", qty_exc)):
                unit = want[key][0]
                priced = its[item_id]["priced_result"]
                assert {k: int(v) for k, v in priced["unit_prices"].items()} == unit, (key, priced["unit_prices"], unit)
                assert int(priced["reference_amounts"]["total"]) == sum(unit.values()) * qty
                for k in total:
                    total[k] += unit[k] * qty
            totals = st["session"]["statement"]["totals"]
            assert totals["materials"] == total["재료비"], (totals, total)
            assert totals["labor"] >= total["노무비"] and totals["expenses"] >= total["경비"]
            return total
        direct = check_amounts(state)
        print("PASS unit prices and amounts match independent calculation, direct cost:", direct)

        # 3. Excel 셀 대조
        def excel(label):
            resp = dialogue_service.export_response(graph, tid)
            return load_workbook(BytesIO(tables.build_xlsx(resp)), data_only=True)

        def check_excel(book, direct, qty_loader, qty_exc, contract):
            bill = list(book["내역서"].iter_rows(values_only=True))
            for item_id, key, qty in ((loader_id, "loader", qty_loader), (exc_id, "exc", qty_exc)):
                row = next(x for x in bill if x[1] == item_id)
                unit = want[key][0]
                assert row[3] == qty and row[2] == "m3", row
                assert row[4:10] == (unit["재료비"], unit["재료비"] * qty, unit["노무비"], unit["노무비"] * qty,
                                    unit["경비"], unit["경비"] * qty), row
            total_row = next(x for x in bill if x[0] and str(x[0]).startswith("합계"))
            assert (total_row[5], total_row[7], total_row[9], total_row[11]) == (direct["재료비"], direct["노무비"], direct["경비"], sum(direct.values())), total_row
            assert next(x for x in book["견적서"].iter_rows(values_only=True) if x[0] == "합 계")[1] == contract
            cost = {x[0]: x for x in book["원가계산서"].iter_rows(values_only=True) if x[0]}
            assert cost["도급액"][2] == contract
            unit_rows = {x[0]: x for x in book["일위대가"].iter_rows(values_only=True) if x[0]}
            assert unit_rows["로더 기계손료"][3] == "hr" and unit_rows["굴착기 운전원"][3] == "hr"
            assert unit_rows["로더 기계손료"][10] == 9711 and unit_rows["굴착기 기계손료"][10] == 9397, unit_rows["로더 기계손료"]
        contract = state["session"]["statement"]["totals"]["contract_amount"]
        check_excel(excel("1"), direct, 1000, 500, contract)
        print("PASS Excel quantity/unit/unit price/amount, direct total, 견적서·원가계산서 contract amount match:", contract)

        # 4. 새로고침 복원
        restored = dialogue_service.restore(graph, tid)
        assert restored["statement"]["totals"]["contract_amount"] == contract
        print("PASS restore after reload keeps the statement total")

        # 5. 로더 물량 수정: 장비시간·금액·합계 갱신, 굴착기 보존
        before_exc = json.dumps(items[exc_id]["computed_result"], sort_keys=True)
        before_exc_rev = items[exc_id]["input_revision"]
        r = dialogue_service.respond(graph, tid, {"message": "로더 상차 2000㎥로 바꿔"})
        r = drive(graph, tid, r)
        state2 = dialogue.load(graph, tid)
        its = state2["session"]["items"]
        assert its[loader_id]["conditions"]["quantity"] == "2000", its[loader_id]["conditions"].get("quantity")
        assert its[exc_id]["input_revision"] == before_exc_rev
        assert json.dumps(its[exc_id]["computed_result"], sort_keys=True) == before_exc
        direct2 = check_amounts(state2, qty_loader=2000)
        contract2 = state2["session"]["statement"]["totals"]["contract_amount"]
        assert contract2 > contract
        check_excel(excel("2"), direct2, 2000, 500, contract2)
        print("PASS loader quantity 1000->2000 updates amounts, total and Excel; excavator result unchanged:", contract, "->", contract2)


if __name__ == "__main__":
    main()
