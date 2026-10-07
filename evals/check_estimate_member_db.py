"""새 공종별 견적의 회원 저장을 실제 로컬 PostgreSQL(로컬 Supabase 127.0.0.1:54322)로 검사한다.

운영 Supabase·유료 API를 쓰지 않는다. 회원은 로컬 auth.users에 만든 검사용 계정이며,
토큰 검증(chat_storage.identity)과 Auth 관리자 삭제 HTTP만 흉내 낸다(기존 check_chat_storage·
check_account_deletion과 같은 방식). 서버 재시작은 단계마다 별도 Python 프로세스로 실행해 확인한다.

사용: CHAT_DATABASE_URL=postgresql://postgres:postgres@127.0.0.1:54322/postgres python evals/check_estimate_member_db.py
"""

from __future__ import annotations

import io
import json
import os
import subprocess
import sys
from pathlib import Path
from unittest.mock import patch
from uuid import uuid4

ROOT = Path(__file__).resolve().parents[1]
STATE = ROOT / "tmp/estimate-member-db-state.json"
BASIS = "2026-10-06"
REBAR = "철근구조물 레미콘 인력운반 타설 150㎥"
PLAIN = "무근구조물 레미콘 인력운반 타설 50㎥"
READY = {"scattered_small_volume": False, "concrete_supply": "관급"}


def _environment() -> None:
    if "127.0.0.1:54322" not in os.environ.get("CHAT_DATABASE_URL", ""):
        raise SystemExit("CHAT_DATABASE_URL을 로컬 Supabase(127.0.0.1:54322)로만 지정하세요")
    os.environ.update({"AGENT_OFFLINE": "1", "AGENT_LLM": "off", "ESTIMATE_PLAN_INPUT": "local"})
    for name in ("INDEX_CONFIG", "APP_ENV", "ESTIMATE_MAX_ITEMS"):
        os.environ.pop(name, None)
    sys.path.insert(0, str(ROOT))


def _answers(response: dict) -> dict:
    out = {}
    for question in response["questions"]:
        field = question["name"].split("@")[0].split(":")[1]
        out[question["name"]] = next(c for c in question["choices"] if "6-1-1" in c) if field == "work" else READY[field]
    return out


def _context():
    from fastapi.testclient import TestClient
    import backend.api.main as api
    from backend.api import chat_storage
    state = json.loads(STATE.read_text(encoding="utf-8")) if STATE.exists() else {}

    def verified(token):
        tokens = {"Bearer fixture-a": state.get("a"), "Bearer fixture-b": state.get("b")}
        if token is None:
            return None
        if token in tokens:
            return tokens[token]
        from fastapi import HTTPException
        raise HTTPException(401, "invalid")
    patcher = patch.object(chat_storage, "identity", side_effect=verified)
    patcher.start()
    return api, chat_storage, TestClient(api.app, raise_server_exceptions=False), state


def _sql(query: str, args=()):
    from backend.api import chat_storage
    with chat_storage.connection() as conn:
        return conn.execute(query, args).fetchall()


def _usage(user_id: str) -> int:
    from backend.api import usage_limits
    subject = usage_limits.subject(user_id, "testclient")
    rows = _sql("SELECT used FROM agent_state.daily_chat_usage WHERE subject=%s AND usage_day=(now() AT TIME ZONE 'Asia/Seoul')::date", (subject,))
    return rows[0]["used"] if rows else 0


def _threads(thread_id: str) -> list[str]:
    rows = _sql("SELECT DISTINCT thread_id FROM agent_state.checkpoints WHERE thread_id IN (%s, %s)", (thread_id, thread_id + "#estimate"))
    return sorted(row["thread_id"] for row in rows)


def _report(results: list) -> None:
    previous = json.loads(STATE.read_text(encoding="utf-8")).get("results", []) if STATE.exists() else []
    state = json.loads(STATE.read_text(encoding="utf-8"))
    state["results"] = previous + results
    STATE.write_text(json.dumps(state, ensure_ascii=False), encoding="utf-8")


def phase_start() -> None:
    """프로세스 1: 검사 회원 준비, 견적 시작 → 질문 하나 답 → 질문 대기 상태로 끝냄."""
    from backend.api import chat_storage
    chat_storage.setup()
    a, b, conv = str(uuid4()), str(uuid4()), str(uuid4())
    with chat_storage.connection() as conn:
        for user_id in (a, b):
            conn.execute("INSERT INTO auth.users(id,aud,role,email) VALUES (%s,'authenticated','authenticated',%s)",
                         (user_id, f"estimate-db-{user_id}@example.invalid"))
    STATE.write_text(json.dumps({"a": a, "b": b, "conv": conv, "results": []}), encoding="utf-8")
    api, chat_storage, client, state = _context()
    headers = {"Authorization": "Bearer fixture-a"}
    first = client.post("/api/chat", headers=headers, json={"conversation_id": conv, "request_id": str(uuid4()),
                                                            "estimate_plan": {"basis_date": BASIS, "items": [
                                                                {"request_text": REBAR}, {"request_text": PLAIN}]}})
    second = client.post("/api/chat", headers=headers, json={"conversation_id": conv, "request_id": str(uuid4()),
                                                             "thread_id": conv, "answers": _answers(first.json())})
    body = second.json()
    state.update(pending_names=sorted(q["name"] for q in body["questions"]), usage_after_start=_usage(a))
    STATE.write_text(json.dumps({**json.loads(STATE.read_text(encoding="utf-8")), **state}), encoding="utf-8")
    _report([("DB1 회원 견적 시작·질문 답(실제 DB 저장)", first.status_code == 200 and second.status_code == 200
              and body["estimate_status"] == "PENDING" and _threads(conv) == [conv + "#estimate"],
              f"질문 {len(body['questions'])}개 대기, 저장 키 {_threads(conv)}")])


def phase_resume() -> None:
    """프로세스 2(재시작): 질문 대기 복원 → 완료 → 중복 요청·다른 회원 접근·Excel·공통 조건 변경."""
    api, chat_storage, client, state = _context()
    a, conv = state["a"], state["conv"]
    headers = {"Authorization": "Bearer fixture-a"}
    restored = client.get(f"/api/conversations/{conv}", headers=headers).json()
    last = restored["messages"][-1]["payload"]
    from backend.agent.estimate.flow import build_estimate_graph
    from backend.api import estimate_service
    from langgraph.checkpoint.postgres import PostgresSaver
    with chat_storage.connection() as conn:
        snap = estimate_service.snapshot(build_estimate_graph(PostgresSaver(conn)), conv)
        stored_names = sorted(f"{q['question_id']}@{q['version']}" for q in snap.values["estimate"]["pending_questions"].values())
    results = [("DB2 재시작 후 질문 대기 상태 복원(메시지·체크포인트)",
                bool(snap.next) and stored_names == state["pending_names"] == sorted(q["name"] for q in last["questions"])
                and len(restored["messages"]) == 4, f"메시지 {len(restored['messages'])}개, 대기 질문 {len(stored_names)}개")]
    request_id = str(uuid4())
    done = client.post("/api/chat", headers=headers, json={"conversation_id": conv, "request_id": request_id, "thread_id": conv,
                                                           "answers": _answers(last)})
    body = done.json()
    usage_once, messages_once = _usage(a), len(client.get(f"/api/conversations/{conv}", headers=headers).json()["messages"])
    with chat_storage.connection() as conn:
        revision_once = estimate_service.snapshot(build_estimate_graph(PostgresSaver(conn)), conv).values["estimate"]["revision"]
    again = client.post("/api/chat", headers=headers, json={"conversation_id": conv, "request_id": request_id, "thread_id": conv,
                                                            "answers": _answers(last)})
    with chat_storage.connection() as conn:
        revision_twice = estimate_service.snapshot(build_estimate_graph(PostgresSaver(conn)), conv).values["estimate"]["revision"]
    results.append(("DB3 견적 완료와 실제 DB 다시 읽기", done.status_code == 200 and body["estimate_status"] == "PARTIAL"
                    and body["statement"]["totals"]["contract_amount"] == 25423718
                    and client.get(f"/api/conversations/{conv}", headers=headers).json()["messages"][-1]["payload"]["statement"]["totals"]["contract_amount"] == 25423718,
                    f"도급액 {body['statement']['totals']['contract_amount']:,}"))
    results.append(("DB4 같은 request_id 재요청: 중복 실행·차감·메시지 저장 없음",
                    again.status_code == 200 and again.json()["answer_id"] == body["answer_id"] and _usage(a) == usage_once
                    and len(client.get(f"/api/conversations/{conv}", headers=headers).json()["messages"]) == messages_once
                    and revision_twice == revision_once and usage_once == state["usage_after_start"] + 1,
                    f"사용량 {state['usage_after_start']} → {usage_once} → {_usage(a)}, 메시지 {messages_once}"))
    other = {"Authorization": "Bearer fixture-b"}
    read_b = client.get(f"/api/conversations/{conv}", headers=other)
    post_b = client.post("/api/chat", headers=other, json={"conversation_id": conv, "request_id": str(uuid4()), "thread_id": conv,
                                                          "conditions": {"contractor_type": "전문건설업"}})
    export_b = client.get(f"/api/export/{conv}.xlsx", headers=other)
    results.append(("DB5 다른 회원의 대화 읽기·요청·Excel 차단", read_b.status_code == post_b.status_code == export_b.status_code == 404,
                    f"{read_b.status_code}/{post_b.status_code}/{export_b.status_code}"))
    export = client.get(f"/api/export/{conv}.xlsx", headers=headers)
    from openpyxl import load_workbook
    text = " ".join(str(c) for row in load_workbook(io.BytesIO(export.content))["견적서"].iter_rows(values_only=True) for c in row if c) \
        if export.status_code == 200 else ""
    changed = client.post("/api/chat", headers=headers, json={"conversation_id": conv, "request_id": str(uuid4()), "thread_id": conv,
                                                              "conditions": {"contractor_type": "전문건설업"}}).json()
    results.append(("DB6 회원 Excel·공통 조건 변경(실제 DB 체크포인트)", "₩25,423,718" in text
                    and changed["statement"]["totals"]["contract_amount"] == 25409021, f"Excel {export.status_code}, 변경 후 25,409,021"))
    state_all = json.loads(STATE.read_text(encoding="utf-8"))
    state_all["contract_after_change"] = changed["statement"]["totals"]["contract_amount"]
    STATE.write_text(json.dumps(state_all, ensure_ascii=False), encoding="utf-8")
    _report(results)


def phase_finish() -> None:
    """프로세스 3(재시작): 완료 결과 복원, 실패 시 트랜잭션, 대화 삭제, 회원 탈퇴, 비회원 가져오기."""
    api, chat_storage, client, state = _context()
    a, b, conv = state["a"], state["b"], state["conv"]
    headers = {"Authorization": "Bearer fixture-a"}
    from backend.agent.estimate.flow import build_estimate_graph
    from backend.api import estimate_service
    from langgraph.checkpoint.postgres import PostgresSaver
    restored = client.get(f"/api/conversations/{conv}", headers=headers).json()
    export = client.get(f"/api/export/{conv}.xlsx", headers=headers)
    results = [("DB7 재시작 후 완료 결과 복원(메시지·Excel)",
                restored["messages"][-1]["payload"]["statement"]["totals"]["contract_amount"] == state["contract_after_change"]
                and export.status_code == 200, f"마지막 도급액 {state['contract_after_change']:,}")]

    def latest():
        with chat_storage.connection() as conn:
            snap = estimate_service.snapshot(build_estimate_graph(PostgresSaver(conn)), conv)
            return snap.config["configurable"]["checkpoint_id"], snap.values["estimate"]["statement"]["totals"]["contract_amount"]
    before_checkpoint = latest()
    before_messages = len(restored["messages"])
    with patch.object(chat_storage, "append_pair", side_effect=RuntimeError("메시지 저장 실패 흉내")):
        failed = client.post("/api/chat", headers=headers, json={"conversation_id": conv, "request_id": str(uuid4()), "thread_id": conv,
                                                                 "conditions": {"contractor_type": "종합건설업"}})
    after_messages = len(client.get(f"/api/conversations/{conv}", headers=headers).json()["messages"])
    results.append(("DB8 메시지 저장 실패 시 견적 체크포인트도 함께 취소", failed.status_code == 500 and latest() == before_checkpoint
                    and after_messages == before_messages, f"응답 {failed.status_code}, 체크포인트·도급액 유지 {before_checkpoint[1]:,}"))
    deleted = client.delete(f"/api/conversations/{conv}", headers=headers)
    results.append(("DB9 대화 삭제: 기존 키·#estimate 키 체크포인트 모두 삭제", deleted.status_code == 204 and _threads(conv) == [],
                    f"삭제 후 남은 키 {_threads(conv)}"))

    # 회원 탈퇴: 기존 흐름 대화 + 새 견적 대화를 가진 회원 A의 탈퇴 작업을 실제 DB에서 처리(Auth 관리자 HTTP만 흉내)
    legacy_conv, estimate_conv = str(uuid4()), str(uuid4())
    client.post("/api/chat", headers=headers, json={"conversation_id": legacy_conv, "request_id": str(uuid4()), "message": REBAR})
    client.post("/api/chat", headers=headers, json={"conversation_id": estimate_conv, "request_id": str(uuid4()),
                                                    "estimate_plan": {"basis_date": BASIS, "items": [{"request_text": REBAR}, {"request_text": PLAIN}]}})
    before = (_threads(legacy_conv), _threads(estimate_conv))
    import httpx
    from backend.api import accounts

    def admin_delete(method, url, **kwargs):
        with chat_storage.connection() as conn:
            conn.execute("DELETE FROM auth.users WHERE id=%s", (url.rsplit("/", 1)[1],))
        return httpx.Response(200, request=httpx.Request(method, url), json={})
    with chat_storage.connection() as conn:
        job = conn.execute("INSERT INTO agent_state.account_deletions(user_id,conversation_ids,feedback_ids) VALUES (%s,%s,%s) RETURNING *",
                           (a, [legacy_conv, estimate_conv], [])).fetchone()
        with patch.dict(os.environ, {"SUPABASE_SECRET_KEY": "local-fixture-secret", "SUPABASE_URL": "http://127.0.0.1:54321"}), \
                patch.object(httpx, "request", side_effect=admin_delete):
            accounts._finish(conn, job)
    after = (_threads(legacy_conv), _threads(estimate_conv))
    conversations_left = _sql("SELECT count(*) AS n FROM public.conversations WHERE user_id=%s", (a,))[0]["n"]
    results.append(("DB10 회원 탈퇴: 기존 흐름 대화·새 견적 대화의 두 키 체크포인트 모두 삭제",
                    before[0] == [legacy_conv] and before[1] == [estimate_conv + "#estimate"] and after == ([], []) and conversations_left == 0,
                    f"탈퇴 전 {before}, 후 {after}"))

    # 비회원 대화 가져오기 → 회원으로 이어서 처리
    guest_session = "guest-import-fixture-" + "y" * 32
    guest = client.post("/api/chat", headers={"X-Guest-Session": guest_session},
                        json={"estimate_plan": {"basis_date": BASIS, "items": [{"request_text": REBAR}, {"request_text": PLAIN}]}}).json()
    guest_thread, target = guest["thread_id"], str(uuid4())
    imported = client.post("/api/conversations/import-guest", headers={"Authorization": "Bearer fixture-b", "X-Guest-Session": guest_session},
                           json={"thread_id": guest_thread, "conversation_id": target})
    memory_left = list(api.GRAPH.checkpointer.list({"configurable": {"thread_id": guest_thread + "#estimate"}}))
    member_b = {"Authorization": "Bearer fixture-b"}
    step = client.post("/api/chat", headers=member_b, json={"conversation_id": target, "request_id": str(uuid4()), "thread_id": target,
                                                            "answers": _answers(guest)}).json()
    final = client.post("/api/chat", headers=member_b, json={"conversation_id": target, "request_id": str(uuid4()), "thread_id": target,
                                                             "answers": _answers(step)}).json()
    results.append(("DB11 비회원 견적을 로그인 후 가져와 회원으로 이어서 완료",
                    imported.status_code == 200 and _threads(target) == [target + "#estimate"] and not memory_left
                    and guest_thread not in api._GUESTS and final.get("estimate_status") == "PARTIAL"
                    and final["statement"]["totals"]["contract_amount"] == 25423718
                    and len(client.get(f"/api/conversations/{target}", headers=member_b).json()["messages"]) == 6,
                    f"가져온 키 {_threads(target)}, 이어서 도급액 {final.get('statement', {}).get('totals', {}).get('contract_amount')}"))
    client.delete(f"/api/conversations/{target}", headers=member_b)
    _report(results)


def cleanup() -> None:
    state = json.loads(STATE.read_text(encoding="utf-8"))
    from backend.api import chat_storage, usage_limits
    with chat_storage.connection() as conn:
        for user_id in (state["a"], state["b"]):
            conn.execute("DELETE FROM agent_state.daily_chat_usage WHERE subject=%s", (usage_limits.subject(user_id, "testclient"),))
            conn.execute("DELETE FROM agent_state.account_deletions WHERE user_id=%s", (user_id,))
            conn.execute("DELETE FROM auth.users WHERE id=%s", (user_id,))


def main() -> int:
    _environment()
    phase = sys.argv[1] if len(sys.argv) > 1 else ""
    if phase:
        {"start": phase_start, "resume": phase_resume, "finish": phase_finish, "cleanup": cleanup}[phase]()
        return 0
    STATE.unlink(missing_ok=True)
    try:
        for name in ("start", "resume", "finish"):
            run = subprocess.run([sys.executable, __file__, name], env={**os.environ, "PYTHONIOENCODING": "utf-8"},
                                 capture_output=True, text=True, encoding="utf-8", errors="replace")
            if run.returncode != 0:
                print(f"[{name} 단계 오류]\n{run.stderr[-2500:]}")
                break
    finally:
        if STATE.exists():
            subprocess.run([sys.executable, __file__, "cleanup"], env=os.environ, capture_output=True)
    results = json.loads(STATE.read_text(encoding="utf-8")).get("results", []) if STATE.exists() else []
    for name, ok, detail in results:
        print(("PASS " if ok else "FAIL ") + name + (f"  — {detail}" if detail else ""))
    passed = sum(bool(ok) for _, ok, _ in results)
    print(f"통과 {passed} / 전체 11 (실행된 {len(results)}개)")
    STATE.unlink(missing_ok=True)
    return 0 if passed == 11 else 1


if __name__ == "__main__":
    raise SystemExit(main())
