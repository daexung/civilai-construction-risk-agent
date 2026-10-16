"""One bounded loopback HTTP reproduction for the local item-labor answer flow."""
from __future__ import annotations

import faulthandler
import io
import json
import os
import subprocess
import sys
import time
import urllib.error
import urllib.request
import zipfile
from pathlib import Path

faulthandler.enable()
faulthandler.dump_traceback_later(10, repeat=True)

QUERY = "시멘트 액체방수를 바닥에 300㎡ 바르려면 인력 품은 어떻게 잡아야 해?"
PORT = 8012


def _request(opener, method: str, url: str, body: dict | None = None,
             headers: dict[str, str] | None = None):
    data = json.dumps(body, ensure_ascii=False).encode("utf-8") if body is not None else None
    request = urllib.request.Request(url, data=data, headers=headers or {}, method=method)
    try:
        with opener.open(request, timeout=5) as response:
            return response.status, response.read(), dict(response.headers)
    except urllib.error.HTTPError as error:
        return error.code, error.read(), dict(error.headers)


def _json(body: bytes) -> dict:
    return json.loads(body.decode("utf-8"))


def _choices(analysis: dict) -> list[dict]:
    actions = {"requested": "include_work", "background": "keep_background",
               "condition": "keep_condition", "excluded": "keep_excluded"}
    decisions = []
    for span in analysis["scope_spans"]:
        action = actions.get(span["classification"], "exclude_work")
        decision = {"span_id": span["span_id"], "action": action}
        if span["classification"] == "unconfirmed" and span.get("work_ids"):
            decision.update(action="include_work", work="방수")
        decisions.append(decision)
    return decisions


def main() -> int:
    root = Path(__file__).resolve().parents[1]
    base = f"http://127.0.0.1:{PORT}"
    opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))
    deadline = time.monotonic() + 52
    stdout_path = root / ".trial-answer-http.server.stdout.log"
    stderr_path = root / ".trial-answer-http.server.stderr.log"
    server_code = (
        "import faulthandler; faulthandler.enable(); faulthandler.dump_traceback_later(10, repeat=True); "
        "from backend.shared import embedding; from backend.knowledge import embed_index; "
        "defn=lambda *a,**k: (_ for _ in ()).throw(AssertionError('external embedding call')); "
        "embedding.embed_texts=defn; embed_index.make_client=defn; "
        f"import uvicorn; uvicorn.run('backend.api.item_labor_trial:app', host='127.0.0.1', port={PORT}, log_level='info')"
    )
    env = os.environ.copy()
    env["LOCAL_TRIAL_ITEM_LABOR"] = "on"
    env["AGENT_OFFLINE"] = "1"
    env["AGENT_LLM"] = "off"
    with stdout_path.open("wb") as server_stdout, stderr_path.open("wb") as server_stderr:
        server = subprocess.Popen([sys.executable, "-B", "-u", "-c", server_code],
                                  cwd=root, env=env, stdout=server_stdout, stderr=server_stderr)
        try:
            ready = False
            while time.monotonic() < deadline and server.poll() is None:
                try:
                    status, body, _ = _request(opener, "GET", base + "/api/item-labor/health")
                    if status == 200:
                        ready = True
                        print("HEALTH_HTTP 200", flush=True)
                        break
                except (OSError, TimeoutError):
                    time.sleep(0.25)
            if not ready:
                print(f"SERVER_NOT_READY pid={server.pid} exit={server.poll()}", flush=True)
                return 10

            status, body, _ = _request(opener, "POST", base + "/api/item-labor/analyze",
                                       {"text": QUERY}, {"Content-Type": "application/json"})
            print("ANALYZE_HTTP", status, flush=True)
            if status != 200:
                print(body.decode("utf-8", "replace")[:1200], flush=True)
                return 11
            trial = _json(body)
            headers = {"Content-Type": "application/json", "X-Trial-Owner": trial["owner_token"]}
            status, body, _ = _request(opener, "POST", f"{base}/api/item-labor/{trial['trial_id']}/confirm",
                                       {"decisions": _choices(trial["analysis"])}, headers)
            print("CONFIRM_HTTP", status, flush=True)
            if status != 200:
                print(body.decode("utf-8", "replace")[:1200], flush=True)
                return 12
            item = _json(body)["items"][0]
            print("CONFIRMED_ITEM", json.dumps({"item_id": item["item_id"],
                  "request_text": item["request_text"], "source_spans": item["source_spans"]}, ensure_ascii=False), flush=True)

            status, body, _ = _request(opener, "POST",
                f"{base}/api/item-labor/{trial['trial_id']}/items/{item['item_id']}/search", {}, headers)
            print("SEARCH_HTTP", status, flush=True)
            if status != 200:
                print(body.decode("utf-8", "replace")[:1200], flush=True)
                return 13
            item = _json(body)["items"][0]
            print("SEARCH_RESULT", json.dumps({"status": item["search_status"],
                "search_info": item.get("search_info"),
                "candidates": [{k: c.get(k) for k in ("section_no", "section", "spec_id")}
                               for c in item.get("candidates", [])],
                "questions": [{k: q.get(k) for k in ("question_id", "version", "field")}
                              for q in item.get("questions", [])]}, ensure_ascii=False), flush=True)
            if item["search_status"] != "COMPLETE" or not (item.get("search_info") or {}).get("query_cache_hit"):
                print("STOP: exact query cache miss; answer not submitted", flush=True)
                return 14
            candidate = next((c for c in item.get("candidates", [])
                              if str(c.get("section_no", "")) == "6-4-1"), None)
            question = next((q for q in item.get("questions", []) if q.get("field") == "work"), None)
            if candidate is None or question is None:
                print("STOP: 6-4-1 choice or work question missing; no forced selection", flush=True)
                return 15
            answer_body = {"question_id": question["question_id"], "version": question["version"],
                           "value": candidate["section"]}
            print("ANSWER_REQUEST", json.dumps(answer_body, ensure_ascii=False), flush=True)
            status, body, _ = _request(opener, "POST",
                f"{base}/api/item-labor/{trial['trial_id']}/items/{item['item_id']}/answer",
                answer_body, headers)
            print("ANSWER_HTTP", status, flush=True)
            assert status == 200, f"answer HTTP status {status}: {body[:1200]!r}"
            if status != 200:
                print("ANSWER_ERROR_BODY", body.decode("utf-8", "replace")[:1500], flush=True)
                return 16
            item = _json(body)["items"][0]
            assert item.get("selected_specification", {}).get("section_no") == "6-4-1"
            assert item.get("status") == "CALCULATED"
            print("ANSWER_SUCCESS", json.dumps({"selected_specification": item.get("selected_specification"),
                "status": item.get("status"), "questions": item.get("questions", []),
                "search_info": item.get("search_info")}, ensure_ascii=False), flush=True)

            if time.monotonic() >= deadline:
                print("FOLLOWUP_SKIPPED: time budget reached", flush=True)
                return 0
            estimate_body = {"basis_date": time.strftime("%Y-%m-%d"), "common_conditions": {}}
            status, body, _ = _request(opener, "POST",
                f"{base}/api/item-labor/{trial['trial_id']}/estimate", estimate_body, headers)
            print("ESTIMATE_HTTP", status, flush=True)
            assert status == 200, f"estimate HTTP status {status}: {body[:1200]!r}"
            if status == 200:
                summary = _json(body)
                assert summary.get("status") == "PARTIAL", summary.get("status")
                print("ESTIMATE_RESULT", json.dumps({"status": summary.get("status"),
                    "basis_date": summary.get("basis_date"), "item_count": len(summary.get("items", []))},
                    ensure_ascii=False), flush=True)
                status, xlsx, xheaders = _request(opener, "GET",
                    f"{base}/api/item-labor/{trial['trial_id']}/estimate.xlsx", headers=headers)
                assert status == 200, f"xlsx HTTP status {status}: {xlsx[:1200]!r}"
                assert zipfile.is_zipfile(io.BytesIO(xlsx)), "xlsx response is not a ZIP archive"
                xlsx_path = root / ".trial" / "multi-trade-answer.xlsx"
                xlsx_path.parent.mkdir(parents=True, exist_ok=True)
                xlsx_path.write_bytes(xlsx)
                with zipfile.ZipFile(io.BytesIO(xlsx)) as archive:
                    assert archive.testzip() is None, "xlsx ZIP member failed CRC validation"
                from openpyxl import load_workbook

                workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
                try:
                    assert workbook.sheetnames, "xlsx contains no sheets"
                    sheetnames = workbook.sheetnames
                finally:
                    workbook.close()
                print("XLSX_OPENPYXL_OK", "SHEETS", json.dumps(sheetnames, ensure_ascii=False),
                      "BYTES", len(xlsx), "PATH", str(xlsx_path), flush=True)
            else:
                print("ESTIMATE_BODY", body.decode("utf-8", "replace")[:1200], flush=True)
            return 0
        finally:
            if server.poll() is None:
                server.terminate()
                try:
                    server.wait(timeout=5)
                except subprocess.TimeoutExpired:
                    server.kill()
                    server.wait(timeout=2)
            faulthandler.cancel_dump_traceback_later()


if __name__ == "__main__":
    raise SystemExit(main())
