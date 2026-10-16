"""Stage 4 local-only multi-trade aggregate and workbook checks."""
from __future__ import annotations
import os, sys
from pathlib import Path
from io import BytesIO
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(1, str(ROOT / "evals"))
os.environ["AGENT_OFFLINE"] = "1"
os.environ["AGENT_LLM"] = "off"
os.environ["LOCAL_TRIAL_ITEM_LABOR"] = "on"
import check_item_labor_trial as base
from openpyxl import load_workbook
from openpyxl import Workbook
from backend.agent.rules.specs import load_specs
from backend.api.tables import _estimate_sheet
from backend.agent.estimate.state import new_session, set_common
import backend.agent.estimate.aggregate as aggregate_module
from backend.api import item_labor_trial as api
from unittest.mock import patch

def main():
    api._TRIALS.clear()
    text = "수성페인트 롤러칠 2회 벽면100㎡ + 실내 방수모르타르10㎜ 바닥20㎡"
    with base.LoopbackClient() as client:
        started, confirmed = base.confirm_text(client, text)
        for item in confirmed["items"]:
            spec_no = "11-2-2" if item["task"] == "도장" else "6-4-3"
            spec = next(s for s in load_specs().values() if s["section_no"] == spec_no)
            base.choose_work_and_compute(client, started["trial_id"], item, [base.hit(spec)])
        url = f"/api/item-labor/{started['trial_id']}/estimate"
        payload = {"basis_date": "2026-10-16", "common_conditions": {}}
        with patch.object(aggregate_module, "calculate_cost_statement",
                          wraps=aggregate_module.calculate_cost_statement) as statement_spy:
            response = client.post(url, json=payload)
        assert response.status_code == 200, response.text
        result = response.json()
        assert result["status"] == "PARTIAL", result
        assert result["aggregate"]["direct_total"] == 1117900, result["aggregate"]
        assert result["statement"]["totals"]["contract_amount"] == 2100289
        assert result["aggregate"]["included"] == ["w2_5", "w25_27"]
        assert all(item["priced"]["reference_amounts"]["subtotals"]["재료비"] is None for item in result["items"])
        assert statement_spy.call_count == 1
        assert statement_spy.call_args.args[0]["reference_amounts"]["total"] == 1117900
        original_labor_version = result["items"][0]["rate_version"]["id"]
        changed_date = client.post(url, json={"basis_date": "2026-01-01", "common_conditions": {}}).json()
        assert changed_date["basis_date"] == "2026-01-01"
        assert changed_date["items"][0]["rate_version"]["id"] != original_labor_version
        result = client.post(url, json=payload).json()
        again = client.post(url, json=payload).json()
        assert again["aggregate"]["direct_total"] == 1117900
        export = client.get(f"/api/item-labor/{started['trial_id']}/estimate.xlsx")
        assert export.status_code == 200, export.text
        example = ROOT / "docs" / "reviews" / "multi-trade-summary-example.xlsx"
        example.write_bytes(export.content)
        book = load_workbook(BytesIO(export.content), data_only=True)
        rows = list(book["공종별상세"].iter_rows(min_row=5, values_only=True))
        assert [row[0] for row in rows] == ["w2_5", "w25_27"], rows
        assert [row[7] for row in rows] == [727000, 390900], rows
        assert [row[4] for row in rows] == ["미산정", "미산정"], rows
        cost_rows = {row[0].value: row[2].value for row in book["원가계산서"].iter_rows(min_row=5)
                     if row[0].value in ("순공사원가", "총원가", "도급액")}
        assert cost_rows == {"순공사원가": 1537323, "총원가": 1909354, "도급액": 2100289}, cost_rows
        unit_values = [str(row[0].value) for row in book["일위대가"].iter_rows(min_row=6)
                       if row[0].value is not None]
        assert any("[w2_5]" in value for value in unit_values)
        assert any("[w25_27]" in value for value in unit_values)
        basis_values = [str(row[0].value) for row in book["산출근거"].iter_rows(min_row=5)
                        if row[0].value is not None]
        assert any("[w2_5]" in value for value in basis_values)
        assert any("[w25_27]" in value for value in basis_values)
        basis_rows = list(book["산출근거"].iter_rows(min_row=5, values_only=True))
        for item_id in ("w2_5", "w25_27"):
            item_unit_basis = [row for row in basis_rows
                               if str(row[0]).startswith(f"[{item_id}]") and "당 (" in str(row[0])]
            assert item_unit_basis and item_unit_basis[0][2] and item_unit_basis[0][3], item_unit_basis
        bill_rows = list(book["내역서"].iter_rows(min_row=6, values_only=True))
        assert len(bill_rows) >= 3 and "[w2_5]" in bill_rows[0][0] and "[w25_27]" in bill_rows[1][0]
        assert [bill_rows[0][2], bill_rows[0][3], bill_rows[1][2], bill_rows[1][3]] == ["㎡", 100, "㎡", 20]
        assert bill_rows[0][5] == bill_rows[1][5] == "미산정"
        assert bill_rows[2][11] == result["aggregate"]["direct_total"]
        headers = [str(row[0].value) for row in book["일위대가"].iter_rows(min_row=6)
                   if row[0].value and str(row[0].value).startswith("제 ")]
        assert len(headers) == 2 and "[w2_5]" in headers[0] and "㎡ 당" in headers[0]
        assert "[w25_27]" in headers[1] and "㎡ 당" in headers[1]
        rate_names = [str(row[1].value) for row in book["단가대비표"].iter_rows(min_row=5)
                      if row[1].value is not None]
        assert any(name.startswith("[w2_5]") for name in rate_names)
        assert any(name.startswith("[w25_27]") for name in rate_names)
        rate_rows = list(book["단가대비표"].iter_rows(min_row=5, values_only=True))
        assert any(str(row[1]).startswith("[w2_5]") and row[5] and row[6] for row in rate_rows)
        assert any(str(row[1]).startswith("[w25_27]") and row[5] and row[6] for row in rate_rows)
        version_rows = [row for row in basis_rows if "기준 버전" in str(row[0])]
        assert any(row[1] == "2026H2" for row in version_rows)
        assert any(row[1] == "260413" for row in version_rows)
        assert "복수 공종 견적" in book["견적서"].cell(2, 1).value
        assert "부분 산정" in book["견적서"].cell(4, 1).value
        assert "부문 기본값" in book["견적서"].cell(4, 1).value
        assert "공사 종류" in book["견적서"].cell(4, 1).value and "work_category" not in book["견적서"].cell(4, 1).value
        assert "\n" in book["견적서"].cell(4, 1).value and book["견적서"].row_dimensions[4].height >= 90
        assert book["견적서"].cell(6, 1).value.startswith("부분 산정 금액(미산정 항목 제외)")
        assert "포함된 재료비 0원 · 별도 공종 재료 미산정" in book["견적서"].cell(9, 3).value
        cost_materials = [row for row in book["원가계산서"].iter_rows(min_row=5)
                          if row[0].value == "재료비" and row[1].value == "직접재료비"]
        assert cost_materials and "별도 공종 재료 미산정" in cost_materials[0][4].value
        print("PASS independent direct subtotal matched aggregate: 1,117,900원; PARTIAL retained")
        print("PASS repeated aggregate and workbook preserve item order, IDs, and per-item totals")
        print("PASS workbook generated with multi-item detail sheet")
        print("Saved sanitized local workbook example:", example)
        print("Statement totals:", result["statement"]["totals"])
        trial = api._get_trial(started["trial_id"], started["owner_token"])
        previous_keys = {item_id: trial["items"][item_id]["session"]["items"][item_id].get("price_key")
                         for item_id in trial["item_order"]}
        with patch("backend.agent.estimate.tools.price_item_for_aggregate",
                   side_effect=[None, ValueError("simulated item price failure")]):
            failed_estimate = client.post(url, json=payload)
        assert failed_estimate.status_code == 422
        assert "summary" not in trial
        assert previous_keys == {item_id: trial["items"][item_id]["session"]["items"][item_id].get("price_key")
                                 for item_id in trial["item_order"]}
        rebuilt = client.post(url, json=payload)
        assert rebuilt.status_code == 200
        print("PASS item failure clears old summary and does not commit a half-updated child state")
        # A quantity change invalidates the aggregate/export but keeps the other item's labor result.
        paint_id, water_id = "w2_5", "w25_27"
        water_labor = trial["items"][water_id]["session"]["items"][water_id]["computed_result"]
        paint_state = trial["items"][paint_id]["session"]["items"][paint_id]
        edit = client.post(f"/api/item-labor/{started['trial_id']}/items/{paint_id}/edit", json={
            "input_revision": paint_state["input_revision"], "quantity": {"value": "101", "unit": "㎡"}})
        assert edit.status_code == 200, edit.text
        invalidated = client.get(f"/api/item-labor/{started['trial_id']}/estimate.xlsx")
        assert invalidated.status_code == 409
        assert trial["items"][water_id]["session"]["items"][water_id]["computed_result"] == water_labor
        print("PASS quantity edit invalidates aggregate export and preserves the other item's labor result")

    # Same-trade repetition keeps separate quantities and stable IDs through pricing.
    api._TRIALS.clear()
    with base.LoopbackClient() as client:
        same_started, same = base.confirm_text(client, "내부 도장 100㎡ 그리고 외부 도장 20㎡")
        same_ids = [item["item_id"] for item in same["items"]]
        paint = next(s for s in load_specs().values() if s["section_no"] == "11-2-2")
        for item in same["items"]:
            base.choose_work_and_compute(client, same_started["trial_id"], item, [base.hit(paint)])
        same_result = client.post(f"/api/item-labor/{same_started['trial_id']}/estimate",
                                  json={"basis_date": "2026-10-16"}).json()
        assert len(set(same_ids)) == 2 and same_result["status"] == "PARTIAL"
        assert same_result["aggregate"]["direct_total"] == sum(
            int(item["priced"]["reference_amounts"]["total"]) for item in same_result["items"])
        print("PASS repeated same-trade instances keep separate stable IDs, quantities, and direct totals")

    # A supported sibling remains visible when another item is still pending or unsupported.
    api._TRIALS.clear()
    with base.LoopbackClient() as client:
        pending_start, pending_trial = base.confirm_text(client, "도장 100㎡ 그리고 방수 20㎡")
        supported = next(item for item in pending_trial["items"] if item["task"] == "도장")
        base.choose_work_and_compute(client, pending_start["trial_id"], supported, [base.hit(paint)])
        pending = client.post(f"/api/item-labor/{pending_start['trial_id']}/estimate",
                              json={"basis_date": "2026-10-16"}).json()
        assert pending["status"] == "PENDING" and any(item.get("priced") for item in pending["items"])
        assert client.get(f"/api/item-labor/{pending_start['trial_id']}/estimate.xlsx").status_code == 409
        pending_item = next(item for item in pending_trial["items"] if item["task"] == "방수")
        water_spec = next(s for s in load_specs().values() if s["section_no"] == "6-4-3")
        searched = client.search(pending_start["trial_id"], pending_item["item_id"], [base.hit(water_spec)])
        assert searched.status_code == 200
        pending_item_state = next(item for item in searched.json()["items"] if item["item_id"] == pending_item["item_id"])
        assert pending_item_state["status"] == "NEEDS_INPUT" and pending_item_state["questions"]
        pending = client.post(f"/api/item-labor/{pending_start['trial_id']}/estimate",
                              json={"basis_date": "2026-10-16"}).json()
        assert pending["status"] == "PENDING"
        assert pending["partial_estimate"]["direct_total"] > 0
        assert pending["partial_estimate"]["complete_contract_amount"] is None
        assert pending["partial_estimate"]["unpriced_by_category"].get("재료비")
        assert any(item["status"] == "NEEDS_INPUT" and item["questions"] for item in pending["items"])
        unsupported = pending_item
        unsupported_hit = {"section_no": "99-99-99", "division": "건축", "rank": 1,
                           "section": "99-99-99 mock", "chunk_id": "mock-none", "score": 1}
        unsupported_response = client.search(pending_start["trial_id"], unsupported["item_id"], [unsupported_hit])
        assert unsupported_response.status_code == 200
        partial = client.post(f"/api/item-labor/{pending_start['trial_id']}/estimate",
                              json={"basis_date": "2026-10-16"}).json()
        assert partial["status"] == "PARTIAL" and any(item["status"] == "UNSUPPORTED" for item in partial["items"])
        assert partial["aggregate"]["contract_amount"] is None
        assert partial["aggregate"]["complete_contract_amount"] is None
        assert partial["aggregate"]["partial_contract_amount"] is not None
        print("PASS NEEDS_INPUT and unsupported items retain status/questions; priced-only subtotal is separate")

    api._TRIALS.clear()
    with base.LoopbackClient() as client:
        unavailable_start, unavailable_trial = base.confirm_text(client, "도장 100㎡")
        unsupported_hit = {"section_no": "99-99-99", "division": "건축", "rank": 1,
                           "section": "99-99-99 mock", "chunk_id": "mock-none", "score": 1}
        base.choose_work_and_compute(client, unavailable_start["trial_id"], unavailable_trial["items"][0],
                                     [unsupported_hit])
        unavailable = client.post(f"/api/item-labor/{unavailable_start['trial_id']}/estimate",
                                  json={"basis_date": "2026-10-16"}).json()
        assert unavailable["status"] == "NO_RESULT"
        assert client.get(f"/api/item-labor/{unavailable_start['trial_id']}/estimate.xlsx").status_code == 409
        print("PASS all-unavailable items return NO_RESULT and no zero-value XLSX")

    # Construct settled item sessions to isolate common-condition and rate-version behavior.
    def parent_fixture(second_section="1-3-1", second_rate="2026H2"):
        specs = load_specs()
        building = next(s for s in specs.values() if s["section_no"] == "11-2-2")
        civil = next((s for s in specs.values() if s["section_no"] == second_section
                      and s["division"] == "토목"), None)
        if civil is None:
            civil = next(s for s in specs.values() if s["division"] == "토목")
        parent = new_session({"basis_date": "2026-10-16", "items": [
            {"item_id": "a", "request_text": "building"}, {"item_id": "b", "request_text": "civil"}]})
        for item_id, spec, rate in (("a", building, "2026H2"), ("b", civil, second_rate)):
            state = parent["items"][item_id]
            state.update(selected_spec_id=spec["id"], selection={"confirmed": True}, status="PRICED",
                         input_revision=0, result_revision=0,
                         priced_result={"reference_amounts": {"subtotals": {"재료비": None, "노무비": 100, "경비": 0},
                                                              "total": 100, "volume": "1"},
                                        "rate_version": {"id": rate},
                                        "equipment_rate_version": {"version": "2026"}})
        return parent

    common_parent = parent_fixture()
    common_result = aggregate_module.aggregate(common_parent)
    assert common_result["aggregate_result"]["status"] == "NEEDS_COMMON"
    assert "common:work_category" in common_result["pending_questions"]
    api._TRIALS["common-case"] = {"trial_id": "common-case", "summary": common_result}
    with base.LoopbackClient() as client:
        answered = client.post("/api/item-labor/common-case/common-answer",
                               json={"field": "work_category", "value": "주택 외 건축"})
        assert answered.status_code == 200, answered.text
        assert answered.json()["status"] in ("COMPLETE", "PARTIAL")
        assert answered.json()["common_conditions"]["work_category"]["source"] == "answer"
    set_common(common_result, "work_category", "주택 외 건축")
    with patch.object(aggregate_module, "calculate_cost_statement",
                      wraps=aggregate_module.calculate_cost_statement) as spy:
        final = aggregate_module.aggregate(common_result)
    assert final["aggregate_result"]["status"] == "PARTIAL"
    assert final["aggregate_result"]["direct_total"] == 200
    assert spy.call_count == 1 and spy.call_args.args[0]["reference_amounts"]["total"] == 200
    conflict = parent_fixture(second_rate="2026H1")
    set_common(conflict, "work_category", "주택 외 건축")
    assert aggregate_module.aggregate(conflict)["aggregate_result"]["status"] == "RATE_VERSION_CONFLICT"
    missing_version = parent_fixture()
    set_common(missing_version, "work_category", "주택 외 건축")
    missing_version["items"]["b"]["priced_result"]["rate_version"] = None
    assert aggregate_module.aggregate(missing_version)["aggregate_result"]["status"] == "RATE_VERSION_CONFLICT"
    unavailable = new_session({"items": [{"item_id": "x", "request_text": "unsupported"}]})
    unavailable["items"]["x"]["status"] = "NOT_FOUND"
    assert aggregate_module.aggregate(unavailable)["aggregate_result"]["status"] == "NO_RESULT"
    print("PASS common-condition conflict prompts; one shared statement call receives 200원 combined direct cost")
    print("PASS conflicting or missing rate versions block aggregation; no priced items produce NO_RESULT")

    # Guard the existing single-trade cover wording while the multi-trade label changes.
    single = {"tables": {}, "statement": {"totals": {"contract_amount": 123456}},
              "result": {"review_status": "AI 초안 · 검토 전"}, "conditions": [],
              "work": {"title": "단일 공종"}, "basis_date": "2026-10-16"}
    single_book = Workbook()
    _estimate_sheet(single_book, single)
    assert single_book["견적서"]["A6"].value.startswith("견적금액 :")
    assert "AI 초안 · 검토 전" in single_book["견적서"]["A4"].value
    print("PASS single-trade cover wording and review badge remain unchanged")
    return 0

if __name__ == "__main__":
    raise SystemExit(main())
