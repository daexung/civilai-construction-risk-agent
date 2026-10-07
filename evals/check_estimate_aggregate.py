"""복수 공종 합산·통합 원가계산서·Excel(단계 C1)을 구조화된 입력으로 검사한다.

오프라인 단어 검색·LLM 꺼짐. 운영 DB·유료 API·운영 라우트는 쓰지 않는다.
단일 공종 결과는 main 9a2af9f 기준 파일의 원가계산서·Excel 숫자와 비교한다.
"""

from __future__ import annotations

import copy
import io
import json
import os
import sys
from pathlib import Path

os.environ["AGENT_OFFLINE"] = "1"
os.environ["AGENT_LLM"] = "off"
os.environ.pop("INDEX_CONFIG", None)
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from langgraph.checkpoint.memory import MemorySaver  # noqa: E402
from langgraph.types import Command  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

import backend.agent.estimate.flow as flow  # noqa: E402
from backend.agent.estimate.aggregate import aggregate  # noqa: E402
from backend.agent.estimate.state import new_session, set_common, set_quantity  # noqa: E402
from backend.agent.tools.calc.cost_statement import calculate_cost_statement  # noqa: E402
from backend.api.estimate_output import EstimateNotReady, build_estimate_xlsx, build_output  # noqa: E402

BASIS = "2026-10-06"
BASELINE = json.loads((ROOT / "evals/fixtures/baseline_main_9a2af9f.json").read_text(encoding="utf-8"))["results"]
REBAR = "철근구조물 레미콘 인력운반 타설 150㎥"
PLAIN = "무근구조물 레미콘 인력운반 타설 50㎥"
PUMP = "철근콘크리트 벽체 260㎥ 32m 펌프차로 타설 비용"
READY = {"work": "6-1-1", "scattered_small_volume": False, "concrete_supply": "관급"}


def finish(session, picks: dict, rounds: int = 6):
    """질문이 없어질 때까지 picks로 답한다. picks[(item_id, field)]가 picks[field]보다 우선."""
    for _ in range(rounds):
        session = aggregate(flow.run_items(session))
        if not session["pending_questions"]:
            return session
        answers = []
        for question in session["pending_questions"].values():
            field, item_id = question["field"], question["item_id"]
            value = picks.get((item_id, field), picks.get(field))
            choices = question["allowed_values"] if isinstance(question["allowed_values"], list) else []
            if field == "work" and value is not None:
                value = next((c for c in choices if value in c), value)
            if value is None:
                value = choices[0] if choices else "10"
            answers.append({"question_id": question["question_id"], "version": question["version"], "value": value})
        session = flow.apply_answers(session, {"answers": answers})
    return aggregate(flow.run_items(session))


def excel_cells(data: bytes) -> dict:
    book = load_workbook(io.BytesIO(data))
    return {name: [[c for c in row if c not in (None, "")] for row in book[name].iter_rows(values_only=True)
                   if any(c not in (None, "") for c in row)] for name in book.sheetnames}


def numbers(sheets: dict) -> dict:
    return {name: [[c for c in row if isinstance(c, (int, float))] for row in rows] for name, rows in sheets.items()}


def main_sent(scenario_id: str) -> dict:
    sent = {}
    for step in BASELINE[scenario_id]:
        sent.update((step.get("op") or {}).get("answers") or {})
    return sent


def main() -> int:
    checks: list[tuple[str, bool, str]] = []
    compute_calls = {"n": 0}
    real_compute = flow.compute

    def counting(state):
        compute_calls["n"] += 1
        return real_compute(state)
    flow.compute = counting

    # 1) 단일 공종: main과 같은 원가계산서·Excel 숫자
    for scenario_id, request in (("single-readymix-rebar", REBAR), ("single-pump", PUMP)):
        steps = BASELINE[scenario_id]
        final = [s["response"] for s in steps if "response" in s][-1]
        main_excel = next(s["excel"] for s in steps if "excel" in s)
        session = finish(new_session({"basis_date": BASIS, "items": [{"request_text": request}]}), main_sent(scenario_id))
        data, filename, output = build_estimate_xlsx(session)
        ours = excel_cells(data)
        same_statement = ([[l["name"], l["amount"], l["status"]] for l in session["statement"]["lines"]] == final["statement"]["lines"]
                          and session["statement"]["totals"] == final["statement"]["totals"])
        same_excel = numbers(ours) == numbers(main_excel)
        text_diff = [name for name in ours if [[c for c in r if not isinstance(c, (int, float))] for r in ours[name]]
                     != [[c for c in r if not isinstance(c, (int, float))] for r in main_excel.get(name, [])]]
        checks.append((f"S1 {scenario_id}: 원가계산서·Excel 숫자 main과 같음", same_statement and same_excel and output["status"] == final["status"],
                       f"도급액 {session['statement']['totals']['contract_amount']:,} (main {final['statement']['totals']['contract_amount']:,}), "
                       f"Excel 글자 차이 시트={text_diff or '없음'}"))

    # 2) 두 공종 합산
    two = finish(new_session({"basis_date": BASIS, "items": [{"request_text": REBAR}, {"request_text": PLAIN}]}), READY)
    agg, statement = two["aggregate_result"], two["statement"]
    lines = {line["name"]: line for line in statement["lines"]}
    rebar_ref = two["items"]["i1"]["priced_result"]["reference_amounts"]
    plain_ref = two["items"]["i2"]["priced_result"]["reference_amounts"]
    direct_labor = int(rebar_ref["subtotals"]["노무비"]) + int(plain_ref["subtotals"]["노무비"])
    independent = calculate_cost_statement(
        {"reference_amounts": {"subtotals": {"재료비": 0, "노무비": direct_labor,
                                              "경비": int(rebar_ref["subtotals"]["경비"]) + int(plain_ref["subtotals"]["경비"])},
                               "total": int(rebar_ref["total"]) + int(plain_ref["total"]), "volume": None},
         "rate_version": {}, "unpriced": [], "excluded": []},
        {name: entry["value"] for name, entry in two["common_conditions"].items()}, BASIS)
    checks.append(("T1 두 공종: 직접비 합산 후 간접비 1회", agg["included"] == ["i1", "i2"] and lines["직접노무비"]["amount"] == direct_labor
                   and sum(line["name"] == "간접노무비" for line in statement["lines"]) == 1
                   and statement["totals"]["contract_amount"] == independent["totals"]["contract_amount"]
                   and agg["direct_total"] == int(rebar_ref["total"]) + int(plain_ref["total"]),
                   f"직접비 {agg['direct_total']:,} / 도급액 {statement['totals']['contract_amount']:,}"))
    defaults = {name: entry["source"] for name, entry in two["common_conditions"].items()}
    checks.append(("T2 공통 조건 기본값도 값·출처를 세션에 저장", defaults == {"work_category": "default_division", "duration": "default",
                                                              "contractor_type": "default", "project_scale": "default"}
                   and two["common_conditions"]["work_category"]["value"] == "기타 토목공사", str(defaults)))
    data, filename, output = build_estimate_xlsx(two)
    cells = excel_cells(data)
    bill_rows = [row for row in cells["내역서"] if row and row[0] == "합계"]
    estimate_text = " ".join(str(c) for row in cells["견적서"] for c in row)
    cost_numbers = {c for row in cells["원가계산서"] for c in row if isinstance(c, int)}
    checks.append(("T3 화면용 출력과 Excel이 같은 집계", [b["total"] for b in output["tables"]["bills"]] == [rebar_ref["total"], plain_ref["total"]]
                   and bill_rows and bill_rows[0][-1] == agg["direct_total"]
                   and f"₩{statement['totals']['contract_amount']:,}" in estimate_text
                   and {line["amount"] for line in statement["lines"] if line["amount"]} <= cost_numbers
                   and [item["item_id"] for item in output["items"]] == ["i1", "i2"],
                   f"{filename} 내역서 합계 {bill_rows[0][-1] if bill_rows else None:,}"))

    # 3) 일부 미산정: 명세 없는 공종은 합산하지 않고 이유와 함께 표시
    part = flow.run_items(new_session({"basis_date": BASIS, "items": [{"request_text": REBAR}, {"request_text": PLAIN}]}))
    no_spec = next(c for c in part["items"]["i2"]["candidates"] if not c["has_spec"])
    part = finish(part, {**READY, ("i2", "work"): no_spec["section"]})
    single = finish(new_session({"basis_date": BASIS, "items": [{"request_text": REBAR}]}), READY)
    data, _, output = build_estimate_xlsx(part)
    missing = " ".join(str(c) for row in excel_cells(data)["원가계산서"] for c in row)
    checks.append(("P1 일부 미산정: 부분 견적, 미산정 공종과 이유 표시, 0원으로 넣지 않음",
                   part["aggregate_result"]["status"] == "PARTIAL" and part["aggregate_result"]["included"] == ["i1"]
                   and part["aggregate_result"]["excluded"][0]["status"] == "UNSUPPORTED"
                   and part["statement"]["totals"] == single["statement"]["totals"]
                   and any(entry["name"].startswith("i2 ") for entry in part["statement"]["unpriced"])
                   and "i2 " in missing and len(output["tables"]["bills"]) == 1,
                   f"제외 {part['aggregate_result']['excluded'][0]['reason']} / 도급액 {part['statement']['totals']['contract_amount']:,}"))

    # 4) 전부 미산정, 질문 대기
    none = flow.run_items(new_session({"basis_date": BASIS, "items": [{"request_text": "아까 거 다시 계산해줘 100㎥"}]}))
    none = aggregate(none)
    try:
        build_estimate_xlsx(none)
        none_excel = "만들어짐"
    except EstimateNotReady as exc:
        none_excel = str(exc)
    checks.append(("N1 전부 미산정: 0원 견적·Excel을 만들지 않음", none["aggregate_result"]["status"] == "NO_RESULT"
                   and none["statement"] is None and "NO_RESULT" in none_excel and flow.session_complete(none), none_excel))
    pending = aggregate(flow.run_items(new_session({"basis_date": BASIS, "items": [{"request_text": REBAR}]})))
    try:
        build_estimate_xlsx(pending)
        pending_excel = "만들어짐"
    except EstimateNotReady as exc:
        pending_excel = str(exc)
    checks.append(("N2 질문 대기: 견적 미확정, Excel 없음", pending["aggregate_result"]["status"] == "PENDING"
                   and pending["statement"] is None and not flow.session_complete(pending) and "PENDING" in pending_excel,
                   f"남은 질문 {list(pending['pending_questions'])}"))

    # 5) 수정 범위: 물량 → 그 공종 계산 + 집계, 공통 조건 → 원가계산서만
    edited = copy.deepcopy(two)
    set_quantity(edited["items"]["i1"], "200", "㎥", "answer")
    stale = aggregate(edited)   # 다시 계산하기 전 집계
    try:
        build_estimate_xlsx(stale)
        stale_excel = "만들어짐"
    except EstimateNotReady:
        stale_excel = "거부"
    checks.append(("R1 낡은 결과는 합계·Excel에 남지 않음", stale["aggregate_result"]["status"] == "PENDING" and stale["statement"] is None
                   and stale_excel == "거부" and stale["aggregate_result"]["waiting"][0]["item_id"] == "i1", str(stale["aggregate_result"]["waiting"])))
    before = compute_calls["n"]
    edited = aggregate(flow.run_items(edited))
    checks.append(("R2 물량 수정: 해당 공종만 다시 계산, 다른 공종 결과 유지", compute_calls["n"] - before == 1
                   and edited["items"]["i2"]["priced_result"] == two["items"]["i2"]["priced_result"]
                   and edited["items"]["i1"]["priced_result"]["reference_amounts"]["volume"] == "200"
                   and edited["statement"]["totals"]["contract_amount"] != two["statement"]["totals"]["contract_amount"],
                   f"계산 {compute_calls['n'] - before}회, 도급액 {two['statement']['totals']['contract_amount']:,} → {edited['statement']['totals']['contract_amount']:,}"))
    common = copy.deepcopy(two)
    set_common(common, "contractor_type", "전문건설업")
    before = compute_calls["n"]
    common = aggregate(common)
    checks.append(("R3 공통 조건 수정: 공종 계산 유지, 원가계산서만 갱신", compute_calls["n"] == before
                   and common["items"] == two["items"] and common["statement"]["conditions"]["contractor_type"] == "전문건설업"
                   and common["statement"]["totals"]["contract_amount"] != two["statement"]["totals"]["contract_amount"],
                   f"계산 0회, 도급액 → {common['statement']['totals']['contract_amount']:,}"))

    # 6) 공종 부문이 달라 공사 종류를 정할 수 없으면 확인
    mixed = new_session({"basis_date": BASIS, "items": [{"request_text": REBAR}, {"request_text": "자동문 3개소 설치"}]})
    item_picks = {**READY, ("i2", "work"): "10-1-7"}
    for _ in range(4):   # 공종 질문에만 답하고 공통 조건 질문은 남겨 둔다
        mixed = aggregate(flow.run_items(mixed))
        item_questions = [q for q in mixed["pending_questions"].values() if q["scope"] == "item"]
        if not item_questions:
            break
        mixed = flow.apply_answers(mixed, {"answers": [
            {"question_id": q["question_id"], "version": q["version"],
             "value": next(c for c in q["allowed_values"] if item_picks.get((q["item_id"], "work"), item_picks["work"]) in c)
             if q["field"] == "work" else item_picks[q["field"]]} for q in item_questions]})
    common_question = mixed["pending_questions"].get("common:work_category") or {}
    resolved = aggregate(flow.run_items(flow.apply_answers(mixed, {"answers": [
        {"question_id": "common:work_category", "version": common_question.get("version"), "value": "주택 외 건축"}]})))
    checks.append(("C1 부문 충돌(공통 6-1-1 + 건축 10-1-7): 첫 공종 기준으로 정하지 않고 확인",
                   mixed["aggregate_result"]["status"] == "NEEDS_COMMON" and mixed["statement"] is None and common_question
                   and resolved["aggregate_result"]["status"] in ("COMPLETE", "PARTIAL")
                   and resolved["common_conditions"]["work_category"] == {"value": "주택 외 건축", "source": "answer", "explicit": True},
                   f"질문 이유: {common_question.get('reason', '')[:40]} → 도급액 {(resolved['statement'] or {}).get('totals', {}).get('contract_amount')}"))

    # 7) 0%가 아닌 할증: 6-1-5 에폭시 4∼6층 + 비계 사용
    epoxy = {"type": "신구-콘크리트 접착제바르기", "ceiling_applied": True, "scaffold_used": True, "floor_level": "4∼6층",
             "floor_level_19_plus": 19, "thickness_adjusted": False, "thickness": "1"}
    results = {}
    for choice in ("예", "아니오"):
        session = finish(new_session({"basis_date": BASIS, "items": [
            {"request_text": "에폭시 콘크리트 접착제 바르기 100㎡ 비용", "conditions": epoxy}]}), {"work": "6-1-5", "apply_adj_2": choice})
        item = session["items"]["i1"]
        painter = next(line["applied"] for line in item["computed_result"]["unit_lines"] if line["name"] == "도장공")
        results[choice] = (painter, item["priced_result"]["total"], session["statement"]["totals"]["contract_amount"])
    checks.append(("J1 0%가 아닌 할증: 예/아니오에 따라 품·금액이 달라짐", results["예"][0] == "0.15" and results["아니오"][0] == "0.144"
                   and results["예"][2] > results["아니오"][2],
                   f"예: 도장공 {results['예'][0]}인 1단위 {results['예'][1]} 도급액 {results['예'][2]:,} / "
                   f"아니오: 도장공 {results['아니오'][0]}인 1단위 {results['아니오'][1]} 도급액 {results['아니오'][2]:,}"))

    # 8) 부모 그래프 끝에서 집계까지
    saver = MemorySaver()
    graph = flow.build_estimate_graph(saver)
    config = {"configurable": {"thread_id": "estimate-agg"}}
    state = graph.invoke({"plan": {"basis_date": BASIS, "items": [{"request_text": REBAR}, {"request_text": PLAIN}]}}, config)
    for _ in range(6):
        if not graph.get_state(config).next:
            break
        session = graph.get_state(config).values["estimate"]
        reply = {"answers": [{"question_id": q["question_id"], "version": q["version"],
                              "value": next(c for c in q["allowed_values"] if "6-1-1" in c) if q["field"] == "work" else READY[q["field"]]}
                             for q in session["pending_questions"].values()]}
        state = graph.invoke(Command(resume=reply), config)
    checks.append(("G1 부모 그래프 종료 시 통합 원가계산서", state["estimate"]["aggregate_result"]["status"] == "PARTIAL"
                   and state["estimate"]["statement"]["totals"]["contract_amount"] == statement["totals"]["contract_amount"],
                   f"도급액 {state['estimate']['statement']['totals']['contract_amount']:,}"))

    flow.compute = real_compute
    for name, ok, detail in checks:
        print(("PASS " if ok else "FAIL ") + name + (f"  — {detail}" if detail else ""))
    passed = sum(bool(ok) for _, ok, _ in checks)
    print(f"통과 {passed} / 전체 {len(checks)}")
    return 0 if passed == len(checks) else 1


if __name__ == "__main__":
    raise SystemExit(main())
