"""Real /api/chat flow for formwork height facts, edits, totals, and XLSX export."""
from __future__ import annotations

import os
import unittest
from contextlib import nullcontext
from io import BytesIO
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.api.main import app


class FormworkChatIntegration(unittest.TestCase):
    def test_chat_answers_edit_aggregate_and_export(self):
        with patch.dict(os.environ, {"AGENT_MODE": "tools", "AGENT_LLM": "off", "AGENT_OFFLINE": "1"}), \
             patch("backend.api.usage_limits.processing", return_value=nullcontext((None, "fixture", False))), \
             patch("backend.api.usage_limits.consume", return_value={}), \
             patch("backend.api.usage_limits.status", return_value={}):
            client = TestClient(app, headers={"X-Guest-Session": "formwork-integration-" + "x" * 32})
            response = client.post("/api/chat", json={
                "message": "합판거푸집 100㎡ 비용 견적, 높이 40m, 양중장비 미사용"
            })
            self.assertEqual(response.status_code, 200)
            first = response.json()
            self.assertEqual(first["status"], "MISSING_INFO")
            self.assertTrue(first["questions"])
            self.assertFalse(any(q["name"] in {"height_deduction_percent", "height_deduction_source"}
                                 for q in first["questions"]))

            answer_values = {"complexity": "복잡", "height": "40", "is_roof_slab": False,
                             "crane_used": False, "lifting_equipment_used": False}
            current = first
            for _ in range(8):
                questions = current.get("questions", [])
                if not questions:
                    break
                answers = {q["name"]: answer_values[q["name"]]
                           for q in questions if q["name"] in answer_values}
                self.assertTrue(answers, questions)
                refs = {q["name"]: q["ref"] for q in questions
                        if q["name"] in answers and q.get("ref")}
                response = client.post("/api/chat", json={"thread_id": first["thread_id"],
                    "answers": answers, "refs": refs})
                self.assertEqual(response.status_code, 200, response.text)
                current = response.json()
            self.assertIn(current["status"], ("PARTIAL", "OK", "COMPUTED"), current)
            daily = current["result"]["daily_volume"]
            self.assertEqual(daily["value"], "10.631060024624333494473")
            self.assertEqual(daily["formula"], "기준 30")
            self.assertTrue(any(row["name"] == "volume" and row["value"] == "100"
                                for row in current["inputs"]))
            self.assertEqual(daily["adjustments"][0]["after"], daily["value"])
            self.assertIn("System assumption", daily["adjustments"][0]["reason"])
            self.assertIn("successively", daily["adjustments"][0]["reason"])
            self.assertIn("PDF p.193 (printed p.137)", daily["adjustments"][0]["source"])
            self.assertEqual(current["result"]["calculation_module"]["module_id"], "formwork-adjustments-v1")
            self.assertEqual(current["result"]["calculation_module"]["version"], "1.0.0")
            self.assertEqual(current["statement"]["status"], "PARTIAL")
            original_total = current["statement"]["totals"]["contract_amount"]
            self.assertGreater(original_total, 0)

            edited_response = client.post("/api/chat", json={"thread_id": first["thread_id"],
                "message": "물량을 120㎡로 변경해서 다시 계산해 주세요"})
            self.assertEqual(edited_response.status_code, 200, edited_response.text)
            edited = edited_response.json()
            self.assertIn(edited["status"], ("PARTIAL", "OK", "COMPUTED"), edited)
            self.assertTrue(any(row["name"] == "volume" and row["value"] == "120"
                                for row in edited["inputs"]), edited["inputs"])
            self.assertEqual(edited["result"]["daily_volume"]["value"], daily["value"])
            edited_total = edited["statement"]["totals"]["contract_amount"]
            self.assertNotEqual(edited_total, original_total)
            self.assertEqual(edited["statement"]["status"], "PARTIAL")

            exported = client.get(f"/api/export/{first['thread_id']}.xlsx")
            self.assertEqual(exported.status_code, 200)
            workbook = load_workbook(BytesIO(exported.content), data_only=False)
            self.assertIn("원가계산서", workbook.sheetnames)
            basis = [" ".join(str(cell) for cell in row if cell is not None)
                     for row in workbook["산출근거"].iter_rows(values_only=True)]
            self.assertTrue(any("10.631060024624333494473" in row for row in basis), basis)
            self.assertTrue(any("System assumption" in row and "successively" in row for row in basis), basis)
            self.assertTrue(any("PDF p.193 (printed p.137)" in row for row in basis), basis)
            self.assertTrue(any("formwork-adjustments-v1" in row and "1.0.0" in row
                                and "source SHA256" in row for row in basis), basis)
            statement_rows = list(workbook["원가계산서"].iter_rows(values_only=True))
            self.assertTrue(any(str(edited_total) in " ".join(map(str, row)) for row in statement_rows))
            # The UI consumes the same endpoint result shape and renders these fields under basis details.
            self.assertTrue(edited["result"]["daily_volume"]["adjustments"])


if __name__ == "__main__":
    unittest.main()
