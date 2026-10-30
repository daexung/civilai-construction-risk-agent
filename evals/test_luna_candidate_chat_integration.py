"""Exercise the saved, unregistered Luna candidate through /api/chat and XLSX."""
from __future__ import annotations

import json
import os
import unittest
from contextlib import nullcontext
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.api.main import app
from backend.agent.tools.calc import generated

ROOT = Path(__file__).resolve().parents[1]
CANDIDATE_PATH = ROOT / "evals/runtime_validation/luna_local_2026-10-18/assembled-source-repaired-v7-module.json"


class SavedCandidateChatIntegration(unittest.TestCase):
    def test_source_verified_legacy_migration_without_registration(self):
        from backend.agent.tools.calc.policies import policy_applies
        from pipeline.calculation_modules import load_case_file, run_candidate_case_detailed

        approved_path = ROOT / "data/calculation_modules/approved/formwork-adjustments-v1.json"
        approved_bytes = approved_path.read_bytes()
        legacy = json.loads(approved_bytes)
        migrated = json.loads((ROOT / "evals/runtime_validation/luna_local_2026-10-18/"
                               "formwork-adjustments-v1-provenance-migration-candidate.json").read_text(encoding="utf-8"))
        policy_id = "formwork-6-3-1-height-started-interval-compound-v1"
        case = load_case_file(ROOT / "evals/calculation_modules/formwork_authoring_cases.json")[0]
        self.assertFalse(policy_applies(policy_id, legacy, case["spec"]))
        self.assertTrue(policy_applies(policy_id, migrated, case["spec"]))
        self.assertEqual(migrated["review_status"], "candidate_unreviewed")
        for evaluation_case in load_case_file(ROOT / "evals/calculation_modules/formwork_authoring_cases.json"):
            detail = run_candidate_case_detailed(migrated,
                {**evaluation_case, "allowed_system_policies": [policy_id]})
            self.assertEqual(detail["status"], "passed", evaluation_case["id"])
        self.assertEqual(approved_path.read_bytes(), approved_bytes)

    def test_candidate_answer_edit_aggregate_and_export(self):
        candidate = json.loads(CANDIDATE_PATH.read_text(encoding="utf-8"))
        original_lookup = generated.registered_module_for

        def candidate_lookup(spec):
            if spec.get("division") == "공통" and spec.get("section_no") == "6-3-1":
                spec.setdefault("quantity_model", {}).setdefault("params", {})["module_id"] = candidate["id"]
                return candidate
            return original_lookup(spec)

        with patch.dict(os.environ, {"AGENT_MODE": "tools", "AGENT_LLM": "off", "AGENT_OFFLINE": "1"}), \
             patch("backend.api.usage_limits.processing", return_value=nullcontext((None, "fixture", False))), \
             patch("backend.api.usage_limits.consume", return_value={}), \
             patch("backend.api.usage_limits.status", return_value={}), \
             patch("backend.api.answer_runs.persist_start", return_value=None), \
             patch("backend.api.answer_runs.persist_finish", return_value=None), \
             patch.object(generated, "registered_module_for", side_effect=candidate_lookup):
            client = TestClient(app, headers={"X-Guest-Session": "luna-candidate-chat-" + "x" * 32})
            first_response = client.post("/api/chat", json={
                "message": "합판거푸집 100제곱미터 비용 견적, 수직고 40미터, 양중장비는 사용하지 않음"})
            self.assertEqual(first_response.status_code, 200, first_response.text)
            first = first_response.json()
            answers_by_name = {"complexity": "복잡", "height": "40", "is_roof_slab": False,
                               "crane_used": False, "lifting_equipment_used": False}
            current = first
            asked_names = set()
            for _ in range(8):
                questions = current.get("questions", [])
                if not questions:
                    break
                asked_names.update(question["name"] for question in questions)
                answers = {question["name"]: answers_by_name[question["name"]]
                           for question in questions if question["name"] in answers_by_name}
                for question in questions:
                    if question["name"] == "work":
                        answers["work"] = next(choice for choice in question.get("choices", [])
                                               if "6-3-1" in choice)
                self.assertTrue(answers, questions)
                refs = {question["name"]: question["ref"] for question in questions
                        if question["name"] in answers}
                response = client.post("/api/chat", json={"thread_id": first["thread_id"],
                    "answers": answers, "refs": refs})
                self.assertEqual(response.status_code, 200, response.text)
                current = response.json()
            self.assertIn("lifting_equipment_used", asked_names)
            self.assertIn(current["status"], ("PARTIAL", "OK", "COMPUTED"), current)
            self.assertEqual(current["result"]["daily_volume"]["value"], "10.631060024624333494473", current)
            self.assertEqual(current["result"]["calculation_module"]["module_id"], candidate["id"])
            adjustment = current["result"]["daily_volume"]["adjustments"][0]
            self.assertEqual(adjustment["official_source"]["pdf_page"], 193)
            self.assertEqual(adjustment["official_source"]["printed_page"], 137)
            self.assertIn("9%까지", adjustment["official_source"]["text"])
            self.assertIn("System assumption:", adjustment["reason"])
            self.assertEqual(adjustment["policy_version"], "1")
            self.assertTrue(current["statement"]["totals"]["contract_amount"] > 0)
            before_total = current["statement"]["totals"]["contract_amount"]

            edited_response = client.post("/api/chat", json={"thread_id": first["thread_id"],
                "message": "물량을 120제곱미터로 변경해 다시 계산해 주세요"})
            self.assertEqual(edited_response.status_code, 200, edited_response.text)
            edited = edited_response.json()
            self.assertIn(edited["status"], ("PARTIAL", "OK", "COMPUTED"), edited)
            self.assertTrue(any(row["name"] == "volume" and row["value"] == "120" for row in edited["inputs"]))
            self.assertEqual(edited["result"]["daily_volume"]["value"], current["result"]["daily_volume"]["value"])
            after_total = edited["statement"]["totals"]["contract_amount"]
            self.assertNotEqual(before_total, after_total)

            export = client.get(f"/api/export/{first['thread_id']}.xlsx")
            self.assertEqual(export.status_code, 200, export.text)
            workbook = load_workbook(BytesIO(export.content), data_only=False)
            values = [" ".join(str(cell) for cell in row if cell is not None)
                      for sheet in workbook.worksheets for row in sheet.iter_rows(values_only=True)]
            evidence = "\n".join(values)
            self.assertIn("10.631060024624333494473", evidence)
            self.assertIn("System assumption:", evidence)
            self.assertIn("Intervals 11", evidence)
            self.assertIn("9%까지", evidence)
            self.assertIn("PDF p.193 (printed p.137)", evidence)
            self.assertIn("formwork-6-3-1-height-started-interval-compound-v1", evidence)
            self.assertIn(str(after_total), evidence)


if __name__ == "__main__":
    unittest.main()
