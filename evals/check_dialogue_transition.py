"""AGENT_MODE 전환: 기존 방식으로 저장된 대화는 tools 모드를 켠 뒤에도 기존 방식으로 조회·이어가기·Excel을 지원한다.

처리 방식은 대화에 저장된 상태로 정한다(새 흐름 상태가 있으면 tools, 기존 그래프 상태만 있으면 기존). 기존 견적을
초기화하거나 새 흐름 상태로 덮어쓰지 않는지 본다. 외부 호출 없음(AGENT_OFFLINE, AGENT_LLM=off).
회원 검사는 로컬 Supabase(CHAT_DATABASE_URL=127.0.0.1:54322)일 때만 한다.

사용: CHAT_DATABASE_URL=postgresql://postgres:postgres@127.0.0.1:54322/postgres python evals/check_dialogue_transition.py
"""

from __future__ import annotations

import io
import os
import sys
from contextlib import nullcontext
from pathlib import Path
from unittest.mock import patch
from uuid import uuid4

os.environ.update(AGENT_OFFLINE="1", AGENT_LLM="off", AGENT_MODE="")
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from fastapi import HTTPException  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from langgraph.checkpoint.postgres import PostgresSaver  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

import backend.api.main as api_main  # noqa: E402
from backend.agent.estimate import dialogue  # noqa: E402
from backend.api import chat_storage  # noqa: E402

START = {"message": "철근콘크리트 벽체 260㎥ 펌프차로 타설 비용", "basis_date": "2026-10-01"}
ANSWERS = {"work": "6-1-4", "pump_size": "32m", "slump_band": "15㎝", "facility_type": "Type-Ⅱ", "site_type": "Type-Ⅱ",
           "placement": "붐", "vibrator_used": True, "reset_status": "없음", "concrete_supply": "관급"}
CHANGE = {"work_category": "주택 외 건축", "duration": "13~36개월", "contractor_type": "전문건설업"}
SECRET = "transition-guest-session-" + "x" * 32


def amounts(content: bytes) -> list:
    book = load_workbook(io.BytesIO(content), data_only=True)
    return [cell.value for sheet in book for row in sheet.iter_rows() for cell in row if isinstance(cell.value, int)]


def legacy_reply(response: dict) -> bool:
    """기존 그래프 응답 모양(새 흐름 표시 없음)."""
    return "estimate_current" not in response and all("ref" not in q for q in response.get("questions") or [])


def flow(post, export, has_dialogue_state, label: str) -> list[tuple[str, bool]]:
    checks = []
    os.environ["AGENT_MODE"] = ""
    done_start = post(START)
    done = post({"thread_id": done_start["thread_id"], "answers": ANSWERS}, done_start)
    waiting = post(START)  # 질문 대기 중인 기존 대화
    before = done["statement"]["totals"]["contract_amount"]
    before_export = export(done["thread_id"])

    os.environ["AGENT_MODE"] = "tools"
    after_export = export(done["thread_id"])
    checks.append((f"{label} T1 전환 후 기존 대화 Excel: 200·같은 도급액",
                   before_export.status_code == after_export.status_code == 200
                   and before in amounts(after_export.content)))
    resumed = post({"thread_id": waiting["thread_id"], "answers": ANSWERS}, waiting)
    checks.append((f"{label} T2 전환 후 대기 중 기존 대화 이어가기: 기존 방식으로 같은 견적",
                   legacy_reply(resumed) and resumed["status"] == done["status"]
                   and resumed["statement"]["totals"]["contract_amount"] == before))
    changed = post({"thread_id": done["thread_id"], "conditions": CHANGE}, done)
    checks.append((f"{label} T3 전환 후 기존 대화 공사 조건 변경: 기존 견적을 유지한 채 재계산",
                   legacy_reply(changed) and changed["statement"] is not None
                   and changed["statement"]["totals"]["contract_amount"] != before
                   and changed["statement"]["totals"]["contract_amount"] in amounts(export(done["thread_id"]).content)))
    followup = post({"thread_id": done["thread_id"], "message": "300㎥로 바꾸면?"}, done)
    checks.append((f"{label} T4 전환 후 기존 대화에 새 질문: 새 흐름으로 덮어쓰지 않음(새 흐름 상태 없음)",
                   legacy_reply(followup) and not has_dialogue_state(done["thread_id"])
                   and not has_dialogue_state(waiting["thread_id"])))
    fresh = post({"message": "콘크리트 타설 1세제곱미터 품셈 알려줘"})
    checks.append((f"{label} T5 전환 후 새 대화는 tools 방식", "estimate_current" in fresh
                   and has_dialogue_state(fresh["thread_id"])))
    os.environ["AGENT_MODE"] = ""
    rolled_back = post({"thread_id": fresh["thread_id"], "message": "콘크리트공은 1세제곱미터당 몇 명이야?"}, fresh)
    checks.append((f"{label} T7 설정을 다시 끄면: 새 흐름으로 저장된 대화는 새 흐름 유지, 기존 대화는 기존 방식",
                   "estimate_current" in rolled_back and has_dialogue_state(fresh["thread_id"])
                   and legacy_reply(post({"thread_id": waiting["thread_id"], "conditions": CHANGE}, waiting))
                   and export(waiting["thread_id"]).status_code == 200))
    return checks


def main() -> int:
    checks = []
    with patch("backend.api.usage_limits.processing", return_value=nullcontext((None, "fixture", False))), \
            patch("backend.api.usage_limits.consume", return_value={}), \
            patch.object(api_main, "warmup_client", return_value=True), \
            TestClient(api_main.app, headers={"X-Guest-Session": SECRET}) as http:
        def guest_post(body: dict, previous: dict | None = None) -> dict:
            result = http.post("/api/chat", json=body)
            assert result.status_code == 200, (result.status_code, result.text[:300])
            return result.json()

        checks += flow(guest_post, lambda thread: http.get(f"/api/export/{thread}.xlsx"),
                       lambda thread: dialogue.load(api_main.DIALOGUE, thread) is not None, "[비회원]")

        if "127.0.0.1:54322" in os.environ.get("CHAT_DATABASE_URL", ""):
            chat_storage.setup()
            member = str(uuid4())
            with chat_storage.connection() as conn:
                conn.execute("INSERT INTO auth.users(id,aud,role,email) VALUES (%s,'authenticated','authenticated',%s)",
                             (member, f"transition-{member}@example.invalid"))
            auth = {"Authorization": "Bearer transition"}

            def identity(token):
                if token is None:
                    return None
                if token == "Bearer transition":
                    return member
                raise HTTPException(401, "invalid")

            def member_post(body: dict, previous: dict | None = None) -> dict:
                conversation = (previous or {}).get("thread_id") or str(uuid4())
                body = {k: v for k, v in body.items() if k != "thread_id"}
                result = http.post("/api/chat", headers=auth,
                                   json={"conversation_id": conversation, "request_id": str(uuid4()), **body})
                assert result.status_code == 200, (result.status_code, result.text[:300])
                return result.json()

            def member_state(thread: str) -> bool:
                with chat_storage.connection() as conn:
                    return dialogue.load(dialogue.build_dialogue_graph(PostgresSaver(conn)), thread) is not None

            try:
                with patch.object(chat_storage, "identity", side_effect=identity):
                    member_checks = flow(member_post, lambda thread: http.get(f"/api/export/{thread}.xlsx", headers=auth),
                                         member_state, "[회원]")
                    listed = http.get("/api/conversations", headers=auth).json()
                    transcripts = [http.get(f"/api/conversations/{c['id']}", headers=auth) for c in listed]
                member_checks.append(("[회원] T6 전환 후 기존 대화 목록·기록 조회",
                                      len(listed) >= 2 and all(t.status_code == 200 and t.json()["messages"]
                                                               for t in transcripts)))
                checks += member_checks
            finally:
                with chat_storage.connection() as conn:
                    conn.execute("DELETE FROM public.conversations WHERE user_id=%s", (member,))
                    conn.execute("DELETE FROM auth.users WHERE id=%s", (member,))
        else:
            print("회원 검사 건너뜀: CHAT_DATABASE_URL이 로컬 Supabase(127.0.0.1:54322)가 아님")

    for name, passed in checks:
        print(f"{'PASS' if passed else 'FAIL'} {name}")
    print(f"통과 {sum(passed for _, passed in checks)} / 전체 {len(checks)}")
    return 0 if all(passed for _, passed in checks) else 1


if __name__ == "__main__":
    raise SystemExit(main())
