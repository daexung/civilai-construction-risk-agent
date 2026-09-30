"""오프라인으로 채팅 API의 상태 전환과 thread_id 규칙을 검사한다."""

from __future__ import annotations

import os
import sys
from pathlib import Path

os.environ["AGENT_OFFLINE"] = "1"
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from fastapi.testclient import TestClient  # noqa: E402

from api.main import app  # noqa: E402

CLIENT = TestClient(app)

PUMP_ANSWERS = {"pump_size": "32m", "slump_band": "15㎝", "facility_type": "Type-Ⅱ", "site_type": "Type-Ⅱ",
                "placement": "붐", "vibrator_used": True, "reset_status": "없음", "concrete_supply": "관급",
                "work_category": "기타 토목공사", "duration": "1~6개월", "contractor_type": "종합건설업",
                "project_scale": "이 견적만"}


def main() -> int:
    checks = []

    outside = CLIENT.post("/api/chat", json={"message": "오늘 현장 날씨 어때?"}).json()
    checks.append(("A1", outside["status"] == "OUT_OF_SCOPE" and not outside["questions"]))

    evidence = CLIENT.post("/api/chat", json={"message": "합판거푸집 설치 인건비"}).json()
    checks.append(("A2", evidence["status"] == "MISSING_INFO" and bool(evidence["questions"])))

    missing = CLIENT.post("/api/chat", json={"message": "철근콘크리트 벽체 260㎥ 펌프차로 타설 비용",
                                              "basis_date": "2026-10-01"}).json()
    thread_id = missing["thread_id"]
    checks.append(("A3", missing["status"] == "MISSING_INFO" and len(missing["questions"]) == 8
                   and missing["work"]["section_no"] == "6-1-4"))

    computed = CLIENT.post("/api/chat", json={"thread_id": thread_id, "answers": PUMP_ANSWERS}).json()
    computed_lines = {line["name"]: line for line in computed.get("result", {}).get("lines", [])}
    checks.append(("A4", computed["status"] == "PARTIAL" and len(computed["inputs"]) == 14
                   and all(item["source"] in ("질문", "답변", "선택", "기본값") for item in computed["inputs"])
                   and computed_lines.get("콘크리트공", {}).get("value") == "8"
                   and computed_lines.get("콘크리트펌프차", {}).get("kind") == "equipment"
                   and computed_lines.get("콘크리트펌프차", {}).get("value") == "2"
                   and computed["result"]["review_status"] == "미완료"
                   and computed["statement"]["totals"] == {
                       "materials": 178360, "labor": 4922974, "expenses": 2263523,
                       "net_cost": 7364857, "management": 589188, "profit": 1166352,
                       "total_cost": 9120397, "vat": 912039, "contract_amount": 10032436
                   }))

    reused = CLIENT.post("/api/chat", json={"thread_id": thread_id, "message": "새 질문"}).json()
    checks.append(("A5", reused["thread_id"] != thread_id))

    health = CLIENT.get("/api/health").json()
    checks.append(("A6", health == {"status": "ok"}))

    blocked_start = CLIENT.post("/api/chat", json={"message": "철근콘크리트 벽체 260㎥ 펌프차로 타설 비용"}).json()
    blocked_answers = {**PUMP_ANSWERS, "reset_status": "있음"}
    blocked = CLIENT.post("/api/chat", json={"thread_id": blocked_start["thread_id"], "answers": blocked_answers}).json()
    checks.append(("A7", blocked["status"] == "BLOCKED" and blocked["result"]["input"] == "reset_status"
                   and bool(blocked["result"]["source"]) and blocked["message"] == blocked["result"]["reason"]))

    unit_lines = {line["name"]: line for line in computed["result"].get("unit_lines", [])}
    checks.append(("A8", computed["result"].get("unit_basis", {}).get("per") == "1㎥"
                   and computed["result"]["unit_basis"]["places"] == 4
                   and unit_lines.get("콘크리트공", {}).get("applied") == "0.0308"
                   and unit_lines.get("콘크리트펌프차", {}).get("exact") == "4/65"
                   and computed_lines["콘크리트공"]["value"] == "8"))

    checks.append(("A9 응답에 llm_info 포함",
                   computed.get("llm_info", {}).get("provider") == "vertex"
                   and computed["llm_info"].get("model") == "gemini-3.5-flash-lite"
                   and computed["llm_info"].get("attempts") == 0
                   and "VERTEX_API_KEY" not in str(computed["llm_info"])))

    # 조건 없이 계산: 기본값으로 바로 도급액이 나온다.
    start = CLIENT.post("/api/chat", json={"message": "철근콘크리트 벽체 260㎥ 펌프차로 타설 비용",
                                             "basis_date": "2026-10-01"}).json()
    answers = {k: v for k, v in PUMP_ANSWERS.items()
               if k not in ("work_category", "duration", "contractor_type", "project_scale")}
    base = CLIENT.post("/api/chat", json={"thread_id": start["thread_id"], "answers": answers}).json()
    condition_sources = {c["name"]: c["source"] for c in base["conditions"]}
    checks.append(("A10 기본값 도급액", base["status"] == "PARTIAL"
                   and base["statement"]["totals"]["contract_amount"] == 10032436
                   and set(condition_sources.values()) == {"기본값"}
                   and base["conditions"][0]["group"] == "토목"
                   and bool(base["conditions"][0]["help"]["기타 토목공사"])))

    exported = CLIENT.get(f"/api/export/{base['thread_id']}.xlsx")  # 조건을 바꾸기 전 상태
    if exported.status_code == 200:
        (ROOT / "evals" / "results" / "sample_견적서.xlsx").write_bytes(exported.content)
    changed = CLIENT.post("/api/chat", json={"thread_id": base["thread_id"], "conditions": {
        "work_category": "주택 외 건축", "duration": "13~36개월", "contractor_type": "전문건설업"}}).json()
    changed_sources = {c["name"]: c["source"] for c in changed["conditions"]}
    checks.append(("A11 조건 바꾸기", changed["statement"]["totals"]["contract_amount"] == 9947033
                   and changed["status"] == "PARTIAL" and changed_sources["duration"] == "선택"
                   and changed_sources["project_scale"] == "기본값"
                   and changed["result"] == base["result"]
                   and any(i["name"] == "volume" and i["value"] == "260" for i in changed["inputs"])))
    group_only = CLIENT.post("/api/chat", json={"thread_id": base["thread_id"],
                                                  "conditions": {"group": "건축"}}).json()
    bad = CLIENT.post("/api/chat", json={"thread_id": base["thread_id"],
                                           "conditions": {"duration": "없는 기간"}})
    checks.append(("A12 묶음만 바꾸기·잘못된 값", group_only["conditions"][0]["value"] == "주택 외 건축"
                   and bad.status_code == 422 and CLIENT.post("/api/chat", json={
                       "thread_id": "nothing", "conditions": {"duration": "1~6개월"}}).status_code == 404))

    from io import BytesIO
    from openpyxl import load_workbook
    book = load_workbook(BytesIO(exported.content))
    first = book["원가계산서"] if "원가계산서" in book.sheetnames else None
    cost_rows = list(first.iter_rows(values_only=True)) if first else []
    cost_header = next((row for row in cost_rows if row[0] == "비목"), None)
    contract = next((row for row in cost_rows if row[0] == "도급액"), None)
    net_cost = next((row for row in cost_rows if row[0] == "순공사원가"), None)
    excluded_rows = [row for row in cost_rows if isinstance(row[4], str) and row[4].startswith("제외 ·")]
    footer_present = any(isinstance(cell.value, str) and "검토 전 참고용" in cell.value
                         for row in first.iter_rows() for cell in row) if first else False
    checks.append(("A13 엑셀 시트·열", exported.status_code == 200
                   and len(book.sheetnames) == 6
                   and book.sheetnames == ["견적서", "원가계산서", "내역서", "일위대가", "단가대비표", "산출근거"]
                   and cost_header is not None and footer_present
                   and CLIENT.get("/api/export/nothing.xlsx").status_code == 404))
    checks.append(("A14 원가계산서 고정 셀", contract is not None and contract[2] == 10032436
                   and net_cost is not None and net_cost[2] == 7364857))
    checks.append(("A15 제외 줄 열 정렬", bool(excluded_rows)
                   and all(row[2] is None and row[4].startswith("제외 ·") for row in excluded_rows)
                   and all(len(row) == 5 for row in cost_rows)))
    unit_sheet = book["일위대가"] if "일위대가" in book.sheetnames else None
    unit_values = list(unit_sheet.iter_rows(values_only=True)) if unit_sheet else []
    unit_headers = next((row for row in unit_values if row[4] == "계"), None)
    unit_title = next((row for row in unit_values if isinstance(row[0], str)
                       and row[0].startswith("제 1호표 ")), None)
    checks.append(("A16 일위대가 호표 합계·열순서", unit_headers is not None and unit_title is not None
                   and unit_title[5] == 21734 and unit_headers.index("노무비") < unit_headers.index("재료비")))

    from urllib.parse import unquote
    disposition = exported.headers.get("content-disposition", "")
    encoded_name = disposition.partition("filename*=UTF-8''")[2]
    decoded_name = unquote(encoded_name)
    checks.append(("A17 공종명 파일명 헤더", 'filename="estimate.xlsx"; filename*=UTF-8\'\'' in disposition
                   and decoded_name.endswith("_견적서.xlsx")
                   and "/" not in decoded_name and "\\" not in decoded_name))

    bill_sheet = book["내역서"]
    bill_numeric = all(cell.value is None or isinstance(cell.value, (int, float))
                       for row in bill_sheet.iter_rows(min_row=6, min_col=5, max_col=12)
                       for cell in row)
    rate_sheet = book["단가대비표"]
    rate_numeric = all(cell.value is None or isinstance(cell.value, (int, float)) or cell.value == "-"
                       for row in rate_sheet.iter_rows(min_row=5, min_col=5, max_col=5)
                       for cell in row)
    basis_sheet = book["산출근거"]
    basis_numeric = all(cell.value is None or isinstance(cell.value, (int, float))
                        or cell.value in ("미산정", "제외")
                        for row in basis_sheet.iter_rows(min_row=5, min_col=2, max_col=2)
                        for cell in row)
    checks.append(("A18 수량·금액 숫자 셀", bill_numeric and rate_numeric and basis_numeric))

    bill_values = list(bill_sheet.iter_rows(values_only=True))
    bill_total = next((row for row in bill_values if row[0] == "합계"), None)
    checks.append(("A19 내역서 합계 금액 칸", bill_total is not None
                   and [bill_total[index] for index in (4, 6, 8, 10)] == [None, None, None, None]
                   and [bill_total[index] for index in (5, 7, 9, 11)] == [178360, 4133480, 1339000, 5650840]))

    missing_title_row = next((index for index, row in enumerate(cost_rows) if row[0] == "빠진 항목"), None)
    rate_title_row = next((index for index, row in enumerate(cost_rows) if row[0] == "적용 요율 출처"), None)
    missing_items = set()
    missing_rows = []
    if missing_title_row is not None:
        missing_end = rate_title_row if rate_title_row is not None else len(cost_rows)
        missing_rows = [row for row in cost_rows[missing_title_row + 2:missing_end] if row[0]]
        missing_items = {row[0] for row in missing_rows}
    checks.append(("A20 원가계산서 빠진 항목 분리", contract is not None and missing_title_row is not None
                   and missing_title_row > cost_rows.index(contract)
                   and all(not row[0] for row in cost_rows[cost_rows.index(contract) + 1:missing_title_row])
                   and any("1-2-9" in item for item in missing_items)
                   and any("살수 양생" in item for item in missing_items)
                   and all(any(expected in item for item in missing_items) for expected in
                           ("레미콘", "펌프차 연료", "퇴직공제", "산업안전보건관리비",
                            "공사이행보증", "건설기계대여대금"))))

    contract_row = cost_rows.index(contract) + 1 if contract else 0
    checks.append(("A21 원가계산서 요약 행 병합", contract_row > 0
                   and f"A{contract_row}:B{contract_row}" in {str(rng) for rng in first.merged_cells.ranges}
                   and "-" in next((row[2] for row in cost_rows if row[1] == "간접재료비(-)"), None)))

    metadata_ok = all(isinstance(sheet["A2"].value, str) and sheet["A2"].value.startswith("공사명:")
                      and isinstance(sheet["A3"].value, str) and sheet["A3"].value.startswith("기준일:")
                      for sheet in book.worksheets if sheet.title != "견적서")
    checks.append(("A22 시트 공통 메타데이터", metadata_ok
                   and "금액 0.1원 미만 버림(품셈 1-2-2)" in [cell.value for row in unit_sheet.iter_rows() for cell in row]))
    labor_rows = [row for row in unit_sheet.iter_rows(min_row=6)
                  if isinstance(row[12].value, str) and "노임 코드" in row[12].value]
    driver_rows = [row for row in unit_sheet.iter_rows(min_row=6) if row[0].value == "펌프차 운전원"]
    driver_pay = [cell for row in driver_rows for cell in row[4:12] if cell.value == 35415.375]
    checks.append(("A23 일위대가 인력 규격·운전원 노임", bool(labor_rows) and bool(driver_pay)
                   and all(row[1].value in (None, "") for row in labor_rows)
                   and driver_pay[0].number_format == "#,##0.0"
                   and "일 283,323원 ÷ 8시간" in driver_rows[0][12].value))

    all_cells = [cell.value for sheet in book.worksheets for row in sheet.iter_rows() for cell in row]
    checks.append(("A24 내부 데이터 경로 비노출", not any(isinstance(value, str) and "data/raw/" in value
                                                        for value in all_cells)))

    rate_values = list(rate_sheet.iter_rows(values_only=True))
    driver_rate = next((row for row in rate_values if row[1] == "펌프차 운전원"), None)
    supplied = next((row for row in rate_values if row[0] == "자재"), None)
    wage_rows = [row for row in rate_values[4:] if row[0] == "노임"]
    machine_rows = [row for row in rate_values[4:] if row[0] == "기계경비"]
    checks.append(("A25 단가대비표 실무 열과 단가", rate_values[3] ==
                   ("구분", "품명", "규격", "단위", "적용 단가", "단가 근거", "적용 기간", "비고")
                   and driver_rate is not None and driver_rate[3] == "원/인" and driver_rate[4] == 283323
                   and "35,415.4원" in driver_rate[7]
                   and bool(wage_rows) and all(row[2] in (None, "") and row[3] == "원/인" for row in wage_rows)
                   and bool(machine_rows) and all(row[3] == "원/hr" for row in machine_rows)
                   and supplied is not None and supplied[4] == "-" and supplied[7] == "관급(발주처 지급)"
                   and all(row[0] != "제비율" for row in rate_values[4:])))

    basis_values = list(basis_sheet.iter_rows(values_only=True))
    checks.append(("A26 산출근거 범위와 적용 기준", not any(value in ("간접노무비", "산재보험료")
                   for row in basis_values for value in row)
                   and any(row[0] == "공구손료 및 경장비 기준액 (원)" and row[1] == 4922974
                           and "× 5%" in row[2] for row in basis_values)
                   and sum(row[0] == "적용 기준" for row in basis_values) == 1
                   and any(row[0] == "공사 규모 판정" for row in basis_values)
                   and any(row[0] == "안전관리비 기초액" for row in basis_values)))

    source_rows = []
    source_headers = None
    if rate_title_row is not None:
        source_headers = cost_rows[rate_title_row + 1]
        for row in cost_rows[rate_title_row + 2:]:
            if row[0] == "표준품셈 기준 금액 · 검토 전 참고용 · 품셈AI":
                break
            if row[0]:
                source_rows.append(row)
    source_names = {row[0] for row in source_rows}
    indirect = next((row for row in source_rows if row[0] == "간접노무비"), None)
    profit_rate = next((row for row in source_rows if row[0] == "이윤"), None)
    vat_row = next((row for row in source_rows if row[0] == "부가가치세"), None)
    checks.append(("A27 적용 요율 출처의 실제 적용 행", source_headers is not None
                   and source_headers[:4] == ("항목", "요율", "적용 구간", "출처")
                   and indirect is not None and indirect[1] == "19.1%"
                   and indirect[2] == "10억 미만·6개월 이하"
                   and indirect[3] == "조달청 제비율 2026.4.13., 토목"
                   and profit_rate is not None and profit_rate[2] == "50억 미만"
                   and vat_row is not None and vat_row[1] == "10%" and vat_row[3] == "부가가치세법"
                   and {"노임", "기계경비"} <= source_names
                   and not any(any(term in name for term in ("레미콘", "연료", "1-2-9", "살수 양생"))
                               for name in source_names)))

    missing_by_name = {row[0]: row[2] for row in missing_rows}
    reasons_ok = (any("레미콘" in name and "관급" in reason for name, reason in missing_by_name.items())
                  and any("연료" in name and "8-1-7" in reason for name, reason in missing_by_name.items())
                  and any("1-2-9" in name and "이번 계산 범위 밖" in reason
                          for name, reason in missing_by_name.items())
                  and any("살수 양생" in name and "6-1-4 사." in reason
                          for name, reason in missing_by_name.items()))
    checks.append(("A28 빠진 항목별 실제 근거", reasons_ok))
    required_reasons = (("퇴직공제", "추정금액 1억 미만"),
                        ("산업안전보건관리비", "총 공사금액 2천만원 미만"),
                        ("공사이행보증", "지방계약법 시행령 제51조"),
                        ("건설기계대여대금", "산정 방법 확인 필요"))
    reasons_present = all(any(keyword in name and token in reason
                              for name, reason in missing_by_name.items())
                          for keyword, token in required_reasons)
    checks.append(("A29 제외·미산정 사유", reasons_present))

    estimate = book["견적서"] if "견적서" in book.sheetnames else None
    estimate_values = {row[0]: row[1] for row in estimate.iter_rows(min_row=1, values_only=True)
                       if row and isinstance(row[0], str)} if estimate else {}
    estimate_items = {estimate.cell(row, 1).value: estimate.cell(row, 2).value
                      for row in range(1, estimate.max_row + 1)} if estimate else {}
    estimate_total = estimate_items.get("합 계")
    estimate_components = sum(estimate_items.get(name) or 0 for name in
                              ("재료비", "노무비", "경비", "일반관리비", "이윤", "부가가치세"))
    estimate_amount_line = estimate["A6"].value if estimate else ""
    checks.append(("A30 견적서 요약 시트·금액", estimate is not None
                   and estimate_total == 10032436 and estimate_items.get("공급가액") == 9120397
                   and estimate_total == estimate_components
                   and "일금 일천삼만이천사백삼십육원정 (₩10,032,436)" in estimate_amount_line
                   and estimate["A2"].value.startswith("공사명 :")
                   and estimate["A3"].value.startswith("기 준 일 :")
                   and estimate.page_setup.orientation == "portrait"
                   and estimate.page_setup.fitToWidth == 1 and estimate.page_setup.fitToHeight == 1))
    from api.tables import won_in_korean
    checks.append(("A31 한글 금액 단위", [won_in_korean(value) for value in
                   (10032436, 9947033, 10319789, 100000000, 1005)] ==
                   ["일천삼만이천사백삼십육", "구백구십사만칠천삼십삼", "일천삼십일만구천칠백팔십구",
                    "일억", "일천오"]))

    first_sagup = CLIENT.post("/api/chat", json={"message": "철근콘크리트 벽체 260㎥ 펌프차로 타설 비용",
                                                  "basis_date": "2026-10-01"}).json()
    sagup_answers = {**PUMP_ANSWERS, "concrete_supply": "사급"}
    sagup_question = CLIENT.post("/api/chat", json={"thread_id": first_sagup["thread_id"],
                                                       "answers": sagup_answers}).json()
    ready_mix_question = next((question for question in sagup_question.get("questions", [])
                               if question["name"] == "ready_mix_price"), None)
    sagup_result = CLIENT.post("/api/chat", json={"thread_id": first_sagup["thread_id"],
                                                    "answers": {"ready_mix_price": "90000"}}).json()
    sagup_line = next((line for line in sagup_result.get("priced", {}).get("supply_lines", [])
                       if line.get("name") == "레미콘(사급)"), {})
    checks.append(("A32 사급 단가 조건 질문과 API 계산", ready_mix_question is not None
                   and ready_mix_question["choices"] == ["모름"]
                   and sagup_result["status"] == "PARTIAL"
                   and sagup_line.get("quantity") == "1.01"
                   and sagup_line.get("unit_price") == "90000"
                   and sagup_line.get("amount") == "90900.0"
                   and sagup_result["priced"]["unit_prices"] == {
                       "재료비": "91586", "노무비": "15898", "경비": "5150"}
                   and sagup_result["priced"]["total"] == "112634"
                   and sagup_result["statement"]["totals"]["contract_amount"] == 41684599
                   and any(citation.get("image_url") == "/api/source/p78.png"
                           for citation in sagup_line.get("citations", []))
                   and CLIENT.get("/api/source/p78.png").status_code == 200))
    sagup_export = CLIENT.get(f"/api/export/{sagup_result['thread_id']}.xlsx")
    if sagup_export.status_code == 200:
        (ROOT / "evals" / "results" / "sample_견적서_사급.xlsx").write_bytes(sagup_export.content)
    sagup_book = load_workbook(BytesIO(sagup_export.content)) if sagup_export.status_code == 200 else None
    sagup_unit = sagup_book["일위대가"] if sagup_book else None
    ready_unit = next((row for row in sagup_unit.iter_rows(values_only=True)
                       if row[0] == "레미콘(사급)"), None) if sagup_unit else None
    sagup_rate = sagup_book["단가대비표"] if sagup_book else None
    rate_values = list(sagup_rate.iter_rows(values_only=True)) if sagup_rate else []
    ready_rate = next((row for row in rate_values if row[0] == "자재" and row[1] == "레미콘(사급)"), None)
    estimate_sheet = sagup_book["견적서"] if sagup_book else None
    estimate_material_note = next((estimate_sheet.cell(row, 3).value for row in
                                    range(1, estimate_sheet.max_row + 1)
                                    if estimate_sheet.cell(row, 1).value == "재료비"), "") if estimate_sheet else ""
    sagup_excel_checks = {
        "export": sagup_export.status_code == 200,
        "unit row": ready_unit is not None,
        "unit spec": ready_unit is not None and ready_unit[1] == "(사급)",
        "unit quantity": ready_unit is not None and ready_unit[2] == 1.01 and ready_unit[3] == "㎥",
        "unit price and amount": ready_unit is not None and ready_unit[8] == 90000 and ready_unit[9] == 90900,
        "unit note": ready_unit is not None and "1-3-1 p78" in str(ready_unit[12]),
        "rate row": ready_rate is not None,
        "rate price": ready_rate is not None and ready_rate[2] == "(사급)"
        and ready_rate[3] == "원/㎥" and ready_rate[4] == 90000,
        "rate source": ready_rate is not None and ready_rate[5] == "사용자 입력(부가세 제외)",
        "rate note": ready_rate is not None and ready_rate[7] == "할증 1%(품셈 1-3-1)",
        "estimate note": estimate_material_note == "레미콘 사급(사용자 입력 90,000원/㎥)",
        "sample saved": (ROOT / "evals" / "results" / "sample_견적서_사급.xlsx").is_file(),
    }
    checks.append(("A33 사급 엑셀 세 시트 표시와 샘플 저장", all(sagup_excel_checks.values())))
    if not all(sagup_excel_checks.values()):
        print("A33 상세:", [name for name, passed in sagup_excel_checks.items() if not passed])

    first_unknown = CLIENT.post("/api/chat", json={"message": "철근콘크리트 벽체 260㎥ 펌프차로 타설 비용",
                                                    "basis_date": "2026-10-01"}).json()
    unknown_question = CLIENT.post("/api/chat", json={"thread_id": first_unknown["thread_id"],
                                                        "answers": sagup_answers}).json()
    unknown_result = CLIENT.post("/api/chat", json={"thread_id": first_unknown["thread_id"],
                                                      "answers": {"ready_mix_price": "모름"}}).json()
    unknown_line = next((line for line in unknown_result.get("priced", {}).get("supply_lines", [])
                         if "레미콘" in line.get("name", "")), {})
    checks.append(("A34 사급 단가 모름 및 관급 미질문", unknown_result["status"] == "PARTIAL"
                   and unknown_question.get("questions", [{}])[0].get("name") == "ready_mix_price"
                   and unknown_line.get("status") == "미산정" and unknown_line.get("amount") is None
                   and unknown_line.get("reason") == "사급 레미콘 단가 미입력"
                   and unknown_result["priced"]["unit_prices"] == {
                       "재료비": "686", "노무비": "15898", "경비": "5150"}
                   and not any(question["name"] == "ready_mix_price" for question in missing.get("questions", []))))

    draft_cases = [
        ("R1", "레디믹스트 콘크리트 100㎥ 인력운반 타설 비용", "6-1-1",
         {"placement_method": "인력운반 타설", "structure": "철근구조물", "volume": "100",
          "scattered_small_volume": False, "concrete_supply": "관급"}, "68693", "6869300", 13150196),
        ("M1", "현장비빔타설 기계비빔 철근구조물 100㎥ 비용", "6-1-2",
         {"mixing_type": "기계비빔타설", "structure": "수량 철근구조물", "volume": "100"},
         "164403", "16440300", 31648540),
        ("S1", "콘크리트 표면 마무리 200㎡ 비용", "6-1-3",
         {"area": "200"}, "948", "189600", 364982),
    ]
    for label, query, section, values, unit_total, direct_total, contract in draft_cases:
        first = CLIENT.post("/api/chat", json={"message": query, "basis_date": "2026-10-01"}).json()
        answers = {**values, **({"work": section} if any(q["name"] == "work" for q in first["questions"]) else {})}
        result = CLIENT.post("/api/chat", json={"thread_id": first["thread_id"],
                                                  "answers": answers}).json()
        priced = result.get("priced") or {}
        statement = result.get("statement") or {}
        ok = (first["status"] == "MISSING_INFO" and result.get("work", {}).get("section_no") == section
              and result.get("status") in ("OK", "PARTIAL")
              and result.get("result", {}).get("review_status") == "AI 초안 · 검토 전"
              and "AI가 품셈 원문으로 만든 계산 초안(검토 전)" in result.get("answer", "")
              and priced.get("total") == unit_total
              and (priced.get("reference_amounts") or {}).get("total") == direct_total
              and (statement.get("totals") or {}).get("contract_amount") == contract)
        if section == "6-1-2":
            structure_question = next((q for q in first.get("questions", [])
                                       if q["name"] == "structure"), {})
            citation_urls = [citation.get("image_url") for line in result.get("result", {}).get("unit_lines", [])
                             for citation in line.get("citations", [])]
            ok = ok and structure_question.get("labels", {}).get("수량 철근구조물") == "철근구조물" \
                and "/api/source/p185-t1.png" in citation_urls \
                and CLIENT.get("/api/source/p185-t1.png").status_code == 200
        checks.append((f"A35 {label} 초안 전체 흐름과 손계산", ok))
        if not ok:
            print("A35 상세:", label, first.get("status"), first.get("work"),
                  [q["name"] for q in first.get("questions", [])], result.get("status"),
                  result.get("work"), [q["name"] for q in result.get("questions", [])],
                  priced.get("total"), priced.get("reference_amounts"), statement.get("totals"),
                  result.get("message"))
        if label == "S1" and result.get("status") in ("OK", "PARTIAL"):
            export = CLIENT.get(f"/api/export/{result['thread_id']}.xlsx")
            book = load_workbook(BytesIO(export.content)) if export.status_code == 200 else None
            checks.append(("A36 초안 엑셀 모든 시트 표시", book is not None
                           and all("AI 초안 · 검토 전" in str(sheet["A3"].value) for sheet in
                                   list(book.worksheets)[1:])
                           and "AI 초안 · 검토 전" in str(book["견적서"]["A4"].value)
                           and (result.get("tables") or {}).get("bill", {}).get("unit") == "㎡"))

    epoxy_query = "에폭시 콘크리트 접착제 바르기 100㎡ 비용"
    epoxy_inputs = {"area": "100", "type": "신구-콘크리트 접착제바르기",
                    "ceiling_applied": False, "scaffold_used": False,
                    "floor_level": "지하층 및 1∼3층", "floor_level_19_plus": 19,
                    "thickness_adjusted": False, "thickness": "1"}
    for label, ceiling, scaffold, expected_unit, expected_contract in (
            ("E0", False, False, "32830", 6284801),
            ("E1", True, False, "39396", 7541741),
            ("E2", True, True, None, None)):
        first = CLIENT.post("/api/chat", json={"message": epoxy_query,
                                                "basis_date": "2026-10-01"}).json()
        result = CLIENT.post("/api/chat", json={"thread_id": first["thread_id"],
                                                  "answers": {**epoxy_inputs,
                                                              "ceiling_applied": ceiling,
                                                              "scaffold_used": scaffold}}).json()
        if label == "E2":
            question = next((item for item in result.get("questions", [])
                             if item["name"] == "apply_adj_2"), {})
            ok = (result["status"] == "MISSING_INFO" and question.get("choices") == ["예", "아니오"]
                  and not question.get("default")
                  and any(citation["internal_id"] == "p188-x6"
                          for citation in question.get("citations", [])))
            answered = CLIENT.post("/api/chat", json={"thread_id": first["thread_id"],
                                                      "answers": {"apply_adj_2": "예"}}).json()
            ok = ok and answered["status"] == "PARTIAL" and answered["priced"]["total"] == "39396"
        else:
            priced = result.get("priced") or {}
            statement = result.get("statement") or {}
            lines = {line["name"]: line for line in result.get("result", {}).get("unit_lines", [])}
            ok = (result["status"] == "PARTIAL" and priced.get("total") == expected_unit
                  and (statement.get("totals") or {}).get("contract_amount") == expected_contract
                  and lines.get("Epoxy신구-콘크리트접착제", {}).get("applied") == "1.2"
                  and lines.get("시너", {}).get("applied") == "0.2")
        checks.append((f"A37 {label} 에폭시", ok))
        if label == "E1" and ok:
            exported = CLIENT.get(f"/api/export/{result['thread_id']}.xlsx")
            book = load_workbook(BytesIO(exported.content)) if exported.status_code == 200 else None
            unit_names = [row[0] for row in book["일위대가"].values] if book else []
            basis_text = [" ".join(str(cell) for cell in row if cell is not None)
                          for row in book["산출근거"].values] if book else []
            checks.append(("A38 E1 엑셀 재료와 할증 근거", book is not None
                           and "Epoxy신구-콘크리트접착제" in unit_names
                           and "시너" in unit_names
                           and any("도장공 가산" in row and "20%" in row for row in basis_text)))

    for name, passed in checks:
        print(f"{'PASS' if passed else 'FAIL'} {name}")
    print(f"통과 {sum(passed for _, passed in checks)} / 전체 {len(checks)}")
    if not all(passed for _, passed in checks):
        print("A3:", missing)
        print("A4:", computed)
        print("A7:", blocked)
    return 0 if all(passed for _, passed in checks) else 1


if __name__ == "__main__":
    raise SystemExit(main())
